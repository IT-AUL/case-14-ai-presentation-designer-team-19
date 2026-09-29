"""Design DNA из пакета PPTX/POTX: declared vs observed (contract D1).

Всё вычисляется из самого пакета — ни имён, ни индексов, ни цветов
конкретных шаблонов здесь нет; на «невиданном» шаблоне, где макетов
почти нет, поля либо заполняются эвристикой по слайдам, либо остаются
честно пустыми.

* **declared** — что шаблон декларирует: макеты (имя, тип, плейсхолдеры
  с bbox в долях слайда), мастера и их макеты, палитры и шрифты тем;
* **observed** — что на слайдах реально: шкала кеглей, шрифты с
  контекстом, цвета, сетка и поля, фоны;
* **anchors** — повторяющиеся служебные элементы (логотип, колонтитул,
  номер, дата) с долей слайдов, где они есть;
* **slide_roles / exemplars** — роль каждого слайда (title, section, toc,
  content, contact, closing, blank) с доказательством и уверенностью;
* **capacities** — сколько текста шаблон реально вмещает;
* **unsupported_features** — что компилятор не переписывает.

Единицы: bbox — доли слайда [0, 1] (как ``AuditIssue.bbox``), отступы —
EMU (как в ``schemas/design-dna.schema.json``).
"""

from __future__ import annotations

import asyncio
import logging
import math
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pydantic import BaseModel, Field

from deckdna.contracts import design_dna as dna
from deckdna.providers.base import ModelGateway
from deckdna.template import autopsy
from deckdna.template.meta import is_meta_xml

logger = logging.getLogger(__name__)

A = autopsy.A
P = autopsy.P
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A, "p": P}

EMU_PER_PT = 12700
_SIZE_STEP = 0.5  # дробные кегли автоподбора (8.12, 6.75) приводятся к шагу 0.5pt
_MAX_FONT_SIZES = 16
_MAX_COLORS = 20
_MAX_ANCHORS = 8
_MIN_ANCHOR_COVERAGE = 0.15
_MIN_LOGO_COVERAGE = 0.3
_MAX_LOGOS = 3
_NEUTRALS = {"000000", "FFFFFF"}

# Шрифты, которые есть в любой офисной системе: их отсутствие во
# вложенных шрифтах не считается проблемой (стандартный набор ОС).
_SYSTEM_FONTS = frozenset(
    f.lower()
    for f in (
        "Arial", "Calibri", "Cambria", "Times New Roman", "Courier New",
        "Verdana", "Tahoma", "Georgia", "Segoe UI", "Helvetica", "Consolas",
        "Trebuchet MS", "Century Gothic", "Comic Sans MS", "Impact",
        "Palatino Linotype", "Garamond", "Symbol", "Wingdings", "Cambria Math",
        "Calibri Light", "Arial Narrow", "Lucida Sans", "Book Antiqua",
        "Franklin Gothic Book", "Gill Sans MT", "Rockwell", "Candara",
    )
)

_LAYOUT_TYPES = {
    "title": "title",
    "secHead": "section",
    "blank": "blank",
    "titleOnly": "title_only",
    "twoObj": "two_content",
    "twoTxTwoObj": "comparison",
    "objTx": "picture",
    "picTx": "picture",
    "tx": "content",
    "obj": "content",
    "objOnly": "content",
    "tbl": "table",
    "chart": "chart",
    "dgm": "diagram",
    "fourObj": "content",
    "vertTx": "content",
}

# Слова-признаки ролей. Это лексика разметки презентаций, а не имена
# шаблонов; для языков, которых здесь нет, роль определится по
# структуре и позиции — с меньшей уверенностью.
_TOC_WORDS = re.compile(
    r"содержание|оглавление|повестка|agenda|contents|table of contents|план", re.I
)
_CLOSING_WORDS = re.compile(
    r"спасибо|благодар|вопрос|thank|questions|q\s*&\s*a|the end", re.I
)
_CONTACT_WORDS = re.compile(r"контакт|связаться|contact|обратная связь", re.I)
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE = re.compile(r"(?:\+?\d[\d\s().-]{8,}\d)")
_URL = re.compile(r"(?:https?://|www\.)\S+|\b[\w-]+\.(?:ru|com|org|io|net)\b", re.I)


# --------------------------------------------------------------------------
# Утилиты
# --------------------------------------------------------------------------


def _step(pt: float) -> float:
    return round(pt / _SIZE_STEP) * _SIZE_STEP


def _percentile(values: list[float], q: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    k = max(0, min(len(ordered) - 1, math.ceil(q * len(ordered)) - 1))
    return ordered[k]


def _ph_type(shape: Any) -> str | None:
    """Тип плейсхолдера как в OOXML (``title``, ``body``, ``pic``, …); у
    ``p:ph`` без ``type`` — ``obj`` (умолчание стандарта). None — не плейсхолдер."""
    found = shape._element.xpath(".//p:nvPr/p:ph")
    if not found:
        return None
    return found[0].get("type") or "obj"


def _ph_idx(shape: Any) -> int | None:
    found = shape._element.xpath(".//p:nvPr/p:ph")
    if not found or found[0].get("idx") is None:
        return None
    try:
        return int(found[0].get("idx"))
    except ValueError:
        return None


def _iter_shapes(shapes: Any):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _frac_bbox(shape: Any, sw: int, sh: int) -> dna.Bbox | None:
    try:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
    except (AttributeError, ValueError):
        return None
    if None in (left, top, width, height) or not sw or not sh:
        return None
    return dna.Bbox(x=left / sw, y=top / sh, w=width / sw, h=height / sh)


def _hex(el: Any) -> str | None:
    """Цвет из a:srgbClr / a:sysClr(lastClr) внутри заливки."""
    if el is None:
        return None
    for child in el:
        tag = etree.QName(child).localname
        if tag == "srgbClr" and child.get("val"):
            return child.get("val").upper()
        if tag == "sysClr" and child.get("lastClr"):
            return child.get("lastClr").upper()
    return None


# --------------------------------------------------------------------------
# Темы, мастера, макеты (declared)
# --------------------------------------------------------------------------


def _theme_colors(zf: zipfile.ZipFile, part: str) -> dict[str, str]:
    root = etree.fromstring(zf.read(part))
    scheme = root.find(f".//{{{A}}}clrScheme")
    colors: dict[str, str] = {}
    if scheme is None:
        return colors
    for slot in scheme:
        value = _hex(slot)
        if value:
            colors[etree.QName(slot).localname] = value
    return colors


def _clr_map(prs: Any) -> dict[str, str]:
    """clrMap мастера: bg1→lt1, tx1→dk1, … (для резолва schemeClr)."""
    try:
        master = prs.slide_masters[0]
    except IndexError:
        return {}
    node = master.element.find(f"{{{P}}}clrMap")
    return dict(node.attrib) if node is not None else {}


def _resolve_scheme(name: str, clr_map: dict[str, str], palette: dict[str, str]) -> str | None:
    slot = clr_map.get(name, name)
    return palette.get(slot) or palette.get(name)


_NAME_HINTS = (
    ("title", re.compile(r"титул|title slide|cover|обложк", re.I)),
    ("closing", re.compile(r"спасибо|thank|финал|завершен|closing|the end", re.I)),
    ("section", re.compile(r"раздел|section|divider|глава|chapter", re.I)),
    ("blank", re.compile(r"пуст|blank|чист", re.I)),
)


def _layout_type(layout: Any, ph_types: list[str]) -> str:
    raw = layout.element.get("type")
    if raw in _LAYOUT_TYPES:
        return _LAYOUT_TYPES[raw]
    # cust / отсутствует: сначала объявленное имя макета (декларация
    # шаблона), затем состав плейсхолдеров
    name = layout.name or ""
    for kind, pattern in _NAME_HINTS:
        if pattern.search(name):
            return kind
    if "ctrTitle" in ph_types:
        return "title"
    if not [t for t in ph_types if t not in {"dt", "ftr", "sldNum"}]:
        return "blank"
    non_title = [t for t in ph_types if t not in {"title", "dt", "ftr", "sldNum"}]
    if not non_title:
        return "title_only"
    return "content"


def _placeholders(shapes: Any, sw: int, sh: int) -> list[dna.Placeholder]:
    out: list[dna.Placeholder] = []
    for shape in shapes:
        ptype = _ph_type(shape)
        if ptype is None:
            continue
        out.append(
            dna.Placeholder(type=ptype, idx=_ph_idx(shape), bbox=_frac_bbox(shape, sw, sh))
        )
    return out


def _declared(prs: Any, zf: zipfile.ZipFile, parts: dict[str, list[str]], sw: int, sh: int):
    themes = []
    theme_fonts = _theme_fonts(zf, parts["themes"])
    for part in parts["themes"]:
        major, minor = theme_fonts[part]
        themes.append(
            dna.Theme(
                part=part,
                major_font=major or "unspecified",
                minor_font=minor or "unspecified",
                colors=_theme_colors(zf, part),
            )
        )

    masters: list[dna.Master] = []
    layouts: list[dna.Layout] = []
    for master in prs.slide_masters:
        master_part = str(master.part.partname).lstrip("/")
        layout_parts = [str(layout.part.partname).lstrip("/") for layout in master.slide_layouts]
        masters.append(
            dna.Master(
                part=master_part,
                layout_parts=layout_parts,
                placeholders=_placeholders(master.placeholders, sw, sh),
            )
        )
        for layout in master.slide_layouts:
            placeholders = _placeholders(layout.placeholders, sw, sh)
            ph_types = [p.type for p in placeholders]
            layouts.append(
                dna.Layout(
                    part=str(layout.part.partname).lstrip("/"),
                    name=layout.name or "unnamed",
                    type=_layout_type(layout, ph_types),
                    matching_name=layout.element.get("matchingName") or None,
                    placeholders=placeholders,
                )
            )
    return dna.Declared(themes=themes, masters=masters, layouts=layouts), theme_fonts


_ThemeFonts = dict[str, tuple[str | None, str | None]]


def _theme_fonts(zf: zipfile.ZipFile, theme_parts: list[str]) -> _ThemeFonts:
    out: _ThemeFonts = {}
    for part in theme_parts:
        root = etree.fromstring(zf.read(part))
        fonts: dict[str, str | None] = {}
        for tag in ("majorFont", "minorFont"):
            node = root.find(f".//{{{A}}}fontScheme/{{{A}}}{tag}/{{{A}}}latin")
            fonts[tag] = node.get("typeface") if node is not None else None
        out[part] = (fonts["majorFont"], fonts["minorFont"])
    return out


# --------------------------------------------------------------------------
# Признаки слайдов
# --------------------------------------------------------------------------


@dataclass
class _SlideFeatures:
    index: int
    part: str
    layout_part: str | None
    layout_type: str
    text_chars: int = 0
    text_shapes: int = 0
    pictures: int = 0
    has_table: bool = False
    has_chart: bool = False
    has_diagram: bool = False
    title: str = ""
    max_size: float = 0.0
    paragraphs: list[str] = field(default_factory=list)
    all_text: str = ""
    table_dims: list[tuple[int, int]] = field(default_factory=list)
    body_paragraph_counts: list[int] = field(default_factory=list)
    words_per_paragraph: list[int] = field(default_factory=list)
    title_chars: int | None = None
    body_chars: int = 0


def _run_size(rpr: Any) -> float | None:
    if rpr is None or not rpr.get("sz"):
        return None
    return int(rpr.get("sz")) / 100


def _slide_features(slide: Any, index: int, layout_types: dict[str, str]) -> _SlideFeatures:
    layout_part = str(slide.slide_layout.part.partname).lstrip("/") if slide.slide_layout else None
    feat = _SlideFeatures(
        index=index,
        part=str(slide.part.partname).lstrip("/"),
        layout_part=layout_part,
        layout_type=layout_types.get(layout_part or "", "content"),
    )
    title_shape = None
    biggest: tuple[float, str] = (0.0, "")
    texts: list[str] = []
    for shape in _iter_shapes(slide.shapes):
        ptype = _ph_type(shape)
        if ptype in {"dt", "ftr", "sldNum"}:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            feat.pictures += 1
        if getattr(shape, "has_table", False) and shape.has_table:
            feat.has_table = True
            feat.table_dims.append((len(shape.table.rows), len(shape.table.columns)))
        if getattr(shape, "has_chart", False) and shape.has_chart:
            feat.has_chart = True
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        feat.text_shapes += 1
        feat.text_chars += len(re.sub(r"\s", "", text))
        texts.append(text)
        paragraphs = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
        feat.paragraphs.extend(paragraphs)
        size = 0.0
        for run in shape.text_frame._txBody.iter(f"{{{A}}}r"):
            s = _run_size(run.find(f"{{{A}}}rPr"))
            if s:
                size = max(size, s)
        if size > biggest[0]:
            biggest = (size, text)
        feat.max_size = max(feat.max_size, size)
        if ptype in {"title", "ctrTitle"} and title_shape is None:
            title_shape = shape
            feat.title = text
            feat.title_chars = len(text)
        else:
            feat.body_chars += len(text)
            if len(paragraphs) >= 2 or ptype in {"body", "obj"}:
                feat.body_paragraph_counts.append(len(paragraphs))
            feat.words_per_paragraph.extend(len(p.split()) for p in paragraphs)
    if not feat.title and biggest[1]:
        feat.title = biggest[1].splitlines()[0]
    feat.all_text = "\n".join(texts)
    return feat


# --------------------------------------------------------------------------
# Роли слайдов
# --------------------------------------------------------------------------


def _classify(feat: _SlideFeatures, total: int, median_size: float) -> tuple[str, float, str]:
    """(роль, уверенность, доказательство) слайда. Только по признакам
    самого слайда и его макета."""
    text = feat.all_text
    last = feat.index >= total - 1
    if feat.text_chars == 0 and not feat.has_table and not feat.has_chart:
        if feat.pictures:
            return "content", 0.6, "слайд только с изображениями, без текста"
        return "blank", 0.9, "на слайде нет текста, таблиц, диаграмм и изображений"
    head = (feat.title or text)[:120]
    if _TOC_WORDS.search(head) and len(feat.paragraphs) >= 4:
        return "toc", 0.85, f"заголовок «{head[:40]}» и {len(feat.paragraphs)} пунктов списка"
    if feat.layout_type == "title" and feat.index <= max(1, total // 10):
        return "title", 0.95, "макет-титул в начале колоды"
    if feat.index == 0 and feat.text_shapes <= 4 and not feat.has_table:
        return "title", 0.8, "первый слайд с короткой подачей"
    contacts = bool(_EMAIL.search(text) or _PHONE.search(text) or _URL.search(text))
    if contacts and (last or _CONTACT_WORDS.search(text)) and feat.text_chars < 600:
        return "contact", 0.85, "e-mail, телефон или адрес сайта рядом с концом колоды"
    if _CONTACT_WORDS.search(head) and feat.text_chars < 600:
        return "contact", 0.7, f"заголовок «{head[:40]}»"
    # во второй половине, а не «среди трёх последних»: у шаблонов с
    # приложениями (инструкции, ресурсы) финал стоит задолго до конца файла
    if _CLOSING_WORDS.search(text) and feat.text_chars < 200 and feat.index >= total // 2:
        return "closing", 0.9, "слова благодарности или вопросов во второй половине колоды"
    if feat.layout_type == "section":
        return "section", 0.9, "макет-разделитель"
    if feat.layout_type == "closing":
        return "closing", 0.9, "макет финального слайда шаблона"
    big = feat.max_size >= 1.5 * median_size if median_size else False
    if big and feat.text_chars <= 80 and feat.text_shapes <= 3 and feat.index > 0:
        return "section", 0.7, "крупный короткий заголовок почти без другого текста"
    if last and feat.text_chars <= 120 and feat.pictures == 0:
        return "closing", 0.6, "последний слайд с минимумом текста"
    if feat.has_table or feat.has_chart or feat.pictures or feat.text_chars > 0:
        return "content", 0.7, "слайд с основным содержанием"
    return "content", 0.5, "роль не определена по признакам — содержание по умолчанию"


def _roles(
    features: list[_SlideFeatures],
    median_size: float,
    ids: dict[int, str],
    meta_indices: frozenset[int] = frozenset(),
):
    # «конец колоды» — последний не служебный слайд: у шаблонов с
    # маркетплейсов за финалом идут 20 слайдов инструкций и каталогов
    # иконок, и «Спасибо!» 35-го из 55 не попадало в «последние три»
    content_total = 1 + max(
        (f.index for f in features if f.index not in meta_indices), default=len(features) - 1
    )
    classified = [
        (_classify(f, content_total if f.index < content_total else len(features), median_size), f)
        for f in features
    ]
    by_role: dict[str, list[tuple[float, str, str]]] = {}
    for (role, conf, evidence), f in classified:
        by_role.setdefault(role, []).append((conf, evidence, ids[f.index]))
    roles: list[dna.SlideRole] = []
    for role, items in sorted(by_role.items(), key=lambda kv: -len(kv[1])):
        evidence = Counter(e for _, e, _ in items).most_common(1)[0][0]
        roles.append(
            dna.SlideRole(
                role=role,
                slide_ids=[sid for _, _, sid in items],
                confidence=round(sum(c for c, _, _ in items) / len(items), 3),
                evidence=evidence,
            )
        )
    return roles, {f.index: c for c, f in classified}


def _clusters(features: list[_SlideFeatures], role_of: dict[int, str]) -> dict[int, str]:
    """cluster_id — структурная подпись слайда: роль + состав (картинки,
    таблица, диаграмма) + число текстовых блоков; одинаковые подписи —
    один кластер."""
    def bucket(n: int) -> str:
        return "0" if n == 0 else "1" if n == 1 else "2-3" if n <= 3 else "4-6" if n <= 6 else "7+"

    seen: dict[tuple, str] = {}
    out: dict[int, str] = {}
    for f in features:
        sig = (
            role_of[f.index],
            f.pictures > 0,
            f.has_table,
            f.has_chart,
            bucket(f.text_shapes),
        )
        if sig not in seen:
            seen[sig] = f"c{len(seen) + 1}"
        out[f.index] = seen[sig]
    return out


# --------------------------------------------------------------------------
# Observed: кегли, шрифты, цвета
# --------------------------------------------------------------------------


def _observed_runs(prs: Any, theme_fonts: dict[str, tuple[str | None, str | None]],
                   palette: dict[str, str], clr_map: dict[str, str]):
    """Проход по всем ранам слайдов: шрифты, кегли, цвета с контекстом."""
    minor = next((m for _, m in theme_fonts.values() if m), None)
    major = next((mj for mj, _ in theme_fonts.values() if mj), None)
    sizes: Counter[float] = Counter()
    size_ctx: dict[float, Counter[str]] = {}
    fonts: Counter[str] = Counter()
    font_ctx: dict[str, Counter[str]] = {}
    colors: Counter[str] = Counter()
    color_ctx: dict[str, Counter[str]] = {}
    runs_total = 0
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            ptype = _ph_type(shape)
            if ptype in {"dt", "ftr", "sldNum"}:
                continue
            # цвета заливки и линий фигуры
            sppr = shape._element.find(f"{{{P}}}spPr")
            if sppr is not None:
                fill = sppr.find(f"{{{A}}}solidFill")
                value = _fill_color(fill, palette, clr_map)
                if value:
                    colors[value] += 1
                    color_ctx.setdefault(value, Counter())["fill"] += 1
                line = sppr.find(f"{{{A}}}ln/{{{A}}}solidFill")
                value = _fill_color(line, palette, clr_map)
                if value:
                    colors[value] += 1
                    color_ctx.setdefault(value, Counter())["line"] += 1
            if not shape.has_text_frame:
                continue
            for run in shape.text_frame._txBody.iter(f"{{{A}}}r"):
                text = run.find(f"{{{A}}}t")
                if text is None or not (text.text or "").strip():
                    continue
                runs_total += 1
                rpr = run.find(f"{{{A}}}rPr")
                size = _run_size(rpr)
                is_title = ptype in {"title", "ctrTitle"}
                if size:
                    st = _step(size)
                    sizes[st] += 1
                    size_ctx.setdefault(st, Counter())["title" if is_title else "run"] += 1
                if rpr is not None:
                    latin = rpr.find(f"{{{A}}}latin")
                    face = latin.get("typeface") if latin is not None else None
                    if face in {"+mn-lt", "+mn-ea"}:
                        face = minor
                    elif face in {"+mj-lt", "+mj-ea"}:
                        face = major
                    if face:
                        fonts[face] += 1
                        font_ctx.setdefault(face, Counter())["title" if is_title else "run"] += 1
                    value = _fill_color(rpr.find(f"{{{A}}}solidFill"), palette, clr_map)
                    if value:
                        colors[value] += 1
                        color_ctx.setdefault(value, Counter())["text"] += 1
    return sizes, size_ctx, fonts, font_ctx, colors, color_ctx, runs_total


def _fill_color(fill: Any, palette: dict[str, str], clr_map: dict[str, str]) -> str | None:
    if fill is None:
        return None
    for child in fill:
        tag = etree.QName(child).localname
        if tag == "srgbClr" and child.get("val"):
            return child.get("val").upper()
        if tag == "sysClr" and child.get("lastClr"):
            return child.get("lastClr").upper()
        if tag == "schemeClr" and child.get("val"):
            return _resolve_scheme(child.get("val"), clr_map, palette)
    return None


def _size_roles(
    sizes: Counter[float], size_ctx: dict[float, Counter[str]]
) -> dict[float, list[str]]:
    """Роль кегля — по его месту относительно основного текста колоды:
    основной — самый частый; ≥1.6× — заголовочный; ≤0.8× — подпись."""
    if not sizes:
        return {}
    body = sizes.most_common(1)[0][0]
    roles: dict[float, list[str]] = {}
    for size in sizes:
        found: list[str] = []
        if size_ctx.get(size, Counter()).get("title") or size >= body * 1.6:
            found.append("title")
        elif size <= body * 0.8:
            found.append("caption")
        else:
            found.append("body")
        roles[size] = found
    return roles


# --------------------------------------------------------------------------
# Observed: сетка и поля, фоны
# --------------------------------------------------------------------------


def _slide_boxes(slide: Any, sw: int, sh: int) -> list[tuple[int, int, int, int]]:
    boxes = []
    for shape in slide.shapes:  # верхний уровень: группы считаем целиком
        if _ph_type(shape) in {"dt", "ftr", "sldNum"}:
            continue
        try:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
        except (AttributeError, ValueError):
            continue
        if None in (left, top, width, height) or width <= 0 or height <= 0:
            continue
        if width * height >= 0.8 * sw * sh:  # фоновая плашка не задаёт сетку
            continue
        boxes.append((left, top, width, height))
    return boxes


def _spacing(prs: Any, sw: int, sh: int) -> dna.Spacing:
    margins: Counter[int] = Counter()
    lines_x: Counter[int] = Counter()
    lines_y: Counter[int] = Counter()
    gaps: Counter[int] = Counter()
    snap = EMU_PER_PT  # 1pt

    def q(v: int) -> int:
        return int(round(v / snap) * snap)

    for slide in prs.slides:
        boxes = _slide_boxes(slide, sw, sh)
        for left, top, width, height in boxes:
            for v in (left, top, sw - (left + width), sh - (top + height)):
                if 0 < v < 0.2 * max(sw, sh):
                    margins[q(v)] += 1
            lines_x[q(left)] += 1
            lines_x[q(left + width)] += 1
            lines_y[q(top)] += 1
            lines_y[q(top + height)] += 1
        ordered = sorted(boxes)
        for i, (l1, t1, w1, h1) in enumerate(ordered):
            for l2, t2, _w2, h2 in ordered[i + 1 : i + 6]:
                gap_x = l2 - (l1 + w1)
                overlap_y = min(t1 + h1, t2 + h2) - max(t1, t2)
                if 0 < gap_x < 0.15 * sw and overlap_y > 0.5 * min(h1, h2):
                    gaps[q(gap_x)] += 1
        by_y = sorted(ordered, key=lambda b: b[1])
        for i, (l1, t1, w1, h1) in enumerate(by_y):
            for l2, t2, w2, _h2 in by_y[i + 1 : i + 6]:
                gap_y = t2 - (t1 + h1)
                overlap_x = min(l1 + w1, l2 + w2) - max(l1, l2)
                if 0 < gap_y < 0.15 * sh and overlap_x > 0.5 * min(w1, w2):
                    gaps[q(gap_y)] += 1

    def top(counter: Counter[int], n: int, min_count: int = 3) -> list[int] | None:
        items = [v for v, c in counter.most_common(n) if c >= min_count and v > 0]
        return items or None

    return dna.Spacing(
        common_margins_emu=top(margins, 4),
        common_gaps_emu=top(gaps, 4),
        alignment_lines_x=top(lines_x, 6, min_count=4),
        alignment_lines_y=top(lines_y, 6, min_count=4),
    )


def _background_of(element: Any, palette: dict[str, str], clr_map: dict[str, str]):
    """(kind, value) явного p:bg элемента (слайд / макет / мастер) или None."""
    bg = element.find(f"{{{P}}}cSld/{{{P}}}bg")
    if bg is None:
        return None
    bg_pr = bg.find(f"{{{P}}}bgPr")
    if bg_pr is not None:
        for child in bg_pr:
            tag = etree.QName(child).localname
            if tag == "solidFill":
                return dna.Kind.solid, _fill_color(child, palette, clr_map)
            if tag == "gradFill":
                stop = child.find(f".//{{{A}}}gs")
                first = _fill_color(stop, palette, clr_map) if stop is not None else None
                return dna.Kind.gradient, first
            if tag == "blipFill":
                blip = child.find(f"{{{A}}}blip")
                rid = blip.get(f"{{{R}}}embed") if blip is not None else None
                return dna.Kind.image, rid
            if tag == "pattFill":
                return dna.Kind.pattern, child.get("prst")
    ref = bg.find(f"{{{P}}}bgRef")
    if ref is not None:
        return dna.Kind.solid, _fill_color(ref, palette, clr_map)
    return None


def _backgrounds(
    prs: Any, palette: dict[str, str], clr_map: dict[str, str]
) -> list[dna.Background] | None:
    counter: Counter[tuple[str, str | None]] = Counter()
    for slide in prs.slides:
        # эффективный фон: слайд → макет → мастер
        found = _background_of(slide.element, palette, clr_map)
        if found is None:
            layout = slide.slide_layout
            found = _background_of(layout.element, palette, clr_map) if layout else None
        if found is None and slide.slide_layout is not None:
            found = _background_of(slide.slide_layout.slide_master.element, palette, clr_map)
        if found is None:
            counter[(dna.Kind.inherit.value, None)] += 1
        else:
            counter[(found[0].value, found[1])] += 1
    if not counter:
        return None
    return [
        dna.Background(kind=dna.Kind(kind), value=value, frequency=n)
        for (kind, value), n in counter.most_common()
    ]


# --------------------------------------------------------------------------
# Якоря
# --------------------------------------------------------------------------


def _anchors(prs: Any, sw: int, sh: int) -> list[dna.Anchor]:
    total = len(prs.slides)
    if not total:
        return []
    found: dict[tuple, dict[str, Any]] = {}

    def add(kind: dna.Kind1, bbox: dna.Bbox, slide_index: int, source: str) -> None:
        key = (kind, round(bbox.x, 2), round(bbox.y, 2), round(bbox.w, 2), round(bbox.h, 2))
        entry = found.setdefault(key, {"bbox": bbox, "slides": set(), "parts": set()})
        entry["slides"].add(slide_index)
        entry["parts"].add(source)

    ph_kinds = {
        "ftr": dna.Kind1.footer,
        "sldNum": dna.Kind1.page_number,
        "dt": dna.Kind1.date,
    }
    for index, slide in enumerate(prs.slides):
        layout = slide.slide_layout
        for shape in slide.shapes:
            ptype = _ph_type(shape)
            box = _frac_bbox(shape, sw, sh)
            if box is None:
                continue
            if ptype in ph_kinds and shape.has_text_frame:
                add(ph_kinds[ptype], box, index, str(slide.part.partname).lstrip("/"))
            elif _is_logo_like(shape, box):
                add(dna.Kind1.logo, box, index, str(slide.part.partname).lstrip("/"))
        # логотипы, размещённые на макете/мастере, видны на каждом слайде
        # этого макета (если макет не скрывает shapes мастера)
        if layout is not None:
            sources = [(layout, str(layout.part.partname).lstrip("/"))]
            if layout.element.get("showMasterSp") != "0":
                master = layout.slide_master
                sources.append((master, str(master.part.partname).lstrip("/")))
            for owner, part in sources:
                for shape in owner.shapes:
                    if _ph_type(shape) is not None:
                        continue
                    box = _frac_bbox(shape, sw, sh)
                    if box is not None and _is_logo_like(shape, box):
                        add(dna.Kind1.logo, box, index, part)
    anchors = []
    for (kind, *_), entry in found.items():
        coverage = len(entry["slides"]) / total
        # служебные плейсхолдеры (номер, колонтитул, дата) значимы и при
        # редком появлении; «логотип» без повторяемости — просто значок
        floor = _MIN_LOGO_COVERAGE if kind == dna.Kind1.logo else _MIN_ANCHOR_COVERAGE
        if coverage < floor:
            continue
        anchors.append(
            dna.Anchor(
                kind=kind,
                bbox=entry["bbox"],
                slide_coverage=round(min(coverage, 1.0), 4),
                source_parts=sorted(entry["parts"])[:5],
            )
        )
    anchors.sort(key=lambda a: -a.slide_coverage)
    kept: list[dna.Anchor] = []
    logos = 0
    for anchor in anchors:
        if anchor.kind == dna.Kind1.logo:
            # не больше _MAX_LOGOS, и не пересекающиеся: ряд одинаковых
            # значков — это оформление, а не несколько логотипов
            if logos >= _MAX_LOGOS or any(
                k.kind == dna.Kind1.logo and _overlap(k.bbox, anchor.bbox) > 0.3 for k in kept
            ):
                continue
            logos += 1
        kept.append(anchor)
    return kept[:_MAX_ANCHORS]


def _overlap(a: dna.Bbox, b: dna.Bbox) -> float:
    """Доля меньшего бокса, покрытая другим."""
    iw = min(a.x + a.w, b.x + b.w) - max(a.x, b.x)
    ih = min(a.y + a.h, b.y + b.h) - max(a.y, b.y)
    if iw <= 0 or ih <= 0:
        return 0.0
    return iw * ih / min(a.w * a.h, b.w * b.h)


def _is_logo_like(shape: Any, box: dna.Bbox) -> bool:
    """Знак/логотип: малый компактный элемент внутри слайда у его края,
    без собственного текста (картинка, векторный знак, группа) либо с
    «logo» в имени. Растянутые полосы и элементы за краем — оформление,
    а не знак. Повторяемость на слайдах проверяет вызывающий: единичный
    значок якорем не станет."""
    name = (shape.name or "").lower()
    named = "logo" in name or "логотип" in name
    has_text = bool(shape.has_text_frame and shape.text_frame.text.strip())
    area = box.w * box.h
    inside = (
        box.x >= -0.001
        and box.y >= -0.001
        and box.x + box.w <= 1.001
        and box.y + box.h <= 1.001
    )
    ratio = max(box.w / box.h, box.h / box.w) if box.w > 0 and box.h > 0 else math.inf
    if not inside or ratio > 8 or not (0.0004 <= area <= 0.05):
        return False
    if has_text and not named:
        return False
    cx, cy = box.x + box.w / 2, box.y + box.h / 2
    near_edge = cx < 0.15 or cx > 0.85 or cy < 0.12 or cy > 0.85
    return named or near_edge


# --------------------------------------------------------------------------
# Ёмкости и неподдерживаемое
# --------------------------------------------------------------------------


def _capacities(
    prs: Any, features: list[_SlideFeatures], chart_series: int | None
) -> dna.Capacities:
    content = [f for f in features if f.text_chars]
    titles = [float(f.title_chars) for f in content if f.title_chars]
    bodies = [float(f.body_chars) for f in content if f.body_chars]
    bullets = [float(n) for f in content for n in f.body_paragraph_counts]
    words = [float(n) for f in content for n in f.words_per_paragraph]
    tables = [d for f in features for d in f.table_dims]

    def as_int(v: float | None) -> int | None:
        return int(round(v)) if v is not None else None

    occupancy: dna.OccupancyRange | None = None
    try:
        from deckdna.audit.basic import _Ctx, _slide_occupancy
        from deckdna.audit.config import default_audit_config

        ctx = _Ctx(prs, "dna", 0, default_audit_config())
        values = sorted(_slide_occupancy(slide, ctx)[0] for slide in prs.slides)
        if values:
            lo = values[int(0.1 * (len(values) - 1))]
            hi = values[int(math.ceil(0.9 * (len(values) - 1)))]
            occupancy = dna.OccupancyRange(min=round(lo, 3), max=round(hi, 3))
    except Exception:  # noqa: BLE001 — метрика факультативна: честно None
        occupancy = None

    return dna.Capacities(
        max_title_chars=as_int(_percentile(titles, 0.9)),
        max_body_chars=as_int(_percentile(bodies, 0.9)),
        max_bullets=as_int(_percentile(bullets, 0.9)),
        max_words_per_bullet=as_int(_percentile(words, 0.9)),
        max_table_rows=max((r for r, _ in tables), default=None),
        max_table_cols=max((c for _, c in tables), default=None),
        max_chart_series=chart_series,
        occupancy_range=occupancy,
    )


def _chart_series(zf: zipfile.ZipFile) -> int | None:
    best: int | None = None
    for name in zf.namelist():
        if re.match(r"^ppt/charts/chart\d+\.xml$", name):
            try:
                n = len(etree.fromstring(zf.read(name)).findall(f".//{{{'http://schemas.openxmlformats.org/drawingml/2006/chart'}}}ser"))
            except etree.XMLSyntaxError:
                continue
            best = max(best or 0, n)
    return best


def _locations(indices: list[int]) -> str:
    shown = ", ".join(str(i + 1) for i in indices[:5])
    more = f" (+{len(indices) - 5})" if len(indices) > 5 else ""
    return f"slides {shown}{more}"


def _unsupported(
    prs: Any,
    fonts: Counter[str],
    forensics: Any,
    theme_fonts: _ThemeFonts,
) -> list[dna.UnsupportedFeature]:
    out: list[dna.UnsupportedFeature] = []
    if getattr(forensics, "charts", 0):
        out.append(
            dna.UnsupportedFeature(
                feature="chart",
                location=f"ppt/charts/ ({forensics.charts} parts)",
                strategy=dna.Strategy.preserve,
                disclosure="detected by package census; render fidelity not verified",
            )
        )
    if getattr(forensics, "embeddings", 0):
        out.append(
            dna.UnsupportedFeature(
                feature="embedded_object",
                location=f"ppt/embeddings/ ({forensics.embeddings} parts)",
                strategy=dna.Strategy.preserve,
                disclosure="detected by package census; render fidelity not verified",
            )
        )

    smartart_slides: list[int] = []
    media_slides: list[int] = []
    anim_slides: list[int] = []
    ink_slides: list[int] = []
    model3d_slides: list[int] = []
    for index, slide in enumerate(prs.slides):
        blob = etree.tostring(slide.element)
        rel_types = " ".join(rel.reltype for rel in slide.part.rels.values())
        if "/diagramData" in rel_types or b"dgm:relIds" in blob:
            smartart_slides.append(index)
        if (
            b"videoFile" in blob
            or b"audioFile" in blob
            or "/video" in rel_types
            or "/audio" in rel_types
        ):
            media_slides.append(index)
        timing = slide.element.find(f"{{{P}}}timing")
        if timing is not None and len(timing):
            anim_slides.append(index)
        if "/ink" in rel_types or b"inkml" in blob or b"contentPart" in blob:
            ink_slides.append(index)
        if b"model3d" in blob:
            model3d_slides.append(index)

    def add(feature: str, slides: list[int], strategy: dna.Strategy, disclosure: str) -> None:
        if slides:
            out.append(
                dna.UnsupportedFeature(
                    feature=feature,
                    location=_locations(slides),
                    strategy=strategy,
                    disclosure=disclosure,
                )
            )

    add("smartart", smartart_slides, dna.Strategy.preserve,
        "SmartArt копируется как есть; текст внутри не подменяется компилятором")
    add("media", media_slides, dna.Strategy.preserve,
        "аудио/видео сохраняются в слайде без изменений; в PDF и HTML не воспроизводятся")
    add("animation", anim_slides, dna.Strategy.preserve,
        "анимации сохраняются в PPTX; в PDF и HTML не отображаются")
    add("ink", ink_slides, dna.Strategy.immutable_decoration,
        "рукописные элементы сохраняются как неизменяемое оформление")
    add("model_3d", model3d_slides, dna.Strategy.preserve,
        "3D-модели сохраняются без изменений; при рендере заменяются запасным изображением")

    embedded: set[str] = set()
    pres = prs.part._element
    for node in pres.findall(f".//{{{P}}}embeddedFontLst/{{{P}}}embeddedFont/{{{P}}}font"):
        if node.get("typeface"):
            embedded.add(node.get("typeface").lower())
    declared = {f for pair in theme_fonts.values() for f in pair if f}
    for face, count in fonts.most_common():
        low = face.lower()
        if count < 3 or low in _SYSTEM_FONTS or low in embedded or face.startswith("+"):
            continue
        out.append(
            dna.UnsupportedFeature(
                feature="non_embedded_font",
                location=f"font: {face} ({count} runs)"
                + (" — шрифт темы" if face in declared else ""),
                strategy=dna.Strategy.preserve,
                disclosure=(
                    "шрифт не встроен в файл: на машине без него текст отрисуется "
                    "запасным шрифтом, метрики могут отличаться"
                ),
            )
        )
    return out


# --------------------------------------------------------------------------
# Конфликты declared vs observed
# --------------------------------------------------------------------------


@dataclass
class Conflict:
    kind: str
    detail: str
    resolution: str

    def to_dict(self) -> dict[str, str]:
        return {"kind": self.kind, "detail": self.detail, "resolution": self.resolution}


def _conflicts(fonts: Counter[str], theme_fonts: dict[str, tuple[str | None, str | None]],
               runs_total: int, colors: Counter[str], palette: dict[str, str]) -> list[Conflict]:
    out: list[Conflict] = []
    declared = [f for pair in theme_fonts.values() for f in pair if f]
    if fonts and runs_total:
        top, count = fonts.most_common(1)[0]
        share = count / runs_total
        if declared and top not in declared and share >= 0.25:
            out.append(
                Conflict(
                    "font",
                    f"тема декларирует «{declared[-1]}», но {share:.0%} текста набрано «{top}»",
                    f"основным считаем «{top}»: он реально используется на слайдах; "
                    f"шрифт темы — запасной",
                )
            )
    if colors and palette:
        from deckdna.audit.basic import _delta_e, _srgb_to_lab

        lab_palette = [_srgb_to_lab(_rgb(v)) for v in palette.values()]
        total = sum(colors.values())
        for value, count in colors.most_common(6):
            if value in _NEUTRALS or count / total < 0.08:
                continue
            lab = _srgb_to_lab(_rgb(value))
            nearest = min(_delta_e(lab, p) for p in lab_palette)
            if nearest > 8.0:
                out.append(
                    Conflict(
                        "color",
                        f"цвет #{value} ({count / total:.0%} применений) не входит в палитру темы "
                        f"(ΔE до ближайшего {nearest:.0f})",
                        "считаем брендовым: цвет часто используется на слайдах, "
                        "но не объявлен в теме",
                    )
                )
    return out


def _rgb(value: str) -> tuple[int, int, int]:
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


# --------------------------------------------------------------------------
# Сборка
# --------------------------------------------------------------------------


@dataclass
class DnaBuild:
    design_dna: dna.DesignDNA
    # индекс слайда → (роль, заголовок) — для списка слайдов шаблона
    slide_meta: dict[int, dict[str, Any]]
    conflicts: list[Conflict]
    confidence: dict[str, float]
    # ADR-018: caller-optional -- describe_exemplars(build.features, gateway)
    # runs the LLM content_description enrichment separately, AFTER this
    # (deterministic, unchanged) build finishes; exposed here so the
    # caller doesn't have to re-parse the package to get them.
    features: list[_SlideFeatures] = field(default_factory=list)


def build_design_dna(
    data: bytes,
    forensics: Any,
    template_id: str,
    analysis_id: str,
    slide_ids: dict[int, str],
    created_at: datetime,
    schema_version: str,
) -> DnaBuild:
    """Design DNA пакета. ``slide_ids`` — индекс слайда → id записи в
    ``GET /templates/{id}/slides``: роли ссылаются на них."""
    with zipfile.ZipFile(BytesIO(data)) as zf:
        prs = Presentation(BytesIO(data))
        sw, sh = int(prs.slide_width), int(prs.slide_height)
        parts = {
            "themes": sorted(n for n in zf.namelist() if re.match(r"^ppt/theme/theme\d+\.xml$", n)),
            "slides": [str(s.part.partname).lstrip("/") for s in prs.slides],
        }
        declared, theme_fonts = _declared(prs, zf, parts, sw, sh)
        palette = declared.themes[0].colors if declared.themes else {}
        clr_map = _clr_map(prs)
        layout_types = {lay.part: lay.type for lay in declared.layouts}

        features = [
            _slide_features(slide, i, layout_types) for i, slide in enumerate(prs.slides)
        ]
        (sizes, size_ctx, fonts, font_ctx, colors, color_ctx, runs_total) = _observed_runs(
            prs, theme_fonts, palette, clr_map
        )
        median_size = sizes.most_common(1)[0][0] if sizes else 0.0
        size_roles = _size_roles(sizes, size_ctx)

        meta_indices = frozenset(
            i for i, slide in enumerate(prs.slides) if is_meta_xml(slide.part.blob)
        )
        roles, classified = _roles(features, median_size, slide_ids, meta_indices)
        role_of = {i: c[0] for i, c in classified.items()}
        cluster_of = _clusters(features, role_of)

        observed = dna.Observed(
            # частоты — из forensics (перепись шрифтов пакета, как и раньше);
            # контекст (заголовок / основной текст) — из прохода по ранам
            fonts=[
                dna.Frequency(
                    value=face,
                    frequency=n,
                    contexts=[c for c, _ in font_ctx.get(face, Counter()).most_common()] or None,
                )
                for face, n in (getattr(forensics, "observed_fonts", None) or fonts).items()
            ],
            font_sizes=[
                dna.FontSize(size_pt=size, frequency=n, roles=size_roles.get(size, ["body"]))
                for size, n in sizes.most_common(_MAX_FONT_SIZES)
            ],
            colors=[
                dna.Frequency(
                    value=value,
                    frequency=n,
                    contexts=[c for c, _ in color_ctx.get(value, Counter()).most_common()] or None,
                )
                for value, n in colors.most_common(_MAX_COLORS)
            ],
            spacing=_spacing(prs, sw, sh),
            backgrounds=_backgrounds(prs, palette, clr_map),
        )

        exemplars = []
        for f in features:
            preserved = f.pictures > 0 and f.text_chars == 0 and not f.has_table and not f.has_chart
            exemplars.append(
                dna.Exemplar(
                    slide_index=f.index,
                    part=f.part,
                    role=role_of[f.index],
                    cluster_id=cluster_of[f.index],
                    layout_part=f.layout_part,
                    mutability=(
                        dna.Mutability.preserve_only if preserved else dna.Mutability.fully_editable
                    ),
                )
            )

        anchors = _anchors(prs, sw, sh)
        capacities = _capacities(prs, features, _chart_series(zf))
        unsupported = _unsupported(prs, fonts, forensics, theme_fonts)
        conflicts = _conflicts(fonts, theme_fonts, runs_total, colors, palette)

    design = dna.DesignDNA(
        schema_version=schema_version,
        template_id=template_id,
        analysis_id=analysis_id,
        created_at=created_at,
        slide_size=dna.SlideSize(
            width_emu=sw, height_emu=sh, aspect_ratio=sw / sh if sh else 0.0
        ),
        declared=declared,
        observed=observed,
        anchors=anchors,
        slide_roles=roles,
        exemplars=exemplars,
        components=[],
        capacities=capacities,
        unsupported_features=unsupported,
    )
    slide_meta = {
        f.index: {
            "role": role_of[f.index],
            "title": f.title[:120] or None,
            "mutability": (m.value if (m := exemplars[f.index].mutability) else None),
        }
        for f in features
    }
    confidence = _confidence(design, features, roles, classified)
    return DnaBuild(design, slide_meta, conflicts, confidence, features)


def _confidence(design: dna.DesignDNA, features: list[_SlideFeatures],
                roles: list[dna.SlideRole], classified: dict) -> dict[str, float]:
    """Уверенность по группам 0..1 — из полноты входных данных, не «на глаз»."""
    def share(n: int, d: int) -> float:
        return round(n / d, 3) if d else 0.0

    total_layouts = len(design.declared.layouts)
    named_layouts = sum(1 for lay in design.declared.layouts if lay.placeholders)
    slides = len(features)
    font_runs = sum(f.frequency for f in design.observed.fonts)
    fonts_conf = (
        min(1.0, round(0.5 + 0.5 * share(font_runs, max(1, slides * 5)), 3))
        if design.observed.fonts
        else 0.3
    )
    weighted = (
        sum(r.confidence * len(r.slide_ids) for r in roles) / slides if slides else 0.0
    )
    return {
        "palette": 0.95 if design.declared.themes and design.declared.themes[0].colors else 0.3,
        "fonts": fonts_conf,
        "sizes": min(1.0, share(sum(s.frequency for s in design.observed.font_sizes), 40)),
        "grid": 0.8 if design.observed.spacing.alignment_lines_x else 0.3,
        "anchors": 0.85 if design.anchors else 0.4,
        "layouts": share(named_layouts, total_layouts),
        "roles": round(weighted, 3),
    }


# --- ADR-018: LLM content_description per exemplar, computed once at
# analyze time and cached in Design DNA -- see the ADR and
# exemplar_describe.v1.yaml for the full rationale. Never runs during
# generation; a failure here just leaves content_description=None,
# same fallback discipline as every other model-touching stage in this
# codebase.

EXEMPLAR_DESCRIBE_PROMPT = "exemplar_describe"
_DESCRIBE_BATCH_SIZE = 10
_MAX_CONCURRENT_DESCRIBE_CALLS = 4
_SAMPLE_TEXT_CHARS = 200


class ExemplarDescriptionItem(BaseModel):
    index: int
    description: str = Field(min_length=1)
    # Live bug (28.09): a team-roster card ("Слайд-визитка команды:
    # название команды, имена и должности") kept getting picked for
    # unrelated content (a 5-stage pipeline description) even with a
    # correct content_description in the prompt -- the instruction was
    # only a soft "prefer a different candidate" preference, not a hard
    # rule, and the model didn't reliably follow it. A structured boolean
    # the model must commit to is a much harder signal to ignore than a
    # free-text hint buried in a longer instruction list, and lets
    # layout_fit.py HARD-exclude these candidates for non-team slides
    # instead of just deprioritizing them.
    is_entity_specific: bool = Field(
        description=(
            "True when this exemplar is built for a SPECIFIC named entity "
            "(a person, a team roster, a single dated event) rather than "
            "generic reusable content -- e.g. fields like name/role/photo "
            "repeated per person. False for a generic fact/stat/feature "
            "grid that could hold any topic."
        )
    )


    # Шаблоны с маркетплейсов (Slidesgo и т. п.) несут десятки служебных
    # слайдов: инструкции, титры, палитру, каталоги иконок. Это свойство
    # самого слайда, его видит модель по тексту — без списков ключевых слов.
    is_template_meta: bool = Field(
        default=False,
        description=(
            "True when the slide is service material of the template itself "
            "(usage instructions, credits, colour palette, font list, icon or "
            "illustration catalogue, 'delete this slide' notes), not a design "
            "for real presentation content."
        ),
    )


class ExemplarDescribeBatch(BaseModel):
    slides: list[ExemplarDescriptionItem] = Field(min_length=1)


def can_describe(gateway: object) -> bool:
    """Same mock-gate pattern as content_writer.can_rewrite/layout_fit.
    can_rewrite: an offline MockProvider without an explicit fixture for
    this prompt would synthesize schema-valid but meaningless sentences,
    which is worse than an honest empty content_description."""
    if getattr(gateway, "provider_name", None) == "mock":
        return EXEMPLAR_DESCRIBE_PROMPT in (getattr(gateway, "fixtures", None) or {})
    return gateway is not None


def _sample_text(feature: _SlideFeatures) -> str:
    parts = [feature.title, *feature.paragraphs]
    joined = " | ".join(p for p in parts if p.strip())
    return joined[:_SAMPLE_TEXT_CHARS] or feature.all_text[:_SAMPLE_TEXT_CHARS]


async def _describe_batch(
    gateway: ModelGateway,
    batch: list[_SlideFeatures],
    *,
    language: str,
    sem: asyncio.Semaphore,
) -> dict[int, ExemplarDescriptionItem]:
    """One batch call. Returns {index: item} only for slides the model
    answered validly for -- partial success returned as-is, same
    discipline as structure.py's _run_structure_batch."""
    payload = {
        "slides": [
            {"index": f.index, "sample_text": _sample_text(f)} for f in batch
        ],
        "language": language,
    }
    wanted = {f.index for f in batch}
    async with sem:
        try:
            out = await gateway.text_json(
                EXEMPLAR_DESCRIBE_PROMPT, payload, ExemplarDescribeBatch
            )
        except Exception as exc:  # noqa: BLE001 — provider failure = honest empty batch
            logger.warning(
                "exemplar_describe batch %s failed (%s); no descriptions for these slides",
                sorted(wanted),
                exc,
            )
            return {}
    if not isinstance(out, ExemplarDescribeBatch):
        return {}
    return {
        item.index: item
        for item in out.slides
        if item.index in wanted and item.description.strip()
    }


async def describe_exemplars(
    features: list[_SlideFeatures],
    gateway: ModelGateway,
    *,
    language: str = "ru",
) -> dict[str, ExemplarDescriptionItem]:
    """part -> description item for every exemplar the model successfully
    described. Empty dict (not an exception) on any failure or when the
    gateway shouldn't participate (can_describe) -- caller decides what
    "no descriptions" means, never blocks analyze.

    Batched (not one call per slide) and fully concurrent across
    batches -- this only ever runs once per template, at analyze time,
    so it's not on the 5-minute generation budget at all, but there's no
    reason to make it slower than it needs to be either.
    """
    if not features or not can_describe(gateway):
        return {}
    batches = [
        features[i : i + _DESCRIBE_BATCH_SIZE]
        for i in range(0, len(features), _DESCRIBE_BATCH_SIZE)
    ]
    sem = asyncio.Semaphore(_MAX_CONCURRENT_DESCRIBE_CALLS)
    batch_results = await asyncio.gather(
        *(_describe_batch(gateway, b, language=language, sem=sem) for b in batches)
    )
    by_index: dict[int, ExemplarDescriptionItem] = {}
    for result in batch_results:
        by_index.update(result)
    return {f.part: by_index[f.index] for f in features if f.index in by_index}
