"""Образцы слайдов из макетов шаблона (ADR-031 §3).

Пайплайн построен «от образцов»: вёрстка клонирует слайды шаблона и
видоизменяет их блоки. Типичный корпоративный ``.potx`` — мастер плюс
макеты, часто без единого слайда, и генерация падала с ``template has no
slides to clone``. Если слайдов меньше ``MIN_EXEMPLAR_SLIDES``, каждый
макет с текстовыми плейсхолдерами инстанцируется слайдом: плейсхолдеры
получают образцовый текст (иначе в них нет ни одного ``a:t`` и вёрстке
некуда писать), форматирование они наследуют от макета и мастера. Дальше
такой слайд идёт обычным путём exemplar-first: Design DNA, роли, выбор
образца, подмена текста.

Рисование с нуля (``composition.native_fallback``) сознательно не
используется: суть задачи — блоки шаблона, а они живут в макетах.
"""

from __future__ import annotations

import copy
import io
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

from deckdna.pptx.composing.background import background_layers
from deckdna.pptx.opc.normalize import normalize_template_bytes

MIN_EXEMPLAR_SLIDES = 3

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_SUBTITLE_TYPES = {PP_PLACEHOLDER.SUBTITLE}
_BODY_TYPES = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
}
# медиа-плейсхолдеры пустыми не оставляем: в редакторе это «Вставьте
# рисунок», а картинку/таблицу/график вёрстка ставит своими средствами
_MEDIA_TYPES = {
    PP_PLACEHOLDER.PICTURE,
    PP_PLACEHOLDER.CHART,
    PP_PLACEHOLDER.TABLE,
    PP_PLACEHOLDER.MEDIA_CLIP,
    PP_PLACEHOLDER.ORG_CHART,
    PP_PLACEHOLDER.BITMAP,
}

# Образцовый текст: длина подобрана под пороги «контентного» образца
# (``RunProfile.is_content_like``: длинные run'ы, без мелких подписей).
_SAMPLE_TITLE = "Заголовок слайда с ключевой мыслью"
_SAMPLE_SUBTITLE = "Подзаголовок, который раскрывает тему слайда"
_SAMPLE_BODY = (
    "Первый тезис раздела с коротким пояснением",
    "Второй тезис раздела с подтверждающим фактом",
    "Третий тезис раздела с выводом для аудитории",
)


@dataclass(frozen=True)
class SynthesisReport:
    original_slides: int
    synthesized_slides: int
    layouts_used: tuple[str, ...]

    def to_dict(self) -> dict:
        return {
            "original_slides": self.original_slides,
            "synthesized_slides": self.synthesized_slides,
            "layouts_used": list(self.layouts_used),
        }


def _placeholder_type(ph) -> PP_PLACEHOLDER | None:
    try:
        return ph.placeholder_format.type
    except (AttributeError, ValueError):
        return None


def _is_primary_body(ph, slide_height: int = 0) -> bool:
    """Основное текстовое тело (не мелкий подзаголовок/подпись < 30% высоты слайда)."""
    if _placeholder_type(ph) not in _BODY_TYPES:
        return False
    h = int(getattr(ph, "height", 0) or 0)
    if slide_height > 0 and h > 0 and h < slide_height * 0.30:
        return False
    return True


def _text_capable(layout, slide_height: int = 0) -> tuple[bool, bool]:
    """(есть заголовок, есть основное тело) у макета; subtitle — не body."""
    from deckdna.pptx.composing.text_replace import _has_nonhorizontal_text_flow

    has_title = has_body = has_subtitle = has_media = False
    for ph in layout.placeholders:
        kind = _placeholder_type(ph)
        body = ph._element.find(qn("p:txBody"))
        if (kind in _TITLE_TYPES | _BODY_TYPES | _SUBTITLE_TYPES and body is not None
                and _has_nonhorizontal_text_flow(body)):
            # Synthesized examples are ordinary horizontal prose hosts. Keep
            # original slides intact, but do not manufacture incompatible donors.
            return False, False
        if kind in _MEDIA_TYPES:
            has_media = True
        elif kind in _TITLE_TYPES:
            has_title = True
        elif kind in _BODY_TYPES:
            if _is_primary_body(ph, slide_height):
                has_body = True
            else:
                has_subtitle = True
        elif kind in _SUBTITLE_TYPES:
            has_subtitle = True
    if has_media and not has_body:
        # Макет «Рисунок с подписью»: после удаления пустого медиа-плейсхолдера
        # остаётся только узкая подпись внизу и дыра на месте иллюстрации.
        return False, False
    return has_title or has_subtitle, has_body


def _fill(slide, slide_height: int = 0) -> None:
    for ph in list(slide.placeholders):
        kind = _placeholder_type(ph)
        if kind in _MEDIA_TYPES:
            el = ph._element
            el.getparent().remove(el)
        elif kind in _TITLE_TYPES:
            ph.text_frame.text = _SAMPLE_TITLE
        elif kind in _SUBTITLE_TYPES:
            ph.text_frame.text = _SAMPLE_SUBTITLE
        elif kind in _BODY_TYPES and ph.has_text_frame:
            tf = ph.text_frame
            if _is_primary_body(ph, slide_height):
                tf.text = _SAMPLE_BODY[0]
                for line in _SAMPLE_BODY[1:]:
                    tf.add_paragraph().text = line
            else:
                tf.text = _SAMPLE_SUBTITLE


def _drawn_background(slide, sw: int, sh: int) -> list:
    """Полноразмерные залитые фигуры/группы без контента, в z-порядке."""
    return [el for el in background_layers(slide._element, sw, sh) if el.tag != qn("p:pic")]


def _background_signature(elements: list) -> tuple[bytes, ...]:
    out = []
    for element in elements:
        clone = copy.deepcopy(element)
        for props in clone.iter(qn("p:cNvPr")):
            props.attrib.clear()
        out.append(etree.tostring(clone, method="c14n"))
    return tuple(out)


def _inherit_drawn_background(slide, originals: list, sw: int, sh: int) -> None:
    """Макет не содержит оформление, нарисованное на исходных слайдах.

    Сначала ищем образец того же макета, затем общий фон его мастера.
    Разные фоны исходных слайдов не смешиваем: неоднозначность остаётся
    за исходными образцами. Собственный фон нового макета приоритетнее.
    """
    if (
        slide.slide_layout._element.find("p:cSld/p:bg", slide.slide_layout._element.nsmap)
        is not None
    ):
        return
    donors = [s for s in originals if s.slide_layout == slide.slide_layout]
    if not donors:
        donors = [
            s for s in originals if s.slide_layout.slide_master == slide.slide_layout.slide_master
        ]
    candidates = [(s, _drawn_background(s, sw, sh)) for s in donors]
    if not candidates or any(not layers for _, layers in candidates):
        return
    signatures = {_background_signature(layers) for _, layers in candidates}
    if len(signatures) != 1:
        return
    donor, layers = candidates[0]
    next_id = slide.shapes._next_shape_id
    for index, layer in enumerate(layers):
        clone = copy.deepcopy(layer)
        for props in clone.iter(qn("p:cNvPr")):
            props.set("id", str(next_id))
            next_id += 1
        # Группа фона может содержать картинку. Копируем ссылки через
        # package API, а не оставляем rId, локальные для слайда-донора.
        for node in clone.iter():
            for attr, rid in list(node.attrib.items()):
                if not attr.startswith("{" + donor._element.nsmap["r"] + "}"):
                    continue
                rel = donor.part.rels[rid]
                target = rel.target_ref if rel.is_external else rel.target_part
                node.set(
                    attr, slide.part.relate_to(target, rel.reltype, is_external=rel.is_external)
                )
        slide.shapes._spTree.insert(2 + index, clone)

    # Светлые авторские шрифты поверх тёмного рисунка не наследуются
    # стандартными ph новых макетов. Переносим цвет и семейство, сохраняя
    # размеры и геометрию целевого макета.
    text_shapes = [s for s in donor.shapes if s.has_text_frame and s.text.strip()]
    if not text_shapes:
        return
    title = donor.shapes.title
    if title is None:
        title = min(text_shapes, key=lambda s: s.top)
    body = max((s for s in text_shapes if s is not title), key=lambda s: len(s.text), default=title)
    for ph in slide.placeholders:
        if not ph.has_text_frame:
            continue
        source = title if _placeholder_type(ph) in _TITLE_TYPES else body
        style = next(
            (
                r._r.rPr
                for p in source.text_frame.paragraphs
                for r in p.runs
                if r.text.strip() and r._r.rPr is not None
            ),
            None,
        )
        if style is None:
            continue
        for para in ph.text_frame.paragraphs:
            for run in para.runs:
                props = run._r.get_or_add_rPr()
                for child in style:
                    if etree.QName(child).localname in {
                        "solidFill",
                        "gradFill",
                        "latin",
                        "ea",
                        "cs",
                    }:
                        existing = props.find(child.tag)
                        if existing is not None:
                            props.remove(existing)
                        props.append(copy.deepcopy(child))


def ensure_exemplar_slides(
    data: bytes, min_slides: int = MIN_EXEMPLAR_SLIDES
) -> tuple[bytes, SynthesisReport | None]:
    """Добавить образцы из макетов, если слайдов в шаблоне меньше *min_slides*.

    Возвращает (байты пакета, отчёт). Шаблон с достаточным числом слайдов
    возвращается теми же байтами и ``None``. Синтезируются слайды для
    неиспользованных макетов и макетов с пустым основным телом. Макет
    «заголовок + тело» с наибольшей площадью тела добавляется дважды,
    чтобы он стал доминирующим макетом пула образцов (как в обычной
    колоде, где контентных слайдов
    больше, чем титульных)."""
    prs = Presentation(io.BytesIO(data))
    original = len(prs.slides)
    if original >= min_slides:
        return data, None
    used_layouts = {s.slide_layout.part.partname for s in prs.slides}
    originals = list(prs.slides)
    sh = int(prs.slide_height or 0)
    content_layouts = []
    other_layouts = []
    for layout in prs.slide_layouts:
        if layout.part.partname in used_layouts:
            bodies = [
                ph
                for s in originals
                if s.slide_layout == layout
                for ph in s.placeholders
                if _is_primary_body(ph, sh) and ph.has_text_frame
            ]
            if not bodies or any(ph.text_frame.text.strip() for ph in bodies):
                continue
            # A referenced but empty main body is not an observed exemplar.
            # Create a filled donor without writing sample text into the original.
        has_title, has_body = _text_capable(layout, sh)
        if not (has_title or has_body):
            continue  # пустой макет: образцу нечего подменять
        (content_layouts if has_title and has_body else other_layouts).append(layout)
    if not content_layouts and not other_layouts:
        return data, None
    # Layout order is authored metadata, not evidence of prose capacity.
    # Read effective geometry through python-pptx (including master inheritance).
    content_layouts.sort(
        key=lambda layout: max(
            (
                int(ph.width or 0) * int(ph.height or 0)
                for ph in layout.placeholders
                if _is_primary_body(ph, sh)
            ),
            default=0,
        ),
        reverse=True,
    )
    order = [*content_layouts[:1], *content_layouts, *other_layouts]
    names: list[str] = []
    for layout in order:
        slide = prs.slides.add_slide(layout)
        _fill(slide, sh)
        _inherit_drawn_background(slide, originals, prs.slide_width, prs.slide_height)
        names.append(layout.name or str(layout.part.partname))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), SynthesisReport(
        original_slides=original,
        synthesized_slides=len(order),
        layouts_used=tuple(names),
    )


def is_synthesized_placeholder_text(text: str) -> bool:
    """Образцовый текст синтеза — для аудита/тестов: на выходе его быть не должно."""
    text = text.strip()
    return text in {_SAMPLE_TITLE, _SAMPLE_SUBTITLE, *_SAMPLE_BODY}


def prepare_template_bytes(data: bytes) -> tuple[bytes, SynthesisReport | None]:
    """Шаблон, готовый к вёрстке: ``.potx`` → тип ``.pptx`` (ADR-031 §2) и
    образцы из макетов при нехватке слайдов (§3). Битый пакет возвращается
    как есть — его типизированно отвергнет ``OpcPackage.open``."""
    data = normalize_template_bytes(data)
    try:
        return ensure_exemplar_slides(data)
    except Exception:  # noqa: BLE001 — не наш формат: решит OpcPackage.open
        return data, None


def prepared_template_path(
    template_path: Path, work_dir: Path
) -> tuple[Path, SynthesisReport | None]:
    """Путь к шаблону, готовому к вёрстке. Если подготовка ничего не меняет
    (обычный ``.pptx`` с образцами) — исходный путь; иначе копия в *work_dir*."""
    try:
        data = template_path.read_bytes()
    except OSError:
        # отсутствующий файл — типизированную ошибку со своей стадией даст
        # первый настоящий читатель шаблона
        return template_path, None
    prepared, report = prepare_template_bytes(data)
    if prepared is data:
        return template_path, None
    work_dir.mkdir(parents=True, exist_ok=True)
    target = work_dir / "template.prepared.pptx"
    target.write_bytes(prepared)
    return target, report


__all__ = [
    "MIN_EXEMPLAR_SLIDES",
    "prepare_template_bytes",
    "prepared_template_path",
    "SynthesisReport",
    "ensure_exemplar_slides",
    "is_synthesized_placeholder_text",
]
