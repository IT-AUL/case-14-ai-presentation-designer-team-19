"""Измеримые метрики Quality Passport, которых не было в v0 (contract D3).

Все три считаются детерминированно по итоговой колоде и уже собранным
audit-issues — без модели и без выдуманных чисел:

- ``style_fidelity`` — доля слайдов колоды БЕЗ нарушений соответствия
  шаблону по четырём осям (палитра, шрифты, происхождение макета,
  якоря). Единица измерения — слайд, поэтому 1.0 значит «ни одного
  слайда с нарушением», а не «ни одной фигуры»;
- ``numbers`` — numeric presence: глобальное совпадение нормализованного
  значения в колоде и источнике. Оно не подтверждает связь числа с
  сущностью, утверждением или арифметикой;
- ``usage`` — вызовы и токены модели за прогон (нули без модели).
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from deckdna.audit.issues import AuditIssue

# --------------------------------------------------------------------------
# style_fidelity
# --------------------------------------------------------------------------

_STYLE_AXES: dict[str, tuple[str, ...]] = {
    "palette_compliance": ("template.color_palette",),
    "font_compliance": ("template.font_family", "template.font_scale"),
    "layout_origin_compliance": ("template.layout_origin",),
    "anchor_compliance": ("template.anchor_position",),
}


def style_fidelity(
    audit_issues: Iterable[AuditIssue], slide_count: int
) -> dict[str, float] | None:
    """Доли слайдов без нарушений по осям соответствия шаблону (0..1).

    None — если в колоде нет слайдов (делить не на что)."""
    if slide_count <= 0:
        return None
    issues = list(audit_issues)
    out: dict[str, float] = {}
    for axis, rules in _STYLE_AXES.items():
        bad = {
            i.slide_index
            for i in issues
            if i.rule_code in rules and i.slide_index is not None
        }
        out[axis] = round(1 - min(len(bad), slide_count) / slide_count, 4)
    return out


def style_fidelity_score(fidelity: Mapping[str, float] | None) -> float | None:
    """Одно число 0..1 для карточки варианта — среднее по осям."""
    if not fidelity:
        return None
    return round(sum(fidelity.values()) / len(fidelity), 4)


# --------------------------------------------------------------------------
# numbers
# --------------------------------------------------------------------------

# 1 234 567 / 12,5 / 30% / 2.5 млн — число с необязательной единицей.
_NUMBER_RE = re.compile(
    r"(?<![\w.,])(\d{1,3}(?:[   ]\d{3})+|\d+)(?:[.,](\d+))?"
    r"(?:\s?(%|млрд|млн|тыс|руб|₽|[xх×]))?",
    re.IGNORECASE,
)
_UNIT_ALIASES = {"х": "x", "×": "x", "руб": "₽"}


def _normalise(match: re.Match[str]) -> str:
    whole = re.sub(r"[   ]", "", match.group(1))
    frac = (match.group(2) or "").rstrip("0")
    unit = (match.group(3) or "").lower()
    unit = _UNIT_ALIASES.get(unit, unit)
    return f"{whole}{'.' + frac if frac else ''}{unit}"


def extract_numbers(text: str) -> set[str]:
    """Нормализованные значимые числа текста.

    Значимое — многозначное (≥2 цифр), дробное или с единицей: одиночные
    «1», «2» — нумерация пунктов, а не факты, и в сверку не идут."""
    found: set[str] = set()
    for m in _NUMBER_RE.finditer(text):
        digits = len(m.group(1).replace(" ", "").replace(" ", "").replace(" ", ""))
        significant = digits >= 2 or m.group(2) or m.group(3)
        if significant:
            found.add(_normalise(m))
    return found


# Skip a leading quote mark ("Рилсы» is a legitimate sentence start) by
# checking the first LETTER, not the first character.
_FIRST_LETTER_RE = re.compile(r"[a-zA-Zа-яА-ЯёЁ]")


def starts_with_capital(text: str) -> bool:
    """False only when the first real letter near the start is lowercase
    -- a strong, cheap signal that a title/bullet got cut off mid-thought
    rather than genuinely tightened (found live in layout_fit.py's Stage
    3 output, 27.09; shared here so content_writer.py's Stage 2 output
    gets the same check -- Stage 3 doesn't always touch every slide, so
    a Stage-2-only fragment could otherwise ship unchecked)."""
    match = _FIRST_LETTER_RE.search(text[:3])
    if match is None:
        return True  # no letters near the start (e.g. starts with a digit) -- don't flag
    return match.group(0).isupper()


_TRAILING_RANGE_RE = re.compile(
    r"(?<!\w)(\d{1,4}(?:[.,]\d+)?)\s*[–—-]\s*(\d{1,4}(?:[.,]\d+)?)\s*[.!?]?\s*$"
)


def ends_with_unqualified_range(text: str) -> bool:
    """A range at the end usually lost its measured object during rewriting.

    Keep full calendar-year ranges, which are meaningful on their own.
    The live failure was ``Пилот строится на 30–50`` (participants omitted).
    """
    match = _TRAILING_RANGE_RE.search(text)
    if match is None:
        return False
    first, second = match.group(1), match.group(2)
    if first.isdigit() and second.isdigit() and len(first) == len(second) == 4:
        if 1900 <= int(first) <= 2100 and 1900 <= int(second) <= 2100:
            return False
    return True


@dataclass(frozen=True)
class NumbersCheck:
    verified: int
    failed: int
    failed_numbers: tuple[str, ...] = ()


def _bare(number: str) -> str:
    return re.sub(r"[^\d.]", "", number)


def check_numbers(deck_text: str, source_text: str) -> NumbersCheck:
    """Проверить глобальное наличие чисел колоды в тексте источника.

    Совпадение — по нормализованному значению; «30%» в колоде совпадает
    и с «30 %», и с «30» без знака процента в источнике. ``verified`` —
    legacy-имя поля NumbersCheck; это только наличие значения где-либо
    в источнике, не доказательство конкретного утверждения."""
    source = extract_numbers(source_text)
    source_bare = {_bare(n) for n in source}
    verified: list[str] = []
    failed: list[str] = []
    for number in sorted(extract_numbers(deck_text)):
        if number in source or _bare(number) in source_bare:
            verified.append(number)
        else:
            failed.append(number)
    return NumbersCheck(len(verified), len(failed), tuple(failed))


def deck_text(pptx_path: str | Path) -> str:
    """Весь текст слайдов колоды (без заметок), включая таблицы и группы."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes: Iterable[Any]) -> Iterable[str]:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                yield shape.text_frame.text
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            yield cell.text

    prs = Presentation(str(pptx_path))
    return "\n".join(t for slide in prs.slides for t in walk(slide.shapes))


def pack_text(pack: Any) -> str:
    """Весь текст ContentPack: заголовки, блоки, таблицы, подписи."""
    parts: list[str] = []
    for section in getattr(pack, "sections", None) or []:
        parts.append(section.heading or "")
        for block in section.blocks:
            if block.text:
                parts.append(block.text)
            parts.extend(block.items or [])
    for table in getattr(pack, "tables", None) or []:
        parts.append(table.title or "")
        parts.extend(table.headers)
        for row in table.rows:
            parts.extend("" if c is None else str(c) for c in row)
    for chart in getattr(pack, "charts", None) or []:
        parts.append(getattr(chart, "title", None) or "")
        for series in getattr(chart, "series", None) or []:
            parts.extend(str(v) for v in getattr(series, "values", None) or [])
    for asset in getattr(pack, "assets", None) or []:
        parts.append(getattr(asset, "caption", None) or "")
    return "\n".join(parts)


# --------------------------------------------------------------------------
# usage
# --------------------------------------------------------------------------


def usage_snapshot(gateway: object | None) -> dict[str, dict[str, int]]:
    """Текущие счётчики gateway (пусто, если он не ведёт учёт)."""
    from deckdna.providers.profiles import UsageSource

    if isinstance(gateway, UsageSource):
        return {m: dict(c) for m, c in gateway.usage_report().items()}
    return {}


def usage_delta(
    before: Mapping[str, Mapping[str, int]],
    after: Mapping[str, Mapping[str, int]],
) -> dict[str, Any]:
    """Расход модели между двумя снимками — нули, если моделей не было.

    Один gateway живёт весь прогон генерации (все варианты), поэтому
    вариант получает разность снимков, а не накопленный итог."""
    per_model: list[dict[str, Any]] = []
    total_tokens = 0
    calls = 0
    for model in sorted(after):
        prev = before.get(model, {})
        d_calls = after[model].get("calls", 0) - prev.get("calls", 0)
        d_in = after[model].get("input_tokens", 0) - prev.get("input_tokens", 0)
        d_out = after[model].get("output_tokens", 0) - prev.get("output_tokens", 0)
        if d_calls <= 0 and d_in <= 0 and d_out <= 0:
            continue
        per_model.append(
            {
                "model_id": model,
                "calls": d_calls,
                "input_tokens": d_in,
                "output_tokens": d_out,
            }
        )
        calls += d_calls
        total_tokens += d_in + d_out
    return {"total_tokens": total_tokens, "model_calls": calls, "per_model": per_model}
