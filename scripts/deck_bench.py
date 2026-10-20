"""Золотой HTTP-бенчмарк: финальные PPTX, аудит API, PDF-контакт-листы.

Без модели по умолчанию. --use-llm допускает максимум два выбранных сценария.
См. docs/BENCHMARK_RU.md; сервер и его конфигурация не изменяются.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import math
import re
import statistics
import subprocess
import tempfile
import time
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import UTC, datetime
from decimal import Decimal
from itertools import groupby, zip_longest
from pathlib import Path
from typing import NamedTuple

import httpx
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
METRIC_VERSION = "2"
FORBIDDEN = [
    "Предположение:",
    "условн",
    "Направление A",
    "Thanks!",
    "CREDITS",
    "Slidesgo",
    "viz_reason",
]
NUMBER = re.compile(r"(?<![\w])[-+−]?(?:\d{1,3}(?:[ \u00a0\u202f]\d{3})+|\d+)(?:[.,]\d+)?")
SLUG = re.compile(r"\b[a-z]+(?:[_-][a-z0-9]+)+\b", re.I)
LABEL = re.compile(r"\bпо\s+(?:месяцам|годам|кварталам|регионам|категориям|каналам)\b", re.I)


def _number_values(text: str) -> set[Decimal]:
    result = set()
    for token in NUMBER.findall(text):
        clean = re.sub(r"[ \u00a0\u202f]", "", token).replace("−", "-").replace(",", ".")
        result.add(Decimal(clean))
    return result


def numbers(text: str) -> set[str]:
    """Сравнивать числа с сохранением знака, без зависимости от десятичного разделителя."""
    return {format(value.normalize(), "f") if value else "0" for value in _number_values(text)}


def strings(value):
    """Текст/значения JSON без чисел из имён служебных полей."""
    if isinstance(value, dict):
        for item in value.values():
            yield from strings(item)
    elif isinstance(value, list):
        for item in value:
            yield from strings(item)
    elif value is not None:
        yield str(value)


def _sniff_delimiter(text: str) -> str:
    """Разделитель CSV: sniffer, иначе — наиболее консистентное число колонок."""
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=";,\t|").delimiter
    except csv.Error:
        best, best_score = ",", -1.0
        for delim in (";", ",", "\t", "|"):
            counts = [len(r.split(delim)) for r in text.splitlines() if r.strip()][:50]
            if not counts:
                continue
            score = min(counts) + (100 if len(set(counts)) == 1 and counts[0] > 1 else 0)
            if score > best_score:
                best, best_score = delim, score
        return best


def csv_cells(path: Path) -> list[list[str]]:
    """CSV/TSV по ячейкам: «2025-01,30.3» как сырой текст склеивало бы
    значение с датой («01,30» → 1.3) и теряло само 30.3."""
    raw = path.read_text(encoding="utf-8-sig")
    return [row for row in csv.reader(io.StringIO(raw), delimiter=_sniff_delimiter(raw)) if row]


def source_text(path: Path) -> str:
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            return "\n".join(
                str(v) for sheet in workbook for row in sheet.values for v in row if v is not None
            )
        finally:
            workbook.close()
    if path.suffix == ".json":
        return "\n".join(strings(json.loads(path.read_text(encoding="utf-8"))))
    if path.suffix in (".csv", ".tsv"):
        # ячейка на строку: и «2025-01» даёт 2025 и 1, и «30.3» даёт 30.3
        return "\n".join(cell for row in csv_cells(path) for cell in row)
    return path.read_text(encoding="utf-8-sig")


_CELL_NUMBER = re.compile(r"^[-+−]?(?:\d{1,3}(?:[ \u00a0\u202f]\d{3})+|\d+)(?:[.,]\d+)?[%％]?$")


def _cell_number(value) -> Decimal | None:
    """Число, если ячейка целиком числовая (допускается % на конце)."""
    text = str(value).strip()
    if not _CELL_NUMBER.match(text):
        return None
    clean = re.sub(r"[ \u00a0\u202f%％]", "", text).replace("−", "-").replace(",", ".")
    try:
        return Decimal(clean)
    except ArithmeticError:
        return None


def _numeric_runs(rows: list[list], min_share=0.6) -> list[list[Decimal]]:
    """Числовые колонки таблицы: ≥2 значений и ≥60% непустых ячеек."""
    columns: list[list[Decimal]] = []
    for col in range(max((len(r) for r in rows), default=0)):
        nonempty = [r[col] for r in rows if col < len(r) and str(r[col]).strip() != ""]
        nums = [n for v in nonempty if (n := _cell_number(v)) is not None]
        if len(nums) >= 2 and len(nums) >= min_share * len(nonempty):
            columns.append(nums)
    return columns


def source_columns(path: Path) -> list[list[Decimal]]:
    """Числовые колонки источника — для производных чисел (доли, итоги)."""
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv"):
        return _numeric_runs(csv_cells(path))
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            return [
                col for sheet in workbook for col in _numeric_runs([list(r) for r in sheet.values])
            ]
        finally:
            workbook.close()
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, dict) and "columns" in payload and "data" in payload:
            return _numeric_runs(payload["data"])
        if isinstance(payload, dict):
            rows = [[value] if not isinstance(value, list) else value for value in payload.values()]
            return _numeric_runs([list(map(list, zip_longest(*rows)))]) if rows else []
        if isinstance(payload, list):
            if payload and all(isinstance(r, dict) for r in payload):
                keys = list(dict.fromkeys(k for r in payload for k in r))
                return _numeric_runs([[r.get(k) for k in keys] for r in payload])
            return _numeric_runs(payload)
    return []


_YEAR = re.compile(r"(?<![\d.,])((?:19|20)\d{2})(?![\d.,])")


def _histogram_bins(values: list[float]) -> tuple[list[float], list[float]]:
    """Бины Фридмана–Диакониса — зеркало backend/analysis/pipeline.py:
    границы и высоты бинов — легальные производные числа колоды."""
    values = sorted(values)
    n = len(values)
    lo, hi = values[0], values[-1]
    if hi == lo:
        return [lo], [float(n)]
    q1, q3 = values[int(n * 0.25)], values[int(n * 0.75)]
    iqr = q3 - q1
    width = 2 * iqr / (n ** (1 / 3)) if iqr > 0 else (hi - lo) / 10
    bins = int(math.ceil((hi - lo) / width)) if width > 0 else 10
    bins = max(5, min(20, bins))
    width = (hi - lo) / bins
    counts = [0.0] * bins
    for v in values:
        counts[min(int((v - lo) / width), bins - 1)] += 1
    edges = [lo + k * width for k in range(bins + 1)]
    return edges, counts


def _pearson(xs: list[float], ys: list[float]) -> float | None:
    n = len(xs)
    if n < 3:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    cov = sum((x - mx) * (y - my) for x, y in zip(xs, ys, strict=True))
    vx = sum((x - mx) ** 2 for x in xs)
    vy = sum((y - my) ** 2 for y in ys)
    if vx == 0 or vy == 0:
        return None
    return cov / math.sqrt(vx * vy)


def derived_numbers(columns: list[list[Decimal]]) -> set[Decimal]:
    """Числа, которые код конвейера выводит из данных: суммы, доли итога,
    типичное значение, края ряда и темпы, разброс, выбросы, гистограмма,
    корреляция, счётчики. Всё это — не выдумки, даже если дословно
    отсутствуют в сырых значениях."""
    out: set[Decimal] = set()

    def put(x):
        if x is not None and math.isfinite(float(x)):
            out.add(Decimal(str(round(float(x), 9))))

    floats = [[float(v) for v in col] for col in columns]
    for vals in floats:
        n = len(vals)
        ordered = sorted(vals)
        total = sum(vals)
        q1, q3 = ordered[int(n * 0.25)], ordered[int(n * 0.75)]
        iqr = q3 - q1
        first, last = vals[0], vals[-1]
        outliers = [v for v in vals if v > q3 + 1.5 * iqr or v < q1 - 1.5 * iqr]
        for x in (
            total,
            total / n,
            statistics.median(ordered),
            q1,
            q3,
            iqr,
            ordered[0],
            ordered[-1],
            first,
            last,
            last - first,
            len(outliers),
            n,
            len(set(vals)),
        ):
            put(x)
        if first:
            put(last / first)
            put((last / first - 1) * 100)
        if total and all(v >= 0 for v in vals):
            # доли частей и накопленные доли топ-k («топ-3 — 41,6% итога»)
            shares = sorted((v / total * 100 for v in vals), reverse=True)
            for s in shares:
                put(s)
            for k in range(2, min(5, n) + 1):
                put(sum(shares[:k]))
                put(n - k)  # «остальные N»
        if n >= 5:
            edges, counts = _histogram_bins(vals)
            for x in (*edges, *counts):
                put(x)
    for xs, ys in ((a, b) for i, a in enumerate(floats) for b in floats[i + 1 :]):
        if len(xs) == len(ys):
            put(_pearson(xs, ys))
    return out


class NumberEvidence(NamedTuple):
    """Что считать подтверждённым числом: значения корпуса и веб-фактов,
    их детерминированные производные, годы из текста источников."""

    values: frozenset
    years: frozenset


def _matches(value: Decimal, pool: frozenset[Decimal], tol: Decimal) -> bool:
    """Совпадение с допуском округления: кандидат, округлённый до точности
    наблюдаемого, совпадает с ним (700 ← 700.1; 19,7 ← 19.7254)."""
    for candidate in pool:
        diff = abs(value - candidate)
        if diff <= tol or diff <= abs(candidate) * Decimal("1e-9"):
            return True
    return False


def number_supported(value: Decimal, evidence: NumberEvidence, slides: int) -> bool:
    decimals = max(0, -value.as_tuple().exponent)
    base_tol = Decimal(1).scaleb(-decimals) / 2
    if _matches(value, evidence.values, base_tol):
        return True
    # проценты ↔ доли и масштабирование оси (тыс./млн/млрд); допуск
    # масштабируется вместе со значением (0,197 доли ↔ 19,7 %)
    for factor in (
        Decimal(100),
        Decimal("0.01"),
        Decimal(10) ** 3,
        Decimal(10) ** 6,
        Decimal(10) ** 9,
        Decimal(10) ** -3,
        Decimal(10) ** -6,
        Decimal(10) ** -9,
    ):
        if _matches(value * factor, evidence.values, base_tol * abs(factor)):
            return True
    if value == value.to_integral_value():
        # порядковые номера слайдов/разделов — не факты
        if 1 <= value <= max(slides, 1):
            return True
        if 1900 <= value <= 2100 and int(value) in evidence.years:
            return True
    return False


def build_number_evidence(texts, files) -> NumberEvidence:
    values = set().union(*(_number_values(t) for t in texts)) if texts else set()
    for path in files:
        values |= derived_numbers(source_columns(Path(path)))
    years = {int(y) for t in texts for y in _YEAR.findall(t)}
    return NumberEvidence(frozenset(values), frozenset(years))


def web_evidence(sources: list[dict], pack: dict) -> list[str]:
    """Only source text and grounded web blocks, never numbers in URLs/IDs/dates."""
    texts = []
    source_ids = {s["id"] for s in sources if s.get("id")}
    for source in sources:
        for field in ("title", "text", "quote", "content"):
            if isinstance(source.get(field), str):
                texts.append(source[field])
    grounded_ids = {
        s["id"]
        for s in pack.get("external_sources", [])
        if s.get("origin") == "web" and s.get("id") in source_ids
    }
    for section in pack.get("sections", []):
        # Same bibliography headings as the API content-fill contract.
        if (
            section.get("heading", "")
            .strip()
            .casefold()
            .startswith(
                (
                    "источники",
                    "список источников",
                    "литература",
                    "references",
                    "bibliography",
                    "sources",
                )
            )
        ):
            continue
        for block in section.get("blocks", []):
            if (block.get("source_ref") or {}).get("artifact_id") in grounded_ids:
                texts.extend(t for t in [block.get("text"), *(block.get("items") or [])] if t)
    return texts


def shapes_tree(shapes):
    for shape in shapes:
        if hasattr(shape, "shapes"):
            yield from shapes_tree(shape.shapes)
        else:
            yield shape


def native_kind(chart) -> str:
    name = chart.chart_type.name
    for prefix, kind in (
        ("DOUGHNUT", "doughnut"),
        ("PIE", "pie"),
        ("LINE", "line"),
        ("XY_SCATTER", "scatter"),
        ("AREA", "area"),
    ):
        if name.startswith(prefix):
            return kind
    if name.startswith("BAR_"):
        return "bar_horizontal"
    if name.startswith("COLUMN_"):
        if "100" in name:
            return "bar_stacked_100"
        return "bar_stacked" if "STACKED" in name else "bar"
    return name.lower()


def slide_content(slide) -> dict:
    shapes = list(shapes_tree(slide.shapes))
    text_shapes = [s for s in shapes if s.has_text_frame and s.text.strip()]
    texts = [s.text for s in text_shapes]
    for shape in shapes:
        if shape.has_table:
            texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        if shape.has_chart:
            # Include titles, series, categories and plotted values from chart parts.
            # Deduplicate repeated category caches across series in the same chart.
            texts.extend(
                dict.fromkeys(
                    node.text
                    for node in shape.chart._chartSpace.xpath(
                        ".//a:t | .//c:strCache/c:pt/c:v | .//c:numCache/c:pt/c:v"
                        " | .//c:strLit/c:pt/c:v | .//c:numLit/c:pt/c:v"
                    )
                    if node.text
                )
            )
    charts = [native_kind(s.chart) for s in shapes if s.has_chart]
    table = any(s.has_table for s in shapes)
    # Placeholder is authoritative when present; otherwise use the top text shape.
    title_shape = slide.shapes.title
    if title_shape is None or not title_shape.text.strip():
        title_shape = min(text_shapes, key=lambda s: (s.top, s.left)) if text_shapes else None
    title = title_shape.text.strip() if title_shape is not None else ""
    bullet_paragraphs = 0
    for shape in text_shapes:
        if shape == title_shape:
            continue
        paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
        explicit = sum(
            bool(re.match(r"\s*[•●▪–-]\s", p.text))
            or bool(p._p.xpath("./a:pPr/a:buChar|./a:pPr/a:buAutoNum"))
            for p in paragraphs
        )
        bullet_paragraphs += max(explicit, len(paragraphs) if len(paragraphs) >= 3 else 0)
    is_list = bullet_paragraphs >= 3 and not (charts or table)
    archetype = (
        "chart"
        if charts
        else "table"
        if table
        else "list"
        if is_list
        else "picture"
        if any(s.shape_type == 13 for s in shapes)
        else "text"
    )
    return {
        "title": title,
        "text": "\n".join(texts),
        "charts": charts,
        "table": table,
        "list": is_list,
        "archetype": archetype,
        "layout": str(slide.slide_layout.part.partname),
    }


def measure_deck(
    path: Path,
    *,
    metadata=(),
    allowed_text=(),
    expected=None,
    forbidden=None,
    evidence: NumberEvidence | None = None,
) -> dict:
    """Метрики финального PPTX. Индексы metadata — API 1-based, а не индекс плана."""
    slides = [slide_content(s) for s in Presentation(path).slides]
    by_index = {m["index"]: m for m in metadata}
    expected = expected or {}
    data = [
        s
        for i, s in enumerate(slides, 1)
        if by_index.get(i, {}).get("purpose") == "data" or s["charts"] or s["table"]
    ]
    all_text = "\n".join(s["text"] for s in slides)
    observed = _number_values(all_text)
    evidence = evidence or build_number_evidence(allowed_text, ())
    unsupported = sorted(
        (v for v in observed if not number_supported(v, evidence, len(slides))),
        key=Decimal,
    )
    unsupported = [format(v.normalize(), "f") if v else "0" for v in unsupported]
    counts = {
        f: len(re.findall(re.escape(f), all_text, re.I))
        for f in (FORBIDDEN if forbidden is None else forbidden)
    }
    # Avoid counting viz_reason twice; other Latin service slugs remain candidates.
    slugs = [m.group() for m in SLUG.finditer(all_text) if m.group().lower() != "viz_reason"]
    counts["english_slug"] = len(slugs)
    checks = []
    for dataset, kinds in expected.items():
        observed_kinds = set()
        for i, slide in enumerate(slides, 1):
            meta = by_index.get(i, {})
            if Path(meta.get("dataset") or "").name != dataset:
                continue
            actual = set(slide["charts"])
            if slide["table"]:
                actual.add("table")
            # Histogram is emitted as editable column bins by the compiler.
            if meta.get("planned_type") == "histogram" and "bar" in actual:
                actual.add("histogram")
            # KPI cards have no distinct OOXML chart class: require visible numbers.
            if meta.get("planned_type") == "kpi" and not actual and numbers(slide["text"]):
                actual.add("kpi")
            observed_kinds.update(actual)
        checks.append(
            {
                "dataset": dataset,
                "expected": kinds,
                "observed": sorted(observed_kinds),
                "match": bool(set(kinds) & observed_kinds),
            }
        )
    archetypes = [s["archetype"] for s in slides]
    return {
        "slides": len(slides),
        "data_slides": len(data),
        "data_title_number_share": (
            sum(bool(numbers(s["title"])) for s in data) / len(data)
            if data
            else 0.0
            if expected
            else None
        ),
        "label_title_count": sum(bool(LABEL.search(s["title"])) for s in data),
        "unsupported_numbers": unsupported,
        "unsupported_number_count": len(unsupported),
        "unsupported_number_share": len(unsupported) / len(observed) if observed else 0.0,
        "numbers_total": len(observed),
        "number_check": "raw_inputs_derived_stats_web_years_and_page_numbers",
        "archetype_count": len(set(archetypes)),
        "archetypes": dict(Counter(archetypes)),
        "layout_count": len({s["layout"] for s in slides}),
        "max_archetype_run": max((len(list(g)) for _, g in groupby(archetypes)), default=0),
        "list_slide_share": sum(s["list"] for s in slides) / len(slides) if slides else 0.0,
        "forbidden_counts": counts,
        "forbidden_total": sum(counts.values()),
        "english_slugs": sorted(set(slugs)),
        "chart_checks": checks,
        "chart_match_share": sum(c["match"] for c in checks) / len(checks) if checks else None,
        "slide_details": [{"index": i, **s} for i, s in enumerate(slides, 1)],
    }


def score_deck(metrics: dict, config: dict) -> tuple[float, dict]:
    n = max(metrics["slides"], 1)
    if not metrics["slides"]:
        return 0.0, {}
    components = {
        "audit": max(0, 1 - (5 * metrics["audit_blockers"] + metrics["audit_errors"]) / n),
        "numeric_titles": metrics["data_title_number_share"],
        "label_titles": max(0, 1 - metrics["label_title_count"] / max(metrics["data_slides"], 1)),
        "unsupported_numbers": 1 - metrics["unsupported_number_share"],
        "diversity": min(1, metrics["archetype_count"] / config["diversity_target"])
        * min(1, config["max_run_target"] / max(metrics["max_archetype_run"], 1)),
        "lists": max(
            0,
            1
            - max(0, metrics["list_slide_share"] - config["list_share_target"])
            / (1 - config["list_share_target"]),
        ),
        "forbidden": max(0, 1 - metrics["forbidden_total"] / n),
        "chart_match": metrics["chart_match_share"],
        "latency": min(
            1,
            config["latency_budget_seconds"] / max(metrics["build_seconds_three_variants"], 0.001),
        ),
    }
    applicable = {
        key: weight for key, weight in config["weights"].items() if components[key] is not None
    }
    score = 100 * sum(components[k] * w for k, w in applicable.items()) / sum(applicable.values())
    return round(score, 3), components


# --- Жёсткие гейты: провал любого → колода получает 0 до общего балла ---

# Плановые типы, которые обязаны материализоваться нативным chart part,
# а не картинкой/надписью. histogram → column chart, kpi → карточки текста.
_CHART_PLANS = {
    "line",
    "bar",
    "bar_horizontal",
    "bar_stacked",
    "bar_stacked_100",
    "pie",
    "doughnut",
    "scatter",
    "area",
    "histogram",
    "waterfall",
    "combo",
}
_TABLE_PLANS = {"table", "reference"}

_SOURCE_RE = re.compile(r"источник|source", re.I)
_UNIT_RE = re.compile(
    r"(%|‰|₽|руб(?:\.|лей|ля|ль|лей)?|млн|млрд|тыс\.?|трлн|шт\.?|балл|ед\.|"
    r"запис|позиц|заказ|клиент|чел\.|сотрудн|услуг|показател|значени|доля|"
    r"раза?\b|\$|€|₸|\bUSD|\bEUR|\bGBP|\bCNY)",
    re.I,
)
# Период — слова календаря/единиц времени; голое число «2052» из данных —
# не период. Год засчитывается только в строке с буквами или внутри
# даты вида «2025-01»/«14.03.2025» — иначе суммы вроде «1901 ₽»
# ложно маркировали бы и слайд, и источник как периодические.
_PERIOD_WORDS = re.compile(
    r"январ|феврал|март|апрел|\bма[йя]\b|июн|июл|август|сентябр|октябр|ноябр|"
    r"декабр|квартал|\bкв\.?\s*\d|\bQ[1-4]\b|недел|месяц|ежемесячн|еженедел|"
    r"ежеднев|ежеквартал|период|\bгод[ауе]?\b|\bлет\b|\bдн(?:ей|я|ь)\b|динамик",
    re.I,
)
_DATE_LIKE = re.compile(r"(?:19|20)\d{2}-\d{1,2}\b|\b\d{1,2}\.(?:0?[1-9]|1[0-2])\.(?:19|20)\d{2}")
_HAS_LETTER = re.compile(r"[A-Za-zА-Яа-яЁё]")
# «1901 ₽»/«2025 шт.» — значение с единицей, а не год. Годом считается
# число 19xx/20xx, за которым не следует единица измерения.
_YEAR_AS_VALUE = re.compile(
    r"(?<![\d.,])(?:19|20)\d{2}(?![\d.,])"
    r"(?!\s*[₽%‰€$₸]|\s*(?:млн|млрд|тыс|трлн|руб(?:\.|лей|ля|ль|лей)?|шт|балл|ед|"
    r"пункт|заказ|клиент|чел|услуг|позиц|запис)\b)",
    re.I,
)


def has_period(text: str) -> bool:
    """Есть ли в тексте указание на период данных (слова, даты, год
    рядом с буквами — «за 2025 год», а не значение «1901 ₽»)."""
    if _PERIOD_WORDS.search(text) or _DATE_LIKE.search(text):
        return True
    return any(
        _YEAR_AS_VALUE.search(line) and _HAS_LETTER.search(line) for line in text.splitlines()
    )


# допуск на границы: ~1 мм — ловит съехавшие объекты, прощает вылеты
# фоновых подложек на долю миллиметра
_BOUNDS_TOL_EMU = 36000


def _group_xfrm(shape):
    """off/ext/chOff/chExt группы — пересчёт дочерних координат в родительские."""
    from pptx.oxml.ns import qn

    xfrm = shape._element.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    if xfrm is None:
        return None

    def point(tag):
        el = xfrm.find(qn(tag))
        return (int(el.get("x", 0)), int(el.get("y", 0))) if el is not None else (0, 0)

    def size(tag):
        el = xfrm.find(qn(tag))
        return (int(el.get("cx", 0)), int(el.get("cy", 0))) if el is not None else (0, 0)

    return (*point("a:off"), *size("a:ext"), *point("a:chOff"), *size("a:chExt"))


def _bounds_findings(slide, width, height):
    """Объекты, выступающие за край слайда больше допуска (с учётом групп
    и поворота вокруг центра)."""
    findings = []

    def emit(shape, x0, y0, x1, y1):
        rot = getattr(shape, "rotation", 0) or 0
        if rot:
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            angle = math.radians(rot)
            ca, sa = math.cos(angle), math.sin(angle)
            pts = [
                (cx + px * ca - py * sa, cy + px * sa + py * ca)
                for px, py in (
                    (x0 - cx, y0 - cy),
                    (x1 - cx, y0 - cy),
                    (x0 - cx, y1 - cy),
                    (x1 - cx, y1 - cy),
                )
            ]
            x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
            x1, y1 = max(p[0] for p in pts), max(p[1] for p in pts)
        overflow = max(-x0, -y0, x1 - width, y1 - height)
        if overflow > _BOUNDS_TOL_EMU:
            findings.append(f"{shape.name} выступает на {round(overflow / 36000, 1)} мм за край")

    def walk(shapes, ox, oy, sx, sy):
        for shape in shapes:
            if hasattr(shape, "shapes"):
                xfrm = _group_xfrm(shape)
                if xfrm is None:
                    continue
                gx, gy, gcx, gcy, chx, chy, ccx, ccy = xfrm
                nsx = sx * gcx / ccx if ccx else sx
                nsy = sy * gcy / ccy if ccy else sy
                walk(shape.shapes, ox + sx * gx - nsx * chx, oy + sy * gy - nsy * chy, nsx, nsy)
                continue
            box = (shape.left, shape.top, shape.width, shape.height)
            if any(v is None for v in box):
                continue  # геометрия наследуется от layout — в границах
            left, top, w, h = box
            emit(shape, ox + sx * left, oy + sy * top, ox + sx * (left + w), oy + sy * (top + h))

    walk(slide.shapes, 0, 0, 1.0, 1.0)
    return findings


def _axis_titles(chart) -> list[str]:
    return [
        n.text
        for n in chart._chartSpace.xpath(".//c:valAx/c:title//a:t | .//c:catAx/c:title//a:t")
        if n.text and n.text.strip()
    ]


def hard_gates(path: Path, *, metadata=(), dataset_periods=None) -> dict:
    """Жёсткие гейты финального PPTX. Любой провал → колода 0, причина —
    в ``failures``. Геометрия считается по реальным частям, без знания
    о конкретных шаблонах."""
    dataset_periods = dataset_periods or {}
    failures = []

    def check(name, fn):
        try:
            for detail in fn():
                failures.append({"gate": name, "detail": detail})
        except Exception as exc:  # сама проверка упала — это провал
            failures.append({"gate": name, "detail": f"{type(exc).__name__}: {exc}"})

    # pptx_reopens: файл открывается и переживает round-trip «сохранить →
    # открыть» без потери слайдов
    try:
        prs = Presentation(path)
    except Exception as exc:
        return {
            "ok": False,
            "failures": [{"gate": "pptx_reopens", "detail": f"{type(exc).__name__}: {exc}"}],
        }

    def _reopens():
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        if len(Presentation(buf).slides) != len(prs.slides):
            yield "перечитанная копия потеряла слайды"

    check("pptx_reopens", _reopens)

    # package_relationships: целостный zip и разрешимые внутренние rels
    def _rels():
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
            if bad:
                yield f"повреждённый член пакета: {bad}"
        for part in prs.part.package.iter_parts():
            for rel in part.rels.values():
                if rel.is_external:
                    continue
                try:
                    _ = rel.target_part
                except Exception:
                    yield f"{part.partname}: битый rels {rel.rId} → {rel.target_ref}"

    check("package_relationships", _rels)

    # native_visuals: каждый chart part имеет embedded workbook; заданные
    # планом графики/таблицы материализованы нативно, а не картинкой
    by_index = {m["index"]: m for m in metadata}

    def _native():
        for i, slide in enumerate(prs.slides, 1):
            shapes = list(shapes_tree(slide.shapes))
            charts = [s for s in shapes if s.has_chart]
            for chart_shape in charts:
                try:
                    _ = chart_shape.chart.part.chart_workbook.xlsx_part
                except Exception:
                    yield f"слайд {i}: chart part без embedded workbook"
            planned = (by_index.get(i) or {}).get("planned_type")
            if planned in _CHART_PLANS and not charts:
                yield f"слайд {i}: запланирован «{planned}», нативного графика нет"
            if planned in _TABLE_PLANS and not any(s.has_table for s in shapes):
                yield f"слайд {i}: запланирована таблица, table object отсутствует"

    check("native_visuals", _native)

    # within_bounds: ничего не выступает за край слайда (допуск ~1 мм)
    def _bounds():
        for i, slide in enumerate(prs.slides, 1):
            for detail in _bounds_findings(slide, prs.slide_width, prs.slide_height):
                yield f"слайд {i}: {detail}"

    check("within_bounds", _bounds)

    # data_context: на слайдах с данными есть источник, единицы/подпись
    # метрики и период (период — если источник его содержал)
    def _context():
        for i, slide in enumerate(prs.slides, 1):
            shapes = list(shapes_tree(slide.shapes))
            meta = by_index.get(i) or {}
            content = slide_content(slide)
            if not (meta.get("purpose") == "data" or content["charts"] or content["table"]):
                continue
            text = content["text"]
            dataset = Path(meta.get("dataset") or "").name
            if not _SOURCE_RE.search(text) and not (dataset and dataset in text):
                yield f"слайд {i}: у слайда с данными нет отметки об источнике"
            units_ok = bool(_UNIT_RE.search(text)) or any(
                _axis_titles(s.chart) for s in shapes if s.has_chart
            )
            if _number_values(text) and not units_ok:
                yield f"слайд {i}: единицы/метрика данных не подписаны"
            if dataset_periods.get(dataset, True) and not has_period(text):
                yield f"слайд {i}: период данных не указан"

    check("data_context", _context)

    # no_empty_duplicate: нет пустых, донорских и повторяющихся слайдов
    # (донорские «Thanks!»/«CREDITS»/«Slidesgo», «Источники» ×2, дубли заголовков)
    donor_re = re.compile(r"(?:Thanks!|\bCREDITS\b|\bSlidesgo\b)", re.I)

    def _empty_dup():
        seen_text, seen_titles = {}, {}
        for i, slide in enumerate(prs.slides, 1):
            content = slide_content(slide)
            has_picture = any(s.shape_type == 13 for s in shapes_tree(slide.shapes))
            if not (
                content["text"].strip() or content["charts"] or content["table"] or has_picture
            ):
                yield f"слайд {i}: пустой слайд"
            donor = donor_re.search(content["text"])
            if donor:
                yield f"слайд {i}: донорский титр шаблона ({donor.group(0)})"
            body = re.sub(r"\s+", " ", content["text"].strip().casefold())
            title = re.sub(r"\s+", " ", content["title"].strip().casefold())
            if body and body in seen_text:
                yield f"слайд {i}: дублирует слайд {seen_text[body]}"
            elif body:
                seen_text[body] = i
            if title and title in seen_titles:
                yield f"слайд {i}: повтор заголовка со слайдом {seen_titles[title]}"
            elif title:
                seen_titles[title] = i

    check("no_empty_duplicate", _empty_dup)
    return {"ok": not failures, "failures": failures}


def trio_title_gates(rows) -> dict[str, str]:
    """Три варианта одного сценария должны нести одни и те же числа
    в заголовках одинаковых позиций (один семантический план)."""
    measured = [r for r in rows if (r.get("metrics") or {}).get("slide_details")]
    if len(measured) < 2:
        return {}
    titles = {
        r["strategy"]: {d["index"]: numbers(d["title"]) for d in r["metrics"]["slide_details"]}
        for r in measured
    }
    strategies = [r["strategy"] for r in measured]
    details = []
    for index in sorted({i for t in titles.values() for i in t}):
        per = {s: tuple(sorted(titles[s].get(index, set()))) for s in strategies}
        if len(set(per.values())) > 1:
            details.append(
                f"слайд {index}: " + "; ".join(f"{s}={list(v) or '—'}" for s, v in per.items())
            )
    if not details:
        return {}
    detail = "; ".join(details[:5])
    return {s: f"variant_title_numbers: {detail}" for s in strategies}


def contact_sheet(pdf: Path, png: Path) -> None:
    """Poppler 40 dpi, numeric page order, one PNG per deck."""
    with tempfile.TemporaryDirectory(prefix="deck-bench-pages-") as temp:
        prefix = Path(temp) / "page"
        subprocess.run(  # noqa: S603
            ["pdftoppm", "-r", "40", "-png", str(pdf), str(prefix)],  # noqa: S607
            check=True,
            capture_output=True,
            timeout=180,
        )
        pages = sorted(Path(temp).glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
        if not pages:
            raise ValueError("PDF has no rendered pages")
        with Image.open(pages[0]) as first:
            w, h = first.size
        sheet = Image.new("RGB", (4 * w, math.ceil(len(pages) / 4) * h), "#dddddd")
        for index, page in enumerate(pages):
            with Image.open(page) as img:
                sheet.paste(img.convert("RGB"), ((index % 4) * w, (index // 4) * h))
        png.parent.mkdir(parents=True, exist_ok=True)
        sheet.save(png)
        sheet.close()


def write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def http_client(base_url: str, timeout: float):
    # Direct connection: the farm must not depend on shell SOCKS/proxy settings.
    return httpx.Client(base_url=base_url.rstrip("/"), timeout=timeout, trust_env=False)


class API:
    def __init__(self, client, timeout=1200, poll=1):
        self.client, self.timeout, self.poll = client, timeout, poll

    def request(self, method, path, **kwargs):
        response = self.client.request(method, path, **kwargs)
        response.raise_for_status()
        return response.json()

    def wait(self, job_id):
        deadline = time.monotonic() + self.timeout
        while time.monotonic() < deadline:
            job = self.request("GET", f"/jobs/{job_id}")
            if job["state"] == "completed":
                return job
            if job["state"] in {"failed", "canceled", "awaiting_user"}:
                raise RuntimeError(f"Job {job_id}: {job['state']}: {job.get('error')}")
            time.sleep(self.poll)
        raise TimeoutError(f"Job {job_id} exceeded {self.timeout}s (not retried)")

    def items(self, path):
        result, cursor = [], None
        while True:
            page = self.request(
                "GET", path, params={"limit": 200, **({"cursor": cursor} if cursor else {})}
            )
            result.extend(page["items"])
            cursor = page.get("next_cursor")
            if not cursor:
                return result

    def download(self, artifact, path):
        # Use artifact ID to avoid duplicating /api/v1 and following external download URLs.
        response = self.client.get(f"/artifacts/{artifact['artifact_id']}/download")
        response.raise_for_status()
        data = response.content
        digest = hashlib.sha256(data).hexdigest()
        if artifact.get("sha256") and digest != artifact["sha256"]:
            raise ValueError("Downloaded artifact checksum mismatch")
        path.write_bytes(data)
        return {"path": str(path), "sha256": digest, "size_bytes": len(data)}


def download_exports(api, export: dict, directory: Path) -> dict:
    filenames = {"pptx": "deck.pptx", "pdf": "deck.pdf", "quality_passport": "passport.json"}
    current = {a["format"]: a for a in export["artifacts"]}
    missing = filenames.keys() - current.keys()
    if missing:
        raise ValueError(f"Missing export formats: {sorted(missing)}")
    if len(current) != len(export["artifacts"]):
        raise ValueError("Duplicate export formats")
    return {fmt: api.download(current[fmt], directory / name) for fmt, name in filenames.items()}


def resolve_inputs(corpus, template_root: Path | None, out: Path):
    templates = {}
    for key, entry in corpus["templates"].items():
        candidates = [template_root / entry["file"]] if template_root else []
        candidates.append(ROOT / entry["fallback"])
        path = next((p for p in candidates if p.is_file()), None)
        if path is None and entry.get("generator"):
            module_path, function = entry["generator"].split(":")
            spec = importlib.util.spec_from_file_location("bench_fixture", ROOT / module_path)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            path = out / "inputs" / entry["file"]
            getattr(module, function)(path)
        if path is None:
            raise FileNotFoundError(f"Template {key}: {candidates}")
        templates[key] = path
    sources = {key: ROOT / entry["path"] for key, entry in corpus["datasets"].items()}
    brand = [ROOT / p for p in corpus["brandbook"]]
    contents = [ROOT / s["content"] for s in corpus["scenarios"] if s.get("content")]
    all_paths = [*templates.values(), *sources.values(), *brand, *contents]
    hashes = {
        str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p): hashlib.sha256(
            p.read_bytes()
        ).hexdigest()
        for p in all_paths
    }
    # The signature uses stable logical keys and bytes, not machine-specific absolute paths.
    logical = {
        **{
            f"template:{k}": hashlib.sha256(p.read_bytes()).hexdigest()
            for k, p in templates.items()
        },
        **{f"data:{k}": hashlib.sha256(p.read_bytes()).hexdigest() for k, p in sources.items()},
        **{
            str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in [*brand, *contents]
        },
    }
    return templates, sources, brand, hashes, logical


def run_scenario(scenario, corpus, templates, sources, brand, args):
    sid = scenario["id"]
    directory = args.out / sid
    directory.mkdir(parents=True, exist_ok=True)
    rows = []
    state = {"scenario": sid, "use_llm": args.use_llm}
    started = time.monotonic()
    with http_client(args.base_url, args.request_timeout) as client:
        api = API(client, args.timeout, args.poll)
        try:
            pid = api.request("POST", "/projects", json={"name": f"bench-v1-{sid}"})["id"]
            state["project_id"] = pid
            paths = [templates[scenario["template"]], *[sources[k] for k in scenario["datasets"]]]
            if scenario.get("content"):
                paths.append(ROOT / scenario["content"])
            if scenario.get("brandbook"):
                paths.extend(brand)
            files = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in paths]
            intake = api.request(
                "POST", f"/projects/{pid}/intake", files=files, data={"idea": scenario["idea"]}
            )
            state["intake"] = intake
            write_json(directory / "run.json", state)
            if intake.get("job_id"):
                api.wait(intake["job_id"])
            rejected = [f for f in intake.get("files", []) if not f["accepted"]]
            if rejected:
                raise ValueError(f"Rejected corpus files: {rejected}")
            if scenario.get("brandbook") and not intake.get("brandbook_id"):
                raise ValueError("Brandbook was not recognized")
            brief = {
                "purpose": scenario["idea"],
                "audience": "руководство и партнёры",
                "language": "ru",
                "target_slide_count": scenario["slides"],
                "content_fill": "shrink",
            }
            if scenario.get("brandbook"):
                api.request(
                    "PATCH", f"/brandbooks/{intake['brandbook_id']}", json={"priority": "brandbook"}
                )
            gen_started = time.monotonic()
            accepted = api.request(
                "POST",
                f"/projects/{pid}/generations",
                json={
                    "template_id": intake["template_id"],
                    "content_pack_id": intake["content_pack_id"],
                    "brief": brief,
                    "use_llm": args.use_llm,
                    "variants": [{"strategy": s} for s in corpus["strategies"]],
                },
            )
            state["accepted"] = accepted
            write_json(directory / "run.json", state)
            generation_error = None
            try:
                api.wait(accepted["job_id"])
            except (RuntimeError, TimeoutError) as exc:
                generation_error = str(exc)
            build_seconds = round(time.monotonic() - gen_started, 3)
            generation = api.request("GET", f"/generations/{accepted['generation_id']}")
            state.update(
                generation=generation,
                build_seconds_three_variants=build_seconds,
                generation_error=generation_error,
            )
            write_json(directory / "run.json", state)
            expected = {
                sources[k].name: corpus["datasets"][k]["expected"] for k in scenario["datasets"]
            }
            allowed = [scenario["idea"], *[source_text(sources[k]) for k in scenario["datasets"]]]
            if scenario.get("content"):
                allowed.append(source_text(ROOT / scenario["content"]))
            data_files = [sources[k] for k in scenario["datasets"]]
            if scenario.get("content"):
                data_files.append(ROOT / scenario["content"])
            base_evidence = build_number_evidence(allowed, data_files)
            dataset_periods = {
                sources[k].name: has_period(source_text(sources[k])) for k in scenario["datasets"]
            }
            web_sources = ((generation.get("content_fill") or {}).get("web") or {}).get(
                "sources", []
            )
            state["web_sources"] = web_sources
            evidence_packs = {}
            for variant in generation["variants"]:
                row = {
                    "scenario": sid,
                    "strategy": variant["strategy"],
                    "score": 0,
                    "status": "failed",
                    "use_llm": args.use_llm,
                    "variant_id": variant["id"],
                    "project_id": pid,
                    "generation_id": generation["id"],
                    "build_seconds_three_variants": build_seconds,
                }
                rows.append(row)
                try:
                    if variant["status"] != "completed":
                        raise RuntimeError(f"Variant {variant['status']}: {generation_error}")
                    variant_dir = directory / variant["strategy"]
                    variant_dir.mkdir(exist_ok=True)
                    exported = api.request(
                        "POST",
                        f"/variants/{variant['id']}/exports",
                        json={"formats": ["pptx", "pdf", "quality_passport"]},
                    )
                    api.wait(exported["job_id"])
                    export = api.request("GET", f"/exports/{exported['export_id']}")
                    artifacts = download_exports(api, export, variant_dir)
                    row["artifacts"] = artifacts
                    passport = json.loads((variant_dir / "passport.json").read_text())
                    pack = {}
                    if web_sources:
                        pack_id = passport["inputs"]["content_pack_id"]
                        if pack_id not in evidence_packs:
                            evidence_packs[pack_id] = api.request(
                                "GET", f"/content-packs/{pack_id}"
                            )
                        pack = evidence_packs[pack_id]
                        write_json(variant_dir / "evidence-pack.json", pack)
                    web_texts = web_evidence(web_sources, pack)
                    write_json(variant_dir / "web-evidence.json", web_texts)
                    slides = api.items(f"/variants/{variant['id']}/slides")
                    metadata = []
                    for slide in slides:
                        explain = api.request("GET", f"/slides/{slide['id']}/explain")
                        metadata.append(
                            {
                                **slide,
                                "dataset": (explain.get("dataset") or {}).get("file"),
                                "planned_type": (explain.get("chosen") or {}).get("type"),
                            }
                        )
                    write_json(variant_dir / "slides.json", metadata)
                    audited = api.request("POST", f"/variants/{variant['id']}/audits", json={})
                    api.wait(audited["job_id"])
                    audit = api.request("GET", f"/audits/{audited['audit_id']}")
                    write_json(variant_dir / "audit.json", audit)
                    if audit.get("deterministic_status") != "completed":
                        raise ValueError("Deterministic audit is incomplete")
                    evidence = NumberEvidence(
                        base_evidence.values | frozenset(_number_values("\n".join(web_texts))),
                        base_evidence.years
                        | frozenset(int(y) for y in _YEAR.findall("\n".join(web_texts))),
                    )
                    metrics = measure_deck(
                        variant_dir / "deck.pptx",
                        metadata=metadata,
                        allowed_text=[*allowed, *web_texts],
                        expected=expected,
                        forbidden=corpus["forbidden"],
                        evidence=evidence,
                    )
                    summary = audit["summary_by_severity"]
                    metrics.update(
                        audit_blockers=summary.get("blocker", 0),
                        audit_errors=summary.get("error", 0),
                        audit_warnings=summary.get("warning", 0),
                        contextual_status=audit.get("contextual_status"),
                        build_seconds_three_variants=build_seconds,
                    )
                    gates = hard_gates(
                        variant_dir / "deck.pptx",
                        metadata=metadata,
                        dataset_periods=dataset_periods,
                    )
                    metrics["gates"] = gates
                    row.update(
                        metrics=metrics,
                        passport_usage=passport.get("metrics", {}).get("usage"),
                        planner=variant.get("planner"),
                    )
                    row["score"], row["score_components"] = score_deck(metrics, corpus["scoring"])
                    png = args.out / f"{sid}__{variant['strategy']}.png"
                    contact_sheet(variant_dir / "deck.pdf", png)
                    row.update(status="completed", contact_sheet=str(png))
                    if not gates["ok"]:
                        detail = "; ".join(
                            f"{g['gate']}: {g['detail']}" for g in gates["failures"][:5]
                        )
                        row.update(status="failed", score=0, error=f"hard gates: {detail}")
                except Exception as exc:  # noqa: BLE001 — preserve other variants on failure
                    row.update(status="failed", score=0, error=f"{type(exc).__name__}: {exc}")
                write_json(directory / "decks.json", rows)
            trio = trio_title_gates(rows)
            for row in rows:
                if row["strategy"] in trio and row["status"] == "completed":
                    row["metrics"].setdefault("gates", {"ok": True, "failures": []})
                    row["metrics"]["gates"]["failures"].append(
                        {"gate": "variant_title_numbers", "detail": trio[row["strategy"]]}
                    )
                    row["metrics"]["gates"]["ok"] = False
                    row.update(status="failed", score=0, error=trio[row["strategy"]])
            if trio:
                write_json(directory / "decks.json", rows)
            if generation_error:
                state["error"] = generation_error
        except Exception as exc:  # noqa: BLE001 — failed cases must not disappear from baseline
            state["error"] = f"{type(exc).__name__}: {exc}"
        finally:
            state["wall_seconds"] = round(time.monotonic() - started, 3)
            write_json(directory / "run.json", state)
    existing = {row["strategy"] for row in rows}
    for strategy in corpus["strategies"]:
        if strategy not in existing:
            rows.append(
                {
                    "scenario": sid,
                    "strategy": strategy,
                    "status": "failed",
                    "score": 0,
                    "use_llm": args.use_llm,
                    "error": state.get("error", "Missing variant"),
                }
            )
    print(f"{sid}: {sum(r['status'] == 'completed' for r in rows)}/{len(rows)} decks", flush=True)
    return rows, state


def markdown(report: dict) -> str:
    lines = [
        "| Сценарий / вариант | Балл | Гейты | 3 варианта, с | Блок./ош. | "
        "Загол. с числом | Ярлыки | Неподтв. числа | Арх./макеты/серия | Списки | "
        "Запр. | Графики |",
        "|---|---:|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]

    def pct(value):
        return "N/A" if value is None else f"{value:.0%}"

    for row in report["decks"]:
        m = row.get("metrics")
        name = f"{row['scenario']} / {row['strategy']}"
        if not m:
            lines.append(f"| {name} | 0 (FAILED) | — | — | — | — | — | — | — | — | — | — |")
            continue
        gates = m.get("gates") or {}
        gate_cell = (
            "✓"
            if gates.get("ok")
            else "✗ " + ", ".join(sorted({g["gate"] for g in gates.get("failures", [])}))
        )
        lines.append(
            f"| {name} | {row['score']:.2f} | {gate_cell} | "
            f"{m['build_seconds_three_variants']:.1f} | "
            f"{m['audit_blockers']}/{m['audit_errors']} | "
            f"{pct(m['data_title_number_share'])} | {m['label_title_count']} | "
            f"{len(m['unsupported_numbers'])} | "
            f"{m['archetype_count']}/{m['layout_count']}/{m['max_archetype_run']} | "
            f"{pct(m['list_slide_share'])} | {m['forbidden_total']} | "
            f"{pct(m['chart_match_share'])} |"
        )
    lines.extend(["", f"Среднее (включая неуспешные): {report['score']:.3f}/100.", ""])
    return "\n".join(lines)


def compare(old: dict, new: dict, tolerance=2) -> tuple[int, str]:
    if old.get("fingerprint") != new.get("fingerprint") or not old.get("fingerprint"):
        return 2, "Incompatible corpus, inputs, scoring or model mode; rerun baseline."

    def indexed(report):
        rows = report["decks"]
        for row in rows:
            score = row["score"]
            if type(score) not in (int, float) or not math.isfinite(score) or not 0 <= score <= 100:
                raise ValueError(f"Invalid deck score: {score}")
        result = {(r["scenario"], r["strategy"]): r for r in rows}
        if len(result) != len(rows) or not rows:
            raise ValueError("Empty or duplicate deck keys")
        return result

    before, after = indexed(old), indexed(new)
    if before.keys() != after.keys():
        return 2, "Missing/extra scenarios or variants; comparison rejected."

    def gate_ids(row):
        return {
            g["gate"] for g in ((row.get("metrics") or {}).get("gates") or {}).get("failures", [])
        }

    failed = False
    lines = [
        "| Колода | Δ балл | Гейты | Δ числовые метрики (new − old) |",
        "|---|---:|---|---|",
    ]
    for key, row in before.items():
        other = after[key]
        delta = round(other["score"] - row["score"], 6)
        # приёмка: гейты не хуже (нет новых провалов), корректность чисел
        # не хуже (неподтверждённых не стало больше), балл не ниже −2,
        # ранее собиравшаяся колода не сломалась
        new_gates = gate_ids(other) - gate_ids(row)
        a, b = row.get("metrics", {}), other.get("metrics", {})
        numbers_worse = (b.get("unsupported_number_count") or 0) > (
            a.get("unsupported_number_count") or 0
        )
        newly_broken = other["status"] != "completed" and row["status"] == "completed"
        failed |= delta < -tolerance or bool(new_gates) or numbers_worse or newly_broken
        gate_cell = (
            "✓"
            if not gate_ids(other)
            else "✗ " + ", ".join(sorted(gate_ids(other))) + (" (новые)" if new_gates else "")
        )
        changes = [
            f"{k}: {b[k] - a[k]:+.4g}"
            for k in sorted(a.keys() & b.keys())
            if isinstance(a[k], (float, int)) and isinstance(b[k], (float, int))
        ]
        if numbers_worse:
            changes.append("⚠ чисел неподтверждённых стало больше")
        lines.append(f"| {' / '.join(key)} | {delta:+.3f} | {gate_cell} | {'; '.join(changes)} |")
    return int(failed), "\n".join(lines) + "\n"


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8088/api/v1")
    parser.add_argument("--corpus", type=Path, default=ROOT / "configs/bench/corpus.v1.json")
    parser.add_argument("--template-root", type=Path)
    parser.add_argument("--out", type=Path, default=ROOT / "bench/latest")
    parser.add_argument("--json", type=Path, help="Итоговый JSON (по умолчанию --out/results.json)")
    parser.add_argument("--scenario", action="append", help="Повторяемый ID сценария")
    parser.add_argument("--use-llm", action="store_true")
    parser.add_argument("--concurrency", type=int, choices=(1, 2), default=1)
    parser.add_argument("--timeout", type=float, default=1800, help="Deadline job, seconds")
    parser.add_argument("--request-timeout", type=float, default=180)
    parser.add_argument("--poll", type=float, default=1)
    parser.add_argument("--compare", nargs=2, type=Path, metavar=("OLD", "NEW"))
    args = parser.parse_args(argv)
    if args.compare:
        try:
            old, new = [json.loads(p.read_text(encoding="utf-8")) for p in args.compare]
            code, table = compare(old, new, old.get("regression_tolerance", 2))
        except (ValueError, KeyError, OSError) as exc:
            code, table = 2, f"Invalid comparison: {exc}"
        print(table)
        return code
    corpus = json.loads(args.corpus.read_text(encoding="utf-8"))
    scenarios = [s for s in corpus["scenarios"] if not args.scenario or s["id"] in args.scenario]
    if not scenarios or (args.scenario and set(args.scenario) - {s["id"] for s in scenarios}):
        parser.error("Unknown/empty scenario selection")
    if args.use_llm and len(scenarios) > 2:
        parser.error("Model budget: select at most two scenarios with --scenario")
    if min(args.timeout, args.request_timeout, args.poll) <= 0:
        parser.error("Timeouts and poll interval must be positive")
    args.out.mkdir(parents=True, exist_ok=True)
    templates, sources, brand, hashes, logical = resolve_inputs(
        corpus, args.template_root, args.out
    )
    fingerprint = hashlib.sha256(
        json.dumps(
            {
                "corpus": corpus,
                "inputs": logical,
                "metric_version": METRIC_VERSION,
                "scenarios": [s["id"] for s in scenarios],
                "use_llm": args.use_llm,
            },
            sort_keys=True,
            ensure_ascii=False,
        ).encode()
    ).hexdigest()
    with http_client(args.base_url, 30) as client:
        capabilities = API(client).request("GET", "/capabilities")
    report = {
        "schema_version": "1",
        "metric_version": METRIC_VERSION,
        "fingerprint": fingerprint,
        "created_at": datetime.now(UTC).isoformat(),
        "base_url": args.base_url,
        "use_llm": args.use_llm,
        "capabilities": capabilities,
        "input_hashes": hashes,
        "scoring": corpus["scoring"],
        "regression_tolerance": corpus["scoring"]["regression_tolerance"],
        "decks": [],
        "runs": [],
        "score": 0,
    }
    output = args.json or args.out / "results.json"
    started = time.monotonic()
    with ThreadPoolExecutor(max_workers=args.concurrency) as executor:
        futures = [
            executor.submit(run_scenario, s, corpus, templates, sources, brand, args)
            for s in scenarios
        ]
        for future in as_completed(futures):
            rows, state = future.result()
            report["decks"].extend(rows)
            report["runs"].append(state)
            order = {s["id"]: i for i, s in enumerate(scenarios)}
            report["decks"].sort(
                key=lambda r: (order[r["scenario"]], corpus["strategies"].index(r["strategy"]))
            )
            report["score"] = round(
                sum(r["score"] for r in report["decks"]) / len(report["decks"]), 3
            )
            report["wall_seconds"] = round(time.monotonic() - started, 3)
            write_json(output, report)
            (args.out / "summary.md").write_text(markdown(report), encoding="utf-8")
    print(markdown(report))
    return 0 if all(r["status"] == "completed" for r in report["decks"]) else 1


if __name__ == "__main__":
    raise SystemExit(main())
