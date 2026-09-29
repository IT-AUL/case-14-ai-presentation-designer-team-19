"""Automated release/acceptance gate for the official submission constraints.

Nothing in the repo currently enforces pass/fail on the official
constraints as a single command: `deckdna benchmark` times `generate()`
and prints a `verdict` column, but always exits `0` — a template that
errors out or blows the time budget does not fail the process, so it
cannot gate a release or run in CI as a real check.

This script runs the real pipeline (`deckdna.generation.pipeline.generate`,
the same function the CLI and the skill runtime use — no new
architecture) for every required (template, strategy) pair with the
LLM/VLM route enabled by default, and enforces every constraint that
actually gates release-readiness:

- all three strategies (faithful/balanced/visual) per template produce
  a deck (OR-007/OR-027);
- duration per deck <= the official 5-minute-per-deck budget (OR-013);
- zero blocker-severity audit issues;
- zero NEW hard-gate deterministic findings — overflow, out-of-bounds,
  unintended overlap, slide-edge clip, font floor, contrast (the required
  "deterministic hard gates"). Findings inherited verbatim from the
  template itself — on protected slides (cloned byte-identical) or
  matching the template baseline on the same exemplar slide the output
  slide was cloned from — are reported separately, never counted as a
  generation regression, and never silently suppressed (they land in
  `hard_gate.inherited` of the report);
- PEI level >= 3 (per pei.py's own rubric, L3 is the first level that
  actually guarantees no raster-only slide -- the required "native
  editable PPTX; no whole-slide rasterization");
- pptx/pdf/html actually produced and structurally valid (HTML rendered
  via the same `render_html` the CLI exporter uses);
- when --llm is on: the contextual (VLM) audit actually ran AND the
  deck plan was produced by the LLM planner — a silent fallback to the
  deterministic planner means the claimed LLM route was never verified
  (MockProvider always falls back by design; run the gate against a
  real endpoint via DECKDNA_PROVIDER_* env to verify the full route).
  `--no-require-llm-planner` narrows the scope to "VLM-path only" and
  marks the report so.

Unexpected (non-DeckDNAError) exceptions are caught per variant and
surfaced as a structured `unexpected_error` failure record — a crashed
run still produces gate-report.json.

Exits 0 only if every check on every (template, strategy) pair passes.
Any failure --  a stage that raised, a budget miss, a blocker or a new
hard-gate issue, a sub-threshold PEI level, a missing export, a planner
fallback -- is a non-zero exit with a structured report; nothing is
ever silently downgraded to a warning.

Usage:
    python3 scripts/release_gate.py                      # official set, --llm on
    python3 scripts/release_gate.py --no-llm              # deterministic-only, faster
    python3 scripts/release_gate.py --templates T1.pptx T2.pptx --out-dir out/gate
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import time
import zipfile
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "backend") not in sys.path:
    sys.path.insert(0, str(ROOT / "backend"))

from deckdna.audit.basic import audit_deck  # noqa: E402
from deckdna.contracts.deck_plan import Brief  # noqa: E402
from deckdna.contracts.variant_spec import Strategy  # noqa: E402
from deckdna.errors import DeckDNAError  # noqa: E402
from deckdna.generation.pipeline import generate  # noqa: E402
from deckdna.planning.story_director import LLM_PLANNER_VERSION  # noqa: E402
from deckdna.pptx.exporting.render import render_html  # noqa: E402
from deckdna.providers.factory import build_gateway  # noqa: E402


def _variant_slug(template: Path, stem_counts: Counter) -> str:
    """Уникальный стабильный slug выходного каталога.

    Одинаковый basename у шаблонов из разных каталогов не должен
    перезаписывать вывод — суффикс из sha1 resolved-пути
    детерминирован и не зависит от порядка аргументов.
    """
    if stem_counts[template.stem] == 1:
        return template.stem
    digest = hashlib.sha1(str(template).encode()).hexdigest()[:8]  # noqa: S324
    return f"{template.stem}-{digest}"


# Official submission set: three organizer benchmark templates.
# Never treated as the supported-template set anywhere else in the
# pipeline -- here specifically because OR-027 requires 3x3=9 decks on
# exactly these three for the intermediate submission.
DEFAULT_TEMPLATES = [
    ROOT / "tests" / "fixtures" / "pptx" / "vk_tech_template.pptx",
    ROOT / "tests" / "fixtures" / "pptx" / "vk_workspace.pptx",
    ROOT / "tests" / "fixtures" / "pptx" / "vk_education.pptx",
]
DEFAULT_CONTENT = ROOT / "tests" / "fixtures" / "content" / "poc_article.md"
DEFAULT_STRATEGIES = (Strategy.faithful, Strategy.balanced, Strategy.visual)

BUDGET_SECONDS = 300.0  # official <=5 min per deck (OR-013)
MIN_PEI_LEVEL = 3  # "no raster-only slide" per pei.py's own rubric (L0-L2 do NOT
# guarantee this -- pei.assess_pptx returns *exactly* level 2, never higher,
# whenever it finds >=1 raster-only slide; L2 for an unrelated reason (an
# oddly-empty slide) is also possible, but every raster-only-slide case is
# provably capped at 2, so >=3 is the correct floor to actually enforce
# "no whole-slide rasterization" -- gating at 2 would silently accept it)

# Deterministic hard gates per the evaluation principles
# (overflow, bounds, collisions, font floor, contrast). integrity.package
# and editability.raster_only already gate via severity=blocker;
# content.source_support is the contextual layer, not deterministic.
HARD_GATE_RULES = frozenset(
    {
        "text.overflow",
        "layout.out_of_bounds",
        "layout.unintended_overlap",
        "text.slide_clip",
        "text.font_floor",
        "accessibility.contrast",
    }
)


def _now_iso() -> str:
    return datetime.now(UTC).isoformat(timespec="seconds")


def _check_pptx(path: Path) -> str | None:
    try:
        with zipfile.ZipFile(path) as zf:
            if zf.testzip() is not None:
                return "pptx zip failed CRC test"
            names = zf.namelist()
            if "[Content_Types].xml" not in names:
                return "pptx missing [Content_Types].xml"
            if not any(n.startswith("ppt/slides/slide") for n in names):
                return "pptx has no slide parts"
    except (OSError, zipfile.BadZipFile) as exc:
        return f"pptx not a valid zip: {exc}"
    return None


def _check_pdf(path: Path) -> str | None:
    try:
        head = path.open("rb").read(5)
    except OSError as exc:
        return f"pdf unreadable: {exc}"
    if head != b"%PDF-":
        return f"pdf missing %PDF- magic bytes (got {head!r})"
    return None


_SLIDE_PNG_RE = re.compile(r"slide-[\w.-]+\.png")


def _check_html(path: Path, expected_slides: int | None = None) -> str | None:
    """index.html exists, non-empty; every slide render it references
    landed on disk; and (when known) the count of rendered slide-*.png
    equals the deck's slide count — a stale or partial render dir
    fails honestly, never silently passes."""
    index = path if path.suffix == ".html" else path / "index.html"
    if not index.exists():
        return f"html index missing: {index}"
    try:
        text = index.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        return f"html index unreadable: {exc}"
    if not text.strip():
        return "html index is empty"
    pngs = {p.name for p in index.parent.glob("slide-*.png")}
    if not pngs:
        return "html export produced no slide-*.png renders"
    referenced = set(_SLIDE_PNG_RE.findall(text))
    if referenced != pngs:
        missing = sorted(referenced - pngs)
        unreferenced = sorted(pngs - referenced)
        return (
            f"index/render mismatch: referenced-but-missing={missing}, "
            f"on-disk-unreferenced={unreferenced}"
        )
    if expected_slides is not None and len(pngs) != expected_slides:
        return f"html export rendered {len(pngs)} slide png(s), expected {expected_slides}"
    return None


# Направление метрик hard-gate правил (measured_value → worse):
# overflow/overhang/clip/overlap — больше = хуже;
# font_floor (min pt) и contrast (min WCAG ratio) — меньше = хуже.
_HIGHER_IS_WORSE = frozenset(
    {
        "text.overflow",
        "layout.out_of_bounds",
        "layout.unintended_overlap",
        "text.slide_clip",
    }
)
_LOWER_IS_WORSE = frozenset(
    {
        "text.font_floor",
        "accessibility.contrast",
    }
)


def _not_worse(rule_code: str, new_value: Any, base_value: Any) -> bool:
    """True ТОЛЬКО если метрика доказуемо не ухудшилась vs baseline.

    Equal-or-better считается наследованным (та же находка, не
    регрессия); ухудшение или нечисловая/отсутствующая метрика —
    доказать non-worsening нельзя → False.
    """
    if not isinstance(new_value, int | float) or isinstance(new_value, bool):
        return False
    if not isinstance(base_value, int | float) or isinstance(base_value, bool):
        return False
    if rule_code in _LOWER_IS_WORSE:
        return new_value >= base_value
    if rule_code in _HIGHER_IS_WORSE:
        return new_value <= base_value
    return False


def _template_baseline(template: Path) -> dict[str, list[dict[str, Any]]]:
    """Hard-gate issues of the pristine template, keyed by slide part.

    audit_deck() numbers slides by ordinal position (enumerate); the
    compose report references exemplars by OPC part name — map one to
    the other through python-pptx part names so no filename/order
    assumption is ever made.
    """
    from pptx import Presentation

    prs = Presentation(str(template))
    part_of_index = {i: str(slide.part.partname).lstrip("/") for i, slide in enumerate(prs.slides)}
    baseline: dict[str, list[dict[str, Any]]] = {}
    for issue in audit_deck(template):
        if issue.rule_code not in HARD_GATE_RULES:
            continue
        d = issue.to_dict()
        part = part_of_index.get(issue.slide_index, "")
        baseline.setdefault(part, []).append(d)
    return baseline


def _classify_hard_gate(
    issues: list[dict[str, Any]],
    compose_report: dict[str, Any] | None,
    baseline: dict[str, list[dict[str, Any]]],
) -> dict[str, Any]:
    """Разделяет hard-gate findings на новые регрессии и наследованные.

    Generic провенанса без привязки к именам/индексам шаблона:
    - слайд помечен protected/verbatim в compose_report → колода
      несёт байт-идентичную копию слайда шаблона, находки на нём
      наследованы по построению;
    - иначе issue сравнивается с baseline-находками конкретного
      exemplar-слайда, из которого выходной слайд склонирован
      (compose_report.slides[i].exemplar.slide_part): тот же набор
      shape_ids И measured_value не хуже по направлению метрики
      правила = та же находка унаследована без регрессии;
    - ухудшенная метрика, иной набор фигур, отсутствие метрики или
      провенансы — честно 'new', регрессии не прячем.
    """
    slides = (compose_report or {}).get("slides") or []
    new: list[dict[str, Any]] = []
    inherited: list[dict[str, Any]] = []
    for issue in issues:
        if issue.get("rule_code") not in HARD_GATE_RULES:
            continue
        if not issue.get("deterministic", True):
            continue
        entry = {
            "rule_code": issue["rule_code"],
            "slide_index": issue.get("slide_index"),
            "severity": issue.get("severity"),
            "message": issue.get("message"),
        }
        idx = issue.get("slide_index")
        srec = slides[idx] if isinstance(idx, int) and 0 <= idx < len(slides) else None
        if srec is None:
            # нет карты на этот слайд — не угадываем, считаем новым
            new.append(entry)
            continue
        if srec.get("protected") or srec.get("verbatim"):
            entry["inherited_via"] = "protected_slide"
            inherited.append(entry)
            continue
        part = (srec.get("exemplar") or {}).get("slide_part")
        sig_ids = frozenset(issue.get("shape_ids") or [])
        matched = False
        for base_issue in baseline.get(part or "", []):
            if base_issue["rule_code"] != issue["rule_code"]:
                continue
            # Тот же набор фигур: shape_ids идентичны (id живёт в XML
            # и переживает exemplar-клонирование). Слайд-уровневые
            # issues — оба набора пустые, совпадает rule_code.
            base_ids = frozenset(base_issue.get("shape_ids") or [])
            if sig_ids != base_ids:
                continue
            # Наследование только при доказанном non-worsening:
            # equal-or-better measured_value по направлению метрики.
            if _not_worse(
                issue["rule_code"],
                issue.get("measured_value"),
                base_issue.get("measured_value"),
            ):
                matched = True
                break
        if matched:
            entry["inherited_via"] = f"exemplar:{part}"
            inherited.append(entry)
        else:
            new.append(entry)
    return {
        "new": new,
        "inherited": inherited,
        "new_by_rule": dict(Counter(i["rule_code"] for i in new)),
        "inherited_by_rule": dict(Counter(i["rule_code"] for i in inherited)),
    }


def evaluate_variant(
    record: dict[str, Any],
    llm: bool,
    budget_seconds: float,
    *,
    require_llm_planner: bool = True,
) -> list[str]:
    """Gate violations for one (template, strategy) run; empty = pass."""
    problems: list[str] = []
    if record["status"] == "error":
        err = record["error"]
        problems.append(f"generate() failed: {err['code']}: {err['message']}")
        return problems  # nothing else to check without artifacts

    if record["duration_seconds"] > budget_seconds:
        problems.append(f"duration {record['duration_seconds']}s exceeds budget {budget_seconds}s")

    blockers = record["issues_by_severity"].get("blocker", 0)
    if blockers:
        problems.append(f"{blockers} blocker-severity audit issue(s)")

    hard_gate = record.get("hard_gate") or {}
    if hard_gate.get("baseline_error"):
        problems.append(
            "template baseline audit unavailable — inherited "
            f"attribution impossible: {hard_gate['baseline_error']}"
        )
    new_hg = hard_gate.get("new") or []
    if new_hg:
        counts = hard_gate.get("new_by_rule") or {}
        detail = ", ".join(f"{k} x{v}" for k, v in sorted(counts.items()))
        problems.append(f"{len(new_hg)} new hard-gate audit issue(s): {detail}")

    dropped = record.get("dropped_units") or {}
    dropped_nz = {k: v for k, v in dropped.items() if v}
    if dropped_nz:
        detail = ", ".join(f"{k} x{v}" for k, v in sorted(dropped_nz.items()))
        problems.append(f"dropped planned content: {detail}")

    ed = record["editability"]
    pei_level = ed.get("pei_level")
    if pei_level is None:
        problems.append("pei_level not computed")
    elif pei_level < MIN_PEI_LEVEL:
        problems.append(f"pei_level {pei_level} < required {MIN_PEI_LEVEL}")
    raster_count = ed.get("raster_only_slides") or 0
    if raster_count:
        problems.append(f"{raster_count} whole-slide-raster slide(s)")

    pptx_problem = _check_pptx(Path(record["pptx_path"]))
    if pptx_problem:
        problems.append(pptx_problem)
    pdf_problem = _check_pdf(Path(record["pdf_path"]))
    if pdf_problem:
        problems.append(pdf_problem)
    if record.get("html_error"):
        problems.append(f"html export failed: {record['html_error']}")
    elif record.get("html_path"):
        html_problem = _check_html(
            Path(record["html_path"]), expected_slides=record.get("slides_out")
        )
        if html_problem:
            problems.append(html_problem)
    else:
        problems.append("html export missing")

    if llm:
        if not record["contextual_audit_ran"]:
            problems.append("--llm requested but contextual (VLM) audit never ran")
        elif record.get("contextual_audit_complete") is not True:
            problems.append(
                "--llm requested but contextual (VLM) audit coverage is incomplete"
            )
        if require_llm_planner and record.get("planner") != LLM_PLANNER_VERSION:
            problems.append(
                "--llm requested but planner fell back to "
                f"{record.get('planner')!r} — LLM route not verified"
            )

    return problems


def run_one(
    template: Path,
    content: Path,
    brief: Brief,
    strategy: Strategy,
    out_dir: Path,
    gateway: Any,
    baseline: dict[str, list[dict[str, Any]]] | None = None,
    baseline_error: str | None = None,
) -> dict[str, Any]:
    started = time.monotonic()
    try:
        report = generate(template, content, brief, out_dir, gateway=gateway, strategy=strategy)
    except DeckDNAError as exc:
        return {
            "template": template.name,
            "strategy": strategy.value,
            "status": "error",
            "error": exc.to_envelope()["error"],
            "duration_seconds": round(time.monotonic() - started, 2),
        }
    except Exception as exc:  # noqa: BLE001 — структурированный FAIL
        # Неожиданное исключение — тоже честная запись в отчёте,
        # не голый traceback и не потерянный прогон.
        return {
            "template": template.name,
            "strategy": strategy.value,
            "status": "error",
            "error": {
                "code": "unexpected_error",
                "message": f"{type(exc).__name__}: {exc}",
            },
            "duration_seconds": round(time.monotonic() - started, 2),
        }

    by_severity: dict[str, int] = {}
    for issue in report["audit_issues"]:
        by_severity[issue["severity"]] = by_severity.get(issue["severity"], 0) + 1

    html_path = ""
    html_error = None
    # Gate-owned html dir: полностью пересоздаём — stale slide-*.png
    # прошлого прогона не могут фальсифицировать счёт рендеров.
    # Cleanup тоже в try: permission/symlink-ошибка → честный
    # html_error в отчёте, не голый traceback.
    html_dir = out_dir / "html"
    try:
        if html_dir.exists():
            shutil.rmtree(html_dir)
        index = render_html(report["artifacts"]["pptx"], html_dir)
        html_path = str(index.parent)
    except Exception as exc:  # noqa: BLE001 — отчёт фиксирует причину
        html_error = f"{type(exc).__name__}: {exc}"

    return {
        "template": template.name,
        "strategy": strategy.value,
        "status": "generated",
        # Wall-clock всего run_one — generate + HTML-экспорт: бюджет
        # <=5min/deck должен покрывать полный путь до артефактов.
        "duration_seconds": round(time.monotonic() - started, 2),
        "generate_seconds": report["duration_seconds"],
        "slides_out": report["slides_out"],
        "dropped_units": dict((report.get("compose_report") or {}).get("dropped_units") or {}),
        "planner": report["planner"],
        "contextual_audit_ran": report["contextual_audit"]["ran"],
        "contextual_audit_complete": report["contextual_audit"].get("complete", False),
        "issues_by_severity": by_severity,
        "hard_gate": {
            **_classify_hard_gate(
                report["audit_issues"], report.get("compose_report"), baseline or {}
            ),
            "baseline_error": baseline_error,
        },
        "editability": report["quality_passport"]["metrics"]["editability"],
        "pptx_path": report["artifacts"]["pptx"],
        "pdf_path": report["artifacts"]["pdf"],
        "html_path": html_path,
        "html_error": html_error,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--templates",
        nargs="+",
        type=Path,
        default=DEFAULT_TEMPLATES,
        help="templates to gate (default: the 3 official organizer benchmark fixtures)",
    )
    parser.add_argument("--content", type=Path, default=DEFAULT_CONTENT)
    parser.add_argument(
        "--strategies",
        nargs="+",
        default=[s.value for s in DEFAULT_STRATEGIES],
        help="default: all three, as the official constraint requires",
    )
    parser.add_argument("--out-dir", type=Path, default=ROOT / "out" / "release_gate")
    parser.add_argument(
        "--llm",
        dest="llm",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="exercise the LLM/VLM route (default: on, per the gate's purpose). "
        "With the default DECKDNA_MOCK_PROVIDER=true this runs fully offline "
        "against MockProvider; point DECKDNA_PROVIDER_BASE_URL/etc at a real "
        "endpoint via env to gate against it instead.",
    )
    parser.add_argument(
        "--require-llm-planner",
        dest="require_llm_planner",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="require the deck plan to come from the LLM planner when --llm "
        "is on (default: yes — a silent deterministic fallback means the "
        "LLM route was never verified). --no-require-llm-planner narrows "
        "the gate scope to the VLM-path only.",
    )
    parser.add_argument("--slides", type=int, default=12)
    parser.add_argument("--budget", type=float, default=BUDGET_SECONDS)
    args = parser.parse_args(argv)

    try:
        strategies = [Strategy(s) for s in args.strategies]
    except ValueError as exc:
        print(f"invalid --strategies value: {exc}", file=sys.stderr)
        return 2

    missing = [t for t in args.templates if not t.exists()]
    if missing:
        print(f"template(s) not found: {[str(m) for m in missing]}", file=sys.stderr)
        return 2
    if not args.content.exists():
        print(f"content not found: {args.content}", file=sys.stderr)
        return 2

    # Коллизии ключей отчёта/каталогов: baselines и summary ключуются по
    # resolved-пути, выходные каталоги — slug(stem)[+hash при дубле].
    # Идентичные resolved-пути — отклоняем: перезапись нечестна.
    resolved = [t.resolve() for t in args.templates]
    dupes = sorted({str(t) for t in resolved if resolved.count(t) > 1})
    if dupes:
        print(f"duplicate template path(s): {dupes}", file=sys.stderr)
        return 2
    stem_counts = Counter(t.stem for t in resolved)

    brief = Brief(
        purpose="Release gate: official 3-variant LLM/VLM route",
        audience="эксперты и жюри",
        language="ru",
        target_slide_count=args.slides,
    )

    gateway = None
    if args.llm:
        try:
            gateway = build_gateway()
        except DeckDNAError as exc:
            print(
                json.dumps({"status": "failed", "error": exc.to_envelope()["error"]}),
            )
            return 1

    args.out_dir.mkdir(parents=True, exist_ok=True)
    started_at = _now_iso()
    t0 = time.monotonic()

    # Baseline hard-gate находки самого шаблона — опора для провенансы
    # inherited issues (generic: по exemplar slide_part, не по имени).
    baselines: dict[str, dict[str, list[dict[str, Any]]]] = {}
    baseline_errors: dict[str, str] = {}
    for template in resolved:
        key = str(template)
        try:
            baselines[key] = _template_baseline(template)
        except Exception as exc:  # noqa: BLE001 — отчёт фиксирует причину
            baselines[key] = {}
            baseline_errors[key] = f"{type(exc).__name__}: {exc}"

    rows: list[dict[str, Any]] = []
    for template in resolved:
        for strategy in strategies:
            variant_dir = args.out_dir / _variant_slug(template, stem_counts) / strategy.value
            record = run_one(
                template,
                args.content,
                brief,
                strategy,
                variant_dir,
                gateway,
                baseline=baselines.get(str(template)),
                baseline_error=baseline_errors.get(str(template)),
            )
            record["template_path"] = str(template)
            record["violations"] = evaluate_variant(
                record,
                args.llm,
                args.budget,
                require_llm_planner=args.require_llm_planner,
            )
            rows.append(record)

    passed = [r for r in rows if not r["violations"]]
    failed = [r for r in rows if r["violations"]]

    widths = (24, 10, 10, 9, 5, 4, 5, 8)
    header = (
        "template".ljust(widths[0])
        + "strategy".rjust(widths[1])
        + "status".rjust(widths[2])
        + "seconds".rjust(widths[3])
        + "pei".rjust(widths[4])
        + "hg".rjust(widths[5])
        + "drop".rjust(widths[6])
        + "verdict".rjust(widths[7])
    )
    lines = [header]
    for r in rows:
        seconds = f"{r['duration_seconds']:.1f}" if "duration_seconds" in r else "-"
        pei = str(r.get("editability", {}).get("pei_level", "-"))
        hg_new = len((r.get("hard_gate") or {}).get("new") or [])
        hg = "-" if r["status"] == "error" else str(hg_new)
        dropped_total = sum((r.get("dropped_units") or {}).values())
        drop = "-" if r["status"] == "error" else str(dropped_total)
        verdict = "PASS" if not r["violations"] else "FAIL"
        lines.append(
            f"{r['template']:<{widths[0]}}"
            f"{r['strategy']:>{widths[1]}}"
            f"{r['status']:>{widths[2]}}"
            f"{seconds:>{widths[3]}}"
            f"{pei:>{widths[4]}}"
            f"{hg:>{widths[5]}}"
            f"{drop:>{widths[6]}}"
            f"{verdict:>{widths[7]}}"
        )
    print("\n".join(lines))
    if failed:
        print("\nFAILED:")
        for r in failed:
            print(f"  {r['template']} / {r['strategy']}:")
            for v in r["violations"]:
                print(f"    - {v}")

    llm_scope = (
        "full LLM/VLM route (LLM planner + VLM audit required)"
        if args.llm and args.require_llm_planner
        else "VLM-path only (LLM planner fallback tolerated)"
        if args.llm
        else "deterministic-only"
    )
    summary = {
        "started_at": started_at,
        "finished_at": _now_iso(),
        "duration_seconds": round(time.monotonic() - t0, 2),
        "budget_seconds": args.budget,
        "min_pei_level": MIN_PEI_LEVEL,
        "llm_enabled": args.llm,
        "require_llm_planner": args.require_llm_planner,
        "gate_scope": llm_scope,
        "hard_gate_rules": sorted(HARD_GATE_RULES),
        "templates": [str(t) for t in args.templates],
        "strategies": [s.value for s in strategies],
        "template_baseline_hard_gate_counts": {
            name: dict(Counter(i["rule_code"] for issues in parts.values() for i in issues))
            for name, parts in baselines.items()
        },
        "baseline_attribution": (
            "hard-gate findings are split into new regressions vs "
            "inherited (protected verbatim slides, or identical shape_ids "
            "with equal-or-better measured_value vs the exemplar slide's "
            "template baseline); inherited findings are reported but do "
            "not fail the gate"
            + (
                f"; LIMITATION: template baseline unavailable for "
                f"{baseline_errors} — all their findings count as new"
                if baseline_errors
                else ""
            )
        ),
        "planner_values": dict(Counter(r.get("planner", "-") for r in rows)),
        "total": len(rows),
        "passed": len(passed),
        "failed": len(failed),
        "status": "pass" if not failed else "fail",
        "rows": rows,
    }
    report_path = args.out_dir / "gate-report.json"
    report_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"\n{len(passed)}/{len(rows)} passed. Report: {report_path}")
    return 0 if not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
