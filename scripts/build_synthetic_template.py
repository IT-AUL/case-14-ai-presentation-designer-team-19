"""Build the synthetic "unseen" template fixture.

Creates ``tests/fixtures/pptx/synthetic_unseen.pptx`` — a deck that is
structurally UNLIKE the three organizer templates (OR-014/OR-026):

- different color scheme (deep-teal canvas, amber/coral accents, set as
  explicit srgbClr fills — not the VK blue palette);
- default python-pptx Office layout set: content slides are manually
  composed on the blank layout, mirroring organizer decks where authored
  slides dominate over placeholder-driven ones;
- real authored content slides: card grids with heading + body text, a
  KPI row, a two-column split — exemplar-clone material, not empty
  placeholders;
- one deliberately decorative slide with fragmented single-glyph runs —
  the is_content_like exemplar gate must reject it;
- one picture (generated PNG, placed at native aspect).

Run: ``.venv/bin/python scripts/build_synthetic_template.py``
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tests" / "fixtures" / "pptx" / "synthetic_unseen.pptx"

# Non-VK palette: deep teal canvas, amber accent, coral secondary.
CANVAS = "1E2A2B"
CARD = "2A3B3C"
CARD_LIGHT = "35504C"
ACCENT = "FFB547"
CORAL = "E8705A"
INK = "F2F0E9"
MUTED = "9FB3AE"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _rgb(hexv: str) -> RGBColor:
    return RGBColor.from_string(hexv)


def _solid(shape, hexv: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hexv)


def _text(shape, lines: list[tuple[str, float, str, bool]]) -> None:
    """lines: (text, size_pt, hexcolor, bold)."""
    tf = shape.text_frame
    tf.word_wrap = True
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
        run.font.name = "Trebuchet MS"


def _box(slide, x, y, w, h, fill_hex=None):
    shape = slide.shapes.add_shape(
        1, Inches(x), Inches(y), Inches(w), Inches(h)  # MSO_SHAPE.RECTANGLE
    )
    if fill_hex:
        _solid(shape, fill_hex)
    else:
        shape.fill.background()
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _canvas(slide) -> None:
    _box(slide, 0, 0, 13.333, 7.5, CANVAS)


def _card(slide, x, y, w, h, heading: str, body: str, accent=CARD) -> None:
    _box(slide, x, y, w, h, accent)
    _box(slide, x, y, 0.09, h, ACCENT)  # левая акцентная полоса
    head = slide.shapes.add_textbox(
        Inches(x + 0.25), Inches(y + 0.2), Inches(w - 0.4), Inches(0.5)
    )
    _text(head, [(heading, 16, INK, True)])
    bodybox = slide.shapes.add_textbox(
        Inches(x + 0.25), Inches(y + 0.75), Inches(w - 0.4), Inches(h - 0.9)
    )
    _text(bodybox, [(body, 11, MUTED, False)])


def _title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.9))
    _text(box, [(text, 28, INK, True)])
    _box(slide, 0.6, 1.25, 1.2, 0.07, CORAL)


def _png(path: Path, w: int, h: int, bg: str) -> None:
    img = Image.new("RGB", (w, h), "#" + bg)
    d = ImageDraw.Draw(img)
    d.ellipse((w // 4, h // 4, 3 * w // 4, 3 * h // 4), fill="#" + ACCENT)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def build(out: Path = OUT) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]  # Blank — авторский контент поверх, как в organizer-деках

    # --- slide 1: титул ---
    s = prs.slides.add_slide(blank)
    _canvas(s)
    _box(s, 0.6, 2.2, 0.12, 2.4, ACCENT)
    t = s.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(10), Inches(1.6))
    _text(t, [("Тёмная сторона данных", 40, INK, True)])
    sub = s.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(10), Inches(0.9))
    _text(
        sub,
        [("Синтетический шаблон, которого пайплайн никогда не видел", 18, MUTED, False)],
    )

    # --- slide 2: сетка карточек 2x3 (основной контентный exemplar) ---
    s = prs.slides.add_slide(blank)
    _canvas(s)
    _title(s, "Шесть опор платформы")
    cells = [
        ("Наблюдаемость",
         "Метрики и трейсы собираются из каждого сервиса без ручной настройки агентов."),
        ("Декларативность",
         "Схема поставки описывается одним манифестом и применяется идемпотентно."),
        ("Изоляция арендаторов",
         "Каждый клиент получает собственный контур секретов и квот на ресурсы."),
        ("Откат за секунды",
         "Любой релиз откатывается атомарно по сохранённому снапшоту состояния."),
        ("Канареечные волны",
         "Трафик нарастает ступенями с автоматической остановкой по метрикам."),
        ("Аудит действий",
         "Все изменения окружения подписываются и попадают в неизменяемый журнал."),
    ]
    for i, (h, b) in enumerate(cells):
        x = 0.6 + (i % 3) * 4.2
        y = 1.6 + (i // 3) * 2.7
        _card(s, x, y, 3.9, 2.4, h, b)

    # --- slide 3: KPI-ряд (контентный exemplar #2) ---
    s = prs.slides.add_slide(blank)
    _canvas(s)
    _title(s, "Год в цифрах")
    kpis = [
        ("99,97%", "доступность платформы по итогам двенадцати месяцев без учёта плановых окон"),
        ("4 200", "деплоев в продакшен выполнили команды за период без единого инцидента отката"),
        ("38 сек",
         "медианное время от обнаружения регресса до полного отката версии на всех узлах"),
        ("12", "регионов присутствия с локальной обработкой данных по требованиям регуляторов"),
    ]
    for i, (num, cap) in enumerate(kpis):
        x = 0.6 + i * 3.2
        _box(s, x, 1.8, 2.9, 3.4, CARD_LIGHT)
        n = s.shapes.add_textbox(Inches(x + 0.2), Inches(2.0), Inches(2.5), Inches(1.0))
        _text(n, [(num, 34, ACCENT, True)])
        c = s.shapes.add_textbox(Inches(x + 0.2), Inches(3.1), Inches(2.5), Inches(2.0))
        _text(c, [(cap, 11, MUTED, False)])

    # --- slide 4: две колонки + картинка (контентный exemplar #3) ---
    s = prs.slides.add_slide(blank)
    _canvas(s)
    _title(s, "Как это работает")
    left = s.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(5.6), Inches(4.8))
    _text(
        left,
        [
            ("Приём шаблона", 16, INK, True),
            (
                "Произвольный файл разбирается до уровня OPC-частей, строится "
                "граф зависимостей и снимается слепок стиля.",
                12,
                MUTED,
                False,
            ),
            ("Перенос контента", 16, INK, True),
            (
                "Текстовые блоки ложатся в слоты exemplar-слайдов с сохранением "
                "всех relationships пакета.",
                12,
                MUTED,
                False,
            ),
        ],
    )
    import tempfile

    img_path = Path(tempfile.gettempdir()) / "deckdna_synthetic_art.png"
    _png(img_path, 640, 480, CARD)
    s.shapes.add_picture(str(img_path), Inches(7.0), Inches(1.7), width=Inches(5.6))

    # --- slide 5: декоративный (побуквенная раскладка) — гейт должен отсеять ---
    s = prs.slides.add_slide(blank)
    _canvas(s)
    word = "SKYLINE"
    for i, ch in enumerate(word):
        b = s.shapes.add_textbox(
            Inches(1.2 + i * 1.5), Inches(2.0 + (i % 2) * 0.35), Inches(1.2), Inches(1.2)
        )
        _text(b, [(ch, 44, CORAL if i % 2 else ACCENT, True)])
    for i in range(30):
        _box(s, 0.4 + i * 0.42, 5.5, 0.25, 0.25 + (i % 4) * 0.15, CARD_LIGHT)

    # --- slide 6: ещё одна карточная (второй кандидат на ротацию) ---
    s = prs.slides.add_slide(blank)
    _canvas(s)
    _title(s, "Дорожная карта релизов")
    rows = [
        ("Q1 — Стабилизация",
         "Закрытие критических регрессий, выравнивание документации и покрытие ядра тестами."),
        ("Q2 — Масштабирование",
         "Горизонтальное расширение брокеров, разделение планов управления и данных."),
        ("Q3 — Экосистема",
         "Публичный SDK, каталог расширений и программа сертификации партнёрских интеграций."),
        ("Q4 — Платформа",
         "Мультитенантная панель, биллинг по потреблению и рынок готовых сценариев."),
    ]
    for i, (h, b) in enumerate(rows):
        _card(s, 0.6, 1.7 + i * 1.35, 12.1, 1.15, h, b, accent=CARD if i % 2 else CARD_LIGHT)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def main() -> int:
    out = build()
    print(f"built {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
