"""Собирает бинарные content-фикстуры (docx/pdf/xlsx) в tests/fixtures/content/.

Та же тема что poc_article.md: заголовки/параграфы/списки + таблица —
проверяем что parse_file реально вытаскивает структуру из каждого
формата, а не только наличие файла. Запуск:

    .venv/bin/python scripts/build_content_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "content"

SECTIONS = [
    (
        "Проблема",
        ["Шаблоны презентаций не масштабируются: раскладка ломается при подстановке контента."],
        ["Ручная вёрстка одной колоды занимает часы", "Растровые генераторы не редактируемы"],
    ),
    (
        "Решение",
        ["Exemplar-first подход: клонируем реальные контентные слайды с графом зависимостей."],
        ["Аутопсия шаблона строит Design DNA", "Story Director распределяет контент"],
    ),
    (
        "Результаты",
        ["Сквозной пайплайн собирает многослайдовую колоду и проходит аудит."],
        ["Ноль ручных правок после генерации", "Quality Passport фиксирует метрики"],
    ),
]
TABLE = [["Метрика", "Значение"], ["Вёрстка", "часы"], ["Правок", "0"]]


def build_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("DeckDNA: компилятор презентаций", 1)
    d.add_paragraph("DeckDNA превращает шаблон и сырой контент в готовую колоду.")
    for heading, paras, items in SECTIONS:
        d.add_heading(heading, 2)
        for p in paras:
            d.add_paragraph(p)
        for item in items:
            d.add_paragraph(item, style="List Bullet")
    t = d.add_table(rows=len(TABLE), cols=2)
    for i, row in enumerate(TABLE):
        for j, value in enumerate(row):
            t.rows[i].cells[j].text = value
    d.save(path)


def _dejavu() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for candidate in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    ):
        if Path(candidate).exists():
            pdfmetrics.registerFont(TTFont("DejaVu", candidate))
            return "DejaVu"
    return "Helvetica"


def build_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _dejavu()
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont(font, 16)
    c.drawString(72, 790, "DeckDNA: компилятор презентаций")
    y = 760
    c.setFont(font, 10)
    c.drawString(72, y, "DeckDNA превращает шаблон и сырой контент в готовую колоду.")
    y -= 30
    for heading, paras, items in SECTIONS:
        c.setFont(font, 13)
        c.drawString(72, y, heading)
        y -= 18
        c.setFont(font, 10)
        for p in paras:
            c.drawString(72, y, p)
            y -= 14
        for item in items:
            c.drawString(84, y, f"• {item}")
            y -= 14
        y -= 10
    # разлинованная таблица — pdfplumber ловит по линиям сетки
    x0, cw, rh = 72, 140, 20
    for i, row in enumerate(TABLE):
        for j, value in enumerate(row):
            c.rect(x0 + j * cw, y - i * rh, cw, rh)
            c.drawString(x0 + j * cw + 4, y - i * rh + 5, value)
    c.showPage()
    c.setFont(font, 10)
    c.drawString(72, 780, "Следующие шаги: semantic slot mapping и inference-контур.")
    c.showPage()
    c.save()


def build_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Метрики"
    for row in TABLE:
        ws.append(row)
    stages = wb.create_sheet("Стадии")
    stages.append(["Стадия", "Статус"])
    for name in ("parse", "plan", "compose", "audit", "repair"):
        stages.append([name, "done"])
    wb.create_sheet("Черновик")  # пустой лист — должен быть пропущен
    wb.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(OUT / "poc_report.docx")
    build_pdf(OUT / "poc_report.pdf")
    build_xlsx(OUT / "poc_tables.xlsx")
    for name in ("poc_report.docx", "poc_report.pdf", "poc_tables.xlsx"):
        print(f"wrote {OUT / name} ({(OUT / name).stat().st_size} bytes)")


if __name__ == "__main__":
    main()
