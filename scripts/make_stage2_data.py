"""Синтетические датасеты второй ступени (ADR-033 §11) → examples/stage2/.

Данные выдуманы и помечены как синтетические в README папки. Генерация
детерминированная (seed), повторный запуск даёт те же файлы.

    uv run python scripts/make_stage2_data.py
"""

from __future__ import annotations

import argparse
import csv
import json
import random
from datetime import date, timedelta
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "examples" / "stage2"

REGIONS = [
    ("Москва", 0.41),
    ("Санкт-Петербург", 0.22),
    ("Казань", 0.13),
    ("Екатеринбург", 0.12),
    ("Новосибирск", 0.12),
]
CATEGORIES = [
    ("Электроника", 0.34, 9000),
    ("Одежда", 0.24, 3200),
    ("Дом", 0.18, 2600),
    ("Спорт", 0.14, 4100),
    ("Книги", 0.10, 900),
]
CHANNELS = [("Сайт", 0.46), ("Приложение", 0.34), ("Маркетплейс", 0.20)]


def _pick(rng: random.Random, items):
    r = rng.random()
    acc = 0.0
    for item in items:
        acc += item[1]
        if r <= acc:
            return item
    return items[-1]


def _ru_money(value: float) -> str:
    return f"{value:.2f}".replace(".", ",")


def purchases(rng: random.Random) -> None:
    """Транзакции интернет-магазина за 2 года: оборот растёт примерно
    вдвое, рост ускоряется во втором году; часть регионов и каналов не
    заполнена."""
    start = date(2024, 10, 1)
    rows = []
    for day in range(730):
        d = start + timedelta(days=day)
        progress = day / 730
        # рост вдвое с ускорением + сезонность декабря
        orders = int(7 + 9 * progress**1.5 + (5 if d.month == 12 else 0) + rng.random() * 3)
        for _ in range(orders):
            region = _pick(rng, REGIONS)[0]
            category, _, base = _pick(rng, CATEGORIES)
            channel = _pick(rng, CHANNELS)[0]
            amount = max(150.0, rng.gauss(base, base * 0.35))
            client = f"C{rng.randint(1, 4200):05d}"
            if rng.random() < 0.04:
                region = ""
            if rng.random() < 0.03:
                channel = ""
            rows.append([d.isoformat(), region, category, channel, _ru_money(amount), client])
    with (OUT / "purchases.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Дата", "Регион", "Категория", "Канал", "Сумма, ₽", "Клиент"])
        w.writerows(rows)


def sales_monthly(rng: random.Random) -> list[list]:
    """Выручка и заказы по месяцам за 2 года (агрегат)."""
    rows = []
    revenue = 11.8
    for i in range(24):
        y, m = 2024 + (9 + i) // 12, (9 + i) % 12 + 1
        revenue *= 1.026 + 0.012 * (i / 23) + rng.uniform(-0.01, 0.012)
        bump = 1.12 if m == 12 else 1.0
        orders = int(revenue * 1000 / 4.3 * bump)
        rows.append([f"{y}-{m:02d}", round(revenue * bump, 2), orders])
    with (OUT / "sales_monthly.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Месяц", "Выручка, млн ₽", "Заказы, шт."])
        w.writerows(rows)
    return rows


def sales_monthly_variants(rows: list[list]) -> None:
    """Тот же датасет в XLSX и JSON — для теста инвариантности по формату."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Продажи"
    ws.append(["Месяц", "Выручка, млн ₽", "Заказы, шт."])
    for r in rows:
        ws.append(r)
    wb.save(OUT / "sales_monthly.xlsx")
    records = [{"Месяц": r[0], "Выручка, млн ₽": r[1], "Заказы, шт.": r[2]} for r in rows]
    (OUT / "sales_monthly.json").write_text(
        json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8"
    )


def regions(rng: random.Random) -> None:
    """Сравнение: выручка по 14 регионам за год (XLSX, шапка под заголовком)."""
    import openpyxl

    names = ["Москва", "Санкт-Петербург", "Казань", "Екатеринбург", "Новосибирск",
             "Нижний Новгород", "Самара", "Краснодар", "Ростов-на-Дону", "Уфа",
             "Пермь", "Воронеж", "Волгоград", "Красноярск"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Регионы"
    ws.append(["Выручка по регионам за 2025 год"])
    ws.append([])
    ws.append(["Регион", "Выручка, млн ₽"])
    value = 410.0
    for name in names:
        ws.append([name, round(value, 1)])
        value *= rng.uniform(0.55, 0.85) if name == "Москва" else rng.uniform(0.82, 0.97)
    wb.save(OUT / "regions.xlsx")


def channels_share() -> None:
    """Доли: структура продаж по каналам (JSON {columns, data})."""
    data = {
        "title": "Структура продаж по каналам, 2025",
        "columns": ["Канал", "Доля, %"],
        "data": [["Сайт", 38.5], ["Приложение", 27.0], ["Маркетплейсы", 21.5],
                 ["Розница", 9.0], ["Партнёры", 4.0]],
    }
    (OUT / "channels_share.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8"
    )


def order_values(rng: random.Random) -> None:
    """Распределение: суммы 600 заказов (одна колонка, правый хвост)."""
    with (OUT / "order_values.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Сумма заказа, ₽"])
        for _ in range(600):
            value = rng.lognormvariate(8.1, 0.55)
            w.writerow([round(value, 2)])


def products_reference(rng: random.Random) -> None:
    """Справочник ≥ 200 строк: товары с артикулом, категорией, ценой, остатком."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["Артикул", "Наименование", "Категория", "Цена, ₽", "Остаток, шт.",
               "Поставщик", "Срок поставки, дн."])
    cats = ["Электроника", "Одежда", "Дом", "Спорт", "Книги", "Красота"]
    for i in range(240):
        cat = cats[i % len(cats)]
        ws.append([f"SKU-{10000 + i}", f"{cat}: товар {i + 1}", cat,
                   round(rng.uniform(300, 45000), 0), rng.randint(0, 900),
                   f"Поставщик {rng.randint(1, 18)}", rng.randint(2, 30)])
    wb.save(OUT / "products_reference.xlsx")


def nps_with_gaps(rng: random.Random) -> None:
    """Пропуски: NPS по месяцам, трёх месяцев нет."""
    with (OUT / "nps_with_gaps.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Месяц", "NPS"])
        nps = 31.0
        for i in range(18):
            y, m = 2025 + i // 12, i % 12 + 1
            nps += rng.uniform(-2.5, 3.2)
            value = "" if i in (5, 6, 11) else round(nps, 1)
            w.writerow([f"{y}-{m:02d}", value])


def report_with_table() -> None:
    """Текст со встроенной таблицей (колонки через пробелы)."""
    text = (
        "Отчёт отдела продаж за 2025 год\n\n"
        "Цель отчёта — показать рост числа клиентов по кварталам и понять, "
        "удалось ли удержать темп после запуска приложения.\n\n"
        "Квартал     Новые клиенты   Ушедшие клиенты\n"
        "1 кв. 2025  1 240           410\n"
        "2 кв. 2025  1 520           395\n"
        "3 кв. 2025  1 980           388\n"
        "4 кв. 2025  2 610           402\n\n"
        "Приложение запущено в начале второго квартала. Отток держится на одном "
        "уровне, а приток растёт каждый квартал.\n"
    )
    (OUT / "report_with_table.txt").write_text(text, encoding="utf-8")


def ad_spend_vs_sales(rng: random.Random) -> None:
    """Рассеяние: расходы на рекламу и продажи по 40 городам."""
    with (OUT / "ad_spend_vs_sales.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Город", "Расходы на рекламу, тыс. ₽", "Продажи, тыс. ₽"])
        for i in range(40):
            spend = rng.uniform(80, 900)
            sales = spend * rng.uniform(3.2, 4.4) + rng.gauss(0, 260) + 400
            w.writerow([f"Город {i + 1}", round(spend, 1), round(sales, 1)])


def kpi() -> None:
    """KPI: четыре ключевых числа года (JSON — массив записей)."""
    data = [
        {"Показатель": "Выручка, млн ₽", "Значение": 412.5},
        {"Показатель": "Клиенты, тыс.", "Значение": 186},
        {"Показатель": "Средний чек, ₽", "Значение": 2217},
        {"Показатель": "NPS", "Значение": 46},
    ]
    (OUT / "kpi.json").write_text(json.dumps(data, ensure_ascii=False, indent=1),
                                   encoding="utf-8")


def main() -> None:
    global OUT
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=OUT)
    OUT = parser.parse_args().out
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(20261016)  # noqa: S311 — синтетика, не криптография
    purchases(rng)
    rows = sales_monthly(rng)
    sales_monthly_variants(rows)
    regions(rng)
    channels_share()
    order_values(rng)
    products_reference(rng)
    nps_with_gaps(rng)
    report_with_table()
    ad_spend_vs_sales(rng)
    kpi()


if __name__ == "__main__":
    main()
