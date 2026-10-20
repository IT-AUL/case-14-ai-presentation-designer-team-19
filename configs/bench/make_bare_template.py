"""Генератор контрольного шаблона «без layouts» (QA-GATES §5).

Валидный PPTX, у которого макеты не несут ни одного placeholder, а
слайды-образцы — обычные надписи без структурной разметки. Конвейер
обязан собирать колоду и на таком шаблоне."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def build(path: Path) -> None:
    prs = Presentation()
    for layout in prs.slide_layouts:
        for ph in list(layout.placeholders):
            ph._element.getparent().remove(ph._element)
    layout = prs.slide_layouts[6]
    for title, note in (
        ("Отчёт о работе", "Собрано вручную"),
        ("Показатели", "Сводка за период"),
    ):
        slide = prs.slides.add_slide(layout)
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1))
        title_box.text_frame.text = title
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(9), Inches(4))
        body.text_frame.text = note
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


if __name__ == "__main__":
    import sys

    build(Path(sys.argv[1]) if len(sys.argv) > 1 else Path("bare_no_layouts.pptx"))
