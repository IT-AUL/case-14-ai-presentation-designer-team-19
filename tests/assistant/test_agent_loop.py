"""ADR-058: цикл помощника без API — действия, наблюдения, честность."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from deckdna.assistant.agent import (
    AgentAction,
    AgentQuestion,
    AgentTurn,
    run_agent,
    slide_number_in,
    ungrounded_numbers,
)
from deckdna.errors import DeckDNAError

ROOT = Path(__file__).resolve().parents[2]


class Gateway:
    def __init__(self, turns):
        self.turns = list(turns)
        self.payloads = []

    async def text_json(self, name, payload, schema):
        assert name == "assistant_agent" and schema is AgentTurn
        self.payloads.append(payload)
        return self.turns.pop(0)


class Workspace:
    def __init__(self, fail=()):
        self.done = []
        self.fail = set(fail)

    def overview(self):
        return {"slides": [{"n": 1, "title": "Обложка"}, {"n": 2, "title": "Рынок"}]}

    def capabilities(self):
        return {"mode": "deck", "web_search": True, "image_generation": False}

    def grounding_corpus(self):
        return ""

    async def execute(self, action):
        self.done.append(action.tool)
        if action.tool in self.fail:
            raise DeckDNAError("invalid_input", "numbers not found: 42")
        if action.tool == "web_research":
            return {"ok": True, "summary": "нашёл фактов: 1", "facts": ["Рынок 10 млрд"],
                    "sources": [{"title": "РБК", "url": "https://example.org"}]}
        return {"ok": True, "summary": "готово", "revision": 2}


def test_numbers_must_come_from_material_or_research():
    assert ungrounded_numbers(["Рынок вырос на 12,5% до 3 400 млн"],
                              "рост 12.5 процента, объём 3400 млн") == []
    assert ungrounded_numbers(["Выручка 42 млн ₽ в 2025"], "выручка выросла") == ["2025", "42"]
    assert ungrounded_numbers(["Шаг 3 из 5"], "") == []


def test_slide_number_is_read_from_free_text():
    assert slide_number_in("Подробнее на 5 слайде") == 5
    assert slide_number_in("добавь аудиторию на слайд 6") == 6
    assert slide_number_in("слайд №12 убери") == 12
    assert slide_number_in("сделай убедительнее") is None


@pytest.mark.asyncio
async def test_research_then_edit_feeds_observations_back_to_model():
    gateway = Gateway([
        AgentTurn(reply="Поищу цифры.", actions=[AgentAction(tool="web_research", slide=2,
                                                             query="рынок 2025")],
                  done=False),
        AgentTurn(reply="Обновил слайд 2 по данным РБК.", actions=[AgentAction(
            tool="rewrite_slide", slide=2, items=["Рынок 10 млрд"], source="web")], done=True),
    ])
    events = []
    ws = Workspace()
    result = await run_agent("Добавь цифры по рынку", workspace=ws, gateway=gateway,
                             emit=lambda e, d: events.append((e, d)))
    assert ws.done == ["web_research", "rewrite_slide"]
    assert gateway.payloads[1]["observations"][0]["facts"] == ["Рынок 10 млрд"]
    assert result.reply.startswith("Обновил") and result.model_calls == 2
    assert [a["tool"] for a in result.actions] == ["web_research", "rewrite_slide"]
    assert result.actions[0]["sources"][0]["title"] == "РБК"
    phases = [d["phase"] for e, d in events if e == "agent.action"]
    assert phases == ["started", "finished", "started", "finished"]


@pytest.mark.asyncio
async def test_failed_action_gets_another_step_and_honest_reply():
    gateway = Gateway([
        AgentTurn(reply="Сделал.", actions=[AgentAction(tool="rewrite_slide", slide=2,
                                                        items=["Выручка 42 млн"])]),
        AgentTurn(reply="Не нашёл 42 млн в материалах, поставил формулировку без числа.",
                  actions=[AgentAction(tool="rewrite_slide", slide=2, items=["Выручка растёт"])]),
    ])
    ws = Workspace(fail={"never"})
    ws.fail = {"rewrite_slide"}
    first = await run_agent("Добавь выручку", workspace=ws, gateway=gateway, max_steps=2)
    assert first.actions[0]["ok"] is False
    assert len(gateway.payloads) == 2
    assert gateway.payloads[1]["observations"][0]["ok"] is False
    assert gateway.payloads[1]["step"]["last"] is True


@pytest.mark.asyncio
async def test_question_returns_without_actions():
    gateway = Gateway([AgentTurn(reply="Уточню.", question=AgentQuestion(
        text="Для кого презентация?", options=["Инвесторы", "Команда"]))])
    ws = Workspace()
    result = await run_agent("Сделай лучше", workspace=ws, gateway=gateway)
    assert ws.done == [] and result.question.options == ["Инвесторы", "Команда"]


@pytest.mark.asyncio
async def test_without_model_free_text_is_answered_honestly_and_quick_commands_work():
    ws = Workspace()
    free = await run_agent("Сделай слайды убедительнее для инвесторов", workspace=ws,
                           gateway=None)
    assert free.mode == "unavailable" and ws.done == []
    quick = await run_agent("сделай короче на слайде 2", workspace=ws, gateway=None)
    assert quick.mode == "local" and ws.done == ["shorter"]
    assert quick.actions[0]["slide"] == 2
    assert quick.actions[0]["revision"] == 2


def test_slide_order_remove_move_insert_keep_other_slides_intact():
    import zipfile

    from deckdna.pptx.composing.slide_order import insert_blank_slide, move_slide, remove_slide
    from deckdna.pptx.composing.slide_swap import slide_parts, swap_slide
    from pptx import Presentation

    deck = (ROOT / "tests/fixtures/pptx/vk_tech_template.pptx").read_bytes()

    def files(data):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return {n: z.read(n) for n in z.namelist()}

    before = files(deck)
    order = slide_parts(before)
    removed = files(remove_slide(deck, 2))
    assert slide_parts(removed) == [order[0], *order[2:]]
    assert order[1] not in removed
    assert all(removed[p] == before[p] for p in slide_parts(removed))
    moved = move_slide(deck, 1, 3)
    assert slide_parts(files(moved))[:3] == [order[1], order[2], order[0]]
    inserted = insert_blank_slide(deck, 2, layout_of=1)
    assert len(Presentation(io.BytesIO(inserted)).slides) == len(order) + 1
    filled = swap_slide(inserted, deck, 2, 3)
    pres = Presentation(io.BytesIO(filled))
    assert len(pres.slides) == len(order) + 1
    out = files(filled)
    assert out[slide_parts(out)[1]] == before[order[2]]
    with pytest.raises(DeckDNAError):
        remove_slide(deck, len(order) + 1)


@pytest.mark.asyncio
async def test_unavailable_model_is_reported_not_papered_over():
    class Down:
        async def text_json(self, name, payload, schema):
            raise DeckDNAError("provider_unavailable", "HTTP 500")

    with pytest.raises(DeckDNAError) as exc:
        await run_agent("Перепиши слайд 2", workspace=Workspace(), gateway=Down())
    assert exc.value.http_status == 503

    class Flaky(Gateway):
        async def text_json(self, name, payload, schema):
            if self.turns:
                return await super().text_json(name, payload, schema)
            raise DeckDNAError("provider_unavailable", "HTTP 500")

    ws = Workspace()
    gateway = Flaky([AgentTurn(reply="Смотрю.", done=False, actions=[
        AgentAction(tool="shorter", slide=2)])])
    result = await run_agent("Короче слайд 2", workspace=ws, gateway=gateway)
    assert ws.done == ["shorter"] and result.reply == "Сделал: готово."
