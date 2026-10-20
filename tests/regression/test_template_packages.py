"""ADR-031 §2–3: шаблоны, которые ломали генерацию на входе.

- ``.potx`` (основная часть ``presentationml.template.main+xml``) падал на
  ``composing.deck``: python-pptx не открывает пакет шаблона;
- шаблон без слайдов (только мастер и макеты) и шаблон с одним слайдом
  падали с ``template has no slides to clone``.

Все прогоны детерминированные, без модели и сети.
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from deckdna.api.app import app
from deckdna.generation.pipeline import generate
from deckdna.pptx.opc.normalize import needs_normalization, normalize_template_bytes
from deckdna.template.synthesize import is_synthesized_placeholder_text
from fastapi.testclient import TestClient
from pptx import Presentation

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
VK_TECH = FIXTURES / "pptx" / "vk_tech.pptx"
CONTENT = FIXTURES / "content" / "poc_article.md"
POTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.template"
A = "/api/v1"

BRIEF = {"purpose": "product", "audience": "руководство", "language": "ru", "target_slide_count": 5}


def _as_potx(pptx: Path) -> bytes:
    src = zipfile.ZipFile(pptx)
    out = io.BytesIO()
    with src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.presentation.main+xml",
                    b"presentationml.template.main+xml",
                )
            dst.writestr(info, data)
    return out.getvalue()


def test_normalize_rewrites_only_the_main_part_type():
    potx = _as_potx(VK_TECH)
    assert needs_normalization(potx)
    fixed = normalize_template_bytes(potx)
    assert not needs_normalization(fixed)
    # python-pptx открывает результат, слайды на месте
    assert len(Presentation(io.BytesIO(fixed)).slides) == len(Presentation(str(VK_TECH)).slides)
    # обычный .pptx возвращается теми же байтами
    plain = VK_TECH.read_bytes()
    assert normalize_template_bytes(plain) is plain


def test_generate_on_potx_template(tmp_path):
    template = tmp_path / "corp.potx"
    template.write_bytes(_as_potx(VK_TECH))
    report = generate(template, CONTENT, BRIEF, tmp_path / "out")
    pptx = Path(report["artifacts"]["pptx"])
    assert pptx.exists()
    assert len(Presentation(str(pptx)).slides) >= 3


def test_api_upload_and_generate_potx(wait_generation):
    client = TestClient(app)
    pid = client.post(f"{A}/projects", json={"name": "potx"}).json()["id"]
    tpl = client.post(
        f"{A}/projects/{pid}/templates",
        files={"file": ("corp.potx", _as_potx(VK_TECH), POTX_MIME)},
    )
    assert tpl.status_code == 201, tpl.text
    assert tpl.json()["package_type"] == "potx"
    analysis = client.post(f"{A}/templates/{tpl.json()['id']}/analyze", json={})
    assert analysis.status_code == 202, analysis.text
    pack = client.post(
        f"{A}/projects/{pid}/content-packs",
        files=[("files", (CONTENT.name, CONTENT.read_bytes(), "text/markdown"))],
        data={"brief": json.dumps({"language": "ru"})},
    )
    assert pack.status_code == 202, pack.text
    gen = client.post(
        f"{A}/projects/{pid}/generations",
        json={
            "template_id": tpl.json()["id"],
            "content_pack_id": pack.json()["content_pack"]["id"],
            "brief": BRIEF,
            "variants": [{"strategy": "balanced"}],
        },
    )
    assert gen.status_code == 202, gen.text
    run = wait_generation(client, gen.json()["generation_id"])
    assert run["state"] == "completed", run


@pytest.mark.parametrize("slides_kept", [0, 1])
def test_layout_only_template_generates_three_variants(tmp_path, slides_kept):
    """ADR-031 §3: образцы синтезируются из макетов шаблона."""
    prs = Presentation()
    for i in range(slides_kept):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = f"Образец {i + 1}"
    template = tmp_path / f"layouts_{slides_kept}.pptx"
    prs.save(str(template))

    for strategy in ("faithful", "balanced", "visual"):
        report = generate(template, CONTENT, BRIEF, tmp_path / strategy, strategy=strategy)
        pptx = Path(report["artifacts"]["pptx"])
        deck = Presentation(str(pptx))
        assert len(deck.slides) >= 3
        codes = {i["rule_code"] for i in report["audit_issues"]}
        assert "integrity.filler_slide" not in codes, (strategy, codes)
        assert "integrity.empty_slide" not in codes, (strategy, codes)
        assert report["template_synthesis"]["synthesized_slides"] > 0
        # образцовый текст синтеза не доживает до колоды, и у каждого слайда
        # есть настоящий текст
        for slide in deck.slides:
            texts = [
                p.text.strip()
                for sh in slide.shapes
                if sh.has_text_frame
                for p in sh.text_frame.paragraphs
                if p.text.strip()
            ]
            assert texts, (strategy, slide.slide_id)
            leftovers = [t for t in texts if is_synthesized_placeholder_text(t)]
            assert not leftovers, (strategy, leftovers)
