"""ADR-032 §2–5: импорт старой презентации, несколько файлов, заметки
докладчика, схемы из текста, график из найденных цифр.

Всё на моках: MockProvider с фикстурами, подменённые поиск и загрузка
страниц, без сети.
"""

from __future__ import annotations

import asyncio
import io
import json
import shutil
from pathlib import Path

import deckdna.api.app as app_module
import pytest
from deckdna.api.app import STORE, app
from deckdna.contracts.deck_plan import Brief
from deckdna.generation.pipeline import generate
from deckdna.ingestion.content_parsers import parse_file, parse_markdown
from deckdna.planning.diagram_detect import detect_diagrams, verify_steps
from deckdna.planning.story_director import plan_deck
from deckdna.planning.talk_track import (
    extractive_track,
    validate_track,
    write_talk_tracks,
)
from deckdna.providers.mock import MockProvider
from deckdna.research.numbers_chart import chart_or_table
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
VK_TECH = FIXTURES / "pptx" / "vk_tech_template.pptx"
VK_EDU = FIXTURES / "pptx" / "vk_education.pptx"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
A = "/api/v1"
client = TestClient(app)

needs_render = pytest.mark.skipif(shutil.which("soffice") is None, reason="needs soffice")


def _old_deck() -> bytes:
    """Старая презентация: обложка, слайд с пунктами и заметками, таблица,
    нативный график."""
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "Итоги пилота"
    cover.placeholders[1].text = "Отчёт для руководства"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Что сделали"
    body = s2.placeholders[1].text_frame
    body.text = "Подключили три отдела продаж к сервису"
    body.add_paragraph().text = "Собрали 120 презентаций за месяц без дизайнера"
    s2.notes_slide.notes_text_frame.text = "Подчеркнуть, что дизайнеры не участвовали."
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Показатели"
    table = s3.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    rows = [("Метрика", "Значение"), ("Колод", "120"), ("Часов сэкономлено", "340")]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Динамика"
    data = CategoryChartData()
    data.categories = ["Июль", "Август", "Сентябрь"]
    data.add_series("Колод", (20, 40, 60))
    s4.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(7),
                        Inches(4), data)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- §2 импорт


def test_pptx_content_keeps_texts_tables_charts_and_notes(tmp_path):
    path = tmp_path / "old.pptx"
    path.write_bytes(_old_deck())
    pack = parse_file(path)
    assert pack.title_hint == "Итоги пилота"
    headings = [s.heading for s in pack.sections]
    assert headings == ["Что сделали", "Показатели", "Динамика"]
    assert pack.tables[0].rows == [["Колод", 120.0], ["Часов сэкономлено", 340.0]]
    assert pack.charts[0].categories == ["Июль", "Август", "Сентябрь"]
    assert pack.charts[0].series[0].values == [20.0, 40.0, 60.0]
    first = pack.sections[0]
    assert all(b.source_ref.page == 2 for b in first.blocks)
    # заметки — в speaker_note плана, не на слайд
    plan = plan_deck(pack, Brief(purpose="Итоги", audience="руководство", language="ru",
                                 target_slide_count=5, content_fill="shrink"))
    texts = " ".join(u.text or "" for s in plan.slides for u in s.content_units)
    assert "Подчеркнуть" not in texts
    assert any("Подчеркнуть" in (s.speaker_note or "") for s in plan.slides)


def test_vk_education_as_content_parses_tables_and_images():
    pack = parse_file(VK_EDU)
    assert len(pack.sections) > 20
    assert len(pack.tables) == 3
    assert pack.assets and all("#ppt/media/" in a.artifact_id for a in pack.assets)


@needs_render
def test_relayout_old_deck_into_vk_tech(tmp_path, wait_generation):
    """Готово, когда: старая презентация (тексты, таблица, график) →
    колода в шаблоне VK Tech с теми же текстами, таблицей и графиком."""
    pid = client.post(f"{A}/projects", json={"name": "relayout"}).json()["id"]
    tpl = client.post(
        f"{A}/projects/{pid}/templates",
        files={"file": (VK_TECH.name, VK_TECH.read_bytes(), PPTX_MIME)},
    ).json()
    pack = client.post(
        f"{A}/projects/{pid}/content-packs",
        files=[("files", ("old.pptx", _old_deck(), PPTX_MIME))],
    )
    assert pack.status_code == 202, pack.text
    gen = client.post(f"{A}/projects/{pid}/generations", json={
        "template_id": tpl["id"],
        "content_pack_id": pack.json()["content_pack"]["id"],
        "brief": {"purpose": "Итоги пилота", "audience": "руководство", "language": "ru",
                  "target_slide_count": 5, "content_fill": "shrink"},
        "variants": [{"strategy": "balanced"}],
    })
    assert gen.status_code == 202, gen.text
    run = wait_generation(client, gen.json()["generation_id"])
    assert run["state"] == "completed", run
    deck_id = run["variants"][0]["deck_artifact_id"]
    deck = Presentation(io.BytesIO(STORE.artifacts[deck_id].data))
    text = " ".join(sh.text_frame.text for s in deck.slides for sh in s.shapes
                    if sh.has_text_frame)
    assert "Подключили три отдела продаж" in text
    tables = [sh for s in deck.slides for sh in s.shapes if getattr(sh, "has_table", False)
              and sh.has_table]
    charts = [sh for s in deck.slides for sh in s.shapes if getattr(sh, "has_chart", False)
              and sh.has_chart]
    assert tables, "таблица из старой колоды должна сохраниться"
    cells = {c.text for t in tables for row in t.table.rows for c in row.cells}
    assert "Часов сэкономлено" in cells
    assert charts, "график из старой колоды должен сохраниться"
    assert list(charts[0].chart.plots[0].categories) == ["Июль", "Август", "Сентябрь"]


def test_multi_file_pack_keeps_separate_sources():
    pid = client.post(f"{A}/projects", json={"name": "multi"}).json()["id"]
    md = "# Доклад\n\n## Проблема\n\nКоманды тратят часы на вёрстку слайдов.\n"
    resp = client.post(
        f"{A}/projects/{pid}/content-packs",
        files=[
            ("files", ("talk.md", md.encode(), "text/markdown")),
            ("files", ("old.pptx", _old_deck(), PPTX_MIME)),
        ],
    )
    assert resp.status_code == 202, resp.text
    pack = resp.json()["content_pack"]
    artifacts = {b["source_ref"]["artifact_id"] for s in pack["sections"] for b in s["blocks"]}
    assert len(artifacts) == 2
    assert [t["id"] for t in pack["tables"]] == ["f2-tbl-0"]
    refs = [b["text"] for s in pack["sections"] for b in s["blocks"] if b["kind"] == "table_ref"]
    assert refs == ["f2-tbl-0"]
    bad = client.post(
        f"{A}/projects/{pid}/content-packs",
        files=[("files", ("a.md", b"# A\n\nx\n", "text/markdown")),
               ("files", ("b.exe", b"MZ", "application/octet-stream"))],
    )
    assert bad.status_code == 422


# ---------------------------------------------------------------- §3 заметки


def _plan_for_notes():
    pack = parse_markdown(
        "# Пилот\n\n## Результаты\n\n- Подготовили 120 колод за месяц\n"
        "- Сэкономили 340 часов дизайнеров\n- Ошибок вёрстки стало меньше\n"
    )
    return plan_deck(pack, Brief(purpose="Итоги пилота", audience="руководство",
                                 language="ru", target_slide_count=4, content_fill="shrink"))


def test_talk_track_rejects_invented_numbers_and_falls_back():
    plan = _plan_for_notes()
    slide = next(s for s in plan.slides if s.purpose.value not in ("title", "thank_you"))
    assert validate_track("Мы подготовили 999 колод. Это рекорд.", slide)
    assert extractive_track(slide)
    gateway = MockProvider(fixtures={"talk_track": {"slides": [
        {"index": s.index, "text": "Мы подготовили 999 колод за месяц. Это рекорд команды."}
        for s in plan.slides
    ]}})
    report = asyncio.run(write_talk_tracks(plan, gateway))
    assert report.rejected >= 1 and report.model_written == 0
    assert all(report.notes.values())


@needs_render
def test_generate_writes_talk_track_into_pptx_notes(tmp_path):
    md = tmp_path / "pilot.md"
    md.write_text(
        "# Пилот\n\n## Результаты\n\n- Подготовили 120 колод за месяц\n"
        "- Сэкономили 340 часов дизайнеров\n\n## Планы\n\n- Подключить ещё два отдела\n"
        "- Добавить шаблоны дочерних компаний\n",
        encoding="utf-8",
    )
    report = generate(VK_TECH, md, {"purpose": "Итоги пилота", "audience": "руководство",
                                    "language": "ru", "target_slide_count": 5,
                                    "content_fill": "shrink"}, tmp_path / "out")
    deck = Presentation(report["artifacts"]["pptx"])
    notes = [s.notes_slide.notes_text_frame.text for s in deck.slides if s.has_notes_slide]
    assert any("120 колод" in n for n in notes)
    metric = report["quality_passport"]["metrics"]["speaker_notes"]
    assert metric["slides_with_notes"] >= 2 and metric["extractive"] >= 2


# ---------------------------------------------------------------- §4 схемы


def test_numbered_steps_become_a_diagram_without_model():
    pack = parse_markdown(
        "# Внедрение\n\n## Как внедряем\n\nВнедрение идёт по этапам.\n\n"
        "1. Аудит\n2. Пилот\n3. Масштабирование\n4. Поддержка\n\n"
        "## Преимущества\n\n- Быстро для любой команды\n- Работает с любым шаблоном\n"
        "- Аудит на каждую колоду\n"
    )
    out, found = asyncio.run(detect_diagrams(pack))
    assert found == 1
    assert out.diagrams[0].steps == ["Аудит", "Пилот", "Масштабирование", "Поддержка"]
    kinds = [b.kind.value for b in out.sections[1].blocks]
    assert kinds == ["paragraph", "diagram_ref"]  # список заменён схемой
    # список аргументов без признаков процесса — не схема
    assert [b.kind.value for b in out.sections[2].blocks] == ["list"]


def test_model_diagram_steps_must_be_verbatim():
    text = "Сначала идёт аудит шаблона, затем пилот в одном отделе и масштабирование."
    assert verify_steps(["аудит шаблона", "пилот в одном отделе", "масштабирование"], text)
    assert not verify_steps(["аудит", "пилотирование", "масштабирование"], text)
    pack = parse_markdown(f"# Внедрение\n\n## План\n\n{text}\n")
    good = MockProvider(fixtures={"diagram_detect": {"diagrams": [
        {"section_id": pack.sections[-1].id, "kind": "process",
         "steps": ["аудит шаблона", "пилот в одном отделе", "масштабирование"]},
    ]}})
    out, found = asyncio.run(detect_diagrams(pack, good))
    assert found == 1 and out.diagrams[0].steps[0] == "Аудит шаблона"
    bad = MockProvider(fixtures={"diagram_detect": {"diagrams": [
        {"section_id": pack.sections[-1].id, "kind": "process",
         "steps": ["Аудит", "Пилотирование", "Масштабирование"]},
    ]}})
    out, found = asyncio.run(detect_diagrams(pack, bad))
    assert found == 0 and not out.diagrams


# ---------------------------------------------------------------- §5 график из цифр


def test_homogeneous_figures_make_a_chart_mixed_make_a_table():
    from deckdna.contracts.content_pack import SourceRef

    src = SourceRef(artifact_id="ext-1")
    chart, table = chart_or_table([
        "В 2022 году рынок ассистентов составил 12,5 млрд рублей.",
        "В 2023 году рынок ассистентов составил 18 млрд рублей.",
        "В 2024 году рынок ассистентов составил 27 млрд рублей.",
    ], "c", src)
    assert table is None and chart.categories == ["2022", "2023", "2024"]
    assert chart.series[0].values == [12.5, 18.0, 27.0]
    assert "млрд ₽" in chart.series[0].name
    chart, table = chart_or_table([
        "Доля компаний с ИИ выросла до 43%.", "Средний чек составил 1 200 рублей.",
    ], "t", src)
    assert chart is None and table.headers == ["Показатель", "Значение"]


@needs_render
def test_numbers_research_puts_a_native_chart_into_the_deck(monkeypatch, wait_generation):
    """Готово, когда: ресёрч цифр → слайд с нативным графиком в итоговом PPTX."""
    import deckdna.research.web as web
    from deckdna.research.config import ResearchConfig, WebResearchConfig
    from deckdna.research.web import EXTRACT_PROMPT, QUERIES_PROMPT, Page, SearchHit

    page_text = (
        "В 2022 году рынок ассистентов составил 12,5 млрд рублей. "
        "В 2023 году рынок ассистентов составил 18 млрд рублей. "
        "В 2024 году рынок ассистентов составил 27 млрд рублей. "
    ) * 3

    class FakeSearx:
        def __init__(self, *args, **kwargs):
            pass

        async def search(self, query, limit, language):
            return [SearchHit(url="https://stats.example.org/market", title="Рынок")]

    async def fake_fetch(url, cfg, transport=None):
        return Page(url=url, title="Рынок ассистентов", text=page_text)

    monkeypatch.setattr(web, "SearxngSearch", FakeSearx)
    monkeypatch.setattr(web, "fetch_page", fake_fetch)
    monkeypatch.setattr(
        app_module, "load_research_config",
        lambda: ResearchConfig(budget_s=20, web=WebResearchConfig(search_url="http://searxng")),
    )
    facts = [
        ("Рынок ассистентов в 2022 году составил 12,5 млрд рублей",
         "В 2022 году рынок ассистентов составил 12,5 млрд рублей"),
        ("Рынок ассистентов в 2023 году составил 18 млрд рублей",
         "В 2023 году рынок ассистентов составил 18 млрд рублей"),
        ("Рынок ассистентов в 2024 году составил 27 млрд рублей",
         "В 2024 году рынок ассистентов составил 27 млрд рублей"),
    ]
    gateway = MockProvider(fixtures={
        QUERIES_PROMPT: {"queries": ["рынок ассистентов"]},
        EXTRACT_PROMPT: {"facts": [{"text": t, "quote": q, "topic": "Рынок"} for t, q in facts]},
    })
    monkeypatch.setattr(app_module, "build_gateway", lambda config=None: gateway)

    pid = client.post(f"{A}/projects", json={"name": "numbers"}).json()["id"]
    tpl = client.post(
        f"{A}/projects/{pid}/templates",
        files={"file": (VK_TECH.name, VK_TECH.read_bytes(), PPTX_MIME)},
    ).json()
    md = ("# Рынок\n\n## Проблема\n\n- Команды тратят часы на вёрстку презентаций.\n"
          "- Фирменный стиль соблюдается не всегда.\n\n## Рынок\n\n"
          "- Рынок ассистентов растёт каждый год.\n- Компании ищут готовые решения.\n")
    cp = client.post(
        f"{A}/projects/{pid}/content-packs",
        files=[("files", ("m.md", md.encode(), "text/markdown"))],
    ).json()["content_pack"]["id"]
    plan_set = client.post(f"{A}/projects/{pid}/plans", json={
        "content_pack_id": cp, "strategies": ["balanced"], "use_llm": False,
        "brief": {"purpose": "Рынок", "audience": "руководство", "language": "ru",
                  "target_slide_count": 4, "content_fill": "shrink"},
    }).json()
    slide = next(s for s in plan_set["plans"][0]["deck_plan"]["slides"]
                 if "Рынок" in s["title_intent"] and s["purpose"] != "title")
    resp = client.post(
        f"{A}/plans/{plan_set['plan_id']}/slides/{slide['id']}/research",
        json={"focus": "numbers", "use_llm": True},
    )
    assert resp.status_code == 200, resp.text
    result = resp.json()
    assert result["research"]["chart_id"]
    assert result["slide"]["desired_visual"] == "chart"
    gen = client.post(f"{A}/projects/{pid}/generations", json={
        "template_id": tpl["id"],
        "content_pack_id": result["plan_set"]["content_pack_id"],
        "deck_plan_id": plan_set["plan_id"],
        "brief": {"purpose": "Рынок", "audience": "руководство", "language": "ru",
                  "target_slide_count": 4},
        "variants": [{"strategy": "balanced"}],
        "use_llm": False,
    })
    assert gen.status_code == 202, gen.text
    run = wait_generation(client, gen.json()["generation_id"])
    assert run["state"] == "completed", run
    deck = Presentation(io.BytesIO(STORE.artifacts[run["variants"][0]["deck_artifact_id"]].data))
    charts = [sh for s in deck.slides for sh in s.shapes
              if getattr(sh, "has_chart", False) and sh.has_chart]
    assert charts, "график из найденных цифр должен быть нативным"
    plot = charts[0].chart.plots[0]
    assert list(plot.categories) == ["2022", "2023", "2024"]
    assert list(plot.series[0].values) == [12.5, 18.0, 27.0]
    json.dumps(result)  # ответ сериализуем
