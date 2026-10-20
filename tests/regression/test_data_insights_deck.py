"""ADR-033, сквозной срез: шаблон + ``purchases.csv`` → колода от выводов.

- слайды с данными несут нативные графики выбранного типа (ряд по месяцам
  — линия, разрезы — столбцы) с подписями осей и строкой «Источник»;
- заголовки — выводы с настоящими числами из данных;
- почему такой график — в заметках докладчика и в паспорте;
- тот же датасет на трёх шаблонах даёт те же типы графиков (ADR-033 §10);
- без ``text.overflow`` и ``chart.metadata`` на слайдах с данными.

Модель не участвует. Рендер-проверка: ``docs/review/assets/adr033_*.png``."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from deckdna.analysis import analyze_pack
from deckdna.generation.pipeline import generate
from deckdna.ingestion.content_parsers import parse_file
from deckdna.pptx.composing.chart_build import chart_type_of
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
STAGE2 = ROOT / "examples" / "stage2"
FIXTURES = ROOT / "tests" / "fixtures" / "pptx"
TEMPLATES = ["vk_tech_template.pptx", "vk_education.pptx", "vk_workspace.pptx"]
BRIEF = {
    "purpose": "Итоги продаж за два года",
    "audience": "руководство",
    "language": "ru",
    "target_slide_count": 7,
}


def _chart_types(pptx_path: str) -> list[list[str]]:
    prs = Presentation(pptx_path)
    out = []
    for slide in prs.slides:
        types = []
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                types.append(chart_type_of(shape.chart.part.blob))
        out.append(types)
    return out


@pytest.fixture(scope="module")
def decks(tmp_path_factory):
    out = {}
    for template in TEMPLATES:
        tmp = tmp_path_factory.mktemp(template.split(".")[0])
        out[template] = generate(
            FIXTURES / template, STAGE2 / "purchases.csv", dict(BRIEF), tmp / "out"
        )
    return out


def test_data_slides_have_typed_native_charts_with_real_numbers(decks):
    report = decks["vk_tech_template.pptx"]
    plan = report["deck_plan"]
    slides = plan["slides"]
    assert re.search(r"\d", slides[1]["title_intent"])
    data = [s for s in slides if s["purpose"] == "data"]
    assert len(data) >= 3
    pack = parse_file(STAGE2 / "purchases.csv")
    [analysis] = analyze_pack(pack, BRIEF["purpose"])
    assert analysis.primary.headline
    line_slide = next(s for s in data if (s.get("chart") or {}).get("type") == "line")
    assert re.search(r"\d", line_slide["title_intent"])
    types = _chart_types(report["artifacts"]["pptx"])
    by_index = {s["index"]: s for s in slides}
    for index, kinds in enumerate(types):
        if by_index[index]["purpose"] == "data":
            unit = next(u for u in by_index[index]["content_units"] if u["kind"] == "chart")
            assert kinds == [unit["chart_spec"]["chart_type"]], (index, kinds)
    assert types[line_slide["index"]] == ["line"]

    prs = Presentation(report["artifacts"]["pptx"])
    first = prs.slides[line_slide["index"]]
    texts = " ".join(sh.text_frame.text for sh in first.shapes if sh.has_text_frame)
    assert "Источник: purchases.csv" in texts
    chart = next(sh for sh in first.shapes if getattr(sh, "has_chart", False)).chart
    xml = chart.part.blob.decode("utf-8")
    assert "Выручка, млн ₽" in xml  # подпись оси с единицей
    assert len(chart.plots[0].categories) == 24


def test_why_this_chart_is_in_notes_and_passport(decks):
    report = decks["vk_tech_template.pptx"]
    prs = Presentation(report["artifacts"]["pptx"])
    plan_slides = report["deck_plan"]["slides"]
    data_index = next(
        s["index"]
        for s in plan_slides
        if s["purpose"] == "data" and (s.get("chart") or {}).get("type") == "line"
    )
    notes = prs.slides[data_index].notes_slide.notes_text_frame.text
    assert notes
    assert "Данные:" not in notes and "Альтернативы" not in notes
    card = plan_slides[data_index]["chart"]
    assert card["type"] == "line" and card["reasons"]
    assert card["alternatives"]
    assert any("24 точки" in r for r in card["reasons"])
    passport = report["quality_passport"]
    [dataset] = passport["analysis"]["datasets"]
    assert dataset["file"] == "purchases.csv"
    shown = [v for v in dataset["views"] if v["slide_index"] is not None]
    assert len(shown) >= 3
    trend = next(v for v in shown if v["slide_index"] == data_index)
    assert trend["chosen"]["type"] == "line" and trend["task"] == "trend"
    assert trend["alternatives"] and trend["rejected"]
    assert trend["insights"][0]["fact"].startswith("Выручка: с ")


def test_same_dataset_same_chart_types_on_every_template(decks):
    """ADR-033 §10: тип графика не зависит от шаблона."""
    signatures = []
    for template, report in decks.items():
        plan = report["deck_plan"]
        planned = [
            u["chart_spec"]["chart_type"]
            for s in plan["slides"]
            for u in s["content_units"]
            if u["kind"] == "chart"
        ]
        rendered = [
            t for kinds in _chart_types(report["artifacts"]["pptx"]) for t in kinds if t is not None
        ]
        assert rendered[: len(planned)] == planned, template
        signatures.append(planned)
    assert signatures[0] == signatures[1] == signatures[2]


@pytest.mark.parametrize("template", TEMPLATES)
def test_data_slides_have_no_overflow_or_chart_metadata_issues(decks, template):
    report = decks[template]
    data = {s["index"] for s in report["deck_plan"]["slides"] if s["purpose"] == "data"}
    bad = [
        (i["slide_index"], i["rule_code"])
        for i in report["audit_issues"]
        if i["slide_index"] in data
        and i["rule_code"]
        in ("text.overflow", "chart.metadata", "layout.out_of_bounds", "integrity.filler_slide")
    ]
    assert not bad, bad
    dropped = [
        s.get("dropped_units")
        for s in report["compose_report"]["slides"]
        if (s.get("dropped_units") or {}).get("chart")
    ]
    assert not dropped
