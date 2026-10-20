"""ADR-031 §5: схема на организаторских шаблонах.

Было (скриншот ``docs/review/assets/diagram_overflow_vk_tech.png``): поток
из четырёх шагов вписывался в узкую строку одной карточки VK Tech,
«Масштабирование» переносилось по буквам (2× ``text.overflow``), слева
оставались пустые маркеры карточек; на VK Education и VK WorkSpace схема
вообще выпадала (место «занимали» демо-картинки и оформление макета).

Стало: хост по ширине (наибольший свободный прямоугольник под заголовком с
учётом оформления макета), кегль по самому длинному слову, оформление
опустевших карточек под схемой удаляется. Шесть шагов с длинными словами
уходят в «змейку» или столбец, а не рвутся.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from deckdna.generation.pipeline import generate
from lxml import etree

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "pptx"
TEMPLATES = ["vk_tech_template.pptx", "vk_education.pptx", "vk_workspace.pptx"]
FLOWS = {
    4: "Аудит -> Пилот -> Масштабирование -> Поддержка",
    6: (
        "Аудит -> Проектирование -> Пилотирование -> Масштабирование -> "
        "Сопровождение -> Совершенствование"
    ),
}
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NEAR = 0.04


def _content(tmp_path: Path, flow: str) -> Path:
    path = tmp_path / "flow.md"
    path.write_text(
        "# Внедрение платформы\n\n## Процесс внедрения\n\n"
        "Внедрение проходит несколько этапов, каждый со своей командой.\n\n"
        f"```diagram id=dg-1\n{flow}\n```\n\n"
        "## Результаты\n\nПилот сократил время подготовки отчётов на 40 процентов.\n",
        encoding="utf-8",
    )
    return path


def _box(el):
    xfrm = el.find(f"{{{P}}}spPr/{{{A}}}xfrm")
    if xfrm is None:
        xfrm = el.find(f"{{{P}}}grpSpPr/{{{A}}}xfrm")
    if xfrm is None or xfrm.find(f"{{{A}}}off") is None:
        return None
    off, ext = xfrm.find(f"{{{A}}}off"), xfrm.find(f"{{{A}}}ext")
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def _gap(a, b) -> int:
    dx = max(0, max(a[0], b[0]) - min(a[0] + a[2], b[0] + b[2]))
    dy = max(0, max(a[1], b[1]) - min(a[1] + a[3], b[1] + b[3]))
    return max(dx, dy)


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(f"{{{A}}}t")).strip()


@pytest.mark.parametrize("steps", sorted(FLOWS))
@pytest.mark.parametrize("template", TEMPLATES)
def test_diagram_fits_without_broken_words_or_orphan_markers(tmp_path, template, steps):
    from pptx import Presentation

    report = generate(
        FIXTURES / template,
        _content(tmp_path, FLOWS[steps]),
        {"purpose": "Показать процесс внедрения", "audience": "руководство",
         "language": "ru", "target_slide_count": 5, "content_fill": "shrink"},
        tmp_path / "out",
    )
    diagram_slides = [
        s for s in report["compose_report"]["slides"]
        if (s.get("diagrams") or {}).get("diagrams_created")
    ]
    assert len(diagram_slides) == 1, "схема должна встать, а не выпасть"
    assert all(not (s.get("dropped_units") or {}).get("diagram")
               for s in report["compose_report"]["slides"])
    pos = report["compose_report"]["slides"].index(diagram_slides[0])

    bad = [
        i for i in report["audit_issues"]
        if i["slide_index"] == pos
        and i["rule_code"] in ("text.overflow", "integrity.filler_slide", "layout.collision")
    ]
    assert not bad, bad

    prs = Presentation(report["artifacts"]["pptx"])
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    slide = prs.slides[pos]
    tree = slide.shapes._spTree
    group = next(
        el for el in tree.iter(f"{{{P}}}grpSp")
        if (el.find(f"{{{P}}}nvGrpSpPr/{{{P}}}cNvPr").get("name") or "").startswith("Diagram")
    )
    assert [_text(sp) for sp in group.iter(f"{{{P}}}sp") if _text(sp)] == FLOWS[steps].split(" -> ")
    # каждое слово шага целиком помещается в ширину шага — без разрыва
    sizes = {int(r.get("sz")) for r in group.iter(f"{{{A}}}rPr") if r.get("sz")}
    assert sizes and min(sizes) >= 900

    # пустые рамки и мелкие значки опустевших карточек не остаются рядом
    # друг с другом (маркер без пункта — «дырявый грид»)
    near = NEAR * min(sw, sh)
    empty_frames = [
        b for el in tree if etree.QName(el).localname == "sp"
        and el.find(f"{{{P}}}txBody") is not None and not _text(el)
        and el.find(f"{{{P}}}nvSpPr/{{{P}}}nvPr/{{{P}}}ph") is None
        and (b := _box(el)) is not None
    ]
    markers = [
        b for el in tree if etree.QName(el).localname in ("sp", "pic")
        and not _text(el) and (b := _box(el)) is not None
        and b[2] * b[3] <= 0.02 * sw * sh
    ]
    orphaned = [m for m in markers if any(_gap(m, f) <= near and m != f for f in empty_frames)]
    assert not orphaned, orphaned
