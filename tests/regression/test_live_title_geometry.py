"""Compile the saved real VK plan, not a rewritten mock storyline."""
import json
from pathlib import Path

from deckdna.audit.basic import audit_deck
from deckdna.contracts.content_pack import Chart, Table
from deckdna.contracts.deck_plan import DeckPlan
from deckdna.generation.pipeline import _auto_fix
from deckdna.pptx.composing.minimal import generate_deck
from pptx import Presentation

ROOT = Path(__file__).parents[2]


def test_saved_vk_plan_compiles_without_title_caption_overlap_and_keeps_charts(tmp_path):
    fixture = json.loads((ROOT / 'tests/fixtures/content/vk_live_title_geometry.json').read_text())
    output = tmp_path / 'deck.pptx'
    generate_deck(
        ROOT / 'tests/fixtures/pptx/vk_workspace.pptx', DeckPlan.model_validate(fixture['plan']),
        output, charts=[Chart.model_validate(x) for x in fixture['charts']],
        tables=[Table.model_validate(x) for x in fixture['tables']],
        exemplar_overrides=fixture['exemplar_overrides'],
    )
    issues, _, _rejected = _auto_fix(output, audit_deck(output), None)
    hard = [i.to_dict() for i in issues if i.severity in {'error', 'blocker'}]
    assert not hard, hard
    assert sum(s.has_chart for slide in Presentation(output).slides for s in slide.shapes) >= 2
