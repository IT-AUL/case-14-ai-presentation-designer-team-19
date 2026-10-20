"""Small native charts need an explicit, configurable geometry warning."""
import pytest
from deckdna.audit.basic import audit_deck
from deckdna.audit.config import AuditConfig
from deckdna.contracts.audit_issue import AuditIssue
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

RULE = 'chart.readability'


def _deck(tmp_path, *, scale=1, wide=.45, high=.3, grouped=False):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10*scale), Inches(6*scale)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['A', 'B', 'C', 'D', 'E', 'F']
    data.add_series('Value', [20, 12, 10, 9, 7, 42])
    shape = slide.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(scale), Inches(scale),
                                   int(prs.slide_width*wide), int(prs.slide_height*high), data)
    if grouped:
        group = slide.shapes.add_group_shape([shape])
        group.width = group.width // 2
        group.height = group.height // 2
    path = tmp_path / 'unseen.pptx'
    prs.save(path)
    return path


@pytest.mark.parametrize('scale', [1, 2])
def test_small_chart_warns_with_geometry_and_no_false_repair_promise(tmp_path, scale):
    issues = [i for i in audit_deck(_deck(tmp_path, scale=scale)) if i.rule_code == RULE]
    assert len(issues) == 1
    issue = issues[0]
    assert issue.severity == 'warning' and issue.deterministic
    assert not issue.repairable
    assert issue.slide_index == 0 and issue.shape_ids
    assert issue.bbox['h'] == pytest.approx(.3)
    assert issue.evidence[0]['kind'] == 'geometry'
    AuditIssue.model_validate(issue.to_dict())


def test_readability_thresholds_are_configurable(tmp_path):
    path = _deck(tmp_path)
    assert any(i.rule_code == RULE for i in audit_deck(path))
    loose = AuditConfig(chart_min_width_ratio=.2, chart_min_height_ratio=.25)
    assert not any(i.rule_code == RULE for i in audit_deck(path, config=loose))


def test_roomy_chart_passes_but_scaled_group_is_measured_on_slide(tmp_path):
    assert not any(i.rule_code == RULE
                   for i in audit_deck(_deck(tmp_path, wide=.6, high=.6)))
    issues = [i for i in audit_deck(_deck(tmp_path, wide=.6, high=.6, grouped=True))
              if i.rule_code == RULE]
    assert len(issues) == 1 and issues[0].bbox['w'] == pytest.approx(.3)
