"""Contrast follows visible master/layout shapes and opaque occlusion."""

import copy

import pytest
from deckdna.audit.basic import RULE_CONTRAST, audit_deck
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


@pytest.mark.parametrize("owner", ["master", "layout"])
@pytest.mark.parametrize("hidden", [False, True])
def test_white_text_on_inherited_dark_panel(tmp_path, owner, hidden):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string("181818")
    target = slide.slide_layout if owner == "layout" else slide.slide_layout.slide_master
    target.shapes._spTree.insert_element_before(copy.deepcopy(panel.element), "p:extLst")
    slide.shapes._spTree.remove(panel.element)
    if hidden:
        slide.element.set("showMasterSp", "0")
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Readable white text"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string("FFFFFF")
    path = tmp_path / "inherited.pptx"
    prs.save(path)
    issues = [i for i in audit_deck(path) if i.rule_code == RULE_CONTRAST]
    assert len(issues) == int(hidden)


@pytest.mark.parametrize("translucent_layer", [False, True])
def test_opaque_panel_above_unknown_fill_restores_resolvable_contrast(tmp_path, translucent_layer):
    import io

    from PIL import Image

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    png = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(png, "PNG")
    png.seek(0)
    slide.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
    if translucent_layer:
        from pptx.oxml.xmlchemy import OxmlElement

        veil = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        veil.fill.solid()
        veil.fill.fore_color.rgb = RGBColor.from_string("181818")
        alpha = OxmlElement("a:alpha")
        alpha.set("val", "50000")
        veil.element.find(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"
        ).append(alpha)
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Invisible white text"
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
    path = tmp_path / "occluded.pptx"
    prs.save(path)
    assert any(i.rule_code == RULE_CONTRAST for i in audit_deck(path))


def test_brandbook_priority_audits_against_changed_background(tmp_path):
    import io
    import json

    from deckdna.template.brandbook import apply_brandbook, build_brandbook

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Brand text"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string("000000")
    buffer = io.BytesIO()
    prs.save(buffer)
    book = build_brandbook(
        [
            (
                "tokens.json",
                json.dumps(
                    {
                        "colors": {"text": "#F8F8F8", "background": "#181818"},
                    }
                ).encode(),
            )
        ],
        "brand",
    )
    branded, _ = apply_brandbook(buffer.getvalue(), book, "brandbook")
    path = tmp_path / "brand.pptx"
    path.write_bytes(branded)
    loaded = Presentation(path)
    assert str(loaded.slides[0].background.fill.fore_color.rgb) == "181818"
    assert not any(i.rule_code == RULE_CONTRAST for i in audit_deck(path))
