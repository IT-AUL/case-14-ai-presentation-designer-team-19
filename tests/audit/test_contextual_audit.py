"""run_contextual_audit: VLM-verdict → AuditIssue (моковый gateway, без сети).

Рендер слайдов идёт через реальный soffice→pdftoppm — тесты помечены
медленными относительно unit-сюиты, но держатся в рамках секунд
(1 слайд на фикстуре).
"""

from __future__ import annotations

import asyncio

import pytest
from deckdna.audit.contextual import (
    PROMPT_NAME,
    SlideChecksResult,
    _vision_json_with_retry,
    run_contextual_audit,
)
from deckdna.errors import DeckDNAError
from deckdna.providers.mock import MockProvider

PPTX = "tests/fixtures/pptx/synthetic_unseen.pptx"


def _verdicts(*checks):
    return {
        "verdicts": [
            {
                "check": str(num),
                "verdict": verdict,
                "rationale": rationale,
                "confidence": confidence,
            }
            for num, verdict, rationale, confidence in checks
        ]
    }


async def test_fail_verdict_becomes_contextual_issue():
    """fail с confidence ≥ порога → issue с frozen rule_code, deterministic=False."""
    gw = MockProvider(
        fixtures={
            PROMPT_NAME: _verdicts(
                ("1", "fail", "Заголовок — тема, а не вывод", 0.92),
                ("2", "pass", "", 0.9),
                ("4", "uncertain", "текста недостаточно", 0.5),
            )
        }
    )
    issues = await run_contextual_audit(PPTX, gw, first=1, last=1)
    assert len(issues) == 1
    issue = issues[0]
    assert issue.rule_code == "content.conclusion_title"
    assert issue.severity == "warning"
    assert issue.deterministic is False
    assert issue.slide_index == 0
    assert issue.confidence == 0.92
    assert issue.evidence[0]["kind"] == "model_verdict"
    # gateway вызван с промптом slide_checks и data-URL картинкой
    assert len(gw.calls) == 1
    assert gw.calls[0]["prompt"] == "slide_checks"
    assert gw.calls[0]["images"][0].startswith("data:image/png;base64,")


async def test_uncertain_and_low_confidence_are_not_issues():
    """uncertain — честная неопределённость, не issue; fail ниже порога — тоже."""
    gw = MockProvider(
        fixtures={
            PROMPT_NAME: _verdicts(
                ("4", "uncertain", "нет evidence на руках", 0.9),
                ("7", "fail", "возможная утечка", 0.3),
            )
        }
    )
    issues = await run_contextual_audit(
        PPTX, gw, first=1, last=1, confidence_threshold=0.5
    )
    assert issues == []


async def test_payload_carries_slide_context():
    """Провайдеру уходит slide_text/title_intent/evidence/language в payload."""
    seen: list[dict] = []

    def _fixture(payload):
        seen.append(payload)
        return _verdicts(("5", "pass", "", 1.0))

    gw = MockProvider(fixtures={PROMPT_NAME: _fixture})
    await run_contextual_audit(PPTX, gw, first=1, last=1, deck_language="ru")
    payload = seen[0]
    assert payload["deck_language"] == "ru"
    assert payload["slide_text"].strip()  # текст слайда реально извлечён
    assert payload["title_intent"] == ""  # deck_plan не передан — честно пусто
    assert payload["evidence_excerpt"] == ""


async def test_all_verdicts_map_to_frozen_rules():
    """Каждая проверка 1-10 → свой frozen rule_code с severity Appendix."""
    gw = MockProvider(
        fixtures={
            PROMPT_NAME: _verdicts(
                *[(str(n), "fail", f"r{n}", 0.9) for n in range(1, 11)]
            )
        }
    )
    issues = await run_contextual_audit(PPTX, gw, first=1, last=1)
    codes = {i.rule_code for i in issues}
    assert codes == {
        "content.conclusion_title",
        "content.title_alignment",
        "content.single_message",
        "content.source_support",
        "content.nonempty",
        "content.visual_relevance",
        "content.prompt_leakage",
        "content.spelling",
        "content.data_relevance",
        "deck.logical_flow",
    }
    by_code = {i.rule_code: i for i in issues}
    assert by_code["content.source_support"].severity == "blocker"
    assert by_code["content.title_alignment"].severity == "error"


async def test_result_type_is_structured():
    """fixture в виде dict валидируется в SlideChecksResult (как настоящий gateway)."""
    gw = MockProvider(fixtures={PROMPT_NAME: _verdicts(("5", "pass", "", 1.0))})
    out = await gw.vision_json(PROMPT_NAME, ["img"], {}, SlideChecksResult)
    assert isinstance(out, SlideChecksResult)


async def test_retries_once_on_structured_output_invalid():
    """Один сбойный ответ не роняет весь аудит — один retry того же запроса."""
    attempts = {"n": 0}

    def _flaky(payload):
        attempts["n"] += 1
        if attempts["n"] == 1:
            raise DeckDNAError(
                code="structured_output_invalid",
                message="provider returned invalid structured output",
                stage="providers",
            )
        return _verdicts(("5", "pass", "", 1.0))

    gw = MockProvider(fixtures={PROMPT_NAME: _flaky})
    issues = await run_contextual_audit(PPTX, gw, first=1, last=1)
    assert issues == []  # pass verdict — no issue, but the call succeeded on retry
    assert attempts["n"] == 2


async def test_second_failure_after_retry_leaves_the_slide_unverified():
    """Retry — не бесконечный: второй подряд structured_output_invalid делает
    слайд непроверенным, а не роняет генерацию (VK Inference, 17.10: модель
    без навязанной схемы ответила своей формой — колода не собралась)."""

    def _always_broken(payload):
        raise DeckDNAError(
            code="structured_output_invalid",
            message="provider returned invalid structured output",
            stage="providers",
        )

    gw = MockProvider(fixtures={PROMPT_NAME: _always_broken})
    diagnostics: dict = {}
    issues = await run_contextual_audit(PPTX, gw, first=1, last=1, diagnostics=diagnostics)
    assert issues == []
    assert diagnostics["unverified_slides"] == [0]


async def test_non_retryable_error_code_is_not_retried():
    """Другие коды ошибок (не structured_output_invalid) не глотаются retry-логикой."""
    attempts = {"n": 0}

    def _auth_failure(payload):
        attempts["n"] += 1
        raise DeckDNAError(
            code="provider_unavailable", message="HTTP 401", stage="providers"
        )

    gw = MockProvider(fixtures={PROMPT_NAME: _auth_failure})
    try:
        await _vision_json_with_retry(gw, PROMPT_NAME, [], {}, SlideChecksResult)
    except DeckDNAError as exc:
        assert exc.code == "provider_unavailable"
    else:
        raise AssertionError("expected DeckDNAError to propagate")
    assert attempts["n"] == 1  # no retry attempted for a non-transient code


async def test_provider_outage_leaves_only_failed_slide_unverified():
    def respond(payload):
        if payload["slide_image"] == 2:
            raise DeckDNAError(
                code="provider_unavailable", message="provider HTTP 500", stage="providers"
            )
        return _verdicts(("4", "fail", "Нет источника утверждения", 0.99))

    gw = MockProvider(fixtures={PROMPT_NAME: respond})
    diagnostics = {}
    issues = await run_contextual_audit(PPTX, gw, first=1, last=3, diagnostics=diagnostics)
    assert diagnostics["requested_slides"] == 3
    assert diagnostics["completed_slides"] == 2
    assert diagnostics["unverified_slides"] == [1]
    assert [(i.slide_index, i.severity) for i in issues] == [(0, "blocker"), (2, "blocker")]
    assert len(gw.calls) == 3  # no additional audit-stage retry on HTTP failure


async def test_unexpected_audit_error_still_propagates():
    def broken(payload):
        raise DeckDNAError(code="internal_error", message="broken audit", stage="audit")

    with pytest.raises(DeckDNAError, match="broken audit"):
        await run_contextual_audit(
            PPTX, MockProvider(fixtures={PROMPT_NAME: broken}), first=1, last=1
        )


class _ConcurrencyTrackingGateway:
    """Fake gateway — измеряет одновременно летящие vision_json-вызовы.

    Каждый вызов реально спит (``asyncio.sleep``), поэтому пересечение
    вызовов во времени наблюдаемо: без конкурентности (последовательный
    ``for``-цикл) наблюдаемый максимум всегда был бы 1.
    """

    def __init__(self, *, delay: float = 0.05) -> None:
        self.delay = delay
        self.in_flight = 0
        self.max_in_flight = 0
        self.call_count = 0

    async def vision_json(self, prompt_name, images, payload, schema):
        self.call_count += 1
        self.in_flight += 1
        self.max_in_flight = max(self.max_in_flight, self.in_flight)
        try:
            await asyncio.sleep(self.delay)
            return SlideChecksResult(verdicts=[])
        finally:
            self.in_flight -= 1


async def test_slides_are_checked_concurrently_not_sequentially():
    """max_concurrent>1 реально даёт пересекающиеся во времени вызовы."""
    gw = _ConcurrencyTrackingGateway()
    issues = await run_contextual_audit(PPTX, gw, first=1, last=4, max_concurrent=4)
    assert issues == []
    assert gw.call_count == 4
    assert gw.max_in_flight > 1, "calls never overlapped — audit is still sequential"


async def test_max_concurrent_bounds_in_flight_calls():
    """max_concurrent реально ограничивает потолок, не просто документация."""
    gw = _ConcurrencyTrackingGateway()
    await run_contextual_audit(PPTX, gw, first=1, last=4, max_concurrent=2)
    assert gw.max_in_flight <= 2


async def test_concurrent_issues_stay_ordered_and_numbered_by_slide():
    """Issue id/seq остаются позиционными по слайду независимо от порядка завершения."""
    gw = MockProvider(
        fixtures={PROMPT_NAME: _verdicts(("5", "fail", "не по делу", 0.9))}
    )
    issues = await run_contextual_audit(PPTX, gw, first=1, last=4, max_concurrent=4)
    assert [i.slide_index for i in issues] == [0, 1, 2, 3]
    assert [i.id.rsplit("-", 1)[-1] for i in issues] == ["1", "2", "3", "4"]


async def test_contextual_issue_matches_frozen_schema():
    import json
    from pathlib import Path

    from jsonschema import validate

    gw = MockProvider(fixtures={PROMPT_NAME: _verdicts(("2", "fail", "missing formula", 0.9))})
    issues = await run_contextual_audit(PPTX, gw, first=2, last=2)
    assert issues[0].slide_index == 1  # same zero-based index as deterministic audit/UI
    schema = json.loads(await asyncio.to_thread(Path('schemas/audit-issue.schema.json').read_text))
    validate(issues[0].to_dict(), schema)


async def test_incomplete_audit_is_disclosed():
    gw = MockProvider(fixtures={PROMPT_NAME: _verdicts(("2", "uncertain", "no evidence", 0.5))})
    diagnostics = {}
    await run_contextual_audit(PPTX, gw, first=1, last=1, diagnostics=diagnostics)
    assert diagnostics == dict(
        expected_slides=1,
        requested_slides=1,
        completed_slides=1,
        unverified_slides=[],
        uncertain_checks=1,
        missing_checks=9,
        duplicate_checks=0,
        invalid_checks=0,
        complete=False,
    )
    diagnostics = {}
    await run_contextual_audit(PPTX, gw, first=1, last=1, budget_s=0, diagnostics=diagnostics)
    assert diagnostics['completed_slides'] == 0
    assert diagnostics['unverified_slides'] == [0]
    assert diagnostics['complete'] is False


async def test_complete_audit_requires_all_unique_checks():
    gw = MockProvider(
        fixtures={
            PROMPT_NAME: _verdicts(
                *[(str(number), "pass", "", 1.0) for number in range(1, 11)]
            )
        }
    )
    diagnostics = {}
    await run_contextual_audit(PPTX, gw, first=1, last=1, diagnostics=diagnostics)
    assert diagnostics["complete"] is True
    assert diagnostics["missing_checks"] == 0
    assert diagnostics["duplicate_checks"] == 0
    assert diagnostics["invalid_checks"] == 0


async def test_duplicate_and_unknown_checks_cannot_complete_audit():
    gw = MockProvider(
        fixtures={
            PROMPT_NAME: _verdicts(
                *[(str(number), "pass", "", 1.0) for number in range(1, 11)],
                ("4", "fail", "duplicate source-support failure", 0.9),
                ("99", "fail", "unknown check", 0.9),
            )
        }
    )
    diagnostics = {}
    issues = await run_contextual_audit(PPTX, gw, first=1, last=1, diagnostics=diagnostics)
    assert issues == [], "duplicate verdict must not manufacture an issue"
    assert diagnostics["missing_checks"] == 0
    assert diagnostics["duplicate_checks"] == 1
    assert diagnostics["invalid_checks"] == 1
    assert diagnostics["complete"] is False


async def test_question_text_instead_of_numbers_maps_by_position():
    """VK Qwen (17.10) писал в `check` текст вопроса, а не номер — все 10
    проверок отбрасывались как invalid, аудит работал впустую."""
    questions = [f"Question text {chr(65 + i)}?" for i in range(10)]

    def _texts(payload):
        return {"verdicts": [
            {"check": q, "verdict": "fail" if i == 0 else "pass",
             "rationale": "r", "confidence": 0.99}
            for i, q in enumerate(questions)
        ]}

    gw = MockProvider(fixtures={PROMPT_NAME: _texts})
    diagnostics: dict = {}
    issues = await run_contextual_audit(PPTX, gw, first=1, last=1, diagnostics=diagnostics)
    assert diagnostics["invalid_checks"] == 0
    assert [i.rule_code for i in issues] == ["content.conclusion_title"]


def test_select_contextual_audit_slides_flagged_vs_sample():
    from deckdna.audit.contextual import select_contextual_audit_slides
    from deckdna.audit.issues import AuditIssue

    # Без детерминированных флагов — выборка ≤ 3 слайдов (первый, средний, последний)
    assert select_contextual_audit_slides(10, [], max_sample_slides=3) == [1, 5, 10]
    assert select_contextual_audit_slides(2, [], max_sample_slides=3) == [1, 2]

    # С детерминированными флагами (warning/error) — только помеченные слайды (1-indexed)
    flagged = [
        AuditIssue(
            id="iss_1",
            audit_run_id="aud_1",
            deck_revision=1,
            slide_index=2,
            rule_code="text.overflow",
            severity="warning",
            deterministic=True,
            message="overflow",
        ),
        AuditIssue(
            id="iss_2",
            audit_run_id="aud_1",
            deck_revision=1,
            slide_index=5,
            rule_code="geometry.overlap",
            severity="error",
            deterministic=True,
            message="overlap",
        ),
        # info не считается флагом предупреждения
        AuditIssue(
            id="iss_3",
            audit_run_id="aud_1",
            deck_revision=1,
            slide_index=7,
            rule_code="style.note",
            severity="info",
            deterministic=True,
            message="info",
        ),
    ]
    assert select_contextual_audit_slides(10, flagged, max_sample_slides=3) == [3, 6]


async def test_contextual_audit_checks_only_flagged_or_sampled_slides(monkeypatch, tmp_path):
    """Parallel pipeline: при полном прогоне проверяются только слайды с детерминированным
    предупреждением или выборка ≤ 3 слайдов, с параллельным ограничением."""
    from pathlib import Path

    from deckdna.audit import contextual as ctx_mod
    from deckdna.audit.issues import AuditIssue
    from pptx import Presentation

    prs = Presentation()
    for i in range(6):
        s = prs.slides.add_slide(prs.slide_layouts[0])
        if s.shapes.title is not None:
            s.shapes.title.text = f"Слайд {i + 1}"
    deck = tmp_path / "six.pptx"
    prs.save(str(deck))

    def fake_render(pptx_path, out_dir, *, dpi=150, first=1, last=None):
        end = last or 6
        out = []
        for idx in range(first, end + 1):
            p = Path(out_dir) / f"slide-{idx:02d}.png"
            p.write_bytes(b"\x89PNG\r\n\x1a\n")
            out.append(p)
        return out

    monkeypatch.setattr(ctx_mod, "render_slides_png", fake_render)

    gw = MockProvider(
        fixtures={
            PROMPT_NAME: _verdicts(
                *[(str(number), "pass", "", 1.0) for number in range(1, 11)]
            )
        }
    )
    # 1. Без детерминированных флагов → выборка ≤ 3 слайдов из 6
    diag_sample: dict = {}
    await run_contextual_audit(
        deck, gw, deterministic_issues=[], diagnostics=diag_sample, max_concurrent=2
    )
    assert diag_sample["requested_slides"] == 3
    assert [c["payload"]["slide_image"] for c in gw.calls] == [1, 3, 6]

    # 2. С флагом только на слайде index=1 (2-й слайд) → ровно 1 вызов VLM
    gw.calls.clear()
    diag_flagged: dict = {}
    det_issue = AuditIssue(
        id="iss_flag",
        audit_run_id="aud_1",
        deck_revision=1,
        slide_index=1,
        rule_code="text.overflow",
        severity="warning",
        deterministic=True,
        message="overflow on slide 1",
    )
    await run_contextual_audit(
        deck,
        gw,
        deterministic_issues=[det_issue],
        diagnostics=diag_flagged,
        max_concurrent=2,
    )
    assert diag_flagged["requested_slides"] == 1
    assert [c["payload"]["slide_image"] for c in gw.calls] == [2]


def test_slide_risk_formula_and_selective_vision_routing():
    """Vision audit gate п.1: risk = 0.30*overflow + 0.20*font_substitution +
    0.20*novel_layout + 0.15*dense_chart + 0.15*low_style_confidence;
    зрением смотрим risk >= 0.45, первый слайд каждого семейства макета
    и слайды с детерминированным предупреждением."""
    from deckdna.audit.contextual import (
        compute_slide_risk,
        load_risk_config,
        select_contextual_audit_slides,
    )
    from deckdna.audit.issues import AuditIssue

    weights, threshold = load_risk_config()
    assert threshold == pytest.approx(0.45)
    assert weights == {
        "overflow": pytest.approx(0.30),
        "font_substitution": pytest.approx(0.20),
        "novel_layout": pytest.approx(0.20),
        "dense_chart": pytest.approx(0.15),
        "low_style_confidence": pytest.approx(0.15),
    }

    # Проверка формулы: overflow (1.0) + dense_chart (1.0) = 0.30 + 0.15 = 0.45
    assert compute_slide_risk(
        {"overflow": 1.0, "dense_chart": 1.0}, weights=weights
    ) == pytest.approx(0.45)
    # font_substitution (1.0) + low_style_confidence (1.0) = 0.20 + 0.15 = 0.35 < 0.45
    assert compute_slide_risk(
        {"font_substitution": 1.0, "low_style_confidence": 1.0}, weights=weights
    ) == pytest.approx(0.35)

    # 6 слайдов:
    # slide 0 (1-indexed 1): семейство A (первый в семействе A -> выбран)
    # slide 1 (1-indexed 2): семейство A, низкий риск, без предупреждений -> пропущен
    # slide 2 (1-indexed 3): семейство B (первый в семействе B -> выбран)
    # slide 3 (1-indexed 4): семейство B, но risk = 0.50 >= 0.45 (novel_layout + overflow) -> выбран
    # slide 4 (1-indexed 5): семейство A, детерминированный warning -> выбран
    # slide 5 (1-indexed 6): семейство B, низкий риск, только info -> пропущен
    compose_report = {
        "slides": [
            {"slide_index": 0, "layout_family": "cards_3", "mode": "cloned"},
            {"slide_index": 1, "layout_family": "cards_3", "mode": "cloned"},
            {"slide_index": 2, "layout_family": "timeline", "mode": "cloned"},
            {
                "slide_index": 3,
                "layout_family": "timeline",
                "mode": "synthesized",
                "overflow": True,
            },
            {"slide_index": 4, "layout_family": "cards_3", "mode": "cloned"},
            {"slide_index": 5, "layout_family": "timeline", "mode": "cloned"},
        ]
    }
    det_issues = [
        AuditIssue(
            id="iss_w",
            audit_run_id="aud_1",
            deck_revision=1,
            slide_index=4,
            rule_code="layout.overlap",
            severity="warning",
            deterministic=True,
            message="warn on slide 4",
        ),
        AuditIssue(
            id="iss_i",
            audit_run_id="aud_1",
            deck_revision=1,
            slide_index=5,
            rule_code="style.info",
            severity="info",
            deterministic=True,
            message="info on slide 5",
        ),
    ]
    selected = select_contextual_audit_slides(
        6, det_issues, compose_report=compose_report
    )
    assert selected == [1, 3, 4, 5]


async def test_unverified_critical_warning_slide_does_not_pass():
    """Vision audit gate п.2: слайд с непроверенным критическим предупреждением
    не считается прошедшим."""
    from deckdna.audit.contextual import slide_audit_passed
    from deckdna.audit.issues import AuditIssue

    crit_issue = AuditIssue(
        id="iss_crit",
        audit_run_id="aud_1",
        deck_revision=1,
        slide_index=0,
        rule_code="text.overflow",
        severity="warning",
        deterministic=True,
        message="critical overflow warning",
    )

    def _broken(payload):
        raise DeckDNAError(
            code="provider_unavailable", message="VLM timeout", stage="providers"
        )

    gw = MockProvider(fixtures={PROMPT_NAME: _broken})
    diagnostics: dict = {}
    await run_contextual_audit(
        PPTX,
        gw,
        first=1,
        last=1,
        deterministic_issues=[crit_issue],
        diagnostics=diagnostics,
    )
    assert diagnostics["unverified_slides"] == [0]
    assert diagnostics["unverified_critical_slides"] == [0]
    assert diagnostics["complete"] is False
    assert slide_audit_passed(0, [crit_issue], diagnostics["unverified_slides"]) is False
    # После успешной проверки VLM слайд без блокеров считается проверенным
    assert slide_audit_passed(0, [crit_issue], []) is True


