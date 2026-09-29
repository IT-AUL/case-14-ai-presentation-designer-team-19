"""AuditConfig: configs/audit.default.yaml реально читается в рантайме.

Пороги правил — конфигурационные данные из yaml, а не константы в коде:
тесты меняют значение в тестовом yaml и проверяют, что поведение
audit_deck реально изменилось.
"""

import yaml
from deckdna.audit.basic import (
    RULE_BULLET_COUNT,
    RULE_EDGE_MARGIN,
    RULE_PLACEHOLDER_TEXT,
    audit_deck,
)
from deckdna.audit.config import (
    DEFAULT_AUDIT_CONFIG_PATH,
    AuditConfig,
    default_audit_config,
    load_audit_config,
)
from pptx import Presentation
from pptx.util import Inches

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _issues_by_rule(issues, code):
    return [i for i in issues if i.rule_code == code]


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _write_config(tmp_path, **overrides):
    """yaml с полным набором порогов + overrides → путь."""
    data = {
        "version": "1",
        "overflow_tolerance": 0.05,
        "aspect_tolerance": 0.03,
        "raster_only_slide_area": 0.90,
        "out_of_bounds_tolerance": 0.005,
        "duplicate_similarity_threshold": 0.9,
        "max_bullets_per_slide": 6,
        "overlap_min_cover": 0.25,
        "overlap_containment": 0.90,
        "overlap_min_shape_area": 0.01,
        "overlap_bg_area": 0.80,
        "edge_margin": 0.03,
        "edge_bleed_span": 0.95,
        "placeholder_patterns": [],
    }
    data.update(overrides)
    path = tmp_path / "audit.yaml"
    path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    return path


def _bullet_deck(tmp_path, bullets):
    prs = Presentation()
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    tf = box.text_frame
    for i in range(bullets):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = f"Пункт {i}"
        ppr = par._p.get_or_add_pPr()
        ppr.append(ppr.makeelement(_A + "buChar", {"char": "•"}))
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def _edge_deck(tmp_path, x):
    """Текстовый бокс с зазором x (доля ширины) от левого края."""
    prs = Presentation()
    slide = _blank_slide(prs)
    sw = prs.slide_width
    box = slide.shapes.add_textbox(
        int(sw * x), Inches(2), int(sw * 0.30), Inches(1)
    )
    box.text_frame.text = "Контент"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def _text_deck(tmp_path, text):
    prs = Presentation()
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(Inches(2), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = text
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_shipped_yaml_loads_with_expected_defaults():
    """Отгружаемый configs/audit.default.yaml парсится и его значения
    совпадают с тем, что было захардкожено в коде до wiring."""
    cfg = load_audit_config(DEFAULT_AUDIT_CONFIG_PATH)
    assert cfg.overflow_tolerance == 0.05
    assert cfg.aspect_tolerance == 0.03
    assert cfg.raster_only_slide_area == 0.90
    assert cfg.out_of_bounds_tolerance == 0.005
    assert cfg.duplicate_similarity_threshold == 0.9
    assert cfg.max_bullets_per_slide == 6
    assert cfg.overlap_min_cover == 0.25
    assert cfg.overlap_containment == 0.90
    assert cfg.overlap_min_shape_area == 0.01
    assert cfg.overlap_bg_area == 0.80
    assert cfg.edge_margin == 0.03
    assert cfg.edge_bleed_span == 0.95
    # roadmap-ключи не потеряны — лежат в raw
    assert cfg.raw["contrast_min_ratio"] == 4.5
    assert cfg.raw["contextual"]["enabled"] is True
    assert "merge_slide" in cfg.raw["repair"]["allowed_actions"]


def test_default_config_is_cached_singleton():
    assert default_audit_config() is default_audit_config()
    assert isinstance(default_audit_config(), AuditConfig)


def test_bullet_threshold_from_yaml_changes_behavior(tmp_path):
    """4 буллета: дефолт max=6 — чисто; yaml max=3 — флагается.

    Доказывает, что порог читается из файла, а не из константы."""
    path = _bullet_deck(tmp_path, bullets=4)
    assert _issues_by_rule(audit_deck(path), RULE_BULLET_COUNT) == []

    cfg = load_audit_config(_write_config(tmp_path, max_bullets_per_slide=3))
    issues = _issues_by_rule(audit_deck(path, config=cfg), RULE_BULLET_COUNT)
    assert len(issues) == 1
    assert issues[0].measured_value == 4
    assert issues[0].threshold == 3


def test_edge_margin_from_yaml_changes_behavior(tmp_path):
    """Бокс с зазором 1% от левого края: дефолт 3% — флагается,
    yaml 0.5% — чисто."""
    path = _edge_deck(tmp_path, x=0.01)
    assert len(_issues_by_rule(audit_deck(path), RULE_EDGE_MARGIN)) == 1

    cfg = load_audit_config(_write_config(tmp_path, edge_margin=0.005))
    assert _issues_by_rule(audit_deck(path, config=cfg), RULE_EDGE_MARGIN) == []


def test_placeholder_patterns_from_yaml_whole_tier(tmp_path):
    """yaml placeholder_patterns — дополнительные WHOLE-ярус паттерны:
    фигура, весь текст которой матчится паттерну, флагается."""
    path = _text_deck(tmp_path, "Секретная заглушка")
    assert _issues_by_rule(audit_deck(path), RULE_PLACEHOLDER_TEXT) == []

    cfg = load_audit_config(
        _write_config(tmp_path, placeholder_patterns=["секретная заглушка"])
    )
    assert len(_issues_by_rule(audit_deck(path, config=cfg), RULE_PLACEHOLDER_TEXT)) == 1


def test_placeholder_patterns_whole_tier_no_substring_fp(tmp_path):
    """Паттерн из yaml матчится только целым текстом — слово внутри
    нормального предложения не флагается (консервативная семантика)."""
    path = _text_deck(tmp_path, "Секретная заглушка важна для анализа")
    cfg = load_audit_config(
        _write_config(tmp_path, placeholder_patterns=["секретная заглушка"])
    )
    assert _issues_by_rule(audit_deck(path, config=cfg), RULE_PLACEHOLDER_TEXT) == []


def test_partial_yaml_falls_back_to_field_defaults(tmp_path):
    """yaml только с одним ключом: остальные пороги — дефолты модели."""
    path = tmp_path / "audit.yaml"
    path.write_text('version: "1"\nedge_margin: 0.10\n', encoding="utf-8")
    cfg = load_audit_config(path)
    assert cfg.edge_margin == 0.10
    assert cfg.max_bullets_per_slide == 6
    assert cfg.raw["edge_margin"] == 0.10
