"""ADR-030: integrity.filler_slide — контентный слайд из одного заголовка.
Роль слайда берётся из отчёта компоновки: обложка и разделитель законно
состоят из заголовка."""

from __future__ import annotations

from deckdna.audit import catalog
from deckdna.audit.basic import RULE_FILLER_SLIDE, audit_filler_slides
from pptx import Presentation
from pptx.util import Inches


def _deck(tmp_path, bodies: list[str]):
    prs = Presentation()
    layout = prs.slide_layouts[5]  # «только заголовок»
    for i, body in enumerate(bodies):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"Заголовок {i}"
        if body:
            box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
            box.text_frame.text = body
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def _report(purposes: list[str]) -> dict:
    return {"slides": [
        {"purpose": p, "title": f"Заголовок {i}", "images": {}} for i, p in enumerate(purposes)
    ]}


def test_title_only_content_slide_is_flagged(tmp_path):
    path = _deck(tmp_path, ["", "", "Содержательный текст слайда про рынок и выводы.", ""])
    issues = audit_filler_slides(
        path, _report(["title", "overview", "solution", "section_divider"]), deck_revision=1
    )
    assert [(i.rule_code, i.slide_index) for i in issues] == [(RULE_FILLER_SLIDE, 1)]
    assert issues[0].severity == "error"
    assert issues[0].deck_revision == 1


def test_mismatched_report_stays_silent(tmp_path):
    path = _deck(tmp_path, ["", ""])
    assert audit_filler_slides(path, _report(["overview"])) == []


def test_rule_is_in_the_catalog():
    assert catalog.RULES[RULE_FILLER_SLIDE].deterministic
