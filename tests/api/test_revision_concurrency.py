"""Concurrent local edits must read the current deck inside the worker lock."""

from __future__ import annotations

import asyncio
import json
import threading
from pathlib import Path

import httpx
import pytest
from deckdna.api import app as api
from deckdna.generation import regenerate as module
from fastapi.testclient import TestClient
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
A = "/api/v1"


@pytest.fixture
def generated(wait_generation):
    client = TestClient(api.app)
    pid = client.post(f"{A}/projects", json={"name": "revision concurrency"}).json()["id"]
    template = ROOT / "tests/fixtures/pptx/vk_tech_template.pptx"
    tid = client.post(f"{A}/projects/{pid}/templates", files={
        "file": (template.name, template.read_bytes(), api._PPTX_MIME),
    }).json()["id"]
    pack = client.post(f"{A}/projects/{pid}/content-packs", files={
        "files": ("article.md", (ROOT / "tests/fixtures/content/poc_article.md").read_bytes(),
                  "text/markdown"),
    }, data={"brief": json.dumps({"language": "ru"})}).json()["content_pack"]["id"]
    accepted = client.post(f"{A}/projects/{pid}/generations", json={
        "template_id": tid, "content_pack_id": pack, "variants": [{"strategy": "balanced"}],
        "brief": {"purpose": "Обзор", "audience": "команда", "language": "ru",
                  "target_slide_count": 5}, "use_llm": False,
    }).json()
    result = wait_generation(client, accepted["generation_id"])
    assert result["state"] == "completed", result
    return api.STORE.runs[result["id"]], accepted["variant_ids"][0]


async def _parallel_edits(variant_id, slides, expected=None):
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=api.app),
                                 base_url="http://test") as client:
        body = {"action": "shorter"}
        if expected is not None:
            body["expected_revision"] = expected
        return await asyncio.gather(*[
            client.post(f"{A}/variants/{variant_id}/slides/{slide.id}/regenerate", json=body)
            for slide in slides
        ])


@pytest.mark.asyncio
async def test_parallel_edits_do_not_compute_from_the_same_deck(generated, monkeypatch):
    run, vid = generated
    slides = api.STORE.slides[vid][:2]
    original = module.regenerate_slide
    entered = threading.Event()
    release = threading.Event()
    bases = []

    def gated(*args, **kwargs):
        bases.append(args[1])
        if len(bases) == 1:
            entered.set()
            assert release.wait(10)
        return original(*args, **kwargs)

    monkeypatch.setattr(module, "regenerate_slide", gated)
    task = asyncio.create_task(_parallel_edits(vid, slides))
    try:
        assert await asyncio.to_thread(entered.wait, 10)
        # Let the second request cross the async handler while the first works.
        await asyncio.sleep(0.2)
    finally:
        release.set()
    responses = await task
    assert [r.status_code for r in responses] == [200, 200], [r.text for r in responses]
    assert sorted(r.json()["deck_revision"] for r in responses) == [2, 3]
    assert len(bases) == 2 and bases[0] != bases[1], "second edit reused stale deck bytes"
    assert run.deck_revision[vid] == 3


@pytest.mark.asyncio
async def test_same_expected_revision_allows_only_one_writer(generated):
    run, vid = generated
    revision = run.deck_revision[vid]
    responses = await _parallel_edits(vid, api.STORE.slides[vid][:2], expected=revision)
    assert sorted(r.status_code for r in responses) == [200, 409]
    assert run.deck_revision[vid] == revision + 1
    assert run.variant_artifacts[vid].keys() == {"pptx"}


def test_latest_deck_still_opens(generated, tmp_path):
    run, vid = generated
    path = tmp_path / "deck.pptx"
    path.write_bytes(api._current_deck(run, vid).data)
    assert len(Presentation(path).slides) == len(api.STORE.slides[vid])


@pytest.mark.asyncio
async def test_export_formats_share_one_revision_while_edit_waits(generated, monkeypatch):
    run, vid = generated
    original = api._ensure_artifact
    entered, release = threading.Event(), threading.Event()
    observed = []

    def gated(*args):
        observed.append(run.deck_revision.get(vid, 1))
        if len(observed) == 1:
            entered.set()
            assert release.wait(10)
        return original(*args)

    monkeypatch.setattr(api, "_ensure_artifact", gated)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=api.app),
                                 base_url="http://test") as client:
        export = asyncio.create_task(client.post(f"{A}/variants/{vid}/exports", json={
            "formats": ["pptx", "pdf", "quality_passport"],
        }))
        edit = None
        try:
            assert await asyncio.to_thread(entered.wait, 10)
            edit = asyncio.create_task(client.post(
                f"{A}/variants/{vid}/slides/{api.STORE.slides[vid][1].id}/regenerate",
                json={"action": "shorter", "expected_revision": 1},
            ))
            await asyncio.sleep(.1)
        finally:
            release.set()
        response = await export
        assert response.status_code == 202, response.text
        assert (await edit).status_code == 200
        saved = api.STORE.exports[response.json()["export_id"]]
        assert saved.deck_revision == 1 and observed == [1, 1, 1]
        assert run.deck_revision[vid] == 2
