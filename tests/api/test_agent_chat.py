"""ADR-058: чат с помощником поверх собранной колоды и плана."""

from __future__ import annotations

import io
import json
import time
import zipfile

import pytest
from deckdna.api import app as api
from deckdna.assistant.agent import AgentAction, AgentQuestion, AgentTurn
from fastapi.testclient import TestClient
from test_revision_concurrency import generated as generated  # noqa: F401

A = "/api/v1"
client = TestClient(api.app)


class Scripted:
    def __init__(self, turns):
        self.turns = list(turns)
        self.payloads = []

    async def text_json(self, name, payload, schema):
        self.payloads.append(payload)
        return self.turns.pop(0) if self.turns else AgentTurn(reply="Готово.")


@pytest.fixture
def model(monkeypatch):
    holder = {}

    def install(*turns):
        gw = Scripted(turns)
        holder["gw"] = gw
        monkeypatch.setattr(api, "_server_provider_ready", lambda: True)
        monkeypatch.setattr(api, "_gateway_for", lambda body, project_id=None: gw)
        return gw

    return install


def _wait(job_id, timeout=180):
    deadline = time.monotonic() + timeout
    while api.STORE.jobs[job_id].state not in api._TERMINAL_JOB_STATES:
        assert time.monotonic() < deadline, "assistant job did not finish"
        time.sleep(.05)
    return api.STORE.jobs[job_id]


def _send(pid, **body):
    response = client.post(f"{A}/projects/{pid}/agent/messages", json=body)
    assert response.status_code == 202, response.text
    return response.json()


def _slides(data):
    from deckdna.pptx.composing.slide_swap import slide_parts

    with zipfile.ZipFile(io.BytesIO(data)) as z:
        files = {n: z.read(n) for n in z.namelist()}
    return [files[p] for p in slide_parts(files)]


def test_free_text_rewrite_changes_only_that_slide(generated, model):  # noqa: F811
    from pptx import Presentation

    run, vid = generated
    pid = run.detail.project_id
    target = api.STORE.slides[vid][2]
    before = _slides(api._current_deck(run, vid).data)
    gw = model(AgentTurn(reply="Переписал слайд 3 короче и по делу.", actions=[AgentAction(
        tool="rewrite_slide", slide=3, title="Главное за минуту",
        items=["Первый тезис без цифр", "Второй тезис без цифр"], source="model")]))
    sent = _send(pid, text="Перепиши третий слайд для питча", variant_id=vid,
                 slide_id=target.slide_plan_id)
    job = _wait(sent["job_id"])
    assert job.state == api.JobState.completed
    assert gw.payloads[0]["selected_slide"] == 3
    assert gw.payloads[0]["deck"]["slides"][2]["n"] == 3
    after_deck = api._current_deck(run, vid).data
    after = _slides(after_deck)
    assert [i for i, (a, b) in enumerate(zip(before, after, strict=True)) if a != b] == [2]
    text = " ".join(sh.text_frame.text for sh in Presentation(io.BytesIO(after_deck)).slides[2]
                    .shapes if sh.has_text_frame)
    assert "Первый тезис" in text
    assert run.deck_revision[vid] == 2
    thread = client.get(f"{A}/projects/{pid}/agent/thread").json()["items"]
    assert [m["role"] for m in thread] == ["user", "assistant"]
    assert thread[1]["actions"][0] == {"tool": "rewrite_slide", "slide": 3, "ok": True,
                                       "summary": "обновил текст слайда", "revision": 2}
    events = [e["event"] for e in api.STORE.job_events[sent["job_id"]]]
    assert "agent.action" in events and "slide.updated" in events
    assert events[-1] == "job.completed" and "agent.message" in events


def test_invented_number_is_rejected_and_deck_kept(generated, model):  # noqa: F811
    run, vid = generated
    pid = run.detail.project_id
    deck_id = api._current_deck(run, vid).id
    gw = model(
        AgentTurn(reply="Добавил.", actions=[AgentAction(
            tool="rewrite_slide", slide=2, items=["Выручка выросла до 987654 млн ₽"])]),
        AgentTurn(reply="Такой цифры нет в материалах — оставил слайд как был."),
    )
    job = _wait(_send(pid, text="Добавь выручку на слайд 2", variant_id=vid)["job_id"])
    assert job.state == api.JobState.completed
    assert api._current_deck(run, vid).id == deck_id
    obs = gw.payloads[1]["observations"][0]
    assert obs["ok"] is False and "987654" in obs["error"]
    reply = client.get(f"{A}/projects/{pid}/agent/thread").json()["items"][-1]
    assert reply["actions"][0]["ok"] is False


def test_add_move_remove_and_undo_slides(generated, model):  # noqa: F811
    from pptx import Presentation

    run, vid = generated
    pid = run.detail.project_id
    count = len(api.STORE.slides[vid])
    model(AgentTurn(reply="Добавил слайд про команду.", actions=[AgentAction(
        tool="add_slide", slide=2, title="Команда проекта",
        items=["Роли и зоны ответственности", "Как мы работаем вместе"], source="model")]))
    assert _wait(_send(pid, text="Добавь слайд про команду после второго",
                       variant_id=vid)["job_id"]).state == api.JobState.completed
    cards = api.STORE.slides[vid]
    assert len(cards) == count + 1 and [c.index for c in cards] == list(range(1, count + 2))
    assert cards[2].title == "Команда проекта"
    deck = api._current_deck(run, vid).data
    assert len(Presentation(io.BytesIO(deck)).slides) == count + 1
    assert all(c.preview_artifact_id for c in cards)
    plan_ids = [s.id for s in run.variant_plans[vid].slides]
    assert plan_ids[2] == cards[2].slide_plan_id

    model(AgentTurn(reply="Переставил.", actions=[AgentAction(tool="move_slide", slide=3, to=4)]))
    assert _wait(_send(pid, text="Перенеси его дальше", variant_id=vid)["job_id"]).state == \
        api.JobState.completed
    assert api.STORE.slides[vid][3].title == "Команда проекта"

    model(AgentTurn(reply="Убрал.", actions=[AgentAction(tool="remove_slide", slide=4)]))
    assert _wait(_send(pid, text="Нет, убери его", variant_id=vid)["job_id"]).state == \
        api.JobState.completed
    assert len(api.STORE.slides[vid]) == count
    assert "Команда проекта" not in [c.title for c in api.STORE.slides[vid]]

    model(AgentTurn(reply="Вернул.", actions=[AgentAction(tool="undo")]))
    assert _wait(_send(pid, text="Верни как было", variant_id=vid)["job_id"]).state == \
        api.JobState.completed
    assert len(api.STORE.slides[vid]) == count + 1
    assert len(Presentation(io.BytesIO(api._current_deck(run, vid).data)).slides) == count + 1
    assert run.deck_revision[vid] == 5


def test_question_and_busy_and_cancel(generated, model):  # noqa: F811
    run, vid = generated
    pid = run.detail.project_id
    model(AgentTurn(reply="Подскажите, для кого презентация?", question=AgentQuestion(
        text="Для кого?", options=["Инвесторы", "Команда"])))
    job = _wait(_send(pid, text="Сделай убедительнее", variant_id=vid)["job_id"])
    assert job.state == api.JobState.completed
    last = client.get(f"{A}/projects/{pid}/agent/thread").json()["items"][-1]
    assert last["question"]["options"] == ["Инвесторы", "Команда"]
    assert run.deck_revision.get(vid, 1) == 1
    api.STORE.jobs["job_busy"] = api.Job(id="job_busy", kind=api.JobKind.assistant,
                                         state=api.JobState.running, project_id=pid,
                                         created_at=api._utcnow(), progress=0)
    busy = client.post(f"{A}/projects/{pid}/agent/messages", json={"text": "ещё"})
    assert busy.status_code == 409
    canceled = client.post(f"{A}/projects/{pid}/agent/jobs/job_busy/cancel")
    assert canceled.status_code == 200
    assert client.delete(f"{A}/projects/{pid}/agent/thread").status_code == 204
    assert client.get(f"{A}/projects/{pid}/agent/thread").json()["items"] == []


def test_without_model_honest_reply_and_quick_command(generated):  # noqa: F811
    run, vid = generated
    pid = run.detail.project_id
    job = _wait(_send(pid, text="добавь аудиторию на 6 слайде", variant_id=vid)["job_id"])
    assert job.state == api.JobState.completed
    reply = client.get(f"{A}/projects/{pid}/agent/thread").json()["items"][-1]
    assert reply["mode"] == "unavailable" and "модель" in reply["text"]
    job = _wait(_send(pid, text="сделай короче на слайде 3", variant_id=vid)["job_id"])
    assert job.state == api.JobState.completed
    reply = client.get(f"{A}/projects/{pid}/agent/thread").json()["items"][-1]
    assert reply["mode"] == "local" and reply["actions"][0]["tool"] == "shorter"


def test_plan_mode_adds_slide_to_plan(model, tmp_path):
    from pathlib import Path

    root = Path(__file__).resolve().parents[2]
    pid = client.post(f"{A}/projects", json={"name": "agent plan"}).json()["id"]
    template = root / "tests/fixtures/pptx/vk_tech_template.pptx"
    tid = client.post(f"{A}/projects/{pid}/templates", files={
        "file": (template.name, template.read_bytes(), api._PPTX_MIME)}).json()["id"]
    pack = client.post(f"{A}/projects/{pid}/content-packs", files={
        "files": ("article.md", (root / "tests/fixtures/content/poc_article.md").read_bytes(),
                  "text/markdown")}, data={"brief": json.dumps({"language": "ru"})}
    ).json()["content_pack"]["id"]
    plan = client.post(f"{A}/projects/{pid}/plans", json={
        "template_id": tid, "content_pack_id": pack, "use_llm": False,
        "brief": {"purpose": "Обзор", "audience": "команда", "language": "ru",
                  "target_slide_count": 5}, "respond_async": False})
    assert plan.status_code in (200, 201), plan.text
    plan_id = plan.json()["plan_id"]
    before = len(api.STORE.plans[plan_id].plans[0].deck_plan.slides)
    model(AgentTurn(reply="Добавил слайд.", actions=[AgentAction(
        tool="add_slide", slide=2, title="Что мы просим", items=["Решение по пилоту"],
        source="model")]))
    job = _wait(_send(pid, text="Добавь слайд с просьбой", plan_id=plan_id)["job_id"])
    assert job.state == api.JobState.completed
    slides = api.STORE.plans[plan_id].plans[0].deck_plan.slides
    assert len(slides) == before + 1 and slides[2].title_intent == "Что мы просим"
    model(AgentTurn(reply="Вернул.", actions=[AgentAction(tool="undo")]))
    _wait(_send(pid, text="верни", plan_id=plan_id)["job_id"])
    assert len(api.STORE.plans[plan_id].plans[0].deck_plan.slides) == before
