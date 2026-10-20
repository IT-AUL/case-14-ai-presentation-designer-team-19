"""Required sections cannot vanish through plan actions or inline generation."""
import copy

import pytest
from deckdna.api import app as api
from fastapi.testclient import TestClient
from pptx import Presentation

A = "/api/v1"
TEXT = "В пилоте проверен один результат. Гарантий повторения пока нет."


@pytest.fixture
def source(monkeypatch, tmp_path):
    monkeypatch.setattr(api, "STORE", api._Store())
    monkeypatch.setattr(api.settings, "state_snapshot", False)
    client = TestClient(api.app)
    pid = client.post(f"{A}/projects", json={"name": "required"}).json()["id"]
    path = tmp_path / "unseen.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(path)
    tid = client.post(f"{A}/projects/{pid}/templates", files={
        "file": (path.name, path.read_bytes(), api._PPTX_MIME),
    }).json()["id"]
    cp = client.post(f"{A}/projects/{pid}/content-packs", files={
        "files": ("author.md", ("# Обзор\n\n## Работа\n\nКоманда подготовила обзор.\n\n"
                                 f"## Риски\n\n{TEXT}\n").encode(), "text/markdown"),
    }).json()["content_pack"]["id"]
    brief = {"purpose": "Обзор", "audience": "команда", "language": "ru",
             "target_slide_count": 4, "mandatory_sections": ["Риски"], "content_fill": "shrink"}
    return client, pid, tid, cp, brief


def build(source):
    client, pid, tid, cp, brief = source
    response = client.post(f"{A}/projects/{pid}/plans", json={
        "content_pack_id": cp, "template_id": tid, "strategies": ["faithful"],
        "brief": brief, "use_llm": False,
    })
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize("entry", ["plans", "generations"])
def test_unknown_section_fails_before_gateway_or_job(source, monkeypatch, entry):
    client, pid, tid, cp, brief = source

    def forbidden(*args, **kwargs):
        raise AssertionError("section preflight must precede gateway creation")

    monkeypatch.setattr(api, "build_gateway", forbidden)
    before = (set(api.STORE.jobs), copy.deepcopy(api.STORE.project_briefs[pid]))
    payload = {
        "content_pack_id": cp, "template_id": tid,
        "brief": {**brief, "mandatory_sections": ["Экономика"]}, "use_llm": True,
    }
    if entry == "generations":
        payload["variants"] = [{"strategy": "faithful"}]
    response = client.post(f"{A}/projects/{pid}/{entry}", json=payload)
    assert response.status_code == 422, response.text
    assert response.json()["error"]["details"]["rule"] == "request.mandatory_section"
    assert (set(api.STORE.jobs), api.STORE.project_briefs[pid]) == before


def test_required_source_is_locked_and_remove_is_atomic(source):
    client, *_ = source
    result = build(source)
    plan = result["plans"][0]["deck_plan"]
    slide = next(s for s in plan["slides"] if any(u.get("text") == TEXT
                                                 for u in s["content_units"]))
    assert slide["mandatory"]
    assert next(u for u in slide["content_units"] if u.get("text") == TEXT)["locked"]
    before = copy.deepcopy(api.STORE.plans[result["plan_id"]])
    response = client.post(f"{A}/plans/{result['plan_id']}/slides/{slide['id']}/actions",
                           json={"action": "remove"})
    assert response.status_code == 422, response.text
    assert api.STORE.plans[result["plan_id"]] == before


@pytest.mark.parametrize("strip_requirement", [False, True])
def test_inline_plan_cannot_remove_required_content(source, strip_requirement):
    client, pid, tid, cp, brief = source
    result = build(source)
    plan = result["plans"][0]["deck_plan"]
    for slide in plan["slides"]:
        for unit in slide["content_units"]:
            if unit.get("text") == TEXT:
                unit["text"] = "Подставленный пересказ другой темы."
                unit["locked"] = False
    if strip_requirement:
        plan["brief"].pop("mandatory_sections", None)
    before = set(api.STORE.jobs)
    response = client.post(f"{A}/projects/{pid}/generations", json={
        "template_id": tid, "content_pack_id": cp, "brief": brief, "deck_plan": plan,
        "variants": [{"strategy": "faithful"}], "use_llm": False,
    })
    assert response.status_code == (409 if strip_requirement else 422), response.text
    assert set(api.STORE.jobs) == before


def test_author_replacement_cannot_silently_remove_a_required_source(source, wait_generation):
    client, pid, tid, cp, brief = source
    result = build(source)
    accepted = client.post(f"{A}/projects/{pid}/generations", json={
        "template_id": tid, "content_pack_id": cp, "brief": brief,
        "deck_plan_id": result["plan_id"], "variants": [{"strategy": "faithful"}], "use_llm": False,
    })
    assert accepted.status_code == 202, accepted.text
    final = wait_generation(client, accepted.json()["generation_id"])
    assert final["state"] == "completed", final
    vid = accepted.json()["variant_ids"][0]
    run = api.STORE.runs[final["id"]]
    original = api._current_deck(run, vid)
    source_slide = next(s.id for s in run.variant_plans[vid].slides
                        if any(u.text == TEXT for u in s.content_units))
    target = next(s for s in api.STORE.slides[vid] if s.slide_plan_id == source_slide)
    response = client.post(f"{A}/variants/{vid}/slides/{target.id}/regenerate", json={
        "action": "replace_text", "text": "Другой материал без обязательного источника.",
        "expected_revision": 1,
    })
    assert response.status_code == 422, response.text
    assert response.json()["error"]["details"]["rule"] == "request.mandatory_section"
    assert api._current_deck(run, vid).id == original.id and run.deck_revision.get(vid, 1) == 1
