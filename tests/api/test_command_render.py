from __future__ import annotations

import time

import pytest
from deckdna.api import app as api
from fastapi.testclient import TestClient
from test_revision_concurrency import generated as generated

A = "/api/v1"
client = TestClient(api.app)


def test_authored_title_and_invalidation_on_new_audience(generated):
    run, vid = generated
    slide = api.STORE.slides[vid][0]
    response = client.post(f"{A}/projects/{run.detail.project_id}/commands", json={
        "text": "заголовок: Итоги проекта", "variant_id": vid,
        "slide_id": slide.slide_plan_id, "expected_revision": 1,
    })
    assert response.status_code == 202, response.text
    record = response.json()["change"]
    deadline = time.monotonic() + 90
    while time.monotonic() < deadline:
        state = api.STORE.jobs[record["job_id"]].state
        if state in api._TERMINAL_JOB_STATES:
            break
        time.sleep(.05)
    assert state == api.JobState.completed
    assert run.variant_plans[vid].slides[0].title_intent == "Итоги проекта"
    assert api.STORE.slides[vid][0].title == 'Итоги проекта'
    assert run.detail.deck_plan.slides[0].title_intent == 'Итоги проекта'
    changed = client.post(f"{A}/projects/{run.detail.project_id}/commands", json={
        "text": "аудитория: инвесторы",
    })
    assert changed.json()["status"] == "applied"
    session = client.get(f"{A}/projects/{run.detail.project_id}/session").json()
    assert session["artifacts"]["generation_id"] is None
    assert vid in api.STORE.slides  # history retained


def test_text_replacement_updates_speaker_notes_and_preserves_other_slide_xml(generated):
    from io import BytesIO

    from deckdna.api import changes
    from pptx import Presentation
    from test_previews_regenerate import _slides_of

    run, vid = generated
    slide = api.STORE.slides[vid][2]
    before = api._current_deck(run, vid).data
    response = client.post(f"{A}/projects/{run.detail.project_id}/commands", json={
        "text": "текст: Авторский вывод по материалам проекта.", "variant_id": vid,
        "slide_id": slide.slide_plan_id, "expected_revision": 1,
    })
    assert response.status_code == 202, response.text
    changes.drain(run.detail.id, vid)
    deadline = time.monotonic() + 90
    job_id = response.json()["change"]["job_id"]
    while api.STORE.jobs[job_id].state not in api._TERMINAL_JOB_STATES:
        assert time.monotonic() < deadline
        time.sleep(.05)
    assert api.STORE.jobs[job_id].state == api.JobState.completed
    after = api._current_deck(run, vid).data
    old, new = _slides_of(before), _slides_of(after)
    assert [name for name in old if old[name] != new[name]] == [list(old)[slide.index-1]]
    actual = Presentation(BytesIO(after)).slides[slide.index-1]
    assert "Авторский вывод" in actual.notes_slide.notes_text_frame.text


@pytest.mark.parametrize('mutation', ['brief', 'delete'])
def test_goal_change_rejects_late_local_edit(generated, monkeypatch, mutation):
    import threading

    from deckdna.generation import regenerate

    run, vid = generated
    before = api._current_deck(run, vid).id
    entered, release = threading.Event(), threading.Event()
    original = regenerate.regenerate_slide

    def paused(*args, **kwargs):
        entered.set()
        assert release.wait(10)
        return original(*args, **kwargs)

    monkeypatch.setattr(regenerate, "regenerate_slide", paused)
    accepted = client.post(f"{A}/projects/{run.detail.project_id}/commands", json={
        "text": "сократи", "variant_id": vid,
        "slide_id": api.STORE.slides[vid][2].slide_plan_id, "expected_revision": 1,
    }).json()["change"]
    try:
        assert entered.wait(10)
        if mutation == 'delete':
            changed = client.delete(f"{A}/projects/{run.detail.project_id}")
            assert changed.status_code == 204
        else:
            changed = client.post(f"{A}/projects/{run.detail.project_id}/commands", json={
                "text": "цель: Новый план выступления",
            })
            assert changed.json()["status"] == "applied"
    finally:
        release.set()
    deadline = time.monotonic() + 60
    while any(i["state"] == "running" for i in run.changes[vid]):
        assert time.monotonic() < deadline
        time.sleep(.05)
    if mutation == 'delete':
        assert vid not in api.STORE.slides
        assert not any(a.project_id == run.detail.project_id for a in api.STORE.artifacts.values())
    else:
        assert api._current_deck(run, vid).id == before
    assert api.STORE.jobs[accepted["job_id"]].state == api.JobState.canceled
    assert not any(e["event"] == "slide.updated" for e in api.STORE.job_events[accepted["job_id"]])
