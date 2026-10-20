"""Hard counts must hold before a regenerated/repaired deck becomes current."""

import io

import pytest
from deckdna.api import app as api
from deckdna.generation import regenerate as module
from fastapi.testclient import TestClient
from pptx import Presentation

A = "/api/v1"


def _extra_page(data):
    prs = Presentation(io.BytesIO(data))
    prs.slides.add_slide(prs.slide_layouts[6])
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


@pytest.fixture
def generated(monkeypatch, tmp_path, wait_generation):
    monkeypatch.setattr(api, "STORE", api._Store())
    monkeypatch.setattr(api.settings, "state_snapshot", False)
    client = TestClient(api.app)
    pid = client.post(f"{A}/projects", json={"name": "hard count"}).json()["id"]
    template = tmp_path / "template.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(template)
    tid = client.post(f"{A}/projects/{pid}/templates", files={
        "file": (template.name, template.read_bytes(), api._PPTX_MIME),
    }).json()["id"]
    cp = client.post(f"{A}/projects/{pid}/content-packs", files={
        "files": ("source.md", b"# Report\n\n## Work\n\n- Findings\n- Limits\n", "text/markdown"),
    }).json()["content_pack"]["id"]
    accepted = client.post(f"{A}/projects/{pid}/generations", json={
        "template_id": tid, "content_pack_id": cp, "use_llm": False,
        "brief": {"purpose": "Report", "audience": "team", "language": "en",
                  "target_slide_count": 3, "slide_count_policy": "maximum"},
        "variants": [{"strategy": "balanced"}],
    }).json()
    result = wait_generation(client, accepted["generation_id"])
    assert result["state"] == "completed", result
    vid = accepted["variant_ids"][0]
    return client, api.STORE.runs[result["id"]], vid


def test_regenerate_cannot_commit_an_extra_page(generated, monkeypatch):
    client, run, vid = generated
    original = module.regenerate_slide
    deck = api._current_deck(run, vid)

    def append(*args, **kwargs):
        result = original(*args, **kwargs)
        result.deck = _extra_page(result.deck)
        return result

    monkeypatch.setattr(module, "regenerate_slide", append)
    response = client.post(f"{A}/variants/{vid}/slides/{api.STORE.slides[vid][1].id}/regenerate",
                           json={"action": "shorter", "expected_revision": 1})
    assert response.status_code == 422, response.text
    assert response.json()["error"]["details"]["rule"] == "request.slide_count"
    assert run.deck_revision.get(vid, 1) == 1
    assert api._current_deck(run, vid).id == deck.id


def test_repair_cannot_commit_an_extra_page(generated, monkeypatch):
    client, run, vid = generated
    audit = api.STORE.audits[api._audit_id_of(vid)]
    issue = api.AuditIssue(
        schema_version="1.0", id="count-probe", audit_run_id=audit.id,
        deck_revision=1, rule_code="layout.edge_margin", deterministic=True,
        severity="warning", message="Repair probe", status="open", repairable=True,
        slide_index=1, slide_id="test",
    )
    api.STORE.issues[audit.id].append(issue)
    original = api.apply_repairs

    def append(source, actions, out, **kwargs):
        report = original(source, actions, out, **kwargs)
        out.write_bytes(_extra_page(out.read_bytes()))
        return report

    monkeypatch.setattr(api, "apply_repairs", append)
    deck = api._current_deck(run, vid)
    response = client.post(f"{A}/audits/{audit.id}/repairs", json={
        "selected_issue_ids": [issue.id], "use_llm": False,
    })
    assert response.status_code == 422, response.text
    assert response.json()["error"]["details"]["rule"] == "request.slide_count"
    assert api._current_deck(run, vid).id == deck.id
    assert run.deck_revision.get(vid, 1) == 1


def test_lazy_pdf_count_conflict_does_not_publish_export(generated, monkeypatch):
    import pymupdf

    client, run, vid = generated
    run.variant_artifacts[vid].pop("pdf", None)
    before = set(api.STORE.exports)

    def missing_page(source, target):
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.new_page()
            pdf.save(target)
        return target

    monkeypatch.setattr(api, "render_pdf", missing_page)
    response = client.post(f"{A}/variants/{vid}/exports", json={"formats": ["pdf"]})
    assert response.status_code == 422, response.text
    assert response.json()["error"]["details"]["rule"] == "request.export_count"
    assert "pdf" not in run.variant_artifacts[vid]
    assert set(api.STORE.exports) == before


def _lock_body(run, vid):
    from deckdna.contracts.deck_plan import Kind
    from deckdna.planning.locked_copy import validate_locked_deck

    plan = run.variant_plans[vid]
    slide = plan.slides[1]
    unit = next(u for u in slide.content_units if u.text and u.kind != Kind.title)
    unit.locked = True
    reports = run.variant_reports[vid]["compose_report"]["slides"]
    validate_locked_deck(plan, api._current_deck(run, vid).data, reports)
    return unit.text


def _corrupt_text(data, old):
    import zipfile

    from lxml import etree

    source, out = io.BytesIO(data), io.BytesIO()
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(out, "w") as zout:
        for entry in zin.infolist():
            payload = zin.read(entry.filename)
            if entry.filename.startswith("ppt/slides/slide") and entry.filename.endswith(".xml"):
                root = etree.fromstring(payload)
                for node in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"):
                    if old in (node.text or ""):
                        node.text = node.text.replace(old, "Changed without author permission")
                payload = etree.tostring(root)
            zout.writestr(entry, payload)
    return out.getvalue()


def test_regenerate_cannot_publish_changed_locked_copy(generated, monkeypatch):
    client, run, vid = generated
    old = _lock_body(run, vid)
    original = module.regenerate_slide
    deck = api._current_deck(run, vid)

    def corrupt(*args, **kwargs):
        result = original(*args, **kwargs)
        result.deck = _corrupt_text(result.deck, old)
        # Even an accidental in-place unlock in a dependency cannot mutate STORE
        # or let the candidate define its own acceptance contract.
        for unit in result.plan.slides[1].content_units:
            unit.locked = False
        return result

    monkeypatch.setattr(module, "regenerate_slide", corrupt)
    response = client.post(f"{A}/variants/{vid}/slides/{api.STORE.slides[vid][1].id}/regenerate",
                           json={"action": "shorter", "expected_revision": 1})
    assert response.status_code == 422, response.text
    assert response.json()["error"]["details"]["rule"] == "request.locked_copy"
    assert api._current_deck(run, vid).id == deck.id and run.deck_revision.get(vid, 1) == 1
    assert any(u.locked for u in run.variant_plans[vid].slides[1].content_units)


def test_repair_cannot_publish_changed_locked_copy(generated, monkeypatch):
    client, run, vid = generated
    old = _lock_body(run, vid)
    audit = api.STORE.audits[api._audit_id_of(vid)]
    issue = api.AuditIssue(
        schema_version="1.0", id="copy-probe", audit_run_id=audit.id,
        deck_revision=1, rule_code="layout.edge_margin", deterministic=True,
        severity="warning", message="Repair probe", status="open", repairable=True,
        slide_index=1, slide_id="test",
    )
    api.STORE.issues[audit.id].append(issue)
    original = api.apply_repairs
    deck = api._current_deck(run, vid)

    def corrupt(source, actions, out, **kwargs):
        report = original(source, actions, out, **kwargs)
        out.write_bytes(_corrupt_text(out.read_bytes(), old))
        return report

    monkeypatch.setattr(api, "apply_repairs", corrupt)
    response = client.post(f"{A}/audits/{audit.id}/repairs", json={
        "selected_issue_ids": [issue.id], "use_llm": False,
    })
    assert response.status_code == 422, response.text
    assert response.json()["error"]["details"]["rule"] == "request.locked_copy"
    assert api._current_deck(run, vid).id == deck.id and run.deck_revision.get(vid, 1) == 1
