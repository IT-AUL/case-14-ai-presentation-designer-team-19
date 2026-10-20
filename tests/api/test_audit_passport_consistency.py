"""B11: an audit upgrade and cached passport must describe the same snapshot."""

import asyncio
import json

import pytest
from deckdna.api import app as api
from fastapi.testclient import TestClient
from pptx import Presentation

A = "/api/v1"


@pytest.fixture
def generated(monkeypatch, tmp_path, wait_generation):
    monkeypatch.setattr(api, "STORE", api._Store())
    monkeypatch.setattr(api.settings, "state_snapshot", False)
    client = TestClient(api.app)
    pid = client.post(f"{A}/projects", json={"name": "audit snapshot"}).json()["id"]
    path = tmp_path / "template.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(path)
    tid = client.post(f"{A}/projects/{pid}/templates", files={
        "file": (path.name, path.read_bytes(), api._PPTX_MIME),
    }).json()["id"]
    text = b"# Report\n\n## Findings\n\n- Measured facts\n- Known limitations\n"
    cp = client.post(f"{A}/projects/{pid}/content-packs", files={
        "files": ("content.md", text, "text/markdown"),
    }).json()["content_pack"]["id"]
    accepted = client.post(f"{A}/projects/{pid}/generations", json={
        "template_id": tid, "content_pack_id": cp, "use_llm": False,
        "brief": {"purpose": "Report", "audience": "team", "language": "en",
                  "target_slide_count": 3}, "variants": [{"strategy": "balanced"}],
    }).json()
    result = wait_generation(client, accepted["generation_id"])
    assert result["state"] == "completed", result
    vid = accepted["variant_ids"][0]
    run = api.STORE.runs[result["id"]]
    audit = api.STORE.audits[api._audit_id_of(vid)]
    audit.contextual_status = "pending"
    run.variant_artifacts[vid].pop("quality_passport", None)
    cached = api._ensure_passport(run, vid)
    assert any(f["strategy"] == "pending" for f in json.loads(cached.data)["fallbacks"])
    monkeypatch.setattr(api, "_require_session", lambda *args: object())
    return client, run, vid, audit, cached


def test_completed_audit_invalidates_only_dependent_passport(generated, monkeypatch):
    client, run, vid, audit, cached = generated
    before = dict(run.variant_artifacts[vid])
    calls = []

    async def contextual(current, *args):
        calls.append(current.deck_revision)
        current.contextual_status = "completed"
        return [api.AuditIssue(
            schema_version="1.0", id="new-contextual", audit_run_id=current.id,
            deck_revision=current.deck_revision, rule_code="content.visual_relevance",
            deterministic=False, severity="warning", message="Review observation",
            status="open", repairable=False, slide_index=0,
        )]

    monkeypatch.setattr(api, "_contextual_audit_issues", contextual)
    response = client.post(f"{A}/variants/{vid}/audits",
                           json={"provider_session_id": "test-session"})
    assert response.status_code == 202, response.text
    fresh = api._ensure_passport(run, vid)
    passport = json.loads(fresh.data)
    assert fresh.id != cached.id
    assert audit.contextual_status == "completed"
    assert api._variant_or_404(vid)[1].contextual_issues == 1
    assert not any(f["strategy"] == "pending" for f in passport["fallbacks"])
    assert passport["issues_summary"]["unresolved"] == audit.issue_count
    assert api.STORE.artifacts[cached.id].data == cached.data  # history is immutable
    assert all(run.variant_artifacts[vid][k] == v for k, v in before.items()
               if k != "quality_passport")
    client.post(f"{A}/variants/{vid}/audits", json={"provider_session_id": "test-session"})
    assert calls == [audit.deck_revision]
    assert api._ensure_passport(run, vid).id == fresh.id


def test_old_audit_cannot_commit_after_edit(generated, monkeypatch):
    client, run, vid, audit, cached = generated
    revision = audit.deck_revision

    async def contextual(current, *args):
        # Emulate another worker committing a new deck while inference is awaited.
        await asyncio.sleep(0)
        with api._variant_lock(vid):
            run.deck_revision[vid] = revision + 1
            audit.deck_revision = revision + 1
        current.contextual_status = "completed"
        return []

    monkeypatch.setattr(api, "_contextual_audit_issues", contextual)
    response = client.post(f"{A}/variants/{vid}/audits",
                           json={"provider_session_id": "test-session"})
    assert response.status_code == 409, response.text
    assert audit.contextual_status == "pending"
    assert audit.deck_revision == revision + 1


def test_partial_audit_invalidates_pending_passport(generated, monkeypatch):
    client, run, vid, audit, cached = generated

    async def contextual(current, *args):
        current.contextual_status = "failed"
        return []

    monkeypatch.setattr(api, "_contextual_audit_issues", contextual)
    response = client.post(f"{A}/variants/{vid}/audits",
                           json={"provider_session_id": "test-session"})
    assert response.status_code == 202
    fresh = api._ensure_passport(run, vid)
    assert fresh.id != cached.id
    status = [f for f in json.loads(fresh.data)["fallbacks"]
              if f["feature"] == "audit.contextual.coverage"]
    assert status and status[-1]["strategy"] == "failed"


async def test_parallel_audit_does_not_duplicate_inference(generated, monkeypatch):
    import httpx

    _, run, vid, audit, _ = generated
    entered, release = asyncio.Event(), asyncio.Event()
    calls = []

    async def contextual(current, *args):
        calls.append(current.deck_revision)
        entered.set()
        await release.wait()
        current.contextual_status = "completed"
        return []

    monkeypatch.setattr(api, "_contextual_audit_issues", contextual)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=api.app),
                                 base_url="http://test") as client:
        first = asyncio.create_task(client.post(f"{A}/variants/{vid}/audits", json={
            "provider_session_id": "test-session",
        }))
        try:
            await asyncio.wait_for(entered.wait(), timeout=5)
            duplicate = await asyncio.wait_for(client.post(f"{A}/variants/{vid}/audits", json={
                "provider_session_id": "test-session",
            }), timeout=5)
            assert duplicate.status_code == 409
            assert calls == [audit.deck_revision]
        finally:
            release.set()
        assert (await first).status_code == 202
    assert audit.contextual_status == "completed"


def test_missing_audit_is_built_from_actual_deck(generated):
    client, run, vid, audit, cached = generated
    del api.STORE.audits[audit.id]
    api._patch_variant(run, vid, metrics=None)
    response = client.post(f"{A}/variants/{vid}/audits", json={})
    assert response.status_code == 202, response.text
    fresh = api.STORE.audits[response.json()["audit_id"]]
    assert fresh.id != audit.id and fresh.contextual_status == "skipped"
    assert fresh.deck_revision == run.deck_revision.get(vid, 1)
    assert fresh.issue_count == len(api.STORE.issues[fresh.id])
    assert api._variant_or_404(vid)[1].metrics.validity == 1.0
    assert api._ensure_passport(run, vid).id != cached.id


def test_author_dismissal_during_audit_is_not_overwritten(generated, monkeypatch):
    client, run, vid, audit, _ = generated
    issue = api.AuditIssue(
        schema_version="1.0", id="author-reviewed", audit_run_id=audit.id,
        deck_revision=audit.deck_revision, rule_code="density.occupancy",
        deterministic=True, severity="warning", message="Sparse slide", status="open",
        repairable=False, slide_index=0,
    )
    api.STORE.issues[audit.id].append(issue)

    async def contextual(current, *args):
        response = client.post(f"{A}/issues/{issue.id}/dismiss",
                               json={"reason": "Intentional space"})
        assert response.status_code == 200
        current.contextual_status = "completed"
        return []

    monkeypatch.setattr(api, "_contextual_audit_issues", contextual)
    response = client.post(f"{A}/variants/{vid}/audits",
                           json={"provider_session_id": "test-session"})
    assert response.status_code == 202
    saved = next(i for i in api.STORE.issues[audit.id] if i.id == issue.id)
    assert saved.status.value == "dismissed"
