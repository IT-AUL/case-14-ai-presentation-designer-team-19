"""Content Ingestion: JSON / Markdown / TXT -> ContentPack."""

import json
from pathlib import Path

import pytest
from deckdna.contracts.content_pack import Kind, Kind1
from deckdna.errors import DeckDNAError
from deckdna.ingestion import parse_file, parse_json_content, parse_markdown, parse_txt
from jsonschema import validate

SCHEMA_PATH = Path(__file__).resolve().parents[2] / "schemas" / "content-pack.schema.json"
SCHEMA = json.loads(SCHEMA_PATH.read_text())

VALID_JSON = {
    "schema_version": "1.0",
    "id": "pack-demo",
    "language": "ru",
    "title_hint": "Демо-продукт",
    "sections": [
        {
            "id": "s1",
            "heading": "Проблема",
            "level": 1,
            "blocks": [
                {
                    "kind": "paragraph",
                    "text": "Ручная сборка колод занимает часы.",
                    "source_ref": {"artifact_id": "src.docx"},
                },
                {
                    "kind": "list",
                    "items": ["3 часа на колоду", "ошибки вёрстки"],
                    "source_ref": {"artifact_id": "src.docx"},
                },
            ],
        }
    ],
    "tables": [],
    "assets": [],
    "warnings": [],
}


class TestJsonParser:
    def test_valid(self):
        pack = parse_json_content(VALID_JSON)
        assert pack.id == "pack-demo"
        assert pack.sections[0].blocks[1].kind is Kind.list
        validate(instance=json.loads(pack.model_dump_json(exclude_none=True)), schema=SCHEMA)

    def test_str_input(self):
        pack = parse_json_content(json.dumps(VALID_JSON))
        assert pack.id == "pack-demo"

    def test_broken_json(self):
        with pytest.raises(DeckDNAError, match="not valid JSON") as exc:
            parse_json_content("{nope")
        assert exc.value.code == "invalid_input"

    def test_schema_violation(self):
        with pytest.raises(DeckDNAError, match="content-pack schema"):
            parse_json_content({"id": "x"})


MARKDOWN = """# Демо-дек

Вводный параграф до первого подзаголовка.

## Проблема

Текст проблемы. Ещё предложение.

- пункт один
- пункт два

> Цитата эксперта

## Метрики

```python
print(1)
```

![Схема](img.png)

### Подробности

Детальный текст.
"""


class TestMarkdownParser:
    def test_sections_and_blocks(self):
        pack = parse_markdown(MARKDOWN, artifact_id="demo.md")
        validate(instance=json.loads(pack.model_dump_json(exclude_none=True)), schema=SCHEMA)
        assert pack.title_hint == "Демо-дек"
        assert pack.language == "ru"
        # sec-0 = "# Демо-дек" с преамбулой; далее Проблема, Метрики, Подробности
        headings = [s.heading for s in pack.sections]
        assert headings == ["Демо-дек", "Проблема", "Метрики", "Подробности"]
        problem = pack.sections[1]
        kinds = [b.kind for b in problem.blocks]
        assert kinds == [Kind.paragraph, Kind.list, Kind.quote]
        assert problem.blocks[1].items == ["пункт один", "пункт два"]
        assert problem.blocks[0].source_ref.heading_path == ["Демо-дек", "Проблема"]

    def test_images_become_assets(self):
        pack = parse_markdown(
            "# Док\n\n![Логотип ВК](pics/logo.png)\n\n![](https://cdn.example.com/x.png)\n",
            artifact_id="doc.md",
        )
        assert len(pack.assets) == 2
        logo = pack.assets[0]
        assert logo.kind is Kind1.image
        assert logo.alt == "Логотип ВК"
        assert logo.artifact_id == "pics/logo.png"
        assert logo.source_url is None
        remote = pack.assets[1]
        assert remote.source_url == "https://cdn.example.com/x.png"
        # image_ref-блок в потоке остаётся как и раньше
        assert any(b.kind is Kind.image_ref for b in pack.sections[0].blocks)

    def test_code_and_image(self):
        pack = parse_markdown(MARKDOWN, artifact_id="demo.md")
        metrics = pack.sections[2]
        assert metrics.blocks[0].kind is Kind.code
        assert metrics.blocks[1].kind is Kind.image_ref
        assert metrics.blocks[1].text == "Схема"

    def test_deterministic_ids(self):
        assert parse_markdown(MARKDOWN).id == parse_markdown(MARKDOWN).id


MARKDOWN_WITH_TABLE = """## Показатели

Текст до таблицы.

| Метрика | Значение | Комментарий |
|---|---|---|
| Выручка | 12.5 | рост |
| Маржа | 30% | |
| | пусто | тест |

Текст после таблицы.
"""


class TestMarkdownTables:
    def test_table_extracted(self):
        pack = parse_markdown(MARKDOWN_WITH_TABLE, artifact_id="tbl.md")
        validate(instance=json.loads(pack.model_dump_json(exclude_none=True)), schema=SCHEMA)
        assert len(pack.tables) == 1
        table = pack.tables[0]
        assert table.headers == ["Метрика", "Значение", "Комментарий"]
        assert table.rows == [
            ["Выручка", 12.5, "рост"],
            ["Маржа", "30%", None],
            [None, "пусто", "тест"],
        ]
        assert table.source_ref.heading_path == ["Показатели"]

    def test_surrounding_content_kept(self):
        pack = parse_markdown(MARKDOWN_WITH_TABLE, artifact_id="tbl.md")
        texts = [b.text for s in pack.sections for b in s.blocks if b.kind is Kind.paragraph]
        assert "Текст до таблицы." in texts
        assert "Текст после таблицы." in texts
        # таблица не остаётся текстовым мусором в блоках
        assert not any("---" in (t or "") or "Выручка" in (t or "") for t in texts)


def _make_docx(tmp_path, with_image: bool = False) -> Path:
    import docx

    document = docx.Document()
    document.add_heading("Отчёт за квартал", 1)
    document.add_paragraph("Вводный текст отчёта.")
    document.add_heading("Показатели", 2)
    document.add_paragraph("пункт первый", style="List Bullet")
    document.add_paragraph("пункт второй", style="List Bullet")
    table = document.add_table(rows=3, cols=2)
    for i, values in enumerate([("Метрика", "Значение"), ("Выручка", "12.5"), ("Маржа", "30%")]):
        for j, value in enumerate(values):
            table.rows[i].cells[j].text = value
    if with_image:
        from io import BytesIO

        from PIL import Image

        img = BytesIO()
        Image.new("RGB", (10, 8), "red").save(img, "PNG")
        img.seek(0)
        document.add_picture(img)
    document.add_paragraph("Заключение.")
    path = tmp_path / "report.docx"
    document.save(path)
    return path


class TestDocxParser:
    def test_sections_lists_tables(self, tmp_path):
        from deckdna.ingestion.content_parsers import parse_docx

        pack = parse_docx(_make_docx(tmp_path).read_bytes(), "report.docx")
        validate(instance=json.loads(pack.model_dump_json(exclude_none=True)), schema=SCHEMA)
        assert pack.title_hint == "Отчёт за квартал"
        assert pack.language == "ru"
        headings = [s.heading for s in pack.sections]
        assert headings == ["Отчёт за квартал", "Показатели"]
        first = pack.sections[0]
        assert first.blocks[0].kind is Kind.paragraph
        assert first.blocks[0].text == "Вводный текст отчёта."
        second = pack.sections[1]
        assert second.blocks[0].kind is Kind.list
        assert second.blocks[0].items == ["пункт первый", "пункт второй"]
        assert second.blocks[0].source_ref.heading_path == ["Отчёт за квартал", "Показатели"]

        assert len(pack.tables) == 1
        table = pack.tables[0]
        assert table.headers == ["Метрика", "Значение"]
        assert table.rows == [["Выручка", 12.5], ["Маржа", "30%"]]
        assert table.source_ref.heading_path == ["Отчёт за квартал", "Показатели"]

    def test_dispatch_docx(self, tmp_path):
        pack = parse_file(_make_docx(tmp_path))
        assert pack.title_hint == "Отчёт за квартал"
        assert len(pack.tables) == 1

    def test_images_become_assets(self, tmp_path):
        pack = parse_file(_make_docx(tmp_path, with_image=True))
        validate(instance=json.loads(pack.model_dump_json(exclude_none=True)), schema=SCHEMA)
        assert len(pack.assets) == 1
        asset = pack.assets[0]
        assert asset.kind is Kind1.image
        assert asset.artifact_id == "report.docx#word/media/image1.png"
        assert asset.width_px == 10
        assert asset.height_px == 8
        # docPr descr не задан в сгенерированном документе -> alt None
        assert asset.alt is None

    def test_images_emit_image_ref_blocks(self, tmp_path):
        """docx-картинка -> image_ref-блок в потоке в точке inline-шейпа
        (мост до Kind.image-юнита плана, как у markdown)."""
        pack = parse_file(_make_docx(tmp_path, with_image=True))
        validate(instance=json.loads(pack.model_dump_json(exclude_none=True)), schema=SCHEMA)
        refs = [
            b for s in pack.sections for b in s.blocks if b.kind is Kind.image_ref
        ]
        assert len(refs) == 1
        assert refs[0].text == "report.docx#word/media/image1.png"
        # позиция в потоке: между абзацем ввода и заключением
        kinds = [b.kind for s in pack.sections for b in s.blocks]
        assert kinds.index(Kind.image_ref) > kinds.index(Kind.paragraph)

    def test_docx_image_reaches_plan_unit(self, tmp_path):
        """image_ref -> ContentUnit(kind=image, asset_ref) в plan_deck."""
        pack = parse_file(_make_docx(tmp_path, with_image=True))
        from deckdna.generation.pipeline import Brief
        from deckdna.planning.story_director import plan_deck

        brief = Brief(
            purpose="Отчёт", audience="жюри", language="ru",
            target_slide_count=10,
        )
        plan = plan_deck(pack, brief)
        units = [
            u for s in plan.slides for u in s.content_units if u.kind.name == "image"
        ]
        assert units and units[0].asset_ref == "report.docx#word/media/image1.png"


def _make_pdf(tmp_path) -> Path:
    """Реальный .pdf: текст + разлинованная таблица (pdfplumber ловит её
    по линиям сетки). Кириллица через DejaVuSans; если TTF нет на
    машине — английский текст (assert'ы ниже это учитывают)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font = "Helvetica"
    cyrillic = False
    for candidate in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/DejaVuSans.ttf",
    ):
        if Path(candidate).exists():
            pdfmetrics.registerFont(TTFont("DejaVu", candidate))
            font = "DejaVu"
            cyrillic = True
            break

    title = "Отчёт за квартал" if cyrillic else "Quarterly report"
    intro = "Вводный текст отчёта." if cyrillic else "Intro paragraph."
    tail = "Заключение второй страницы." if cyrillic else "Second page conclusion."
    headers = ["Метрика", "Значение"] if cyrillic else ["Metric", "Value"]
    rows = (
        [["Выручка", "12.5"], ["Маржа", "30%"]]
        if cyrillic
        else [["Revenue", "12.5"], ["Margin", "30%"]]
    )

    path = tmp_path / "report.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont(font, 14)
    c.drawString(72, 780, title)
    c.setFont(font, 10)
    c.drawString(72, 760, intro)
    # сетка таблицы 3x2 — линии + текст в ячейках
    x0, y0, cw, rh = 72, 640, 120, 20
    for i, row in enumerate([headers] + rows):
        for j, value in enumerate(row):
            c.rect(x0 + j * cw, y0 - i * rh, cw, rh)
            c.drawString(x0 + j * cw + 4, y0 - i * rh + 5, value)
    c.showPage()
    c.setFont(font, 10)
    c.drawString(72, 780, tail)
    c.showPage()
    c.save()
    return path


class TestPdfParser:
    def test_sections_and_tables(self, tmp_path):
        from deckdna.ingestion.content_parsers import parse_pdf

        pack = parse_pdf(_make_pdf(tmp_path).read_bytes(), "report.pdf")
        validate(instance=json.loads(pack.model_dump_json(exclude_none=True)), schema=SCHEMA)
        # крупная строка — заголовок раздела; текст обеих страниц — в нём
        assert pack.title_hint == pack.sections[0].heading
        assert pack.sections[0].heading not in ("", "Page 1")
        first = pack.sections[0].blocks[0]
        assert first.source_ref.heading_path == [pack.sections[0].heading]
        texts = " ".join(
            (b.text or "") + " ".join(b.items or []) for s in pack.sections for b in s.blocks
        )
        assert "12.5" not in texts, "table cells are not duplicated as prose"

        assert len(pack.tables) == 1
        table = pack.tables[0]
        assert len(table.headers) == 2
        assert table.rows[0][1] == 12.5  # число -> float
        assert table.source_ref.heading_path == ["Page 1"]
        # table_ref-блок связывает таблицу с потоком блоков
        assert any(b.kind is Kind.table_ref and b.text == table.id for b in pack.sections[0].blocks)

    def test_dispatch_pdf(self, tmp_path):
        pack = parse_file(_make_pdf(tmp_path))
        assert pack.sections and pack.sections[0].heading
        assert len(pack.tables) == 1

    def test_empty_pdf_keeps_one_section(self, tmp_path):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        path = tmp_path / "blank.pdf"
        c = canvas.Canvas(str(path), pagesize=A4)
        c.showPage()
        c.save()
        pack = parse_file(path)
        assert len(pack.sections) == 1
        assert pack.tables == []


def _make_xlsx(tmp_path) -> Path:
    """Реальный .xlsx через openpyxl: два листа с данными, пустой лист
    и лист из одной строки заголовков без данных."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.append(["Метрика", "Значение"])
    ws.append(["Выручка", 12.5])
    ws.append(["Маржа", "30%"])

    wb.create_sheet("Пустой")  # без данных — должен быть пропущен

    solo = wb.create_sheet("Одна строка")
    solo.append(["Только заголовки", "без данных"])  # <2 строк — пропуск

    sales = wb.create_sheet("Продажи")
    sales.append([])  # ведущая пустая строка — не заголовок
    sales.append(["Регион", "Сумма"])
    sales.append(["Север", 7])
    sales.append([])  # пустая строка внутри — отфильтровывается
    sales.append(["Юг", 3.25])

    path = tmp_path / "report.xlsx"
    wb.save(path)
    return path


class TestXlsxParser:
    def test_sheets_to_tables(self, tmp_path):
        from deckdna.ingestion.content_parsers import parse_xlsx

        pack = parse_xlsx(_make_xlsx(tmp_path).read_bytes(), "report.xlsx")
        validate(instance=json.loads(pack.model_dump_json(exclude_none=True)), schema=SCHEMA)
        assert pack.language == "ru"
        # два листа с данными -> две таблицы и две секции
        assert len(pack.tables) == 2
        assert len(pack.sections) == 2
        assert [s.heading for s in pack.sections] == ["Отчёт", "Продажи"]

        first = pack.tables[0]
        assert first.headers == ["Метрика", "Значение"]
        assert first.rows == [["Выручка", 12.5], ["Маржа", "30%"]]
        assert first.source_ref.heading_path == ["Отчёт"]

        second = pack.tables[1]
        assert second.headers == ["Регион", "Сумма"]
        # внутренняя пустая строка отфильтрована; 7 -> 7.0 (число -> float)
        assert second.rows == [["Север", 7.0], ["Юг", 3.25]]

        # table_ref-блок связывает каждую таблицу с её секцией
        for section, table in zip(pack.sections, pack.tables, strict=True):
            assert any(b.kind is Kind.table_ref and b.text == table.id for b in section.blocks)

    def test_dispatch_xlsx(self, tmp_path):
        pack = parse_file(_make_xlsx(tmp_path))
        assert len(pack.tables) == 2

    def test_empty_workbook_keeps_one_section(self, tmp_path):
        from openpyxl import Workbook

        path = tmp_path / "empty.xlsx"
        Workbook().save(path)
        pack = parse_file(path)
        assert len(pack.sections) == 1
        assert pack.tables == []


class TestTxtParser:
    def test_single_paragraph(self):
        pack = parse_txt("строка раз\nстрока два", artifact_id="note.txt")
        assert len(pack.sections) == 1
        assert len(pack.sections[0].blocks) == 1
        assert pack.sections[0].blocks[0].kind is Kind.paragraph
        assert pack.language == "ru"

    def test_language_detect_en(self):
        pack = parse_txt("plain english text")
        assert pack.language == "en"


class TestParseFile:
    def test_dispatch_md(self, tmp_path):
        p = tmp_path / "doc.md"
        p.write_text("# Заголовок\n\nтекст", encoding="utf-8")
        pack = parse_file(p)
        assert pack.title_hint == "Заголовок"
        assert pack.sections[0].blocks[0].source_ref.artifact_id == "doc.md"

    def test_dispatch_json(self, tmp_path):
        p = tmp_path / "doc.json"
        p.write_text(json.dumps(VALID_JSON), encoding="utf-8")
        assert parse_file(p).id == "pack-demo"

    def test_unsupported(self, tmp_path):
        p = tmp_path / "doc.xyz"
        p.write_bytes(b"garbage")
        with pytest.raises(DeckDNAError, match="unsupported content format"):
            parse_file(p)


class TestGarbageCssIsDropped:
    """Найдено вживую: пользователь вставил текст из HTML-письма (VK,
    шрифт MailSans), стилевой блок дошёл как обычный markdown-параграф и
    попал в презентацию как контент ("body {\\npadding: 0;..."). parse_file
    должен отфильтровывать такие блоки, а не отдавать их как контент."""

    CSS_BLOB = (
        "body {\n"
        "padding: 0;\n"
        "margin: 0;\n"
        "height: 100vh;\n"
        "overflow: hidden;\n"
        "font-family: 'MailSans Roman', sans-serif;\n"
        "}"
    )

    def test_css_paragraph_in_markdown_is_dropped_with_warning(self, tmp_path):
        p = tmp_path / "doc.md"
        p.write_text(
            f"# Заголовок\n\n{self.CSS_BLOB}\n\nНормальный текст про продукт.",
            encoding="utf-8",
        )
        pack = parse_file(p)
        texts = [b.text for b in pack.sections[0].blocks]
        assert "Нормальный текст про продукт." in texts
        assert not any(t and "padding" in t for t in texts)
        assert any(w.code == "content.garbage_text_dropped" for w in pack.warnings)

    def test_css_paragraph_in_txt_is_dropped(self, tmp_path):
        """parse_txt always emits exactly one paragraph block — when that
        block IS the garbage, nothing real is left: the section collapses
        away entirely rather than surviving as an empty husk."""
        p = tmp_path / "doc.txt"
        p.write_text(self.CSS_BLOB, encoding="utf-8")
        pack = parse_file(p)
        assert pack.sections == []
        assert any(w.code == "content.garbage_text_dropped" for w in pack.warnings)

    def test_css_list_item_is_dropped_but_siblings_kept(self, tmp_path):
        css_item = "body { padding: 0; margin: 0; overflow: hidden; }"
        p = tmp_path / "doc.md"
        p.write_text(
            f"- пункт первый\n- {css_item}\n- пункт третий",
            encoding="utf-8",
        )
        pack = parse_file(p)
        list_block = next(b for b in pack.sections[0].blocks if b.kind is Kind.list)
        assert "пункт первый" in list_block.items
        assert "пункт третий" in list_block.items
        assert not any("padding" in item for item in list_block.items)

    def test_normal_text_with_colons_is_not_flagged(self, tmp_path):
        p = tmp_path / "doc.md"
        text = "Итог: выручка выросла. Вывод: план выполнен. Дальше: масштабирование."
        p.write_text(text, encoding="utf-8")
        pack = parse_file(p)
        assert pack.sections[0].blocks[0].text == text
        assert not pack.warnings


def test_display_equations_survive_markdown_escapes_and_underscores():
    from deckdna.ingestion.content_parsers import parse_markdown

    source = r"""# Measurement
## Equation
\[
Index = 100 \times \sigma(0.25z_{active} - 0.75z_{idle})
\]

$$
E = m c^2
$$
"""
    pack = parse_markdown(source)
    texts = [block.text for section in pack.sections for block in section.blocks]
    assert "Index = 100 × σ(0.25z_active - 0.75z_idle)" in texts
    assert "E = m c^2" in texts


def test_pdf_sections_tidy_caps_cover_and_repeated_subheadings():
    from deckdna.contracts.content_pack import Block, Kind, Section, SourceRef
    from deckdna.ingestion.content_parsers import _tidy_pdf_sections

    ref = SourceRef(artifact_id="a.pdf")

    def sec(heading, text=None):
        blocks = [Block(kind=Kind.paragraph, text=text, source_ref=ref)] if text else []
        return Section(id="x", heading=heading, level=2, blocks=blocks)

    sections = [
        sec("ПУБЛИЧНЫЕ ВЫСТУПЛЕНИЯ"),
        sec("КАК ВЫСТРОИТЬ АРГУМЕНТАЦИЮ. ПИРАМИДА МИНТО"),
        sec("ВАРИАНТ 1. ПИРАМИДА МИНТО", "Структура по методу Минто и SCQA."),
        sec("Назначение", "Убедить руководителя."),
        sec("ВАРИАНТ 2. ИСТОРИЯ", "Личный опыт вовлекает."),
        sec("Назначение", "Вовлечь зал."),
        sec("ОШИБКА", "Начинать с деталей."),
    ]
    body = "Структура по методу Минто и SCQA. Убедить руководителя."
    out, title = _tidy_pdf_sections(sections, None, body)
    assert title == "Как выстроить аргументацию. Пирамида Минто"
    assert [s.heading for s in out] == ["Вариант 1. Пирамида Минто", "Вариант 2. История"]
    assert out[0].blocks[1].text == "Назначение: Убедить руководителя."
    assert [b.text for b in out[1].blocks][1:] == [
        "Назначение: Вовлечь зал.",
        "Ошибка: Начинать с деталей.",
    ]
