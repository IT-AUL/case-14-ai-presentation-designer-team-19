"""Quality Passport: реальный прогон compose -> audit -> PEI -> assemble.

Собираем паспорт на настоящем multi-slide выводе generate_deck на
organizer-шаблоне и проверяем его настоящей схемой через canonical
serialize-хелпер.
"""

import hashlib
import json
import zipfile
from pathlib import Path

import pytest
from deckdna.audit.basic import audit_deck
from deckdna.contracts import to_schema_dict
from deckdna.evaluation import assemble_quality_passport
from deckdna.pptx.composing.minimal import generate_deck
from jsonschema import validate
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

FIXTURE = Path("tests/fixtures/pptx/vk_tech_template.pptx")
DECK_PLAN = Path("tests/fixtures/content/poc_deck_plan.json")
SCHEMA = json.loads(
    (Path(__file__).resolve().parents[2] / "schemas" / "quality-passport.schema.json").read_text()
)


@pytest.fixture(scope="module")
def passport(tmp_path_factory):
    if not FIXTURE.exists():
        pytest.skip("organizer fixture not present")
    out = tmp_path_factory.mktemp("qp") / "out.pptx"
    plan = json.loads(DECK_PLAN.read_text(encoding="utf-8"))
    report = generate_deck(FIXTURE, plan, out)
    issues = audit_deck(out)
    return (
        assemble_quality_passport(
            out, report, issues, content_pack_id="pack-poc", duration_seconds=12.5
        ),
        report,
        issues,
        out,
    )


def test_schema_valid(passport):
    qp, *_ = passport
    validate(instance=to_schema_dict(qp), schema=SCHEMA)


def test_honest_fields_computed(passport):
    qp, report, issues, out = passport
    m = qp.metrics
    assert m.validity.opens_cleanly is True
    assert m.editability.pei_level is not None and 0 <= m.editability.pei_level <= 5
    assert m.editability.raster_only_slides == sum(
        1 for i in issues if i.rule_code == "editability.raster_only"
    )
    assert m.editability.native_text_ratio is not None
    assert 0.0 <= m.editability.native_text_ratio <= 1.0
    assert m.readability.overflow_count == sum(
        1 for i in issues if i.rule_code == "text.overflow"
    )
    assert m.content_support.supported_claims == report["totals"]["runs_replaced"]
    assert m.timings.total_seconds == 12.5


def test_unknown_fields_stay_none(passport):
    """Ничего не выдумано: без текста источника числа не сверяются (None),
    без модели usage — честные нули; измеримое по колоде — посчитано."""
    qp, *_ = passport
    assert qp.metrics.content_support.numbers_verified is None
    assert qp.metrics.content_support.numbers_failed is None
    assert not any(
        f.strategy == "numeric_presence_only" for f in qp.fallbacks or []
    )
    assert qp.metrics.usage.model_calls == 0 and qp.metrics.usage.total_tokens == 0
    fid = qp.metrics.style_fidelity
    assert fid is not None and all(
        0 <= v <= 1 for v in fid.model_dump().values()
    )
    assert 0 < qp.metrics.readability.avg_occupancy <= 1
    assert qp.fallbacks is None


def test_numeric_presence_is_disclosed_as_limited_metric(passport):
    _, report, issues, out = passport
    qp = assemble_quality_passport(
        out,
        report,
        issues,
        content_pack_id="pack-poc",
        source_text="Доля возвратов составила 30%. Выручка снизилась.",
    )
    assert qp.metrics.content_support.numbers_verified is not None
    disclosure = next(
        f for f in qp.fallbacks or [] if f.strategy == "numeric_presence_only"
    )
    assert disclosure.feature == "metrics.content_support.numbers_verified"
    assert "do not verify the associated claim" in disclosure.disclosure
    validate(instance=to_schema_dict(qp), schema=SCHEMA)


def test_validity_real_round_trip_and_ooxml(passport):
    """validity: реальный round-trip на сгенерированной колоде и
    структурный sanity — на чистом выводе ooxml_errors == 0."""
    qp, _, issues, _ = passport
    v = qp.metrics.validity
    assert v.opens_cleanly is True
    assert v.round_trip_ok is True
    assert v.ooxml_errors == 0
    assert qp.metrics.readability.contrast_failures == sum(
        1 for i in issues if i.rule_code == "accessibility.contrast"
    )


def test_validity_detects_broken_package(tmp_path):
    """Seeded defect: висячий rel и недекларированная часть →
    ooxml_errors > 0, round_trip отдельно остаётся честным."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "текст"
    good = tmp_path / "good.pptx"
    prs.save(good)

    broken = tmp_path / "broken.pptx"
    with zipfile.ZipFile(good) as zin, zipfile.ZipFile(broken, "w") as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == "ppt/_rels/presentation.xml.rels":
                root = etree.fromstring(data)
                rel = root[0]
                rel.set("Target", "does/not/exist.bin")
                data = etree.tostring(root)
            if item != "[Content_Types].xml":
                zout.writestr(item, data)
            else:
                root = etree.fromstring(data)
                for el in list(root):
                    if el.get("PartName", "").endswith("slideMaster1.xml"):
                        root.remove(el)
                zout.writestr(item, etree.tostring(root))

    report = {"slides": [{}], "template": str(good), "totals": {}}
    qp = assemble_quality_passport(broken, report, [])
    v = qp.metrics.validity
    assert v.ooxml_errors and v.ooxml_errors >= 1

    notzip = tmp_path / "notzip.pptx"
    notzip.write_bytes(b"not a zip")
    qp2 = assemble_quality_passport(
        notzip, {"slides": [], "template": str(good), "totals": {}}, []
    )
    assert qp2.metrics.validity.round_trip_ok is False
    assert qp2.metrics.validity.ooxml_errors >= 1


def test_issues_summary_and_exports(passport):
    qp, _, issues, out = passport
    s = qp.issues_summary
    assert s.unresolved == len(issues)
    assert (s.error or 0) + (s.warning or 0) + (s.blocker or 0) + (s.info or 0) == len(issues)
    assert len(qp.exports) == 1
    assert qp.exports[0].sha256 == hashlib.sha256(out.read_bytes()).hexdigest()
    assert qp.exports[0].format.value == "pptx"


def test_inputs(passport):
    qp, *_ = passport
    assert qp.inputs.template_name == FIXTURE.name
    assert qp.inputs.template_sha256 == hashlib.sha256(FIXTURE.read_bytes()).hexdigest()
    assert qp.inputs.content_pack_id == "pack-poc"
    assert qp.inputs.brief_hash


def test_dropped_units_surfaced(tmp_path):
    """dropped_units из compose_report видны в паспорте, а не теряются:
    сумма по kind плюсуется в unsupported_claims (контент запрошен,
    в вывод не попал), per-kind детализация — в fallbacks с disclosure."""
    out = tmp_path / "out.pptx"
    plan = json.loads(DECK_PLAN.read_text(encoding="utf-8"))
    plan["slides"][0]["content_units"].extend(
        [
            {"role": "body", "kind": "table", "text": "2x2 матрица рисков"},
            {"role": "body", "kind": "image", "text": "схема пайплайна"},
        ]
    )
    report = generate_deck(FIXTURE, plan, out)
    issues = audit_deck(out)
    qp = assemble_quality_passport(out, report, issues, content_pack_id="pack-poc")

    dropped = report["dropped_units"]
    # table/image are structural kinds this exemplar has no slot for at
    # all — always dropped, deterministically. text_unplaced (bullets that
    # lost their slot once exemplar selection also had to satisfy the new
    # table/image needs) is a budget-fit side effect of *which* exemplar
    # gets picked, not what this test exercises — asserted loosely.
    assert dropped["table"] == 1
    assert dropped["image"] == 1
    dropped_total = sum(dropped.values())
    expected_unsupported = report["totals"].get("runs_skipped", 0) + dropped_total
    assert qp.metrics.content_support.unsupported_claims == expected_unsupported
    assert qp.fallbacks is not None
    by_feature = {f.feature: f for f in qp.fallbacks}
    assert by_feature["composing.table_units"].strategy == "drop"
    assert "1 table unit(s)" in by_feature["composing.table_units"].disclosure
    assert "1 image unit(s)" in by_feature["composing.image_units"].disclosure
    validate(instance=to_schema_dict(qp), schema=SCHEMA)


def test_repair_fixes_disclose_stale_content_support(tmp_path):
    """When repair_fixes is non-empty (this passport
    describes a post-repair revision), original compose counts must be
    absent from current-revision metrics until they can be recomputed."""
    out = tmp_path / "out.pptx"
    plan = json.loads(DECK_PLAN.read_text(encoding="utf-8"))
    report = generate_deck(FIXTURE, plan, out)
    issues = audit_deck(out)

    without_repair = assemble_quality_passport(
        out, report, issues, content_pack_id="pack-poc"
    )
    assert not any(
        f.feature == "metrics.content_support"
        for f in without_repair.fallbacks or []
    )
    assert without_repair.metrics.content_support.supported_claims is not None

    with_repair = assemble_quality_passport(
        out,
        report,
        issues,
        content_pack_id="pack-poc",
        repair_fixes={"layout.out_of_bounds": 1},
    )
    by_feature = {f.feature: f for f in with_repair.fallbacks or []}
    fb = by_feature["metrics.content_support"]
    assert fb.strategy == "not_recomputed"
    assert "repair" in fb.disclosure
    assert with_repair.metrics.content_support.supported_claims is None
    assert with_repair.metrics.content_support.unsupported_claims is None
    serialized = to_schema_dict(with_repair)["metrics"]["content_support"]
    assert "supported_claims" not in serialized
    assert "unsupported_claims" not in serialized
    validate(instance=to_schema_dict(with_repair), schema=SCHEMA)


def test_placeholder_residual_surfaced_as_fallback(tmp_path):
    """Остаточный integrity.placeholder_text — сигнал contentless-слотов:
    одна агрегированная fallback-запись (strategy='retain'), но НЕ
    плюсуется в unsupported_claims — там знаменатель юниты плана,
    а тут фигуры, пересечение с runs_skipped непроверяемо."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for y in (1, 3):
        box = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(5), Inches(1))
        box.text_frame.text = "Заголовок"
    out = tmp_path / "stock.pptx"
    prs.save(out)

    issues = audit_deck(out)
    placeholders = [i for i in issues if i.rule_code == "integrity.placeholder_text"]
    assert len(placeholders) == 2

    report = {
        "template": str(out),
        "totals": {"runs_replaced": 4, "runs_skipped": 1},
        "dropped_units": {},
        "deck_plan_id": "synthetic",
    }
    qp = assemble_quality_passport(out, report, issues)

    assert qp.metrics.content_support.supported_claims == 4
    assert qp.metrics.content_support.unsupported_claims == 1  # не раздуто
    by_feature = {f.feature: f for f in qp.fallbacks or []}
    fb = by_feature["audit.integrity.placeholder_text"]
    assert fb.strategy == "retain"
    assert "2 shape(s)" in fb.disclosure
    validate(instance=to_schema_dict(qp), schema=SCHEMA)


def test_provenance_real_skill_version_and_config_hashes(passport):
    """OR-006: паспорт ссылается на реальные версии, использованные прогоном.

    skill_version — version из skill/manifest.yaml; config_hashes —
    sha256 реально загруженных конфигов (generation + audit). LLM-вызовов
    в пайплайне нет — prompt_versions/model_profiles честно пусты.
    """
    import yaml

    qp, *_ = passport
    prov = qp.provenance

    manifest = yaml.safe_load(
        (Path("skill") / "manifest.yaml").read_text(encoding="utf-8")
    )
    assert prov.skill_version == str(manifest["version"])

    assert prov.prompt_versions == {}
    assert prov.model_profiles == []

    expected = {
        "configs/generation.default.yaml",
        "configs/audit.default.yaml",
    }
    assert set(prov.config_hashes) == expected
    for rel, digest in prov.config_hashes.items():
        actual = hashlib.sha256(Path(rel).read_bytes()).hexdigest()
        assert digest == actual

    serialized = to_schema_dict(qp)["provenance"]
    assert serialized["skill_version"] == prov.skill_version
    assert serialized["config_hashes"] == prov.config_hashes
    validate(instance=to_schema_dict(qp), schema=SCHEMA)
