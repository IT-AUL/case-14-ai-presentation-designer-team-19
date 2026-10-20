import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

import json  # noqa: E402
import threading  # noqa: E402
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer  # noqa: E402

import pytest  # noqa: E402


class LocalOpenAIServer:
    """Локальный OpenAI-compatible stub на реальном TCP (stdlib http.server).

    Отвечает на `POST {base_url}/chat/completions` телом вида
    `{"choices": [{"message": {"content": <json>}}]}`, где `<json>` —
    дамп `responses[name]`, а `name` — `response_format.json_schema.name`
    из пришедшего запроса (`DeckPlan` для storyline,
    `SlideChecksResult` для slide_checks). Каждый запрос (path,
    Authorization-заголовок, распарсенное тело) записывается в
    `.requests` — тесты читают её для assertions о модели, сообщениях
    и image_url data-URL'ах.
    """

    def __init__(self) -> None:
        self.responses: dict[str, dict] = {}
        self.requests: list[dict] = []
        outer = self

        class _Handler(BaseHTTPRequestHandler):
            def do_POST(self) -> None:  # noqa: N802 (http.server API)
                length = int(self.headers.get("Content-Length") or 0)
                raw = self.rfile.read(length)
                try:
                    body = json.loads(raw)
                except json.JSONDecodeError:
                    body = {}
                outer.requests.append(
                    {
                        "method": "POST",
                        "path": self.path,
                        "authorization": self.headers.get("Authorization"),
                        "json": body,
                    }
                )
                name = (
                    body.get("response_format", {})
                    .get("json_schema", {})
                    .get("name", "")
                )
                if self.path != "/chat/completions" or name not in outer.responses:
                    self.send_response(404)
                    self.send_header("Content-Length", "0")
                    self.end_headers()
                    return
                payload = {
                    "id": "chatcmpl-local-stub",
                    "object": "chat.completion",
                    "choices": [
                        {
                            "index": 0,
                            "message": {
                                "role": "assistant",
                                "content": json.dumps(
                                    outer.responses[name], ensure_ascii=False
                                ),
                            },
                            "finish_reason": "stop",
                        }
                    ],
                }
                data = json.dumps(payload).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            def log_message(self, *args) -> None:  # тишина в тестах
                pass

        self._server = ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
        self._server.daemon_threads = True
        self._thread = threading.Thread(
            target=self._server.serve_forever, daemon=True
        )
        self._thread.start()

    @property
    def base_url(self) -> str:
        return f"http://127.0.0.1:{self._server.server_port}"

    def close(self) -> None:
        self._server.shutdown()
        self._thread.join(timeout=5)


_DECKDNA_PROVIDER_ENV_KEYS = (
    "DECKDNA_MOCK_PROVIDER",
    "DECKDNA_PROVIDER_BASE_URL",
    "DECKDNA_PROVIDER_API_KEY",
    "DECKDNA_FALLBACK_PROVIDER_BASE_URL",
    "DECKDNA_FALLBACK_PROVIDER_API_KEY",
    "DECKDNA_FALLBACK_MODEL_TEXT",
    "DECKDNA_FALLBACK_MODEL_VISION",
    "DECKDNA_RESEARCH_SEARCH_URL",
    "DECKDNA_MODEL_TEXT",
    "DECKDNA_MODEL_VISION",
    "DECKDNA_MODEL_EMBED",
    "DECKDNA_MODEL_IMAGE",
)


def pytest_configure(config: pytest.Config) -> None:
    """Tests must never depend on the ambient .env's real provider
    credentials. Found live 28.09: this repo's own local .env has real
    Cerebras credentials (mock_provider=false) for the live UI session
    -- with no isolation, ``Settings()`` (pydantic-settings, env_file=
    ".env") picks them up in every test run on this machine, including
    a FRESH ``Settings(mock_provider=False)`` construction that only
    overrides one field and expects the rest to be genuinely empty
    (test_mock.py::test_build_gateway_requires_config_when_flag_off --
    documented all session as a "known pre-existing false positive",
    turns out to be this exact bug, not unrelated). A test asserting
    "no LLM used by default" (ADR-018's test_analyze_default_leaves_
    content_description_empty) failed the same way, and a module-scoped
    fixture elsewhere (test_frontend_handoff.py's ``shared``) that
    relies on "auto" gateway resolution started timing out against a
    real endpoint once the ambient .env stopped being silently broken.

    A per-test (autouse) fixture is the wrong tool here: pytest sets up
    broader-scoped fixtures (module/session) before function-scoped
    ones for the same test, so a function-scoped monkeypatch never
    takes effect before a module-scoped fixture's own setup code runs.
    A ``pytest_configure`` hook runs once, unconditionally, before any
    fixture or test collection -- no scope-ordering hazard.

    Real env vars take precedence over the .env file in pydantic-
    settings' default source order, so setting them here (not just the
    already-constructed `settings` singleton) neutralizes BOTH the
    singleton AND any fresh `Settings()` a test constructs. Individual
    tests that need a specific provider config already do so explicitly
    (local_openai stub + explicit Settings(...) kwargs, or
    DECKDNA_VK_LIVE=1-gated live smoke tests reading their own separate
    env vars) and are unaffected."""
    import os

    for key in _DECKDNA_PROVIDER_ENV_KEYS:
        os.environ[key] = "true" if key == "DECKDNA_MOCK_PROVIDER" else ""

    from deckdna.settings import settings

    settings.mock_provider = True
    settings.provider_base_url = ""
    settings.provider_api_key = ""
    settings.fallback_provider_base_url = ""
    settings.fallback_provider_api_key = ""
    settings.fallback_model_text = ""
    settings.fallback_model_vision = ""
    settings.research_search_url = ""
    settings.model_text = ""
    settings.model_vision = ""
    settings.model_embed = ""
    settings.model_image = ""


@pytest.fixture
def local_openai() -> LocalOpenAIServer:
    """Поднимает LocalOpenAIServer и гасит его после теста."""
    server = LocalOpenAIServer()
    yield server
    server.close()


@pytest.fixture
def wait_generation():
    """Поллит GET /generations/{run_id} до терминального состояния.

    POST /generations отвечает 202 ДО завершения (процесс-локальный async
    lifecycle) — тесты обязаны дождаться queued→running→terminal, а не
    читать мгновенный снимок."""
    import time

    def _wait(client, run_id: str, timeout: float = 180.0) -> dict:
        deadline = time.monotonic() + timeout
        while True:
            run = client.get(f"/api/v1/generations/{run_id}").json()
            if run.get("state") in {"completed", "failed", "canceled"}:
                return run
            if time.monotonic() > deadline:
                raise TimeoutError(
                    f"generation {run_id} still {run.get('state')} after {timeout}s"
                )
            time.sleep(0.2)

    return _wait


@pytest.fixture(autouse=True)
def _isolated_shared_pacers():
    """ADR-031: `shared_pacer()` хранит один пейсер на endpoint на процесс.
    Без сброса тест получает пейсер с `_next_request_at`, сдвинутым
    предыдущим тестом, и ждёт чужой слот."""
    from deckdna.providers import openai_compat

    openai_compat._SHARED_PACERS.clear()
    yield
    openai_compat._SHARED_PACERS.clear()


def squeeze_text_frames(pptx_bytes: bytes, count: int = 3) -> bytes:
    """ADR-031: колода с настоящими repairable-проблемами для repair-тестов.

    Конвейер сам чинит переполнения до ревизии 1 (Q2 auto-fix, text_fit),
    поэтому свежая колода больше не даёт ``text.overflow``. Здесь у *count*
    текстовых рамок (не заголовков) высота урезается до четверти — текст
    переполняет рамку, а места под ней на слайде достаточно, чтобы
    ``resize_shape`` вернул её. Позиция и текст не трогаются."""
    from io import BytesIO

    from pptx import Presentation

    prs = Presentation(BytesIO(pptx_bytes))
    squeezed = 0
    for slide in list(prs.slides)[1:]:
        for shape in slide.shapes:
            if squeezed >= count:
                break
            if (
                not shape.has_text_frame
                or shape.is_placeholder and shape.placeholder_format.type in (1, 3)
                or len(shape.text_frame.text.strip()) < 40
                or not shape.height
                or shape.top + shape.height > prs.slide_height * 0.8
            ):
                continue
            shape.text_frame.auto_size = None
            shape.text_frame.word_wrap = True
            shape.height = max(int(shape.height) // 4, 120000)
            squeezed += 1
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def squeeze_variant_text_frames(variant_id: str) -> None:
    """Give API repair tests a real damaged PPTX, even when generation is clean."""
    import hashlib

    from deckdna.api.app import STORE, _variant_or_404

    _, variant = _variant_or_404(variant_id)
    deck_id = variant.deck_artifact_id
    artifact = STORE.artifacts[deck_id]
    data = squeeze_text_frames(artifact.data)
    STORE.artifacts[deck_id] = artifact.model_copy(update={
        "data": data, "sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data),
    })
    # A pipeline audit describes the pre-damage bytes. The next POST /audits
    # must measure the modified file rather than returning that cached audit.
    for audit_id in [a.id for a in STORE.audits.values() if a.variant_id == variant_id]:
        STORE.audits.pop(audit_id)
        STORE.issues.pop(audit_id, None)
