"""Offline checks for the delivered-deck benchmark, without a server or model."""

import importlib.util
import json
import subprocess
import sys
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches

SCRIPT = Path(__file__).resolve().parents[2] / "scripts/deck_bench.py"
spec = importlib.util.spec_from_file_location("deck_bench", SCRIPT)
bench = importlib.util.module_from_spec(spec)


def load_bench():
    assert SCRIPT.exists(), "deck_bench.py is not implemented"
    spec.loader.exec_module(bench)
    return bench


def make_deck(tmp_path):
    prs = Presentation()
    for title in ("Рост 12,5%", "Продажи по месяцам", "Предположение: 999"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(2))
        body.text = "• Первый\n• Второй\n• Третий"
    cd = CategoryChartData()
    cd.categories = ["Январь", "Февраль"]
    cd.add_series("Продажи", [10, 20])
    prs.slides[0].shapes.add_chart(
        XL_CHART_TYPE.LINE, Inches(1), Inches(4), Inches(5), Inches(2), cd
    )
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_metrics_read_rendered_title_and_native_chart_not_api_title(tmp_path):
    b = load_bench()
    path = make_deck(tmp_path)
    metadata = [
        {
            "index": 1,
            "purpose": "data",
            "title": "API title 55",
            "dataset": "sales.csv",
            "planned_type": "bar",
        },
        {"index": 2, "purpose": "data", "dataset": "missing.csv"},
    ]
    result = b.measure_deck(
        path,
        metadata=metadata,
        allowed_text=["12,5 10 20"],
        expected={"sales.csv": ["line"], "missing.csv": ["pie"]},
    )
    assert result["data_title_number_share"] == 0.5
    assert result["label_title_count"] == 1
    assert result["unsupported_numbers"] == ["999"]
    assert result["unsupported_number_count"] == 1
    assert result["numbers_total"] == 4
    assert result["unsupported_number_share"] == 0.25
    assert result["forbidden_counts"]["Предположение:"] == 1
    assert result["chart_match_share"] == 0.5
    assert result["chart_checks"][0]["observed"] == ["line"]
    assert result["archetype_count"] == 2
    assert result["layout_count"] == 1
    assert result["max_archetype_run"] == 2
    assert result["list_slide_share"] == pytest.approx(2 / 3)


def test_numbers_preserve_sign_thousands_and_decimal_notation():
    b = load_bench()
    assert b.numbers("−1 234,50; 12.50%; 1\u202f000; +7; 0,25") == {
        "-1234.5",
        "12.5",
        "1000",
        "7",
        "0.25",
    }


def test_nested_shapes_and_service_slugs_are_visible(tmp_path):
    b = load_bench()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Thanks! CREDITS Slidesgo viz_reason sales_monthly Направление A условный"
    path = tmp_path / "groups.pptx"
    prs.save(path)
    metrics = b.measure_deck(path)
    assert metrics["forbidden_total"] == 7
    assert metrics["forbidden_counts"]["english_slug"] == 1


def test_missing_expected_dataset_is_failure_not_na(tmp_path):
    b = load_bench()
    metrics = b.measure_deck(make_deck(tmp_path), expected={"absent.csv": ["scatter"]})
    assert metrics["chart_match_share"] == 0
    assert metrics["chart_checks"][0]["observed"] == []


def test_source_numbers_do_not_include_json_structure(tmp_path):
    b = load_bench()
    path = tmp_path / "data.json"
    path.write_text('{"series2": [12.5, 99], "label": "2025"}')
    assert b.numbers(b.source_text(path)) == {"12.5", "99", "2025"}


def test_compare_cli_rejects_regression_and_missing_decks(tmp_path):
    load_bench()
    base = {
        "schema_version": "1",
        "fingerprint": "same",
        "decks": [
            {
                "scenario": "one",
                "strategy": "faithful",
                "score": 80,
                "status": "completed",
                "metrics": {"audit_errors": 0},
            }
        ],
    }
    old, new = tmp_path / "old.json", tmp_path / "new.json"
    old.write_text(json.dumps(base))
    for score, code in ((78, 0), (77.9, 1)):
        base["decks"][0]["score"] = score
        new.write_text(json.dumps(base))
        result = subprocess.run(  # noqa: S603 — fixed interpreter and repository script
            [sys.executable, str(SCRIPT), "--compare", str(old), str(new)],
            capture_output=True,
            text=True,
        )
        assert result.returncode == code, result.stderr
        assert "audit_errors" in result.stdout
    base["decks"] = []
    new.write_text(json.dumps(base))
    result = subprocess.run(  # noqa: S603 — fixed interpreter and repository script
        [sys.executable, str(SCRIPT), "--compare", str(old), str(new)],
        capture_output=True,
        text=True,
    )
    assert result.returncode != 0


def test_contact_sheet_orders_pages_numerically(tmp_path):
    b = load_bench()
    import fitz
    from PIL import Image

    doc = fitz.open()
    for index in range(12):
        page = doc.new_page(width=72, height=72)
        page.draw_rect(page.rect, color=None, fill=(index / 12, 0, 0))
    pdf = tmp_path / "deck.pdf"
    doc.save(pdf)
    doc.close()
    png = tmp_path / "sheet.png"
    b.contact_sheet(pdf, png)
    with Image.open(png) as sheet:
        reds = [sheet.getpixel(((i % 4) * 40 + 20, (i // 4) * 40 + 20))[0] for i in range(12)]
    assert reds == sorted(reds)
    assert len(set(reds)) == 12


def test_direct_http_client_does_not_require_optional_socks_package(monkeypatch):
    b = load_bench()
    monkeypatch.setenv("ALL_PROXY", "socks5://127.0.0.1:1")
    monkeypatch.setenv("HTTP_PROXY", "socks5://127.0.0.1:1")
    with b.http_client("http://127.0.0.1:8088/api/v1", 10) as client:
        assert str(client.build_request("GET", "/capabilities").url) == (
            "http://127.0.0.1:8088/api/v1/capabilities"
        )


def test_score_penalizes_audit_and_renormalizes_only_true_na():
    b = load_bench()
    config = json.loads((SCRIPT.parents[1] / "configs/bench/corpus.v1.json").read_text())["scoring"]
    metrics = {
        "slides": 5,
        "data_slides": 0,
        "audit_blockers": 0,
        "audit_errors": 0,
        "data_title_number_share": None,
        "label_title_count": 0,
        "unsupported_number_share": 0,
        "archetype_count": 4,
        "max_archetype_run": 2,
        "list_slide_share": 0,
        "forbidden_total": 0,
        "chart_match_share": None,
        "build_seconds_three_variants": 100,
    }
    assert b.score_deck(metrics, config)[0] == 100
    metrics["audit_blockers"] = 1
    assert b.score_deck(metrics, config)[0] == 60
    metrics["slides"] = 0
    assert b.score_deck(metrics, config)[0] == 0


def test_wrong_native_chart_is_rejected_even_when_plan_claims_expected(tmp_path):
    b = load_bench()
    metrics = b.measure_deck(
        make_deck(tmp_path),
        metadata=[{"index": 1, "dataset": "data.csv", "planned_type": "pie"}],
        expected={"data.csv": ["pie"]},
    )
    assert metrics["chart_match_share"] == 0
    assert metrics["chart_checks"][0]["observed"] == ["line"]


def test_compare_rejects_incompatible_inputs_and_failed_artifacts():
    b = load_bench()
    old = {
        "fingerprint": "one",
        "decks": [{"scenario": "s", "strategy": "visual", "score": 50, "status": "completed"}],
    }
    assert b.compare(old, {**old, "fingerprint": "two"})[0] == 2
    new = {**old, "decks": [{**old["decks"][0], "status": "failed"}]}
    assert b.compare(old, new)[0] == 1


def test_api_pagination_job_failure_and_download_checksum(tmp_path):
    b = load_bench()
    import httpx

    def serve(request):
        if request.url.path.endswith("/jobs/job1"):
            return httpx.Response(200, json={"state": "failed", "error": "render_failed"})
        if request.url.path.endswith("/download"):
            return httpx.Response(200, content=b"wrong artifact")
        if request.url.params.get("cursor"):
            return httpx.Response(200, json={"items": [{"index": 2}], "next_cursor": None})
        return httpx.Response(200, json={"items": [{"index": 1}], "next_cursor": "next"})

    with httpx.Client(base_url="http://test/api/v1", transport=httpx.MockTransport(serve)) as c:
        api = b.API(c)
        assert api.items("/variants/v1/slides") == [{"index": 1}, {"index": 2}]
        with pytest.raises(RuntimeError, match="render_failed"):
            api.wait("job1")
        with pytest.raises(ValueError, match="checksum"):
            api.download({"artifact_id": "art1", "sha256": "incorrect"}, tmp_path / "bad.pptx")
    assert not (tmp_path / "bad.pptx").exists()


def test_runner_generates_downloads_and_scores_three_decks_offline(tmp_path, monkeypatch):
    b = load_bench()
    from argparse import Namespace

    from deckdna.api.app import app
    from deckdna.settings import settings
    from fastapi.testclient import TestClient

    monkeypatch.setattr(settings, "state_snapshot", False)
    monkeypatch.setattr(settings, "artifact_root", tmp_path / "artifacts")
    monkeypatch.setattr(
        b, "http_client", lambda *_: TestClient(app, base_url="http://testserver/api/v1")
    )
    template = make_deck(tmp_path)
    corpus = json.loads((SCRIPT.parents[1] / "configs/bench/corpus.v1.json").read_text())
    scenario = {
        "id": "offline",
        "template": "generated",
        "datasets": [],
        "slides": 3,
        "idea": "Описание проекта. Сервис проверяет документы. Команда проводит пилот. "
        "Партнёры оценивают результаты. Затем команда выбирает следующий шаг.",
    }
    args = Namespace(
        out=tmp_path / "out",
        base_url="http://testserver/api/v1",
        use_llm=False,
        request_timeout=120,
        timeout=120,
        poll=0.1,
    )
    rows, state = b.run_scenario(scenario, corpus, {"generated": template}, {}, [], args)
    assert len(rows) == 3
    assert {r["strategy"] for r in rows} == {"faithful", "balanced", "visual"}
    assert all(r["status"] == "completed" for r in rows), (rows, state)
    for row in rows:
        assert 0 <= row["score"] <= 100
        assert Path(row["contact_sheet"]).is_file()
        assert set(row["artifacts"]) == {"pptx", "pdf", "quality_passport"}
        assert row["metrics"]["audit_errors"] >= 0


def test_chart_cache_text_and_values_are_audited(tmp_path):
    b = load_bench()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["CREDITS 999"]
    data.add_series("Thanks! 777", [123456789])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(3), data
    )
    path = tmp_path / "chart.pptx"
    prs.save(path)
    result = b.measure_deck(path)
    assert result["forbidden_counts"]["CREDITS"] == 1
    assert result["forbidden_counts"]["Thanks!"] == 1
    assert result["unsupported_numbers"] == ["777", "999", "123456789"]


def test_web_evidence_uses_grounded_blocks_not_ids_dates_or_model_text():
    b = load_bench()
    sources = [
        {
            "id": "web123",
            "url": "https://example.org/900",
            "title": "Отчёт",
            "retrieved_at": "2026-10-19",
            "query": "рынок 555",
        }
    ]
    pack = {
        "external_sources": [
            {"id": "web123", "origin": "web"},
            {"id": "model456", "origin": "model"},
        ],
        "sections": [
            {
                "heading": "Итоги",
                "blocks": [
                    {"source_ref": {"artifact_id": "web123"}, "text": "Доля составила 42%"},
                    {"source_ref": {"artifact_id": "model456"}, "text": "Рост 888%"},
                ],
            }
        ],
    }
    assert b.numbers("\n".join(b.web_evidence(sources, pack))) == {"42"}


def test_incomplete_export_cannot_reuse_previous_files(tmp_path):
    b = load_bench()
    for name in ("deck.pptx", "deck.pdf", "passport.json"):
        (tmp_path / name).write_text("stale")
    # Validation must happen before even the first download or reading old artifacts.
    with pytest.raises(ValueError, match="Missing export formats"):
        b.download_exports(
            None, {"artifacts": [{"format": "pptx", "artifact_id": "new"}]}, tmp_path
        )


@pytest.mark.parametrize("score", [float("nan"), float("inf"), -1, 101])
def test_compare_rejects_nonfinite_or_invalid_scores(score):
    b = load_bench()
    row = {"scenario": "s", "strategy": "visual", "status": "completed", "score": 80}
    old = {"fingerprint": "same", "decks": [row]}
    with pytest.raises(ValueError, match="score"):
        b.compare(old, {**old, "decks": [{**row, "score": score}]})


@pytest.mark.parametrize("heading", ["Источники", "References"])
def test_web_bibliography_is_not_numeric_evidence(heading):
    b = load_bench()
    sources = [{"id": "web1", "title": "Отчёт", "url": "https://example.org/900"}]
    pack = {
        "external_sources": [{"id": "web1", "origin": "web"}],
        "sections": [
            {
                "heading": heading,
                "blocks": [
                    {
                        "source_ref": {"artifact_id": "web1"},
                        "items": [
                            "Отчёт. example.org. URL: https://example.org/900 "
                            "(дата обращения: 19.10.2026)"
                        ],
                    }
                ],
            }
        ],
    }
    assert b.numbers("\n".join(b.web_evidence(sources, pack))) == set()


def test_csv_source_text_splits_cells_not_delimiter_fragments(tmp_path):
    """«2025-01,30.3» сырым текстом давало фантомы 1.3 и 3 и теряло 30.3."""
    b = load_bench()
    path = tmp_path / "nps.csv"
    path.write_text("Месяц,NPS\n2025-01,30.3\n2025-02,31.8\n", encoding="utf-8")
    assert b.numbers(b.source_text(path)) == {"2025", "1", "30.3", "2", "31.8"}


def test_derived_numbers_cover_shares_totals_medians_ratios_and_bins():
    b = load_bench()
    derived = b.derived_numbers([[Decimal(10), Decimal(20), Decimal(30), Decimal(40)]])
    for expected in (
        "100",  # сумма
        "25",  # медиана
        "40",  # максимум
        "10",  # доля 10/100 → 10%
        "70",  # топ-2 доля
        "90",  # топ-3 доля
        "4",  # last/first
        "300",  # рост, %
        "2",  # остальные
    ):
        assert Decimal(expected) in derived, expected
    two_cols = b.derived_numbers(
        [[Decimal(x) for x in (1, 2, 3, 4, 5)], [Decimal(x) for x in (2, 4, 6, 8, 10)]]
    )
    assert Decimal("1") in two_cols  # r = 1
    # бины гистограммы: края и высоты — легальные производные
    bins = b.derived_numbers([[Decimal(x) for x in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10)]])
    assert Decimal("10") in bins and Decimal("1") in bins


def test_number_supported_allows_rounding_scale_years_and_page_numbers(tmp_path):
    b = load_bench()
    source = tmp_path / "d.csv"
    source.write_text("Канал;Доля\nСайт;19.7254\nМаркет;8\n", encoding="utf-8")
    ev = b.build_number_evidence(["Доля 19.7254 за 2025 год"], [source])
    assert b.number_supported(Decimal("19.7"), ev, 13)  # округление
    assert b.number_supported(Decimal("0.197"), ev, 13)  # доля ↔ процент
    assert b.number_supported(Decimal("19725.4"), ev, 13)  # масштаб ×1000
    assert b.number_supported(Decimal("2025"), ev, 13)  # год из источника
    assert b.number_supported(Decimal("3"), ev, 13)  # номер слайда
    assert b.number_supported(Decimal("28.9"), ev, 13)  # доля Маркета 8/27.7
    assert not b.number_supported(Decimal("999"), ev, 13)
    assert not b.number_supported(Decimal("1999"), ev, 13)  # год не из источника
    assert not b.number_supported(Decimal("27"), ev, 13)  # за пределами слайдов


def test_numbers_are_not_supported_when_only_scale_hints_exist():
    b = load_bench()
    ev = b.NumberEvidence(frozenset({Decimal("42")}), frozenset())
    assert not b.number_supported(Decimal("42.6"), b.NumberEvidence(ev.values, ev.years), 5)


def make_gate_deck(tmp_path, name="deck.pptx"):
    """Чистая колода: обложка + слайд данных с источником, единицами и периодом."""
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[5])
    cover.shapes.title.text = "Отчёт по выручке"
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Выручка выросла до 410 млн ₽"
    note = slide.shapes.add_textbox(Inches(1), Inches(1.6), Inches(6), Inches(0.6))
    note.text_frame.text = "Источник: sales.csv, янв–июн 2025. Единицы: млн ₽"
    cd = CategoryChartData()
    cd.categories = ["Январь", "Июнь"]
    cd.add_series("Выручка", [100, 410])
    slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(2.4), Inches(6), Inches(3), cd)
    path = tmp_path / name
    prs.save(path)
    return path


GATE_META = [{"index": 2, "purpose": "data", "dataset": "sales.csv", "planned_type": "line"}]
GATE_PERIODS = {"sales.csv": True}


def gate_names(result):
    return {f["gate"] for f in result["failures"]}


def test_hard_gates_pass_on_clean_deck(tmp_path):
    b = load_bench()
    result = b.hard_gates(
        make_gate_deck(tmp_path), metadata=GATE_META, dataset_periods=GATE_PERIODS
    )
    assert result == {"ok": True, "failures": []}


def test_hard_gates_report_unreadable_pptx_as_reopen_failure(tmp_path):
    b = load_bench()
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"definitely not a zip")
    result = b.hard_gates(broken)
    assert result["ok"] is False
    assert gate_names(result) == {"pptx_reopens"}


def test_hard_gates_flag_chart_without_embedded_workbook(tmp_path):
    b = load_bench()
    source = make_gate_deck(tmp_path)
    stripped = tmp_path / "no_workbook.pptx"
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(stripped, "w") as zout:
        for item in zin.infolist():
            if "embeddings" in item.filename and item.filename.endswith(".xlsx"):
                continue
            zout.writestr(item, zin.read(item.filename))
    result = b.hard_gates(stripped, metadata=GATE_META, dataset_periods=GATE_PERIODS)
    assert result["ok"] is False
    assert "native_visuals" in gate_names(result)


def test_hard_gates_flag_raster_instead_of_planned_chart_and_table(tmp_path):
    b = load_bench()
    from PIL import Image

    image = tmp_path / "pic.png"
    Image.new("RGB", (64, 32), "#808080").save(image)
    for planned in ("pie", "table"):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(3), Inches(2))
        path = tmp_path / f"{planned}.pptx"
        prs.save(path)
        result = b.hard_gates(
            path, metadata=[{"index": 1, "purpose": "data", "planned_type": planned}]
        )
        assert result["ok"] is False, planned
        assert "native_visuals" in gate_names(result)


def test_hard_gates_flag_shapes_beyond_slide_edge_groups_and_rotation(tmp_path):
    b = load_bench()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # далеко за правым краем
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, prs.slide_width - Emu(10000), Inches(1), Inches(2), Inches(1)
    )
    # ребёнок группы с собственным пространством координат
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(Inches(30), Inches(1), Inches(2), Inches(1))
    child.text_frame.text = "x"
    path = tmp_path / "oob.pptx"
    prs.save(path)
    result = b.hard_gates(path)
    assert result["ok"] is False
    details = [f["detail"] for f in result["failures"] if f["gate"] == "within_bounds"]
    assert len(details) == 2, details


def test_hard_gates_flag_missing_source_units_and_period(tmp_path):
    b = load_bench()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = ["A", "B"]
    cd.add_series("S", [1, 2])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(3), cd
    )
    path = tmp_path / "nocontext.pptx"
    prs.save(path)
    result = b.hard_gates(
        path,
        metadata=[{"index": 1, "purpose": "data", "dataset": "d.csv"}],
        dataset_periods={"d.csv": True},
    )
    assert result["ok"] is False
    details = [f["detail"] for f in result["failures"] if f["gate"] == "data_context"]
    assert any("источник" in d for d in details)
    assert any("единиц" in d or "метрик" in d for d in details)
    assert any("период" in d for d in details)


def test_hard_gates_flag_empty_and_duplicate_slides(tmp_path):
    b = load_bench()
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # пустой
    for _ in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Источники"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(2))
        box.text_frame.text = "Отчёт по продажам, 2025"
    donor = prs.slides.add_slide(prs.slide_layouts[5])
    donor.shapes.title.text = "Thanks!"
    dbox = donor.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(2))
    dbox.text_frame.text = "CREDITS: presentation template by Slidesgo"
    path = tmp_path / "dup.pptx"
    prs.save(path)
    result = b.hard_gates(path)
    assert result["ok"] is False
    details = [f["detail"] for f in result["failures"] if f["gate"] == "no_empty_duplicate"]
    assert any("пустой" in d for d in details)
    assert any("дублирует" in d or "повтор" in d for d in details)
    assert any("донорский" in d for d in details)


def test_trio_title_gates_require_same_numbers_across_variants():
    b = load_bench()

    def row(strategy, titles):
        return {
            "strategy": strategy,
            "status": "completed",
            "metrics": {
                "slide_details": [{"index": i, "title": t} for i, t in enumerate(titles, 1)]
            },
        }

    same = [
        row("faithful", ["Обложка", "Рост 10%"]),
        row("balanced", ["Обложка", "Рост 10%"]),
        row("visual", ["Обложка", "Рост 10%"]),
    ]
    assert b.trio_title_gates(same) == {}
    diff = [
        row("faithful", ["Обложка", "Рост 10%"]),
        row("balanced", ["Обложка", "Рост 12%"]),
        row("visual", ["Обложка", "Рост 10%"]),
    ]
    result = b.trio_title_gates(diff)
    assert set(result) == {"faithful", "balanced", "visual"}
    assert "variant_title_numbers" in result["balanced"]


def _deck_row(score=80, gates=None, unsupported=0, status="completed"):
    return {
        "scenario": "s",
        "strategy": "faithful",
        "score": score,
        "status": status,
        "metrics": {
            "gates": {"ok": not gates, "failures": [{"gate": g} for g in (gates or [])]},
            "unsupported_number_count": unsupported,
        },
    }


def test_compare_rejects_new_gate_failures_and_worse_numbers():
    b = load_bench()
    old = {"fingerprint": "same", "decks": [_deck_row()]}
    # те же баллы, но появился новый провал гейта
    new = {"fingerprint": "same", "decks": [_deck_row(gates=["within_bounds"])]}
    assert b.compare(old, new)[0] == 1
    # те же баллы и гейты, но неподтверждённых чисел стало больше
    new = {"fingerprint": "same", "decks": [_deck_row(unsupported=4)]}
    assert b.compare(old, new)[0] == 1
    # тот же набор гейтов, что и раньше, — не регресс
    old2 = {"fingerprint": "same", "decks": [_deck_row(gates=["within_bounds"], score=50)]}
    new2 = {"fingerprint": "same", "decks": [_deck_row(gates=["within_bounds"], score=52)]}
    assert b.compare(old2, new2)[0] == 0
    # меньше неподтверждённых чисел — улучшение, допускается
    old3 = {"fingerprint": "same", "decks": [_deck_row(unsupported=3)]}
    new3 = {"fingerprint": "same", "decks": [_deck_row(unsupported=1)]}
    assert b.compare(old3, new3)[0] == 0
    # completed → failed — регресс при любом балле
    new4 = {"fingerprint": "same", "decks": [_deck_row(status="failed", score=0)]}
    assert b.compare(old, new4)[0] == 1


def test_has_period_requires_calendar_words_or_year_in_context():
    b = load_bench()
    assert b.has_period("Выручка за 2025 год")
    assert b.has_period("NPS по месяцам, янв 2025 – июн 2026")
    assert b.has_period("Динамика, 2025-01 — 2025-06")
    assert b.has_period("Срез на 14.03.2025")
    assert not b.has_period("Заказы\n2667.9\n1901\n")
    assert not b.has_period("Сумма заказа 1901 ₽")
