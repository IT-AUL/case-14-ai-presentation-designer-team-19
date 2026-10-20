"""Do not manufacture vertical prose donors or modify original brand slides."""
import io

from deckdna.template.synthesize import prepare_template_bytes
from pptx import Presentation

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def test_synthesized_prose_is_horizontal_and_existing_vertical_slide_stays():
    prs = Presentation()
    vertical = next(layout for layout in prs.slide_layouts
                    if any(ph._element.find(f'{{{P}}}nvSpPr/{{{P}}}nvPr/{{{P}}}ph').get('orient')
                           == 'vert' for ph in layout.placeholders))
    original = prs.slides.add_slide(vertical)
    original.shapes.title.text = 'Original brand typography'
    out = io.BytesIO()
    prs.save(out)
    data, report = prepare_template_bytes(out.getvalue())
    prepared = Presentation(io.BytesIO(data))
    assert report is not None and len(prepared.slides) > 1
    assert prepared.slides[0].shapes.title.text == 'Original brand typography'
    assert any(ph._element.find(f'{{{P}}}nvSpPr/{{{P}}}nvPr/{{{P}}}ph').get('orient') == 'vert'
               for ph in prepared.slides[0].placeholders)
    for slide in list(prepared.slides)[1:]:
        assert all(ph._element.find(f'{{{P}}}nvSpPr/{{{P}}}nvPr/{{{P}}}ph').get('orient') != 'vert'
                   for ph in slide.placeholders)


def _body_layout(prs):
    from pptx.enum.shapes import PP_PLACEHOLDER

    return next(
        layout
        for layout in prs.slide_layouts
        if any(
            ph.placeholder_format.type in {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
            for ph in layout.placeholders
        )
        and all(
            ph._element.find(f"{{{P}}}nvSpPr/{{{P}}}nvPr/{{{P}}}ph").get("orient") != "vert"
            for ph in layout.placeholders
        )
    )


def test_empty_used_body_layout_is_instantiated_without_changing_original():
    prs = Presentation()
    layout = _body_layout(prs)
    original = prs.slides.add_slide(layout)
    original.shapes.title.text = "Author title"
    out = io.BytesIO()
    prs.save(out)
    data, report = prepare_template_bytes(out.getvalue())
    prepared = Presentation(io.BytesIO(data))
    assert report is not None
    assert prepared.slides[0].shapes.title.text == "Author title"
    assert all(
        not ph.text
        for ph in prepared.slides[0].placeholders
        if ph != prepared.slides[0].shapes.title
    )
    donors = [
        s
        for s in list(prepared.slides)[1:]
        if s.slide_layout.part.partname == prepared.slides[0].slide_layout.part.partname
    ]
    assert donors and any(
        ph.text.strip() for s in donors for ph in s.placeholders if ph != s.shapes.title
    )


def test_cover_subtitle_is_not_the_dominant_prose_host_for_layout_only_template():
    from collections import Counter

    prs = Presentation()
    out = io.BytesIO()
    prs.save(out)
    data, _ = prepare_template_bytes(out.getvalue())
    prepared = Presentation(io.BytesIO(data))
    counts = Counter(s.slide_layout.part.partname for s in prepared.slides)
    dominant = counts.most_common(1)[0][0]
    exemplar = next(s for s in prepared.slides if s.slide_layout.part.partname == dominant)
    bodies = [ph for ph in exemplar.placeholders if ph != exemplar.shapes.title]
    assert any(ph.height > prepared.slide_height * 0.4 for ph in bodies), (
        "a short cover subtitle displaced the primary body layout"
    )


def test_replace_text_runs_prefers_primary_body_over_narrow_subtitle():
    from deckdna.pptx.composing.text_replace import replace_text_runs
    from lxml import etree

    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    slide_xml = f"""
    <p:sld xmlns:p="{P}" xmlns:a="{A}">
      <p:cSld>
        <p:spTree>
          <p:sp>
            <p:nvSpPr>
              <p:cNvPr id="2" name="Title"/><p:cNvSpPr/>
              <p:nvPr><p:ph type="title"/></p:nvPr>
            </p:nvSpPr>
            <p:spPr>
              <a:xfrm><a:off x="457200" y="274638"/><a:ext cx="8229600" cy="800000"/></a:xfrm>
            </p:spPr>
            <p:txBody><a:bodyPr/><a:p><a:r><a:t>Old title</a:t></a:r></a:p></p:txBody>
          </p:sp>
          <p:sp>
            <p:nvSpPr>
              <p:cNvPr id="3" name="Subtitle"/><p:cNvSpPr/>
              <p:nvPr><p:ph type="subTitle" idx="1"/></p:nvPr>
            </p:nvSpPr>
            <p:spPr>
              <a:xfrm><a:off x="457200" y="1150000"/><a:ext cx="8229600" cy="450000"/></a:xfrm>
            </p:spPr>
            <p:txBody><a:bodyPr/><a:p><a:r><a:t>Narrow subtitle strip</a:t></a:r></a:p></p:txBody>
          </p:sp>
          <p:sp>
            <p:nvSpPr>
              <p:cNvPr id="4" name="Main Body"/><p:cNvSpPr/>
              <p:nvPr><p:ph type="body" idx="2"/></p:nvPr>
            </p:nvSpPr>
            <p:spPr>
              <a:xfrm><a:off x="457200" y="1750000"/><a:ext cx="8229600" cy="4200000"/></a:xfrm>
            </p:spPr>
            <p:txBody><a:bodyPr/><a:p><a:r><a:t>Primary body text</a:t></a:r></a:p></p:txBody>
          </p:sp>
        </p:spTree>
      </p:cSld>
    </p:sld>
    """.strip().encode()

    updated, report = replace_text_runs(
        slide_xml,
        ["Основной абзац контентного слайда"],
        title_text="Заголовок слайда",
        slide_size=(9144000, 6858000),
    )
    root = etree.fromstring(updated)
    shapes = {
        sp.find(f"{{{P}}}nvSpPr/{{{P}}}cNvPr").get("id"): "".join(
            t.text or "" for t in sp.iter(f"{{{A}}}t")
        )
        for sp in root.findall(f".//{{{P}}}sp")
    }
    assert shapes["4"] == "Основной абзац контентного слайда"
    assert shapes["3"] == ""
    assert list(report.filled_shape_ids) == ["4"]
