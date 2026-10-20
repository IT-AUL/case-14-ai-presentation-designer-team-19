"""Sparse donors follow usable main bodies rather than captions or layout order."""

import copy
import io
from collections import Counter

from deckdna.template.synthesize import prepare_template_bytes
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches


def _two_body_layouts():
    prs = Presentation()
    layouts = [prs.slide_layouts[1], prs.slide_layouts[3]]
    keep = {layout.part for layout in layouts}
    master = prs.slide_masters[0]
    for entry in list(master._element.sldLayoutIdLst):
        if master.part.related_part(entry.rId) not in keep:
            master._element.sldLayoutIdLst.remove(entry)
    # Both layouts deliberately contain only one body to isolate the decision.
    for layout in layouts:
        bodies = [
            ph for ph in layout.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.OBJECT
        ]
        for ph in bodies[1:]:
            ph._element.getparent().remove(ph._element)
        body = bodies[0]
        body.left, body.top = Inches(1), Inches(2)
        body.width, body.height = Inches(8), Inches(4)
    return prs, layouts


def _prepare(prs):
    out = io.BytesIO()
    prs.save(out)
    data, report = prepare_template_bytes(out.getvalue())
    return Presentation(io.BytesIO(data)), report


def test_populated_caption_does_not_hide_empty_primary_body():
    prs, layouts = _two_body_layouts()
    layout = layouts[0]
    main = next(
        ph for ph in layout.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.OBJECT
    )
    caption = copy.deepcopy(main._element)
    caption.find("p:nvSpPr/p:cNvPr", caption.nsmap).set("id", "99")
    caption.find("p:nvSpPr/p:nvPr/p:ph", caption.nsmap).set("idx", "99")
    caption.find("p:spPr/a:xfrm/a:ext", caption.nsmap).set("cy", str(Inches(0.4)))
    layout.shapes._spTree.insert_element_before(caption, "p:extLst")
    original = prs.slides.add_slide(layout)
    original.shapes.title.text = "Author title"
    original.placeholders[99].text = "Author caption"

    prepared, report = _prepare(prs)
    assert report is not None
    assert prepared.slides[0].placeholders[99].text == "Author caption"
    assert prepared.slides[0].placeholders[main.placeholder_format.idx].text == ""
    donors = [
        slide
        for slide in list(prepared.slides)[1:]
        if slide.slide_layout.part.partname == layout.part.partname
    ]
    assert donors, "a populated caption must not suppress a main-body donor"
    assert donors[0].placeholders[main.placeholder_format.idx].text.strip()


def test_largest_primary_body_dominates_independently_of_layout_order():
    prs, layouts = _two_body_layouts()
    first = next(
        ph for ph in layouts[0].placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.OBJECT
    )
    first.width = Inches(2)
    prepared, report = _prepare(prs)
    assert report is not None
    counts = Counter(slide.slide_layout.part.partname for slide in prepared.slides)
    assert counts[layouts[1].part.partname] == 2
    assert counts[layouts[0].part.partname] == 1


def test_populated_primary_body_does_not_add_redundant_donor():
    prs, layouts = _two_body_layouts()
    original = prs.slides.add_slide(layouts[0])
    original.shapes.title.text = "Author title"
    body = next(
        ph for ph in original.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.OBJECT
    )
    body.text = "Already populated primary body"
    prepared, _ = _prepare(prs)
    matches = [
        slide
        for slide in prepared.slides
        if slide.slide_layout.part.partname == layouts[0].part.partname
    ]
    assert len(matches) == 1
    assert matches[0].placeholders[body.placeholder_format.idx].text == body.text
