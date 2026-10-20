"""Explicit text and background brand roles apply to declared and observed DNA."""

import io
import json

from deckdna.template.brandbook import apply_brandbook, build_brandbook, template_look
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


def test_brandbook_priority_overrides_explicit_neutral_roles_and_fonts():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    box.text = "Brand text"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string("000000")
    data = io.BytesIO()
    prs.save(data)
    book = build_brandbook(
        [
            (
                "tokens.json",
                json.dumps(
                    {
                        "colors": {"text": "#F8F8F8", "background": "#181818"},
                        "fonts": {"heading": "DejaVu Serif", "body": "Liberation Sans"},
                    }
                ).encode(),
            )
        ],
        "brand",
    )
    branded, report = apply_brandbook(data.getvalue(), book, "brandbook")
    result = Presentation(io.BytesIO(branded))
    assert str(result.slides[0].background.fill.fore_color.rgb) == "181818"
    changed = result.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert str(changed.font.color.rgb) == "F8F8F8"
    assert changed.font.name == "Liberation Sans"
    look = template_look(branded)
    assert look.theme_colors["dk1"] == "#F8F8F8"
    assert look.theme_colors["lt1"] == "#181818"
    assert report.applied
    original, report = apply_brandbook(data.getvalue(), book, "template")
    assert original == data.getvalue() and not report.applied


def test_brandbook_style_refresh_replaces_stale_dna_and_preserves_exemplars(tmp_path):
    from datetime import UTC, datetime

    from deckdna.template import autopsy, dna

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    box.text = "Original text"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string("2277BB")
    path = tmp_path / "source.pptx"
    prs.save(path)
    original = dna.build_design_dna(
        path.read_bytes(),
        autopsy.analyze_template(path),
        "tpl",
        "ana",
        {0: "source-slide"},
        datetime.now(UTC),
        "1",
    ).design_dna
    original.exemplars[0].content_description = "Author's semantic description"
    book = build_brandbook(
        [
            (
                "tokens.json",
                json.dumps(
                    {
                        "colors": {"primary": "#AA5522"},
                        "fonts": {"heading": "DejaVu Serif", "body": "Liberation Sans"},
                    }
                ).encode(),
            )
        ],
        "brand",
    )
    branded, _ = apply_brandbook(path.read_bytes(), book, "brandbook")
    refreshed = dna.restyle_design_dna(original, branded)
    assert refreshed.declared.themes[0].major_font == "DejaVu Serif"
    assert refreshed.declared.themes[0].minor_font == "Liberation Sans"
    assert "Arial" not in {f.value for f in refreshed.observed.fonts}
    assert "2277BB" not in {c.value for c in refreshed.observed.colors}
    assert refreshed.exemplars == original.exemplars
    assert refreshed.capacities == original.capacities
    assert "Arial" in {f.value for f in original.observed.fonts}
    assert refreshed.template_id == original.template_id
