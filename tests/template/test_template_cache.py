"""Регрессионные тесты дискового кэша разбора шаблона и детектора иконок (Parallel pipeline)."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from deckdna.providers.mock import MockProvider
from deckdna.template import autopsy, dna, icons
from pptx import Presentation
from pptx.util import Inches


def _make_minimal_pptx(path: Path, title_text: str = "Тестовый шаблон") -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title is not None:
        slide.shapes.title.text = title_text
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    if slide2.shapes.title is not None:
        slide2.shapes.title.text = "Контентный слайд"
    prs.save(str(path))
    return path


def test_autopsy_disk_cache_skips_reparse_and_invalidates_on_version(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setenv("DECKDNA_ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    tpl = _make_minimal_pptx(tmp_path / "tpl.pptx")

    parse_calls = 0
    real_parse = autopsy._parse_template

    def counting_parse(path: Path):
        nonlocal parse_calls
        parse_calls += 1
        return real_parse(path)

    monkeypatch.setattr(autopsy, "_parse_template", counting_parse)

    first = autopsy.analyze_template(tpl)
    assert parse_calls == 1
    assert first.slides == 2

    # Повторный запуск на том же файле (даже по другому пути с тем же sha256)
    tpl_copy = tmp_path / "tpl_copy.pptx"
    tpl_copy.write_bytes(tpl.read_bytes())
    second = autopsy.analyze_template(tpl_copy)
    assert parse_calls == 1
    assert {k: v for k, v in second.to_dict().items() if k != "path"} == {
        k: v for k, v in first.to_dict().items() if k != "path"
    }

    # Смена версии анализатора инвалидирует кэш
    third = autopsy.analyze_template(tpl, analyzer_version="autopsy-test-v2")
    assert parse_calls == 2
    assert third.slides == 2


async def test_design_dna_and_exemplar_descriptions_cached_on_disk(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setenv("DECKDNA_ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    tpl = _make_minimal_pptx(tmp_path / "tpl_dna.pptx")
    raw = tpl.read_bytes()

    forensics = autopsy.analyze_template(tpl)
    slide_ids = {i: f"tsl_{i}" for i in range(forensics.slides)}
    now = datetime.now(UTC)

    res1 = dna.build_design_dna(raw, forensics, "tpl_1", "ana_1", slide_ids, now, "1.0.0")
    assert len(res1.design_dna.exemplars) == 2

    # На втором вызове даже при подмене _build_design_dna_uncached берётся кэш с диска
    monkeypatch.setattr(
        dna,
        "_build_design_dna_uncached",
        lambda *_a, **_kw: (_ for _ in ()).throw(AssertionError("should load from disk cache")),
    )
    res2 = dna.build_design_dna(raw, forensics, "tpl_2", "ana_2", {0: "a", 1: "b"}, now, "1.0.0")
    assert res2.design_dna.template_id == "tpl_2"
    assert sorted(sid for r in res2.design_dna.slide_roles for sid in r.slide_ids) == ["a", "b"]

    # Смена версии DNA-анализатора инвалидирует кэш
    monkeypatch.setattr(dna, "ANALYZER_VERSION", "dna-test-v2")
    assert (
        dna.load_cached_template_analysis(raw, "tpl_3", "ana_3", lambda: "tsl_x", now, "1.0.0")
        is None
    )


async def test_describe_exemplars_cached_by_template_sha256(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("DECKDNA_ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    tpl = _make_minimal_pptx(tmp_path / "tpl_desc.pptx")
    raw = tpl.read_bytes()
    forensics = autopsy.analyze_template(tpl)
    res = dna.build_design_dna(
        raw,
        forensics,
        "tpl_1",
        "ana_1",
        {i: f"tsl_{i}" for i in range(forensics.slides)},
        datetime.now(UTC),
        "1.0.0",
    )

    gateway = MockProvider(
        fixtures={
            dna.EXEMPLAR_DESCRIBE_PROMPT: lambda payload: {
                "slides": [
                    {
                        "index": s["index"],
                        "description": "Описание макета из кэша",
                        "is_entity_specific": False,
                        "is_template_meta": False,
                    }
                    for s in payload["slides"]
                ]
            }
        }
    )
    desc1 = await dna.describe_exemplars(res.features, gateway, package_bytes=raw)
    assert len(desc1) == 2
    calls_after_first = len(gateway.calls)
    assert calls_after_first > 0

    # Второй вызов на тех же байтах шаблона не обращается к модели
    desc2 = await dna.describe_exemplars(res.features, gateway, package_bytes=raw)
    assert desc2 == desc1
    assert len(gateway.calls) == calls_after_first


def test_has_template_icons_distinguishes_plain_template_from_icon_catalog(
    tmp_path: Path,
) -> None:
    plain = _make_minimal_pptx(tmp_path / "plain.pptx")
    assert icons.has_template_icons(plain) is False

    # Шаблон со слайдом-каталогом из >= 12 мелких векторных значков
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(12):
        col = i % 4
        row = i // 4
        shape = slide.shapes.add_shape(
            1,
            Inches(0.5 + col * 1.2),
            Inches(0.5 + row * 1.2),
            Inches(0.4),
            Inches(0.4),
        )
        # Добавим custGeom в XML фигуры, чтобы она распознавалась как векторная пиктограмма
        sp_pr = shape._element.spPr
        for child in list(sp_pr):
            if child.tag.endswith("prstGeom"):
                sp_pr.remove(child)
        from lxml import etree

        cust = etree.SubElement(
            sp_pr, "{http://schemas.openxmlformats.org/drawingml/2006/main}custGeom"
        )
        path_lst = etree.SubElement(
            cust, "{http://schemas.openxmlformats.org/drawingml/2006/main}pathLst"
        )
        p = etree.SubElement(
            path_lst, "{http://schemas.openxmlformats.org/drawingml/2006/main}path"
        )
        for idx in range(4):
            ln = etree.SubElement(p, "{http://schemas.openxmlformats.org/drawingml/2006/main}lnTo")
            etree.SubElement(
                ln,
                "{http://schemas.openxmlformats.org/drawingml/2006/main}pt",
                x=str(idx * 10),
                y=str(idx * 10),
            )
    with_icons = tmp_path / "with_icons.pptx"
    prs.save(str(with_icons))
    assert icons.has_template_icons(with_icons) is True
