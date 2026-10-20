"""Sparse, slide-designed templates keep their artwork on added exemplars."""

import io

import pytest
from deckdna.template.synthesize import prepare_template_bytes
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _template(grouped, color="172C46"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
    rect = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor.from_string(color)
    rect.line.fill.background()
    if grouped:
        icon = io.BytesIO()
        Image.new("RGB", (8, 8), "orange").save(icon, format="PNG")
        shapes.add_picture(
            io.BytesIO(icon.getvalue()), Inches(0.1), Inches(0.1), Inches(0.1), Inches(0.1)
        )
    for y, text, size in [(1, "Authored title", 32), (3, "Authored body in the template", 18)]:
        shape = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(8), Inches(1))
        shape.text = text
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        shape.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("F0F3F8")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("grouped", [False, True])
@pytest.mark.parametrize("color", ["172C46", "692F46"])
def test_synthesized_slides_keep_drawn_background_and_foreground(grouped, color):
    data, report = prepare_template_bytes(_template(grouped, color))
    assert report.synthesized_slides > 0
    prs = Presentation(io.BytesIO(data))
    for slide in list(prs.slides)[1:]:
        background = slide.shapes[0]
        rect = background.shapes[0] if grouped else background
        assert rect.fill.fore_color.rgb == RGBColor.from_string(color)
        assert (background.left, background.top, background.width, background.height) == (
            0,
            0,
            prs.slide_width,
            prs.slide_height,
        )
        # It remains native artwork, underneath all synthetic text.
        assert background.shape_type == (
            MSO_SHAPE_TYPE.GROUP if grouped else MSO_SHAPE_TYPE.AUTO_SHAPE
        )
        ids = [n.get("id") for n in slide._element.iter(f"{{{P}}}cNvPr")]
        assert len(set(ids)) == len(ids)
        if grouped:
            assert background.shapes[1].image.blob
        for ph in slide.placeholders:
            if ph.has_text_frame and ph.text:
                assert ph.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor.from_string(
                    "F0F3F8"
                )
    # Repeated serialization must preserve copied image relationships.
    again = io.BytesIO()
    prs.save(again)
    Presentation(io.BytesIO(again.getvalue()))


def test_partial_panel_is_not_promoted_to_deck_background():
    prs = Presentation(io.BytesIO(_template(False)))
    prs.slides[0].shapes[0].width = prs.slide_width // 2
    buf = io.BytesIO()
    prs.save(buf)
    data, _ = prepare_template_bytes(buf.getvalue())
    for slide in list(Presentation(io.BytesIO(data)).slides)[1:]:
        assert all(sh.is_placeholder for sh in slide.shapes)
