"""Template attribution differs from a presentation about credits or licences."""
import io
from datetime import UTC, datetime

import pytest
from deckdna.template.dna import build_design_dna
from deckdna.template.meta import is_meta_xml
from pptx import Presentation
from pptx.util import Inches


def _slide(prs, *texts):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, text in enumerate(texts):
        slide.shapes.add_textbox(Inches(.5), Inches(i + .2), Inches(8), Inches(.7)).text = text
    return slide


@pytest.mark.parametrize('heading,body', [
    ('CREDITS', 'This presentation template is free for everyone to use thanks to the following:'),
    ('Attribution', 'Please keep this slide to credit the authors of this template.'),
    ('Лицензия шаблона', 'Укажите авторов шаблона при использовании этой презентации.'),
])
def test_explicit_template_attribution_is_not_a_content_donor(heading, body):
    prs = Presentation()
    assert is_meta_xml(_slide(prs, heading, body)._element.xml.encode())


@pytest.mark.parametrize('heading,body', [
    ('Credits', 'Customer credits decreased by 14% during the pilot.'),
    ('Attribution', 'Attribution of sales across channels remains uncertain.'),
    ('Open-source licenses', 'Our product supports templates and lists third-party dependencies.'),
])
def test_topic_words_alone_do_not_exclude_business_or_technical_content(heading, body):
    prs = Presentation()
    assert not is_meta_xml(_slide(prs, heading, body)._element.xml.encode())


def test_meta_first_does_not_steal_title_role():
    prs = Presentation()
    _slide(prs, 'CREDITS',
           'This presentation template is free for everyone to use thanks to the following:')
    _slide(prs, 'Исходный титул', 'Содержательный подзаголовок')
    _slide(prs, 'Описание', 'Основное содержание и аргументы')
    buf = io.BytesIO()
    prs.save(buf)
    dna = build_design_dna(buf.getvalue(), None, 't', 'a', {i: str(i) for i in range(3)},
                           datetime.now(UTC), '1.0').design_dna
    assert dna.exemplars[0].content_is_template_meta is True
    assert dna.exemplars[0].role != 'title'
    assert dna.exemplars[1].role == 'title'
