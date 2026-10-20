import pytest
from deckdna.contracts.deck_plan import Brief, Kind, Purpose
from deckdna.ingestion.content_parsers import parse_file
from deckdna.planning.data_story import plan_data_deck
from openpyxl import Workbook

_STORY = (
    "Узнала о курсе из рассылки, долго сравнивала с другими школами, "
    "купила после второго вебинара и попросила рассрочку на три месяца"
)


def _pack(tmp_path):
    wb = Workbook()
    sales = wb.active
    sales.title = "Продажи"
    sales.append(["Регион", "Выручка, ₽"])
    for name, value in [("Москва", 420), ("Казань", 180), ("Пермь", 95), ("Тула", 60)]:
        sales.append([name, value])
    talks = wb.create_sheet("Интервью")
    talks.append(["№", "Клиент", "Что рассказал"])
    for i in range(1, 6):
        talks.append([i, f"Клиент {i}. {_STORY}", f"{_STORY}. Ответ {i}"])
    blank = wb.create_sheet("Бланк")
    blank.append(["Кто клиент", "Что хотим узнать", "Первый блок вопросов"])
    blank.append(["Образец заполнения — на соседних листах"] * 3)
    path = tmp_path / "mixed.xlsx"
    wb.save(path)
    return parse_file(path, artifact_id="mixed")


@pytest.mark.parametrize("sheet", ["Интервью", "Бланк"])
def test_prose_and_blank_sheets_get_no_table_slide(tmp_path, sheet):
    pack = _pack(tmp_path)
    plan = plan_data_deck(
        pack,
        Brief(purpose="Продажи", audience="Руководство", language="ru", target_slide_count=8),
    )
    talks = next(d for d in pack.datasets if d.source_ref.sheet == sheet)
    talks_ev = f"{pack.id}:{talks.id}"
    data = [s for s in plan.slides if s.purpose == Purpose.data]
    assert data
    for slide in data:
        assert talks_ev not in slide.evidence_ids
        for unit in slide.content_units:
            assert talks_ev not in unit.evidence_ids
            assert unit.kind != Kind.table or _STORY not in str(unit)
