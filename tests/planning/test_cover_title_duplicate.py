"""ADR-031 §9: раздел с тем же названием, что обложка, не даёт слайдов
«Идея — часть 2»: заголовок берётся из темы предложения слайда."""

from deckdna.contracts.deck_plan import Brief
from deckdna.ingestion.content_parsers import parse_markdown
from deckdna.planning.story_director import _topic_title, plan_deck

MD = (
    "# Идея\n\nСервис собирает презентацию из шаблона компании. Пользователь "
    "загружает текст и получает три варианта. Каждый вариант проходит аудит "
    "качества. Время сборки укладывается в пять минут.\n"
)


def test_no_content_slide_repeats_cover_title():
    plan = plan_deck(
        parse_markdown(MD),
        Brief(purpose="Идея", audience="жюри", language="ru", target_slide_count=6,
              content_fill="shrink"),
    )
    cover = plan.slides[0].title_intent
    content = [s for s in plan.slides if s.purpose.value not in ("title", "thank_you")]
    assert content
    for slide in content:
        assert not slide.title_intent.startswith(cover), slide.title_intent


def test_topic_title_heuristic():
    assert _topic_title("Время сборки укладывается в пять минут.", "ru") == "Время сборки"
    assert _topic_title("Пилот сократил время на 40%.", "ru") == "Пилот"
    assert _topic_title("Short English sentence here.", "en") is None
    assert _topic_title("Без глагола только существительные и прилагательные здесь", "ru") is None


def test_technical_slugs_never_leak_into_cover_or_speaker_notes():
    """F-09 / F-18 / F-54: служебные токены presentation/project/balanced не
    должны попадать в заголовки обложки, подзаголовки и заметки спикера."""
    from deckdna.contracts.deck_plan import (
        ContentUnit,
        DensityBudget,
        Kind,
        Purpose,
        SlidePlan,
    )
    from deckdna.planning.talk_track import _author_note, extractive_track

    pack = parse_markdown(MD)
    pack = pack.model_copy(update={"title_hint": "presentation"})
    plan = plan_deck(
        pack,
        Brief(
            purpose="project",
            key_message="balanced",
            audience="руководство",
            language="ru",
            target_slide_count=5,
        ),
    )
    cover = plan.slides[0]
    assert cover.title_intent == "Идея"
    cover_texts = [u.text for u in cover.content_units if u.text]
    assert cover_texts == ["Идея"]

    from deckdna.planning.content_fill import estimate_content_capacity

    sec_with_slug = pack.sections[0].model_copy(update={"heading": "presentation"})
    slug_pack = pack.model_copy(update={"sections": [sec_with_slug]})
    suff = estimate_content_capacity(
        slug_pack,
        Brief(purpose="project", audience="руководство", language="ru", target_slide_count=4),
    )
    assert "presentation" not in suff.covered_topics

    budget = DensityBudget(level="sparse", max_words=60, max_bullets=4)
    slide = SlidePlan(
        id="slide-1",
        index=1,
        purpose=Purpose.problem,
        title_intent="Анализ узких мест",
        key_message="project",
        evidence_ids=["e1"],
        content_units=[
            ContentUnit(
                role="title", kind=Kind.title, text="Анализ узких мест", evidence_ids=["e1"]
            ),
            ContentUnit(
                role="body",
                kind=Kind.bullet,
                text="Контекст: project. Время согласования снизилось на 30% по.",
                evidence_ids=["e1"],
            ),
        ],
        density_budget=budget,
        speaker_note="Контекст: presentation. Примечание: таблица по. Раздел: 1",
    )
    track = extractive_track(slide)
    note = _author_note(slide)
    assert "project" not in track.casefold()
    assert not track.endswith(" по.")
    assert note == "Примечание: таблица."


def test_cover_subtitle_not_duplicated_into_second_column_on_preview(tmp_path):
    """F-69: на обложке с несколькими текстовыми слотами подзаголовок
    записывается ровно в один слот и не дублируется во вторую колонку."""
    from deckdna.contracts.deck_plan import (
        ContentUnit,
        DensityBudget,
        Kind,
        Purpose,
        SlidePlan,
    )
    from deckdna.pptx.composing.minimal import generate_deck
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    # Имитируем двухколоночный макет/обложку с заголовком и двумя текстовыми колонками
    tbox = s.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.0), Inches(1.0))
    tbox.text_frame.text = "Шаблонный заголовок презентации"
    # Сделаем первую фигуру title placeholder через XML
    from lxml import etree
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    nv_pr = tbox._element.find(f"{{{P}}}nvSpPr/{{{P}}}nvPr")
    etree.SubElement(nv_pr, f"{{{P}}}ph", type="title")

    col1 = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(4.0), Inches(2.5))
    col1.text_frame.text = "Первая колонка подзаголовка в шаблоне"
    col2 = s.shapes.add_textbox(Inches(5.2), Inches(2.2), Inches(4.0), Inches(2.5))
    col2.text_frame.text = "Вторая колонка подзаголовка в шаблоне"

    tpl_path = tmp_path / "two_col_cover.pptx"
    prs.save(str(tpl_path))

    subtitle_text = "Единая платформа аналитики и планирования"
    budget = DensityBudget(level="sparse", max_words=60, max_bullets=4)
    base_plan = plan_deck(
        parse_markdown(MD),
        Brief(
            purpose=subtitle_text,
            audience="руководство",
            language="ru",
            target_slide_count=3,
        ),
    )
    plan = base_plan.model_copy(
        update={
            "slides": [
                SlidePlan(
                    id="slide-0",
                    index=0,
                    purpose=Purpose.title,
                    title_intent="Стратегия развития 2026",
                    key_message=subtitle_text,
                    evidence_ids=["test:meta"],
                    content_units=[
                        ContentUnit(
                            role="title",
                            kind=Kind.title,
                            text="Стратегия развития 2026",
                            evidence_ids=["test:meta"],
                        ),
                        ContentUnit(
                            role="subtitle",
                            kind=Kind.subtitle,
                            text=subtitle_text,
                            evidence_ids=["test:meta"],
                        ),
                        ContentUnit(
                            role="body",
                            kind=Kind.bullet,
                            text=subtitle_text,
                            evidence_ids=["test:meta"],
                        ),
                    ],
                    density_budget=budget,
                )
            ]
        }
    )
    out_pptx = tmp_path / "out.pptx"
    generate_deck(tpl_path, plan, out_pptx)
    rendered = Presentation(str(out_pptx))
    slide_texts = [
        sh.text_frame.text.strip()
        for sh in rendered.slides[0].shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    ]
    assert slide_texts.count(subtitle_text) == 1, slide_texts
