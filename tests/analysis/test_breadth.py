"""ADR-033, блок D: все входы, все типы, интерпретация моделью с проверкой
чисел, сложные случаи, слайд-приложение, стиль графика-донора, бенчмарк.

Модель — MockProvider с фикстурами, сети нет."""

from __future__ import annotations

import asyncio
import sys
from pathlib import Path

import pytest
from deckdna.analysis import analyze_pack
from deckdna.analysis.interpret import (
    INTENT_PROMPT,
    INTERPRET_PROMPT,
    classify_intent,
    interpret_pack,
)
from deckdna.analysis.shape import shape_for_type
from deckdna.contracts.content_pack import Chart, ChartSeries, SourceRef
from deckdna.contracts.deck_plan import Brief
from deckdna.ingestion.content_parsers import parse_file
from deckdna.planning.story_director import plan_deck
from deckdna.providers.mock import MockProvider

ROOT = Path(__file__).resolve().parents[2]
STAGE2 = ROOT / "examples" / "stage2"
sys.path.insert(0, str(ROOT / "scripts"))


def _brief(**extra) -> Brief:
    return Brief(
        purpose="Итоги по данным",
        audience="руководство",
        language="ru",
        target_slide_count=extra.pop("target", 7),
        **extra,
    )


def test_benchmark_accuracy():
    from benchmark_viz import run

    report = run()
    assert report["cases"] >= 20
    assert report["task_accuracy"] >= 0.9
    assert report["top1_accuracy"] >= 0.85
    assert report["top3_accuracy"] == 1.0


def test_table_inside_text_becomes_a_dataset_with_context():
    pack = parse_file(STAGE2 / "report_with_table.txt")
    [ds] = pack.datasets
    assert [c.name for c in ds.columns] == ["Квартал", "Новые клиенты", "Ушедшие клиенты"]
    assert "показать рост числа клиентов" in ds.context_text
    # текст вокруг таблицы остался абзацем, без строк таблицы
    paragraph = pack.sections[0].blocks[0].text
    assert "Приложение запущено" in paragraph and "1 кв. 2025" not in paragraph
    [analysis] = analyze_pack(pack)
    assert analysis.primary.task.task == "trend"
    assert len(analysis.primary.view.series) == 2


def test_markdown_pipe_table_with_numbers_becomes_a_dataset(tmp_path):
    md = tmp_path / "report.md"
    md.write_text(
        "# Итоги\n\n## Продажи по регионам\n\nСравнить регионы по выручке.\n\n"
        "| Регион | Выручка, млн ₽ |\n|---|---|\n| Москва | 41 |\n| Казань | 12 |\n"
        "| Пермь | 8 |\n",
        encoding="utf-8",
    )
    pack = parse_file(md)
    [ds] = pack.datasets
    assert "Сравнить регионы" in ds.context_text
    assert analyze_pack(pack)[0].primary.task.task == "comparison"


def test_kpi_cards_and_reference_top_n_with_continuations():
    kpi = parse_file(STAGE2 / "kpi.json")
    plan = plan_deck(kpi, _brief(target=5))
    slide = next(s for s in plan.slides if s.purpose.value == "data")
    # заголовок — вывод с главным числом, а не ярлык «Ключевые показатели»
    assert "Выручка" in slide.title_intent and "412" in slide.title_intent
    # между числом и единицей — неразрывный пробел
    assert [
        u.text.replace("\u00a0", " ") for u in slide.content_units if u.kind.value == "bullet"
    ] == [
        "Выручка: 412 млн ₽",
        "Клиенты: 186 тыс.",
        "Средний чек: 2 217 ₽",
        "NPS: 46",
    ]
    assert not any(s.title_intent == "Ключевые выводы" for s in plan.slides)

    ref = parse_file(STAGE2 / "products_reference.xlsx")
    plan = plan_deck(ref, _brief(target=8))
    first = next(s for s in plan.slides if s.purpose.value == "data")
    table = next(u for u in first.content_units if u.kind.value == "table")
    by_id = {t.id: t for t in ref.tables}
    rows = by_id[table.table_ref].rows
    # топ-N позиций + сводка «Прочие» и «Итого» вместо слайдов-продолжений
    assert len(rows) == 6 + 2
    assert rows[-2][0].startswith("Прочие (")
    assert rows[-1][0] == "Итого"
    assert len(by_id[table.table_ref].headers) <= 5
    assert first.title_intent.startswith("Товары: в таблице топ-6 из 240 позиций")
    assert "240 строк" in first.speaker_note
    assert not any("продолжение" in s.title_intent for s in plan.slides)


def test_percent_shares_are_read_as_shares():
    pack = parse_file(STAGE2 / "channels_share.json")
    [analysis] = analyze_pack(pack)
    assert analysis.primary.rec.chosen.chart_type in ("pie", "doughnut")
    assert analysis.primary.headline == "Сайт — крупнейшая доля, 38,5%"


def test_single_metric_dataset_builds_a_story_not_a_stub():
    """BE3: regions.xlsx давал колоду «титул + один слайд + спасибо».

    У категориальной метрики есть второе честное представление — доля
    каждой категории от итога; планировщик обязан показать оба среза,
    а не схлопывать историю до трёх слайдов."""
    pack = parse_file(STAGE2 / "regions.xlsx")
    [analysis] = analyze_pack(pack)
    derived = [v for v in analysis.views if v.view.derived]
    assert derived, "ожидалось производное представление «доля от итога»"
    # доли посчитаны из исходных значений — новых чисел не появляется
    share = derived[0].view.series[0]
    assert share.unit == "%"
    assert abs(sum(v for v in share.values if v is not None) - 100) < 1
    # при хвосте «прочие» > 25% круговая законно отклоняется — доли
    # показываются горизонтальным баром; pie остаётся для компактных долей
    assert derived[0].rec.chosen.chart_type in (
        "pie", "doughnut", "bar_stacked_100", "bar_horizontal", "bar",
    )

    plan = plan_deck(pack, _brief(target=8))
    data = [s for s in plan.slides if s.purpose.value == "data"]
    assert len(plan.slides) >= 5
    assert len(data) >= 2  # сравнение регионов и их доли — разные слайды
    assert any(s.title_intent == "Ключевые выводы" for s in plan.slides)
    assert all("Данные:" not in (s.speaker_note or "") for s in data)
    # вывод производного среза не дублирует факт на «Ключевых выводах»
    key = next(s for s in plan.slides if s.title_intent == "Ключевые выводы")
    key_ids = {e for u in key.content_units for e in (u.evidence_ids or [])}
    assert not any(f"{pack.id}:{i.id}" in key_ids for v in derived for i in v.insights)


def test_pie_with_many_slices_is_folded_into_others():
    chart = Chart(
        id="c",
        categories=[f"К{i}" for i in range(9)],
        series=[ChartSeries(name="Доля", values=[30, 20, 12, 10, 8, 7, 6, 4, 3])],
        source_ref=SourceRef(artifact_id="x"),
        chart_type="bar_horizontal",
    )
    pie = shape_for_type(chart, "pie")
    assert pie.categories == ["К0", "К1", "К2", "К3", "К4", "Прочие"]
    assert pie.series[0].values[-1] == 20
    many = chart.model_copy(
        update={
            "categories": [f"Г{i}" for i in range(15)],
            "series": [ChartSeries(name="З", values=list(range(15, 0, -1)))],
        }
    )
    bars = shape_for_type(many, "bar_horizontal")
    # хвост 15 = лидеру: агрегат-столбец читался бы как главная категория,
    # поэтому сумма уходит в сноску, а не в отдельный бар
    assert len(bars.categories) == 10
    assert "показаны 10 крупнейших" in bars.source_note
    assert "остальные 5 — всего 15" in bars.source_note


def test_model_interpretation_is_checked_against_the_fact():
    pack = parse_file(STAGE2 / "sales_monthly.csv")
    target = max((i for i in pack.insights if i.fact_type == "growth"), key=lambda i: i.strength)

    def answer(payload):
        items = []
        for item in payload["items"]:
            if item["id"] == target.id:
                items.append(
                    {
                        "id": item["id"],
                        "meaning": f"Продажи выросли в {target.numbers[0]} раза за два года.",
                        "implication": "Стоит поощрить команду продаж.",
                    }
                )
            else:
                items.append(
                    {
                        "id": item["id"],
                        "meaning": "Выросло на 37% за квартал.",
                        "implication": "Нанять ещё 15 менеджеров.",
                    }
                )
        return {"items": items}

    gateway = MockProvider(fixtures={INTERPRET_PROMPT: answer})
    out, report = asyncio.run(interpret_pack(pack, _brief(), gateway))
    by_id = {i.id: i for i in out.insights}
    good = by_id[target.id]
    assert good.meaning == f"Продажи выросли в {target.numbers[0]} раза за два года."
    assert good.implication == "Стоит поощрить команду продаж."
    assert good.needs_confirmation and not good.confirmed
    others = [i for i in out.insights if i.id != target.id and i.meaning]
    for ins in others:
        # «37%» и «15» не из факта: смысл — шаблон, следствие не пишется
        assert "37" not in ins.meaning and ins.implication is None
    assert report.rejected >= 2


def test_without_model_nothing_is_invented():
    pack = parse_file(STAGE2 / "sales_monthly.csv")
    out, report = asyncio.run(interpret_pack(pack, _brief(), None))
    assert out is pack and report.interpreted == 0
    out, _ = asyncio.run(interpret_pack(pack, _brief(), MockProvider()))
    assert all(i.meaning is None and i.implication is None for i in out.insights)


def test_model_intent_adds_weight_when_text_has_no_markers():
    pack = parse_file(STAGE2 / "purchases.csv")
    gateway = MockProvider(fixtures={INTENT_PROMPT: {"task": "composition"}})
    out, task = asyncio.run(classify_intent(pack, "Итоги года", gateway))
    assert task == "composition"
    assert "Цель: доли" in out.datasets[0].context_text
    [analysis] = analyze_pack(out)
    assert analysis.primary.task.task == "composition"
    assert analysis.primary.rec.chosen.chart_type in ("pie", "doughnut")
    # маркеры в тексте есть — модель не спрашиваем
    _, task = asyncio.run(classify_intent(pack, "Покажите рост продаж", gateway))
    assert task is None


def test_explain_slide_flag_adds_the_appendix():
    pack = parse_file(STAGE2 / "purchases.csv")
    plan = plan_deck(pack, _brief(target=8, explain_slide=True))
    appendix = next(s for s in plan.slides if s.title_intent == "Как выбирались визуализации")
    texts = [u.text for u in appendix.content_units if u.kind.value == "bullet"]
    assert any("линия" in t for t in texts) and len(texts) >= 3
    assert plan.slides[-1].purpose.value == "thank_you"


def test_confirmed_implication_closes_the_deck():
    pack = parse_file(STAGE2 / "sales_monthly.csv")
    ins = pack.insights[0].model_copy(
        update={
            "implication": "Закрепить успех: поощрить команду",
            "confirmed": True,
        }
    )
    pack = pack.model_copy(update={"insights": [ins, *pack.insights[1:]]})
    plan = plan_deck(pack, _brief(target=6))
    assert plan.slides[-1].title_intent == "Что предлагаем"
    assert plan.slides[-1].content_units[1].text == "Закрепить успех: поощрить команду"


@pytest.fixture
def chart_template(tmp_path):
    """Шаблон с графиком-донором: своя палитра и кегль подписей."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Пример графика"
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("Ряд", (1, 2, 3))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(4.5), data
    )
    chart = frame.chart
    chart.plots[0].series[0].format.fill.solid()
    chart.plots[0].series[0].format.fill.fore_color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    chart.category_axis.tick_labels.font.size = Pt(14)
    text = prs.slides.add_slide(prs.slide_layouts[1])
    text.shapes.title.text = "Текст"
    text.placeholders[1].text = "Пункт"
    path = tmp_path / "donor.pptx"
    prs.save(str(path))
    return path


def test_donor_chart_style_is_applied_to_any_chart_type(tmp_path, chart_template):
    from deckdna.generation.pipeline import generate
    from pptx import Presentation

    report = generate(
        chart_template,
        STAGE2 / "sales_monthly.csv",
        {"purpose": "Итоги", "audience": "руководство", "language": "ru", "target_slide_count": 4},
        tmp_path / "out",
    )
    prs = Presentation(report["artifacts"]["pptx"])
    xml = next(
        sh.chart.part.blob.decode()
        for s in prs.slides
        for sh in s.shapes
        if getattr(sh, "has_chart", False)
    )
    assert "<c:lineChart>" in xml  # тип — по данным, не от донора
    assert 'srgbClr val="C0392B"' in xml  # палитра донора
    assert 'sz="1400"' in xml  # кегль подписей донора
