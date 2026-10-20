"""End-to-end pipeline: content file -> ContentPack -> DeckPlan -> deck ->
audit -> pdf -> QualityPassport.

Runs the real generation pipeline on the organizer template and a
markdown fixture; verifies every artifact exists and the passport
validates against schemas/quality-passport.schema.json.
"""

import json
import shutil
import zipfile
from pathlib import Path

import pytest
from deckdna.errors import DeckDNAError
from deckdna.generation.pipeline import generate
from jsonschema import validate
from pptx import Presentation

FIXTURE = Path("tests/fixtures/pptx/vk_tech_template.pptx")
CONTENT = Path("tests/fixtures/content/poc_article.md")
SCHEMA = json.loads(
    (Path(__file__).resolve().parents[2] / "schemas" / "quality-passport.schema.json").read_text()
)

BRIEF = {
    "purpose": "Показать сквозной пайплайн DeckDNA",
    "audience": "эксперты VK Tech",
    "language": "ru",
    "target_slide_count": 12,
}

HAS_SOFFICE = shutil.which("soffice") is not None


@pytest.fixture(scope="module")
def report(tmp_path_factory):
    if not FIXTURE.exists():
        pytest.skip("organizer fixture not present")
    if not HAS_SOFFICE:
        pytest.skip("soffice not installed — pdf stage needs it")
    out_dir = tmp_path_factory.mktemp("pipeline")
    return generate(FIXTURE, CONTENT, dict(BRIEF), out_dir), out_dir


def test_all_artifacts_exist(report):
    result, _ = report
    for key in ("pptx", "pdf", "quality_passport"):
        path = Path(result["artifacts"][key])
        assert path.exists(), f"missing artifact: {path}"
    assert zipfile.is_zipfile(result["artifacts"]["pptx"])
    assert Path(result["artifacts"]["pdf"]).read_bytes()[:5] == b"%PDF-"


def test_deck_matches_plan_size(report):
    result, _ = report
    # Короткий материал не дробится ради target; сборка следует финальному плану.
    planned = len(result["deck_plan"]["slides"])
    assert 3 <= planned <= BRIEF["target_slide_count"]
    assert result["slides_out"] == result["compose_report"]["slides_out"] == planned
    assert len(Presentation(result["artifacts"]["pptx"]).slides) == planned


def test_quality_passport_validates_against_schema(report):
    result, _ = report
    validate(result["quality_passport"], SCHEMA)
    metrics = result["quality_passport"]["metrics"]
    assert metrics["validity"]["opens_cleanly"] is True
    assert metrics["editability"]["pei_level"] is not None
    assert metrics["timings"]["total_seconds"] > 0


def test_audit_issues_are_dicts_with_fields(report):
    result, _ = report
    for issue in result["audit_issues"]:
        assert issue["rule_code"]
        assert issue["severity"] in {"blocker", "error", "warning", "info"}


def test_unsupported_content_format_typed(tmp_path):
    weird = tmp_path / "content.xyz"
    weird.write_text("whatever", encoding="utf-8")
    with pytest.raises(DeckDNAError) as excinfo:
        generate(FIXTURE, weird, dict(BRIEF), tmp_path / "out")
    assert excinfo.value.code == "invalid_input"
    assert excinfo.value.stage == "content_ingestion"


def test_bad_slide_count_typed(tmp_path):
    if not FIXTURE.exists():
        pytest.skip("organizer fixture not present")
    # Brief.target_slide_count is ge=3/le=40 (deck_plan.py) -- 3 is a
    # deliberately supported short-deck size now ("any input, any
    # template", 29.09), not below-minimum. Use a value actually outside
    # the contract's bounds instead.
    brief = dict(BRIEF, target_slide_count=2)
    with pytest.raises(DeckDNAError) as excinfo:
        generate(FIXTURE, CONTENT, brief, tmp_path / "out")
    assert excinfo.value.code == "invalid_input"
    assert excinfo.value.stage == "planning"


def test_style_previews_cover_subtitle_not_duplicated_in_two_columns(tmp_path, monkeypatch):
    """F-69: render_style_previews не дублирует подзаголовок обложки во второй
    текстовый слот/колонку на шаблонах с несколькими слотами обложки."""
    import deckdna.pptx.exporting.previews as exp_previews
    import deckdna.pptx.exporting.render as exp_render
    from deckdna.contracts.deck_plan import (
        Brief,
        ContentUnit,
        DensityBudget,
        Kind,
        Purpose,
        SlidePlan,
    )
    from deckdna.generation.preview import render_style_previews
    from deckdna.ingestion.content_parsers import parse_markdown
    from deckdna.planning.story_director import plan_deck
    from lxml import etree
    from pptx.util import Inches

    monkeypatch.setattr(exp_render, "render_pdf", lambda pptx, pdf: pdf)
    monkeypatch.setattr(exp_previews, "pdf_page_pngs", lambda pdf: [b"png"])

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tbox = s.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.0), Inches(1.0))
    tbox.text_frame.text = "Шаблонный заголовок презентации"
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    nv_pr = tbox._element.find(f"{{{p_ns}}}nvSpPr/{{{p_ns}}}nvPr")
    etree.SubElement(nv_pr, f"{{{p_ns}}}ph", type="title")

    col1 = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(4.0), Inches(2.5))
    col1.text_frame.text = "Первая колонка подзаголовка в шаблоне"
    col2 = s.shapes.add_textbox(Inches(5.2), Inches(2.2), Inches(4.0), Inches(2.5))
    col2.text_frame.text = "Вторая колонка подзаголовка в шаблоне"

    tpl_path = tmp_path / "two_col_cover.pptx"
    prs.save(str(tpl_path))

    subtitle_text = "Единая платформа аналитики и планирования"
    pack = parse_markdown("# Стратегия\n\nПлатформа объединяет отчётность и витрины данных.\n")
    base_plan = plan_deck(
        pack,
        Brief(
            purpose=subtitle_text,
            audience="руководство",
            language="ru",
            target_slide_count=3,
        ),
    )
    cover = SlidePlan(
        id="slide-0",
        index=0,
        purpose=Purpose.title,
        title_intent="Стратегия развития 2026",
        key_message=subtitle_text,
        evidence_ids=["test:meta"],
        content_units=[
            ContentUnit(
                role="title",
                kind=Kind.title,
                text="Стратегия развития 2026",
                evidence_ids=["test:meta"],
            ),
            ContentUnit(
                role="subtitle",
                kind=Kind.subtitle,
                text=subtitle_text,
                evidence_ids=["test:meta"],
            ),
            ContentUnit(
                role="body",
                kind=Kind.bullet,
                text=subtitle_text,
                evidence_ids=["test:meta"],
            ),
        ],
        density_budget=DensityBudget(level="sparse", max_words=60, max_bullets=4),
    )
    plan = base_plan.model_copy(update={"slides": [cover, *base_plan.slides[1:]]})

    out_dir = tmp_path / "previews"
    res = render_style_previews(
        tpl_path, pack, plan, out_dir, slide_ids=["slide-0"], strategies=("balanced",)
    )
    assert res["styles"][0]["error"] is None
    rendered = Presentation(str(out_dir / "balanced" / "preview.pptx"))
    slide_texts = [
        sh.text_frame.text.strip()
        for sh in rendered.slides[0].shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    ]
    assert slide_texts.count(subtitle_text) == 1, slide_texts
