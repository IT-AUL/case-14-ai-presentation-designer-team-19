
import pytest
from deckdna.errors import DeckDNAError
from deckdna.generation.local_edit import enlarge_chart
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def test_enlargement_preserves_native_chart_and_content(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.5), Inches(.2), Inches(8), Inches(.5))
    title.text = "Sales"
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Revenue", [5, 9])
    shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(3), Inches(2),
                                   Inches(2), Inches(2), data)
    old_area = shape.width * shape.height
    path = tmp_path / "slide.pptx"
    deck.save(path)
    enlarge_chart(path)
    actual = Presentation(path)
    chart = next(s for s in actual.slides[0].shapes if s.has_chart)
    assert chart.width * chart.height > old_area * 1.08
    assert list(chart.chart.series[0].values) == [5, 9]
    assert title.text in [s.text for s in actual.slides[0].shapes if s.has_text_frame]
    assert chart.left >= 0 and chart.top >= 0
    assert chart.left + chart.width <= actual.slide_width
    assert chart.top + chart.height <= actual.slide_height


def test_no_chart_is_an_explicit_unavailable_operation(tmp_path):
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    path = tmp_path / "slide.pptx"
    deck.save(path)
    before = path.read_bytes()
    with pytest.raises(DeckDNAError, match="one native chart"):
        enlarge_chart(path)
    assert path.read_bytes() == before
