"""B02 count gates use the actual compiled package, not a mutable report."""

import io
from pathlib import Path

import pymupdf
import pytest
from deckdna.contracts.deck_plan import Brief
from deckdna.errors import DeckDNAError
from deckdna.generation import pipeline
from deckdna.planning.request_constraints import validate_deck_count
from pptx import Presentation


def _brief(policy="exact", **extra):
    return Brief(purpose="Обзор работы", audience="команда", language="ru",
                 target_slide_count=3, content_fill="shrink", slide_count_policy=policy, **extra)


@pytest.fixture
def inputs(tmp_path):
    template = tmp_path / "unseen.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(template)
    content = tmp_path / "content.md"
    content.write_text("# Работа команды\n\n## Результаты\n\n"
                       "- Проверили материалы.\n- Согласовали ограничения.\n- Подготовили обзор.\n")
    return template, content


def test_exact_three_is_generated_and_rendered(inputs, tmp_path):
    report = pipeline.generate(*inputs, _brief(), tmp_path / "accepted")
    pptx = Presentation(report["artifacts"]["pptx"])
    assert len(pptx.slides) == report["slides_out"] == 3
    with pymupdf.open(Path(report["artifacts"]["pdf"])) as pdf:
        assert len(pdf) == 3


def test_extra_compiled_page_cannot_hide_behind_old_report(inputs, tmp_path, monkeypatch):
    original = pipeline.generate_deck

    def add_continuation(template, plan, out, **kwargs):
        report = original(template, plan, out, **kwargs)
        prs = Presentation(out)
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(out)
        assert report["slides_out"] == 3
        return report

    def forbidden_render(*args, **kwargs):
        raise AssertionError("count conflict must be caught before final rendering")

    monkeypatch.setattr(pipeline, "generate_deck", add_continuation)
    monkeypatch.setattr(pipeline, "render_pdf", forbidden_render)
    out = tmp_path / "conflict"
    with pytest.raises(DeckDNAError) as error:
        pipeline.generate(*inputs, _brief(), out)
    assert error.value.code == "validation_failed"
    assert error.value.details["actual"] == 4
    assert error.value.details["required"] == 3
    assert (out / "deck.pptx").exists()  # available as diagnostic, never claimed ready


def test_maximum_checks_physical_pages_and_preserves_original_bound():
    prs = Presentation()
    for _ in range(4):
        prs.slides.add_slide(prs.slide_layouts[6])
    data = io.BytesIO()
    prs.save(data)
    validate_deck_count(_brief("maximum", requested_slide_count=5), data.getvalue())
    with pytest.raises(DeckDNAError) as error:
        validate_deck_count(_brief("maximum"), data.getvalue())
    assert error.value.details["actual"] == 4 and error.value.details["required"] == 3


def test_inline_plan_cannot_drop_the_callers_hard_limit(inputs, tmp_path):
    from deckdna.ingestion.content_parsers import parse_file
    from deckdna.planning import plan_deck

    source = parse_file(inputs[1])
    soft = Brief(purpose="Обзор", audience="команда", language="ru", target_slide_count=4)
    plan = plan_deck(source, soft)
    assert len(plan.slides) == 3
    hard = soft.model_copy(update={"slide_count_policy": "exact", "requested_slide_count": 4})
    with pytest.raises(DeckDNAError) as error:
        pipeline.generate(*inputs, hard, tmp_path / "inline", deck_plan=plan, content_pack=source)
    assert error.value.details["actual"] == 3 and error.value.details["required"] == 4


def test_pdf_cannot_hide_a_page_even_below_the_maximum(inputs, tmp_path, monkeypatch):
    def missing_page(source, target):
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.new_page()
            pdf.save(target)
        return target

    monkeypatch.setattr(pipeline, "render_pdf", missing_page)
    with pytest.raises(DeckDNAError) as error:
        pipeline.generate(*inputs, _brief("maximum", requested_slide_count=5), tmp_path / "pdf")
    assert error.value.details == {"rule": "request.export_count", "required": 3, "actual": 2}
