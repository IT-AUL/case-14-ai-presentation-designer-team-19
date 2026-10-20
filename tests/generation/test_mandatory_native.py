"""Required source content is checked in the deck, not merely plan metadata."""
import copy

import pymupdf
import pytest
from deckdna.contracts.deck_plan import Brief, DeckPlan
from deckdna.errors import DeckDNAError
from deckdna.generation import pipeline
from deckdna.ingestion.content_parsers import parse_file
from deckdna.planning.story_director import plan_deck
from pptx import Presentation

COPY = "В пилоте проверен один результат. Гарантий повторения пока нет."


@pytest.fixture
def inputs(tmp_path):
    template = tmp_path / "unseen.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(template)
    source = tmp_path / "source.md"
    source.write_text("# Обзор\n\n## Работа\n\nКоманда подготовила обзор результатов.\n\n"
                      f"## Риски\n\n{COPY}\n")
    pack = parse_file(source)
    brief = Brief(purpose="Обзор", audience="команда", language="ru", target_slide_count=4,
                  mandatory_sections=["Риски"], slide_count_policy="exact", content_fill="shrink")
    return template, source, pack, brief


@pytest.mark.parametrize("heading", ["Риски", "Следующие шаги", "Контакты"])
def test_required_source_is_preserved_in_native_and_pdf(inputs, tmp_path, heading):
    template, source, pack, brief = inputs
    if heading != "Риски":
        source.write_text(source.read_text().replace("## Риски", f"## {heading}"))
        pack = parse_file(source)
        brief = brief.model_copy(update={"mandatory_sections": [heading]})
    report = pipeline.generate(template, source, brief, tmp_path / "accepted", content_pack=pack)
    plan = DeckPlan.model_validate(report["deck_plan"])
    slide = next(s for s in plan.slides if any(u.text == COPY for u in s.content_units))
    assert slide.mandatory and next(u for u in slide.content_units if u.text == COPY).locked
    with pymupdf.open(report["artifacts"]["pdf"]) as pdf:
        assert len(pdf) == 4
        assert COPY in " ".join(pdf[slide.index].get_text().split())
        lines = [line for block in pdf[slide.index].get_text("dict")["blocks"]
                 for line in block.get("lines", [])
                 if COPY.split()[1] in "".join(s["text"] for s in line["spans"])]
        assert lines and all(abs(line["dir"][1]) < .01 and line["dir"][0] > .99
                             for line in lines), "horizontal prose was rotated by a donor layout"


def test_heading_and_empty_report_cannot_hide_required_content_loss(inputs, tmp_path, monkeypatch):
    original = pipeline.generate_deck

    def lose(*args, **kwargs):
        report = original(*args, **kwargs)
        path = args[2]
        prs = Presentation(path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and COPY in shape.text:
                    shape.text = shape.text.replace(COPY, "Иная тема без исходного содержания.")
        prs.save(path)
        report["dropped_units"] = {}
        return report

    def forbidden(*args, **kwargs):
        raise AssertionError("content loss must be checked before final rendering")

    monkeypatch.setattr(pipeline, "generate_deck", lose)
    monkeypatch.setattr(pipeline, "render_pdf", forbidden)
    template, source, pack, brief = inputs
    with pytest.raises(DeckDNAError) as exc:
        pipeline.generate(template, source, brief, tmp_path / "lost", content_pack=pack)
    assert exc.value.details["rule"] == "request.locked_copy"


def test_inline_plan_cannot_disable_the_original_request(inputs, tmp_path):
    template, source, pack, brief = inputs
    soft = brief.model_copy(update={"mandatory_sections": None})
    plan = plan_deck(pack, soft)
    plan = copy.deepcopy(plan)
    for slide in plan.slides:
        for unit in slide.content_units:
            if unit.text == COPY:
                unit.text = "Другой материал вместо обязательного источника."
    # The normal planner may split the source paragraph into title/body; neither
    # a title nor a fragment is a complete substantive source anchor.
    with pytest.raises(DeckDNAError) as exc:
        pipeline.generate(template, source, brief, tmp_path / "inline",
                          content_pack=pack, deck_plan=plan)
    assert exc.value.details["rule"] == "request.mandatory_section"
