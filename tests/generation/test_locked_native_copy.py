"""Native publication checks reject missing, relocated and silently edited copy."""
import io
from pathlib import Path

import pymupdf
import pytest
from deckdna.contracts.deck_plan import Brief, ContentUnit, Kind
from deckdna.errors import DeckDNAError
from deckdna.generation import pipeline
from deckdna.ingestion.content_parsers import parse_file
from deckdna.planning.locked_copy import validate_locked_deck, validate_locked_transition
from deckdna.planning.story_director import plan_deck
from pptx import Presentation

COPY = "Один проверенный результат без гарантий."
TITLE = "Результат."


@pytest.fixture
def inputs(tmp_path):
    template = tmp_path / "unseen.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(template)
    source = tmp_path / "source.md"
    source.write_text(f"# Обзор\n\n## Результат\n\n{COPY}\n")
    pack = parse_file(source)
    brief = Brief(purpose="Обзор", audience="команда", language="ru", target_slide_count=3,
                  content_fill="shrink")
    plan = plan_deck(pack, brief)
    slide = plan.slides[1].model_copy(update={"title_intent": TITLE, "content_units": [
        ContentUnit(role="title", kind=Kind.title, text=TITLE, locked=True),
        ContentUnit(role="body", kind=Kind.paragraph, text=COPY, locked=True),
    ]})
    plan = plan.model_copy(update={"slides": [plan.slides[0], slide, *plan.slides[2:]]})
    return template, source, pack, brief, plan


def _deck(texts):
    prs = Presentation()
    for text in texts:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 8000000, 4000000).text = text
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


@pytest.mark.parametrize("bad", ["Один проверенный результат с гарантиями.",
                                      "Один проверенный результат без гарантий",
                                      "один проверенный результат без гарантий."])
def test_compiled_copy_requires_verbatim_punctuation_and_qualification(inputs, bad):
    *_, plan = inputs
    with pytest.raises(DeckDNAError) as exc:
        validate_locked_deck(plan, _deck(["Обзор", TITLE + "\n" + bad, COPY]))
    assert exc.value.details["rule"] == "request.locked_copy"


def test_copy_on_a_different_slide_does_not_hide_loss(inputs):
    *_, plan = inputs
    with pytest.raises(DeckDNAError):
        validate_locked_deck(plan, _deck([COPY, TITLE, "Спасибо"]))


def test_physical_mapping_handles_inserted_template_slides(inputs):
    *_, plan = inputs
    deck = _deck([COPY, "Обзор", TITLE + "\n" + COPY, "Спасибо"])
    validate_locked_deck(plan, deck, [{"id": "protected:template"},
                                     *({"id": s.id} for s in plan.slides)])
    with pytest.raises(DeckDNAError) as exc:
        validate_locked_deck(plan, deck)
    assert exc.value.details["reason"] == "physical_slide_mapping_unavailable"


def test_unlocking_or_moving_copy_without_author_action_fails(inputs):
    *_, plan = inputs
    unlocked = plan.model_copy(deep=True)
    for unit in unlocked.slides[1].content_units:
        unit.locked = False
    with pytest.raises(DeckDNAError):
        validate_locked_transition(plan, unlocked)
    moved = plan.model_copy(deep=True)
    moved.slides[2].content_units += moved.slides[1].content_units
    moved.slides[1].content_units = []
    with pytest.raises(DeckDNAError):
        validate_locked_transition(plan, moved)
    validate_locked_transition(plan, moved, allow_moves=True)  # explicit split/merge


def test_locked_copy_is_generated_and_survives_pdf(inputs, tmp_path):
    template, source, pack, brief, plan = inputs
    report = pipeline.generate(template, source, brief, tmp_path / "accepted",
                               content_pack=pack, deck_plan=plan)
    validate_locked_deck(plan, Path(report["artifacts"]["pptx"]),
                         report["compose_report"]["slides"])
    with pymupdf.open(report["artifacts"]["pdf"]) as pdf:
        text = " ".join(pdf[1].get_text().split())
        assert TITLE in text and COPY in text


def test_corruption_after_compiler_is_caught_before_render(inputs, tmp_path, monkeypatch):
    original = pipeline.generate_deck

    def corrupt(*args, **kwargs):
        report = original(*args, **kwargs)
        output = args[2]
        prs = Presentation(output)
        for shape in prs.slides[1].shapes:
            if shape.has_text_frame and COPY in shape.text:
                shape.text = shape.text.replace(COPY, "Результат гарантирован.")
        prs.save(output)
        return report

    def forbidden(*args, **kwargs):
        raise AssertionError("locked copy must be checked before rendering")

    monkeypatch.setattr(pipeline, "generate_deck", corrupt)
    monkeypatch.setattr(pipeline, "render_pdf", forbidden)
    template, source, pack, brief, plan = inputs
    with pytest.raises(DeckDNAError) as exc:
        pipeline.generate(template, source, brief, tmp_path / "bad",
                          content_pack=pack, deck_plan=plan)
    assert exc.value.details["rule"] == "request.locked_copy"


def test_hidden_copy_and_speaker_notes_are_not_visible_preservation(inputs):
    *_, plan = inputs
    prs = Presentation(io.BytesIO(_deck(["Обзор", TITLE + "\n" + COPY, "Спасибо"])))
    shape = prs.slides[1].shapes[0]
    shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}nvSpPr/"
                        "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
                        ).set("hidden", "1")
    prs.slides[1].notes_slide.notes_text_frame.text = TITLE + "\n" + COPY
    out = io.BytesIO()
    prs.save(out)
    with pytest.raises(DeckDNAError):
        validate_locked_deck(plan, out.getvalue())


def test_one_native_occurrence_cannot_cover_duplicate_locked_units(inputs):
    *_, plan = inputs
    doubled = plan.model_copy(deep=True)
    doubled.slides[1].content_units.append(doubled.slides[1].content_units[-1])
    with pytest.raises(DeckDNAError):
        validate_locked_deck(doubled, _deck(["Обзор", TITLE + "\n" + COPY, "Спасибо"]))
