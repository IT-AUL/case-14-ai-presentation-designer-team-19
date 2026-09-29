"""OR-014/OR-026: невиданный шаблон должен проходить весь пайплайн.

synthetic_unseen.pptx строится scripts/build_synthetic_template.py —
структурно непохож на organizer-фикстуры: другая палитра (teal/amber,
не VK blue), Office-набор layouts, авторские карточки на blank-layout.
"""

import re
import shutil
import sys
import zipfile
from pathlib import Path

import pytest
from deckdna.audit.basic import audit_deck
from deckdna.generation.pipeline import generate
from deckdna.pptx.cloning.exemplar import select_exemplar_slides
from deckdna.pptx.opc.package import OpcPackage
from deckdna.template.autopsy import analyze_template
from lxml import etree
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from build_synthetic_template import build  # noqa: E402

FIXTURE = ROOT / "tests" / "fixtures" / "pptx" / "synthetic_unseen.pptx"
CONTENT = ROOT / "tests" / "fixtures" / "content" / "poc_article.md"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

BRIEF = {
    "purpose": "Показать сквозной пайплайн DeckDNA",
    "audience": "эксперты",
    "language": "ru",
    "target_slide_count": 12,
}

HAS_SOFFICE = shutil.which("soffice") is not None


def test_fixture_rebuildable(tmp_path):
    """Генератор воспроизводит валидный pptx из кода, без интернета."""
    out = build(tmp_path / "rebuilt.pptx")
    assert zipfile.is_zipfile(out)
    prs = Presentation(str(out))
    assert len(prs.slides) == 6


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    if not FIXTURE.exists():
        pytest.skip("synthetic fixture not present — run scripts/build_synthetic_template.py")
    if not HAS_SOFFICE:
        pytest.skip("soffice not installed — pdf stage needs it")
    out_dir = tmp_path_factory.mktemp("unseen_pipeline")
    return generate(str(FIXTURE), str(CONTENT), dict(BRIEF), out_dir)


def test_autopsy_on_unseen_structure():
    """Autopsy считает чужую структуру без допущений под organizer."""
    if not FIXTURE.exists():
        pytest.skip("synthetic fixture not present")
    f = analyze_template(FIXTURE)
    assert f.slides == 6
    assert f.dominant_layout is not None
    # чужая палитра — ни одного VK-синего 0077FF в observed top
    assert "0077FF" not in f.observed_colors
    assert set(f.observed_colors) & {"FFB547", "2A3B3C", "1E2A2B"}


def test_exemplar_gate_rejects_decorative_slide():
    """Побуквенная декоративная раскладка не попадает в пул exemplar'ов."""
    if not FIXTURE.exists():
        pytest.skip("synthetic fixture not present")
    pkg = OpcPackage.open(FIXTURE)
    choices = select_exemplar_slides(pkg, count=6)
    picked = {c.slide_part for c in choices}
    # slide5 — декоративный SKYLINE с однобуквенными runs
    assert "ppt/slides/slide5.xml" not in picked
    # контентные карточные слайды выбраны
    assert picked & {
        "ppt/slides/slide2.xml",
        "ppt/slides/slide3.xml",
        "ppt/slides/slide6.xml",
    }


def test_full_pipeline_on_unseen_template(generated):
    result = generated
    assert result["slides_out"] == BRIEF["target_slide_count"]
    for key in ("pptx", "pdf", "quality_passport"):
        assert Path(result["artifacts"][key]).exists()
    # паспорт и issues валидной формы
    assert result["quality_passport"]["metrics"]["validity"]["opens_cleanly"]


def test_no_stock_text_leak(generated):
    """Оригинальные строки шаблона не должны остаться в выводе —
    включая одиночные короткие runs ('12'), которые гейт раньше
    считал фрагментарной типографикой и не очищал."""
    stock = {"12", "99,97%", "4 200", "38 сек", "SKYLINE", "Наблюдаемость",
             "Тёмная сторона данных", "Q1 — Стабилизация"}
    zf = zipfile.ZipFile(generated["artifacts"]["pptx"])
    for name in zf.namelist():
        if not re.match(r"ppt/slides/slide\d+\.xml$", name):
            continue
        root = etree.fromstring(zf.read(name))
        texts = {(t.text or "").strip() for t in root.iter(f"{{{A}}}t")}
        leaked = texts & stock
        assert not leaked, f"{name}: утечка стокового текста {leaked}"


def test_audit_runs_on_unseen_deck(generated):
    """Аудит исполняется; issues если есть — валидной формы."""
    issues = audit_deck(generated["artifacts"]["pptx"])
    for i in issues:
        assert i.rule_code and i.severity in {"blocker", "error", "warning", "info"}
